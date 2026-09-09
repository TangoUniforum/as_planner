"""HarvestPlan rows must sum to the week the planner DECIDED.

THE INCIDENT (2026-09-08). A plan that never exceeded the 55,000-fish weekly
processing ceiling reported THREE breaches of it. 2028-W09 wrote

    B55 tank 41   26,593
    B55 tank 46   26,593
    B55 tank 52    1,815
                 --------
                  55,001      against a decision of exactly 55,000

Each event's count was rounded independently (`round(ev.count, 0)`), so three
rounded parts out-summed the whole by one fish. Advisory records the DECISION
and said 55,000; every gate that sums HarvestPlan rows -- the ceiling count,
the board, the Analyze checklist -- saw 55,001 and called it a breach.

That is expensive in exactly the way the other measurement bugs on this project
are: it does not crash, it produces a confident wrong number, and it argues
against a plan that is actually legal. Three phantom breaches were reported to
the operator, on a plan whose real count was zero.

FIX: round the WEEK, not the row. Floor every event, then hand the shortfall to
the largest fractional parts (largest remainder), so the rows always tie to the
week's rounded total. The tiebreak is the tank id -- a set-order tiebreak here
is the bug class that made the whole engine irreproducible earlier the same day
(commit 87ee040).
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date

import openpyxl

from forecast.excel_io import (whole_parts, write_daily_harvest_schedule,
                               write_harvest_plan_output, write_harvest_report)


@dataclass
class _Ev:
    event_date: date
    source_tank_id: int
    batch_id: str
    count: float
    avg_wt_g: float = 3450.0


def _rows(events):
    wb = openpyxl.Workbook()
    write_harvest_plan_output(wb, events, 0.86, {})
    ws = wb["HarvestPlan"]
    out = []
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith("20") and "-W" in r[0]:
            out.append((r[0], r[2], r[3]))
    return out


D = date(2028, 3, 1)


def test_the_incident_three_rows_that_summed_to_one_over():
    """26,592.5 + 26,592.5 + 1,815 = 55,000 exactly. It must not read 55,001."""
    evs = [_Ev(D, 41, "B55", 26592.5), _Ev(D, 46, "B55", 26592.5),
           _Ev(D, 52, "B55", 1815.0)]
    rows = _rows(evs)
    assert sum(r[2] for r in rows) == 55000
    assert all(float(r[2]).is_integer() for r in rows)


def test_every_row_stays_a_whole_fish():
    evs = [_Ev(D, 11, "B1", 100.4), _Ev(D, 12, "B1", 100.4), _Ev(D, 13, "B1", 100.2)]
    rows = _rows(evs)
    assert sum(r[2] for r in rows) == 301
    assert all(float(r[2]).is_integer() for r in rows)


def test_weeks_are_rounded_independently_of_each_other():
    """A shortfall must never be borrowed across a week boundary."""
    evs = [_Ev(date(2028, 3, 1), 11, "B1", 10.5), _Ev(date(2028, 3, 1), 12, "B1", 10.5),
           _Ev(date(2028, 3, 8), 11, "B1", 20.5), _Ev(date(2028, 3, 8), 12, "B1", 20.5)]
    rows = _rows(evs)
    per = {}
    for wk, _tank, cnt in rows:
        per[wk] = per.get(wk, 0) + cnt
    assert sorted(per.values()) == [21, 41]


def test_whole_numbers_are_left_alone():
    """The common case must be untouched — no drift introduced by the fix."""
    evs = [_Ev(D, 11, "B1", 30000.0), _Ev(D, 12, "B1", 25000.0)]
    assert sum(r[2] for r in _rows(evs)) == 55000


def test_a_single_event_week_still_ties():
    assert sum(r[2] for r in _rows([_Ev(D, 11, "B1", 54999.6)])) == 55000


def test_the_split_is_deterministic_across_input_order():
    """Same events, different order in -> same per-tank numbers out."""
    a = [_Ev(D, 41, "B1", 26592.5), _Ev(D, 46, "B1", 26592.5), _Ev(D, 52, "B1", 1815.0)]
    b = [_Ev(D, 52, "B1", 1815.0), _Ev(D, 46, "B1", 26592.5), _Ev(D, 41, "B1", 26592.5)]
    assert sorted(_rows(a)) == sorted(_rows(b))


def test_no_fish_are_invented_or_lost_overall():
    evs = [_Ev(D, t, "B1", 1000.0 + t * 0.37) for t in range(11, 21)]
    exact = round(sum(e.count for e in evs), 0)
    assert sum(r[2] for r in _rows(evs)) == exact


# ---- The other two writers had the SAME bug (2026-09-09) ----
# eae18ae fixed write_harvest_plan_output and left write_harvest_report and
# write_daily_harvest_schedule rounding each row on its own. The audit found
# the phantom still sitting one sheet over: HarvestReport read 55,001 against
# a 55,000 decision at 2028-W09, 8 weeks disagreed with HarvestPlan by a fish,
# and the Daily Harvest Schedule's day rows missed their own Total row in 58
# of 85 weeks. The rounding rule now lives in ONE function, `whole_parts`,
# which all three call.

def test_whole_parts_ties_to_the_whole():
    assert sum(whole_parts([26592.5, 26592.5, 1815.0])) == 55000
    assert sum(whole_parts([10.5, 10.5])) == 21
    assert whole_parts([]) == []
    assert sum(whole_parts([1000.0 + t * 0.37 for t in range(10)])) == round(
        sum(1000.0 + t * 0.37 for t in range(10)), 0)


def test_whole_parts_returns_whole_numbers_only():
    for v in whole_parts([26592.5, 26592.5, 1815.0]):
        assert float(v).is_integer()


def test_whole_parts_is_order_deterministic():
    """Ties break on position, so a sorted input gives a stable answer."""
    a = whole_parts([10.5, 10.5, 10.5, 10.5])
    assert a == whole_parts([10.5, 10.5, 10.5, 10.5])
    assert sum(a) == 42


def test_harvest_report_ties_to_the_week():
    """The incident: this sheet read 55,001 where HarvestPlan read 55,000."""
    evs = [_Ev(D, 41, "B55", 26592.5), _Ev(D, 46, "B55", 26592.5),
           _Ev(D, 52, "B55", 1815.0)]
    wb = openpyxl.Workbook()
    write_harvest_report(wb, evs, 0.86, {})
    total = sum(r[6] for r in wb["HarvestReport"].iter_rows(values_only=True)
                if r and isinstance(r[6], (int, float)))
    assert total == 55000


def test_the_two_harvest_sheets_agree():
    """They are the same events; they must not disagree by a fish."""
    evs = [_Ev(D, 41, "B55", 26592.5), _Ev(D, 46, "B55", 26592.5),
           _Ev(D, 52, "B55", 1815.0)]
    wb1 = openpyxl.Workbook(); write_harvest_plan_output(wb1, evs, 0.86, {})
    wb2 = openpyxl.Workbook(); write_harvest_report(wb2, evs, 0.86, {})
    plan = sum(r[3] for r in wb1["HarvestPlan"].iter_rows(values_only=True)
               if r and isinstance(r[3], (int, float)))
    rep = sum(r[6] for r in wb2["HarvestReport"].iter_rows(values_only=True)
              if r and isinstance(r[6], (int, float)))
    assert plan == rep == 55000


def test_daily_schedule_day_rows_tie_to_their_total():
    """58 of 85 weeks missed their own Total row by 1-2 fish."""
    wb = openpyxl.Workbook()
    write_daily_harvest_schedule(wb, [_Ev(D, 41, "B1", 23387.0)], None, 0.86, {})
    days, total = [], None
    for r in wb["Daily Harvest Schedule"].iter_rows(values_only=True):
        if not r or r[0] is None:
            continue
        if str(r[2]).strip() == "Total":
            total = r[5]
        elif isinstance(r[5], (int, float)):
            days.append(r[5])
    assert total is not None and sum(days) == total
