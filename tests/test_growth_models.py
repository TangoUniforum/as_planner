"""Per-batch seawater growth models (biology.yaml sgr_sw_by_model + batches.yaml
growth_model).

Operator request 2026-10-18: "add an actual growth model to different batches"
instead of only an SGR correction factor. Step 1 is the mechanism: a library of
named SW curves, each batch picks one, blank = the base curve. The tests pin:
blank changes nothing (plans identical), a named curve really is the one grown
on, and a name that does not exist is refused rather than quietly run on the
base curve.
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import math
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
import yaml

from forecast.biology import (BASE_GROWTH_MODEL, _require_growth_models,
                              sgr_pct_per_day, sw_growth_curve)
from forecast.config_io import biology_from_dict, biology_to_dict
from forecast.models import BatchInput, BiologyTables
from forecast.scenario_io import batches_from_list, batches_to_list

ROOT = Path(__file__).resolve().parent.parent
FIXTURE_PR = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"


def _tables(**models):
    return BiologyTables(sgr_size_g=[100.0, 1000.0, 5000.0],
                         sgr_fw_pct_day=[3.0, 2.0, 1.0],
                         sgr_sw_pct_day=[2.0, 1.0, 0.5],
                         sgr_sw_by_model=dict(models))


def _b(bid="B60", **kw):
    d = dict(batch_id=bid, input_date=datetime(2026, 1, 1), input_count=550000,
             tran_sf_date=datetime(2026, 3, 1), tran_og_date=datetime(2026, 12, 1),
             tran_og_count=290000, tran_og_avg_wt_g=370.0, tran_og_cv=16.0,
             fcr_model="FCR_116_Quick", fw_correction=1.0, sgr_correction=1.0)
    d.update(kw)
    return BatchInput(**d)


# --------------------------------------------------------------------------- #
# The growth rate itself
# --------------------------------------------------------------------------- #
def test_blank_growth_model_grows_on_the_base_curve():
    t = _tables(Slow=[1.0, 0.5, 0.25])
    for name in ("", BASE_GROWTH_MODEL):
        b = _b(growth_model=name)
        assert sgr_pct_per_day(1000.0, "SW", b, t) == pytest.approx(1.0)
        assert sw_growth_curve(t, b) is t.sgr_sw_pct_day


def test_named_growth_model_is_the_curve_grown_on_and_corrections_still_layer():
    t = _tables(Slow=[1.0, 0.5, 0.25])
    t.og_sgr_by_week = {"2026-W50": 0.9}
    b = _b(growth_model="Slow", sgr_correction=0.95)
    assert sgr_pct_per_day(1000.0, "SW", b, t) == pytest.approx(0.5 * 0.95)
    assert sgr_pct_per_day(1000.0, "SW", b, t, "2026-W50") == pytest.approx(
        0.5 * 0.95 * 0.9)


def test_growth_model_never_touches_freshwater():
    t = _tables(Slow=[1.0, 0.5, 0.25])
    assert sgr_pct_per_day(1000.0, "FW", _b(growth_model="Slow"), t) == \
        sgr_pct_per_day(1000.0, "FW", _b(), t)


def test_unknown_growth_model_is_refused_by_name():
    t = _tables(Slow=[1.0, 0.5, 0.25])
    with pytest.raises(ValueError, match="'Typo'.*B61"):
        _require_growth_models([_b(), _b("B61", growth_model="Typo")], t)
    with pytest.raises(ValueError, match="Typo"):
        sw_growth_curve(t, _b(growth_model="Typo"))
    _require_growth_models([_b(), _b(growth_model="Slow")], t)   # fine


def test_unsorted_size_rows_permute_the_growth_models_too():
    t = BiologyTables(sgr_size_g=[5000.0, 100.0], sgr_fw_pct_day=[1.0, 3.0],
                      sgr_sw_pct_day=[0.5, 2.0], sgr_sw_by_model={"S": [0.25, 1.0]})
    assert t.sgr_size_g == [100.0, 5000.0]
    assert t.sgr_sw_by_model["S"] == [1.0, 0.25]


# --------------------------------------------------------------------------- #
# YAML round trips stay byte-identical when the feature is unused
# --------------------------------------------------------------------------- #
def test_biology_without_growth_models_writes_no_new_key():
    d = biology_to_dict(_tables())
    assert "sgr_sw_by_model" not in d
    back = biology_from_dict(d)
    assert back.sgr_sw_by_model == {}


def test_biology_growth_models_round_trip_and_wrong_length_is_refused():
    d = biology_to_dict(_tables(Actual=[1.9, 0.9, None]))
    assert biology_from_dict(d).sgr_sw_by_model == {"Actual": [1.9, 0.9, None]}
    d["sgr_sw_by_model"]["Short"] = [1.0, 0.5]
    with pytest.raises(ValueError, match="'Short' has 2 values"):
        biology_from_dict(d)


def test_batches_round_trip_writes_growth_model_only_when_set():
    rows = batches_to_list([_b("B60"), _b("B61", growth_model="Actual")])
    assert "growth_model" not in rows[0]
    assert rows[1]["growth_model"] == "Actual"
    back = batches_from_list(rows)
    assert [b.growth_model for b in back] == ["", "Actual"]
    # A pandas grid gives the unset row NaN, which must read as blank.
    rows[0]["growth_model"] = float("nan")
    assert batches_from_list(rows)[0].growth_model == ""


# --------------------------------------------------------------------------- #
# End to end on the fixture PR
# --------------------------------------------------------------------------- #
PLAN_SHEETS = ("HarvestPlan", "TransferPlan", "BiologyProjection", "WeeklyReport")


def _run(tmp, copy_config, edit_bio=None, edit_batches=None):
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    if edit_bio:
        p = cfg / "biology.yaml"
        d = yaml.safe_load(p.read_text(encoding="utf-8"))
        edit_bio(d)
        p.write_text(yaml.safe_dump(d, sort_keys=False), encoding="utf-8")
    if edit_batches:
        p = scn / "batches.yaml"
        d = yaml.safe_load(p.read_text(encoding="utf-8"))
        edit_batches(d["batches"])
        p.write_text(yaml.safe_dump(d, sort_keys=False), encoding="utf-8")
    inp, out = tmp / "pr.xlsx", tmp / "out.xlsm"
    shutil.copy(FIXTURE_PR, inp)
    from forecast.run import main
    with contextlib.redirect_stdout(io.StringIO()):
        assert main(str(inp), str(out), config_dir=str(cfg),
                    scenario_dir=str(scn), calib_log_path="") == 0
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    h = {}
    for s in PLAN_SHEETS:
        rows = [r for r in wb[s].iter_rows(values_only=True)
                if not (r and isinstance(r[0], str) and r[0].startswith("Generated"))]
        h[s] = hashlib.md5(repr(rows).encode()).hexdigest()
    bi = [tuple(r) for r in wb["BatchInputs"].iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(bi) if r and r[0] == "Batch")
    col = list(bi[hi]).index("Growth model (SW)")
    labels = {}
    for r in bi[hi + 1:]:
        if not r or r[0] is None:
            break
        labels[r[0]] = r[col]
    harvest = sum(r[8] or 0 for r in wb["HarvestPlan"].iter_rows(min_row=5, values_only=True)
                  if r[0] and str(r[0]).startswith("20"))
    wb.close()
    return h, labels, harvest


def _slow(d):
    d["sgr_sw_by_model"] = {"Slow": [None if v is None else round(v * 0.8, 6)
                                     for v in d["sgr_sw_pct_day"]]}


@pytest.mark.skipif(not FIXTURE_PR.exists(), reason="fixture PR missing")
def test_end_to_end_library_unused_is_identical_and_used_changes_growth(
        tmp_path, copy_config):
    base_h, base_lbl, base_hog = _run(tmp_path / "base", copy_config)
    assert set(base_lbl.values()) == {BASE_GROWTH_MODEL}

    # The curve exists but no batch picks it: every plan sheet identical.
    lib_h, _, _ = _run(tmp_path / "lib", copy_config, edit_bio=_slow)
    assert lib_h == base_h

    # Every batch on a curve 20% slower: the plan must move, and the
    # BatchInputs sheet must name the curve. (Negative control for the
    # identity above: proves the hashes can differ.)
    def _all_slow(batches):
        for b in batches:
            b["growth_model"] = "Slow"
    slow_h, slow_lbl, slow_hog = _run(tmp_path / "slow", copy_config,
                                      edit_bio=_slow, edit_batches=_all_slow)
    assert slow_h["BiologyProjection"] != base_h["BiologyProjection"]
    assert set(slow_lbl.values()) == {"Slow"}
    assert slow_hog != base_hog and not math.isnan(slow_hog)
