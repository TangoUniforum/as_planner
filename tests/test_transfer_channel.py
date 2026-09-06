"""Every move carries the code path that emitted it, and TransferPlan exports it.

6N inflow is the sum of several independent channels, and the plan alone cannot
tell them apart. Measured on the live config: 16% more fish enter 6N in 2027
than the move-in controller asked for, and the obvious suspect (out-of-rotation
make-room dumps) turned out to be 0.3% of horizon inflow -- 10,857 fish, one
row. The real excess is grow-out housekeeping choosing a 6N tank as a
destination. None of that was visible before the tag.

DIAGNOSTIC ONLY. Nothing reads `channel` to decide anything; it defaults to
None so construction sites that do not set it need no change.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook

from forecast.events import Transfer
from forecast.excel_io import write_transfer_plan_output
from forecast.placement import TankAllocation

D = date(2026, 11, 11)


def _mv(src, dst, count, channel=None, batch="B45"):
    return Transfer(
        batch_id=batch, event_date=D, source_tank_id=src,
        destinations=[TankAllocation(tank_id=dst, count=count,
                                     avg_wt_g=3100.0, cv_pct=16.0)],
        count_transferred=count, channel=channel,
    )


def _sheet(events):
    wb = Workbook()
    write_transfer_plan_output(wb, events, [])
    ws = wb["TransferPlan"]
    header = [c.value for c in ws[4]]
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
    return header, body


def test_channel_defaults_to_none():
    # The ~30 sites that do not set it must stay constructible unchanged.
    assert _mv(31, 61, 100.0).channel is None


def test_column_is_exported_last():
    header, _ = _sheet([_mv(31, 61, 100.0, "rotation_fill")])
    assert header[-1] == "Channel"
    assert header[:9] == [
        "Week", "Batch", "Type", "From_Tank", "To_Tank",
        "Count (fish)", "Avg_Weight (kg)", "Grade", "CV (%)",
    ]


def test_channel_reaches_the_sheet():
    _, body = _sheet([_mv(31, 61, 100.0, "rotation_fill")])
    assert body[0][-1] == "rotation_fill"


def test_untagged_move_exports_blank_not_a_guess():
    _, body = _sheet([_mv(31, 61, 100.0)])
    assert body[0][-1] == ""


def test_same_week_legs_still_merge_to_one_row():
    """The merge key must NOT include the channel.

    A row is one physical pumping event and the handling budget counts rows, so
    splitting by channel would inflate the gate with moves nobody performs.
    """
    _, body = _sheet([_mv(31, 61, 100.0, "rotation_fill"),
                      _mv(31, 61, 40.0, "rotation_fill")])
    assert len(body) == 1
    assert body[0][5] == 140


def test_merged_legs_from_two_channels_read_mixed():
    # Attribution must not silently credit the row to whichever leg was built
    # first -- that would move fish between channels in the measurement.
    _, body = _sheet([_mv(31, 61, 100.0, "rotation_fill"),
                      _mv(31, 61, 40.0, "_equalize")])
    assert len(body) == 1
    assert body[0][5] == 140
    assert body[0][-1] == "mixed"
