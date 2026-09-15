"""A batch split across freshwater and seawater at the PR close.

OPERATOR DECISIONS (2026-09-11). A batch the ProductionReport holds partly in
freshwater and partly in seawater has its FW part MODELLED automatically:
  WHEN  the scenario tran_og_date, or the first forecast week once it has
        passed (og_entry_week_start(max(tran_og_date, forecast_start)));
  HOW   the PR's FW count through FW biology, culled to the REMAINING target
  MANY  (tran_og_count - the SW part at the PR close; B49 on 8/31 -> 242,257);
        remaining >= FW count -> no cull; remaining <= 0 -> no cull and a loud
        "target already met by the SW part";
  WHERE a top-up of the batch's OWN entry-tier stage-SW tanks, heaviest first
        (the big class to the heavier tank); a class that would breach the
        density target spills into empty entry tanks; never 6N / STARVE / OG3+;
  GROWTH the batch's configured fw_correction.
split_batch_fw: auto | off (default auto). A MANUAL fw_to_og always wins. An
overdue wholly-FW batch (tran_og_date before the PR close) moves in the first
forecast week. A SECOND manual fw_to_og for one batch is refused (today two
events create ~240,000 fish and every gate passes).

These are the no-workbook unit tests; tests/test_split_batch_integration.py
runs the real 8/31 PR.
"""
from __future__ import annotations

import dataclasses
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast.events import TranOGEntry
from forecast.models import BatchInput, BatchWeekState, ControlParams, SizeClassSplit
from forecast.state import FacilityState, TankState

ROOT = Path(__file__).resolve().parent.parent
PR_CLOSE = date(2026, 8, 31)
FS = date(2026, 9, 1)


# ---- fixtures ---------------------------------------------------------------

def _b49(**kw):
    b = BatchInput(
        batch_id="B49", input_date=datetime(2025, 9, 11), input_count=550_000,
        tran_sf_date=datetime(2025, 12, 8), tran_og_date=datetime(2026, 9, 14),
        tran_og_count=290_000, tran_og_avg_wt_g=370.0, tran_og_cv=16.0,
        fcr_model="FCR_116_Quick", fw_correction=0.9517, sgr_correction=0.95)
    return dataclasses.replace(b, **kw) if kw else b


def _b36():
    """B36 from the 2024-11-30 corpus registry (reconstructed there: its
    tran_og_count equals the SW count, i.e. the whole-batch target is met)."""
    return BatchInput(
        batch_id="B36", input_date=datetime(2024, 3, 2), input_count=118_008,
        tran_sf_date=datetime(2024, 6, 1), tran_og_date=datetime(2024, 11, 30),
        tran_og_count=106_207, tran_og_avg_wt_g=296.806, tran_og_cv=16.0,
        fcr_model="FCR_118_Quick", fw_correction=1.0, sgr_correction=1.0)


@pytest.fixture(scope="module")
def tables():
    from forecast.config_io import load_biology_tables
    return load_biology_tables(ROOT / "config")


@pytest.fixture(scope="module")
def control():
    from forecast.config_io import load_control
    c = load_control(ROOT / "config")
    c.forecast_start = datetime(2026, 9, 1)
    return c


def _og(batch, tank, count):
    return SimpleNamespace(batch_id=batch, tank_id=tank, closing_count=count,
                           closing_biomass_kg=count * 0.44, closing_avg_wt_g=440.0)


def _fw(batch, count, kg, unit="PostS.01"):
    return SimpleNamespace(batch_id=batch, unit_label=unit, fw_system="PostS",
                           closing_count=count, closing_biomass_kg=kg)


def _tank(tank_id, system_id, batch=None, count=0.0, wt=0.0, stage="SW",
          vol=1720.0, ttype="OG"):
    t = TankState(location_id=f"{system_id}-{tank_id}", tank_id=tank_id,
                  system_id=system_id, volume_m3=vol, max_density_kg_m3=85.0,
                  max_feed_kg_day_cap=3000.0, type=ttype)
    if batch:
        t.assign(batch_id=batch, count=count, avg_wt_g=wt, cv_pct=16.0, stage=stage)
    return t


def _row(batch, week, ws, stage, *, open_c, close_c, wt, close_wt=None, mort=0.0,
         cull=0.0, feed_w=0.0, feed_d=0.0):
    cw = wt if close_wt is None else close_wt
    n = (open_c + close_c) / 2.0
    return BatchWeekState(
        batch_id=batch, week_label=week, week_start=datetime.combine(ws, datetime.min.time()),
        days_since_input=300, week_from_input=43, count=n, avg_weight_g=(wt + cw) / 2,
        biomass_kg=n * (wt + cw) / 2000.0, feed_kg_day=feed_d, feed_kg_week=feed_w,
        sgr_pct_day=1.0, fcr=1.0, stage=stage, feed_type="F", mortality_pct_weekly=0.1,
        cull_count_week=cull, cull_biomass_kg_week=cull * wt / 1000.0,
        open_count=open_c, open_avg_weight_g=wt, open_biomass_kg=open_c * wt / 1000.0,
        close_count=close_c, close_avg_weight_g=cw, close_biomass_kg=close_c * cw / 1000.0,
        mort_count_week=mort)


# ---- the control key ----------------------------------------------------------

def test_split_batch_fw_defaults_to_auto_and_parses():
    from forecast import split_batch
    from forecast.config_io import control_from_dict, control_to_dict, load_control
    assert ControlParams.__dataclass_fields__["split_batch_fw"].default == "auto"
    d = control_to_dict(load_control(ROOT / "config"))
    d.pop("split_batch_fw", None)
    # a config written before the key existed loads as the default
    assert control_from_dict(d).split_batch_fw == "auto"
    assert control_from_dict({**d, "split_batch_fw": "off"}).split_batch_fw == "off"
    # an unquoted YAML `off` arrives as False
    import yaml
    assert split_batch.mode(SimpleNamespace(
        split_batch_fw=yaml.safe_load("k: off")["k"])) == "off"
    assert split_batch.mode(SimpleNamespace(split_batch_fw="auto")) == "auto"
    assert split_batch.mode(SimpleNamespace(split_batch_fw="OFF")) == "off"
    assert split_batch.mode(SimpleNamespace()) == "auto"
    with pytest.raises(ValueError, match="split_batch_fw"):
        split_batch.mode(SimpleNamespace(split_batch_fw="maybe"))


# ---- classification ------------------------------------------------------------

def test_split_batch_classification():
    from forecast.split_batch import classify
    meta = {"B49": _b49(), "B50": _b49(batch_id="B50", tran_og_date=datetime(2026, 10, 15)),
            "B51": _b49(batch_id="B51", tran_og_date=datetime(2026, 8, 27)),
            "B52": _b49(batch_id="B52", tran_og_date=datetime(2026, 8, 20)),
            "B60": _b49(batch_id="B60")}
    og = [_og("B49", 14, 20_000.0), _og("B49", 24, 27_743.0), _og("B41", 30, 9_000.0),
          _og("B60", 16, 5_000.0), _og("B99", 17, 1_000.0)]
    fw = [_fw("B49", 150_000.0, 55_000.0), _fw("B49", 100_225.0, 37_333.0, "PostS.02"),
          _fw("B50", 250_000.0, 60_000.0), _fw("B51", 200_000.0, 50_000.0),
          _fw("B52", 180_000.0, 40_000.0),
          _fw("B60", 0.0, 0.0),             # a zero-count FW row is ignored
          _fw("B99", 5_000.0, 1_000.0)]     # no registry row: not modellable
    c = classify(og, fw, transferred_fw=set(), batch_by_id=meta, forecast_start=FS)
    assert c.split_ids == ["B49"]
    assert c.sw_count["B49"] == pytest.approx(47_743.0)
    assert c.fw_count["B49"] == pytest.approx(250_225.0)
    assert c.fw_kg["B49"] == pytest.approx(92_333.0)
    # wholly-FW with tran_og_date before the close = overdue; a future one is not
    assert c.overdue_ids == ["B51", "B52"]
    # a manual fw_to_og wins: the batch is not auto-moved (split or overdue)
    c = classify(og, fw, transferred_fw={"B49", "B52"}, batch_by_id=meta, forecast_start=FS)
    assert c.split_ids == [] and c.overdue_ids == ["B51"]


def test_classification_reads_the_pr_not_the_post_window_state():
    """The post-window state has absorbed a manual fw_to_og; the split is a fact
    of the PR, so classify takes the PR's own records (no state argument)."""
    import inspect
    from forecast.split_batch import classify
    assert "state" not in inspect.signature(classify).parameters


# ---- how many fish, and when -----------------------------------------------------

def test_split_remaining_target():
    from forecast.split_batch import remaining_target
    assert remaining_target(290_000, 47_743.0) == pytest.approx(242_257.0)
    assert remaining_target(106_207, 106_207.0) == 0.0          # met -> no cull
    assert remaining_target(100_000, 120_000.0) == 0.0          # exceeded -> no cull
    assert remaining_target(None, 5.0) == 0.0                   # no target -> no cull
    assert remaining_target(0, 5.0) == 0.0


def test_b36_whole_batch_target_would_cull_half_the_fw_part(tables, control):
    """NEGATIVE CONTROL for the remaining-target rule. B36's registry target
    equals its SW count, so applying the whole-batch target to the FW
    remainder (the projector as it stands) culls ~47% of 201,807 fish; the
    remaining rule culls none (handling mortality only)."""
    from forecast.biology import project_in_flight_fw_batch
    from forecast.split_batch import fw_track_batch, remaining_target
    fs = date(2024, 12, 1)
    ctl = dataclasses.replace(control, forecast_start=datetime(2024, 12, 1),
                              horizon_weeks=6)
    b36 = _b36()
    whole = dataclasses.replace(b36, tran_og_date=datetime(2024, 12, 1))
    s_whole = project_in_flight_fw_batch(whole, tables, ctl, 201_807.0, 261.0,
                                         date(2024, 11, 30))
    tb = fw_track_batch(b36, forecast_start=fs,
                        target=remaining_target(b36.tran_og_count, 106_207.0))
    s_rem = project_in_flight_fw_batch(tb, tables, ctl, 201_807.0, 261.0,
                                       date(2024, 11, 30))
    cull_whole = sum(s.cull_count_week for s in s_whole[0])
    cull_rem = sum(s.cull_count_week for s in s_rem[0])
    assert cull_whole > 0.40 * 201_807
    assert cull_rem < 0.01 * 201_807                           # handling mortality only
    assert s_rem[2] and s_rem[2][0].post_cull_count > 199_000


def test_split_effective_date(tables, control):
    """B49 enters on its own scenario date (2026-09-14, forecast week 2); B36's
    date is already past at the close, so it enters in forecast week 0. A split
    is emitted in both cases -- under the raw past date the projector starts in
    seawater and emits none (the defect: B36 had 0 FW weeks)."""
    from forecast.biology import project_in_flight_fw_batch
    from forecast.split_batch import effective_tran_og_date, fw_track_batch
    from forecast.time_grid import og_entry_week_start, week_index
    e49 = effective_tran_og_date(_b49().tran_og_date, FS)
    assert e49 == date(2026, 9, 14)
    assert week_index(og_entry_week_start(e49, FS), FS) == 2
    fs36 = date(2024, 12, 1)
    e36 = effective_tran_og_date(_b36().tran_og_date, fs36)
    assert e36 == fs36 and week_index(og_entry_week_start(e36, fs36), fs36) == 0
    _s, _r, raw = project_in_flight_fw_batch(
        _b36(), tables, dataclasses.replace(control, forecast_start=datetime(2024, 12, 1)),
        201_807.0, 261.0, date(2024, 11, 30))
    assert raw == []                                           # the defect
    _s, _r, sp = project_in_flight_fw_batch(
        fw_track_batch(_b36(), forecast_start=fs36, target=0.0), tables,
        dataclasses.replace(control, forecast_start=datetime(2024, 12, 1)),
        201_807.0, 261.0, date(2024, 11, 30))
    assert len(sp) == 1
    st, _r, sp = project_in_flight_fw_batch(
        fw_track_batch(_b49(), forecast_start=FS, target=242_257.0), tables, control,
        250_225.0, 92_333.0 * 1000 / 250_225.0, PR_CLOSE)
    assert len(sp) == 1 and sp[0].post_cull_count == pytest.approx(242_257.0)
    fw_weeks = [s.week_label for s in st if s.stage == "FW"]
    assert fw_weeks == ["2026-W36", "2026-W37"]


def test_fw_track_batch_keeps_the_configured_growth_and_type():
    from forecast.split_batch import fw_track_batch
    b = _b49()
    t = fw_track_batch(b, forecast_start=FS, target=242_257.0)
    assert t.fw_correction == b.fw_correction                   # no auto-calibration
    assert t.tran_og_count == 242_257
    assert isinstance(t.tran_og_date, datetime) and t.tran_og_date.date() == date(2026, 9, 14)
    assert b.tran_og_count == 290_000                           # the registry is untouched
    past = fw_track_batch(_b49(tran_og_date=datetime(2026, 8, 20)), forecast_start=FS,
                          target=0.0)
    assert past.tran_og_date.date() == FS and not past.tran_og_count


# ---- one SW stream per batch -----------------------------------------------------------

def test_merge_split_states():
    from forecast.split_batch import merge_states
    W = [("2026-W36", date(2026, 9, 1)), ("2026-W37", date(2026, 9, 7)),
         ("2026-W38", date(2026, 9, 14))]
    og = [_row("B49", w, d, "SW", open_c=47_743 - 50 * i, close_c=47_693 - 50 * i,
               wt=440 + 10 * i, feed_w=100.0, feed_d=15.0) for i, (w, d) in enumerate(W)]
    fw = [_row("B49", W[0][0], W[0][1], "FW", open_c=250_225, close_c=250_100, wt=369,
               mort=125, feed_w=900.0, feed_d=130.0),
          _row("B49", W[1][0], W[1][1], "FW", open_c=250_100, close_c=249_993, wt=380,
               mort=107, feed_w=950.0, feed_d=140.0),
          _row("B49", W[2][0], W[2][1], "SW", open_c=249_993, close_c=242_000, wt=458,
               close_wt=470, mort=257, cull=7_736.0, feed_w=2_000.0, feed_d=300.0)]
    sw, fw_phase = merge_states(og, fw)
    assert [s.week_label for s in sw] == [w for w, _d in W]      # one row per week
    assert all(s.stage == "SW" for s in sw)
    assert sw[0] is og[0] and sw[1] is og[1]                     # pre-entry = OG track
    m = sw[2]
    for f in ("count", "biomass_kg", "feed_kg_day", "feed_kg_week", "cull_count_week",
              "cull_biomass_kg_week", "open_count", "open_biomass_kg", "close_count",
              "close_biomass_kg", "mort_count_week"):
        assert getattr(m, f) == pytest.approx(getattr(og[2], f) + getattr(fw[2], f)), f
    assert m.avg_weight_g == pytest.approx(m.biomass_kg * 1000.0 / m.count)
    assert m.close_avg_weight_g == pytest.approx(m.close_biomass_kg * 1000.0 / m.close_count)
    assert m.cull_count_week == pytest.approx(7_736.0)           # the crossing cull rides
    w_og, w_fw = og[2].biomass_kg, fw[2].biomass_kg
    assert m.sgr_pct_day == pytest.approx(
        (og[2].sgr_pct_day * w_og + fw[2].sgr_pct_day * w_fw) / (w_og + w_fw))
    assert [s.stage for s in fw_phase] == ["FW", "FW"]           # returned separately
    assert fw_phase[0] is fw[0]


def test_merge_with_no_fw_rows_is_the_og_track():
    from forecast.split_batch import merge_states
    og = [_row("B49", "2026-W36", date(2026, 9, 1), "SW", open_c=10, close_c=9, wt=400)]
    sw, fwp = merge_states(og, [])
    assert sw == og and fwp == []


def test_fw_addends_extra():
    from forecast.placement import fw_addends_by_week
    base = {"B50": [_row("B50", "2026-W36", date(2026, 9, 1), "FW", open_c=10_000,
                         close_c=9_990, wt=300, feed_d=40.0)]}
    extra = {"B49": [_row("B49", "2026-W36", date(2026, 9, 1), "FW", open_c=250_225,
                          close_c=250_100, wt=369, feed_d=130.0)]}
    assert fw_addends_by_week(base, extra_fw_states=None) == fw_addends_by_week(base)
    b0, f0 = fw_addends_by_week(base)
    b1, f1 = fw_addends_by_week(base, extra_fw_states=extra)
    assert b1["2026-W36"] == pytest.approx(b0["2026-W36"] + extra["B49"][0].biomass_kg)
    assert f1["2026-W36"] == pytest.approx(f0["2026-W36"] + 130.0)


# ---- where the arrival lands ----------------------------------------------------

def _split(big=121_128.5, big_wt=523.0, small=121_128.5, small_wt=405.0):
    return SizeClassSplit(batch_id="B49", tran_og_date=datetime(2026, 9, 14),
                          post_cull_count=big + small,
                          post_cull_avg_wt_g=(big * big_wt + small * small_wt) / (big + small),
                          post_cull_cv_pct=16.0, big_class_count=big,
                          big_class_avg_wt_g=big_wt, small_class_count=small,
                          small_class_avg_wt_g=small_wt)


def test_own_entry_tanks_are_heaviest_first_and_never_6n_starve_or_og3():
    from forecast.split_batch import own_entry_tanks
    st = FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0), _tank(24, "OG2S", "B49", 22_743, 480.0),
        _tank(16, "OG1N", "B49", 5_000, 500.0, stage="STARVE"),
        _tank(30, "OG3N", "B49", 9_000, 900.0), _tank(61, "OG6N", "B49", 4_000, 950.0),
        _tank(15, "OG1N", "B48", 30_000, 600.0), _tank(17, "OG1N")])
    assert [t.tank_id for t in own_entry_tanks(st, "B49")] == [24, 14]
    assert own_entry_tanks(st, "B50") == []


def test_split_topup_prefers_own_entry_tanks():
    """The heavier own tank takes the big class, the lighter one the small; no
    empty tank is needed (a top-up costs 0 moves)."""
    from forecast.split_batch import own_entry_tanks, plan_topup
    st = FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0), _tank(24, "OG2S", "B49", 22_743, 480.0),
        _tank(17, "OG1N")])
    tp = plan_topup(_split(), own_entry_tanks(st, "B49"), [st.tanks_by_id[17]],
                    density_target_pct=0.85)
    by = {(a.tank_id, a.size_class): a for a in tp.allocations}
    assert set(by) == {(24, "big"), (14, "small")}
    assert by[(24, "big")].count == pytest.approx(121_128.5)
    assert by[(24, "big")].avg_wt_g == pytest.approx(523.0)
    assert by[(14, "small")].avg_wt_g == pytest.approx(405.0)
    assert tp.empty_tanks_used == [] and not tp.warnings
    assert sum(a.count for a in tp.allocations) == pytest.approx(242_257.0)


def test_split_topup_spills_over_the_density_target_into_empty_tanks():
    from forecast.split_batch import own_entry_tanks, plan_topup, spill_tanks_needed
    # 500 m3 own tanks: 500 x 85 x 0.85 = 36,125 kg each at the density target
    st = FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0, vol=500.0),
        _tank(24, "OG2S", "B49", 22_743, 480.0, vol=500.0),
        _tank(17, "OG1N"), _tank(18, "OG1N")])
    own = own_entry_tanks(st, "B49")
    assert spill_tanks_needed(_split(), own, density_target_pct=0.85,
                              empty_cap_kg=1720 * 85 * 0.85) == 2
    # the real 1,720 m3 own tanks hold the whole cohort: no empty tank needed
    big_own = own_entry_tanks(FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0), _tank(24, "OG2S", "B49", 22_743, 480.0)]),
        "B49")
    assert spill_tanks_needed(_split(), big_own, density_target_pct=0.85,
                              empty_cap_kg=1720 * 85 * 0.85) == 0
    tp = plan_topup(_split(), own, [st.tanks_by_id[17], st.tanks_by_id[18]],
                    density_target_pct=0.85)
    kg = {}
    for a in tp.allocations:
        kg[a.tank_id] = kg.get(a.tank_id, 0.0) + a.count * a.avg_wt_g / 1000.0
    assert kg[24] + 22_743 * 0.48 == pytest.approx(36_125.0, rel=1e-6)   # filled to target
    assert kg[14] + 25_000 * 0.40 == pytest.approx(36_125.0, rel=1e-6)
    assert tp.empty_tanks_used == [17, 18]
    classes = {a.tank_id: a.size_class for a in tp.allocations if a.tank_id in (17, 18)}
    assert classes == {17: "big", 18: "small"}                  # surplus keeps its class
    assert sum(a.count for a in tp.allocations) == pytest.approx(242_257.0)  # nothing lost


def test_split_topup_with_no_empty_tank_keeps_every_fish_and_says_so():
    from forecast.split_batch import own_entry_tanks, plan_topup
    st = FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0, vol=500.0),
        _tank(24, "OG2S", "B49", 22_743, 480.0, vol=500.0)])
    tp = plan_topup(_split(), own_entry_tanks(st, "B49"), [], density_target_pct=0.85)
    assert sum(a.count for a in tp.allocations) == pytest.approx(242_257.0)
    assert tp.empty_tanks_used == []
    assert tp.warnings and "density target" in tp.warnings[0] and "B49" in tp.warnings[0]


def test_a_single_own_tank_takes_both_classes():
    from forecast.split_batch import own_entry_tanks, plan_topup
    st = FacilityState(today=FS, tanks=[_tank(24, "OG2S", "B49", 22_743, 480.0)])
    tp = plan_topup(_split(), own_entry_tanks(st, "B49"), [], density_target_pct=0.85)
    assert {a.tank_id for a in tp.allocations} == {24}
    assert {a.size_class for a in tp.allocations} == {"big", "small"}


def test_the_topup_lands_through_the_same_cohort_tranog():
    """TranOGEntry.apply is the one door: it merges a same-batch landing and
    refuses 6N / OG3+ / off-feed tanks (test_same_cohort_topup)."""
    from forecast.split_batch import own_entry_tanks, plan_topup
    st = FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0), _tank(24, "OG2S", "B49", 22_743, 480.0)])
    tp = plan_topup(_split(), own_entry_tanks(st, "B49"), [], density_target_pct=0.85)
    ev = TranOGEntry(batch_id="B49", event_date=date(2026, 9, 14), destinations=tp.allocations)
    assert ev.apply(st) == []
    assert ev.count_placed == pytest.approx(242_257.0) and ev.count_refused == 0
    assert st.tanks_by_id[24].count == pytest.approx(22_743 + 121_128.5)


# ---- a second manual fw_to_og ---------------------------------------------------

def _entry_state():
    return FacilityState(today=FS, tanks=[_tank(i, "OG1N") for i in (1, 2, 3, 4, 5, 6)])


def _fw_ev(batch, dests, count=None, week=1):
    from forecast.manual_events import ManualDest, ManualEvent
    return ManualEvent(type="fw_to_og", week=week, batch=batch, count=count,
                       destinations=[ManualDest(tank=t) for t in dests])


def test_second_fw_to_og_refused_in_the_run():
    from forecast.manual_events import apply_events_for_week
    lk = {("B49", "2026-W36"): (250_000.0, 370.0, 16.0),
          ("B50", "2026-W36"): (100_000.0, 300.0, 16.0)}
    evs = [_fw_ev("B49", [1, 2], 240_000), _fw_ev("B49", [3, 4], 240_000),
           _fw_ev("B50", [5, 6])]
    st = _entry_state()
    consumed: dict = {}
    _tr, _hv, tn, warns, fwb = apply_events_for_week(
        st, evs, 1, FS, week_label="2026-W36", fw_lookup=lk, consumed_fw=consumed)
    assert [e.batch_id for e in tn] == ["B49", "B50"]          # different batches both apply
    refused = [w for w in warns if w.startswith("MANUAL EVENT REFUSED")]
    assert len(refused) == 1 and "fw_to_og #2" in refused[0] and "B49" in refused[0]
    assert "already" in refused[0]
    assert fwb["B49"][0] == pytest.approx(250_000.0)           # the balance holds ONE event
    assert set(consumed) == {"B49", "B50"}
    assert st.tanks_by_id[3].is_empty and st.tanks_by_id[4].is_empty
    # the refusal carries across weeks: a week-2 event for B49 is refused too
    lk2 = {("B49", "2026-W37"): (249_000.0, 380.0, 16.0)}
    _tr, _hv, tn2, w2, _f = apply_events_for_week(
        st, [_fw_ev("B49", [3], week=2)], 2, date(2026, 9, 7), week_label="2026-W37",
        fw_lookup=lk2, consumed_fw=consumed)
    assert tn2 == [] and any(w.startswith("MANUAL EVENT REFUSED") for w in w2)


def test_second_fw_to_og_hole_without_the_guard():
    """NEGATIVE CONTROL -- documents the hole the guard closes: with no consumed
    set (the pre-2026-09-11 call) both events place the WHOLE FW part, 480,000
    fish from 250,000, and nothing downstream caught it (the over-production
    guard fires only above input_count x 1.001)."""
    from forecast.manual_events import apply_events_for_week
    lk = {("B49", "2026-W36"): (250_000.0, 370.0, 16.0)}
    st = _entry_state()
    _tr, _hv, tn, _w, _f = apply_events_for_week(
        st, [_fw_ev("B49", [1, 2], 240_000), _fw_ev("B49", [3, 4], 240_000)], 1, FS,
        week_label="2026-W36", fw_lookup=lk)
    assert len(tn) == 2
    assert sum(e.count_placed for e in tn) == pytest.approx(480_000.0)


def test_a_refused_first_event_does_not_consume_the_batch():
    """Only an fw_to_og that actually moved fish consumes the batch."""
    from forecast.manual_events import apply_events_for_week
    lk = {("B49", "2026-W36"): (250_000.0, 370.0, 16.0)}
    st = _entry_state()
    st.tanks_by_id[1].assign(batch_id="B48", count=10.0, avg_wt_g=500.0, cv_pct=16.0,
                             stage="SW")
    consumed: dict = {}
    _tr, _hv, tn, _w, _f = apply_events_for_week(
        st, [_fw_ev("B49", [1], 240_000), _fw_ev("B49", [3, 4], 240_000)], 1, FS,
        week_label="2026-W36", fw_lookup=lk, consumed_fw=consumed)
    assert [e.batch_id for e in tn] == ["B49"] and len(tn) == 1


def test_second_fw_to_og_refused_by_the_validator(tables, control):
    from forecast.manual_events import validate_manual_events
    st = _entry_state()
    res = validate_manual_events(
        st, [_fw_ev("B49", [1, 2], 240_000), _fw_ev("B49", [3, 4], 240_000)],
        batch_by_id={"B49": _b49()}, tables=tables, forecast_start=FS, control=control,
        pr_closing=PR_CLOSE, fw_records=[_fw("B49", 250_225.0, 92_333.0)])
    (i1, ok1, m1), (i2, ok2, m2) = res
    assert ok1, m1
    assert not ok2 and any("already" in m and "B49" in m for m in m2), m2


def test_the_editor_does_not_offer_an_already_scripted_batch(tables, control):
    import app
    ctx = dict(fw_records=[_fw("B49", 250_225.0, 92_333.0)], control=control,
               batch_by_id={"B49": _b49()}, tables=tables, pr_closing=PR_CLOSE)
    labels = ["2026-W36", "2026-W37"]
    assert "B49" in app._mw_fw_avail(ctx, labels)
    assert "B49" not in app._mw_fw_avail(ctx, labels, scripted={"B49"})
    assert "B49" in app._mw_fw_avail(ctx, labels, scripted={"B50"})


# ---- a past-date split can be scripted manually -----------------------------------

def test_manual_past_date_split_is_offered_week_1_at_the_pr_state(tables, control):
    """_build_fw_lookup indexed FW-stage weeks only, and a batch whose
    tran_og_date has passed is projected starting in SEAWATER -- no FW week, so
    no fw_to_og could ever be scripted for it (B36 on 2024-11-30). Fallback:
    week 1 carries the PR-measured FW state."""
    from forecast.manual_window import _build_fw_lookup
    past = _b49(tran_og_date=datetime(2026, 8, 20))
    ev = [_fw_ev("B49", [1, 2])]
    lk = _build_fw_lookup(ev, [_fw("B49", 250_225.0, 92_333.0)], control, PR_CLOSE,
                          tables, {"B49": past})
    assert lk[("B49", "2026-W36")] == pytest.approx((250_225.0, 369.0, 16.0), rel=1e-3)
    # a future date keeps the projected FW weeks (no fallback involved): week 2
    # is offered too. Each week offers the fish at its START
    # (manual_window.fw_week_start_states), so week 1 is the PR's own count and
    # week 2 carries a week of freshwater losses.
    lk = _build_fw_lookup(ev, [_fw("B49", 250_225.0, 92_333.0)], control, PR_CLOSE,
                          tables, {"B49": _b49()})
    assert lk[("B49", "2026-W36")][0] == pytest.approx(250_225.0)
    assert lk[("B49", "2026-W37")][0] < 250_225.0


def test_the_fallback_never_offers_eggs(tables, control):
    """A batch still in EGGS at week 1 has no FW state and gets no fallback."""
    from forecast.manual_window import _build_fw_lookup
    eggs = _b49(batch_id="B56", input_date=datetime(2026, 8, 1),
                tran_sf_date=datetime(2026, 11, 6), tran_og_date=datetime(2027, 8, 1))
    lk = _build_fw_lookup([_fw_ev("B56", [1])], [_fw("B56", 500_000.0, 0.0)], control,
                          PR_CLOSE, tables, {"B56": eggs})
    assert ("B56", "2026-W36") not in lk


# ---- the new ValidationLog lines ------------------------------------------------

def test_auto_transfer_line_names_batch_count_week_and_rule():
    from forecast.split_batch import auto_transfer_line
    ln = auto_transfer_line("B49", fw_count=250_225.0, week_label="2026-W38",
                            placed=242_257.0, tran_og_date=date(2026, 9, 14),
                            pr_closing=PR_CLOSE, tran_og_count=290_000, sw_count=47_743.0)
    assert ln.startswith("SPLIT BATCH B49: 250,225 FW fish auto-transferred 2026-W38 "
                         "per scenario tran_og_date")
    assert "script an fw_to_og to override" in ln
    assert "242,257" in ln and "47,743" in ln and "290,000" in ln
    past = auto_transfer_line("B36", fw_count=201_807.0, week_label="2024-W48",
                              placed=200_800.0, tran_og_date=date(2024, 11, 30),
                              pr_closing=date(2024, 11, 30), tran_og_count=106_207,
                              sw_count=106_207.0)
    assert "target already met by the SW part" in past
    assert "already past" in past and "first forecast week" in past
    over = auto_transfer_line("B50", fw_count=254_135.0, week_label="2026-W36",
                              placed=250_000.0, tran_og_date=date(2026, 8, 27),
                              pr_closing=PR_CLOSE, tran_og_count=260_000, sw_count=0.0,
                              overdue=True)
    assert over.startswith("OVERDUE FW BATCH B50: 254,135 FW fish auto-transferred 2026-W36")
    assert "script an fw_to_og to override" in over


def test_the_new_lines_land_in_their_own_validation_log_categories():
    from forecast import excel_io
    from forecast.split_batch import auto_transfer_line, not_placed_line
    lines = [
        auto_transfer_line("B49", fw_count=250_225.0, week_label="2026-W38",
                           placed=242_257.0, tran_og_date=date(2026, 9, 14),
                           pr_closing=PR_CLOSE, tran_og_count=290_000, sw_count=47_743.0),
        auto_transfer_line("B50", fw_count=254_135.0, week_label="2026-W36",
                           placed=250_000.0, tran_og_date=date(2026, 8, 27),
                           pr_closing=PR_CLOSE, tran_og_count=260_000, sw_count=0.0,
                           overdue=True),
        not_placed_line("B49", expected=242_257.0, placed=0.0, week_label="2026-W38"),
        "SPLIT BATCH AT PR CLOSE - B40: 1 fish (0 kg) were in freshwater",
    ]
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(wb, invariant_warnings=lines)
    cats = [r[1] for r in wb["ValidationLog"].iter_rows(values_only=True)
            if r and isinstance(r[0], int)]
    assert cats == ["WARNING - Split batch FW part auto-transferred",
                    "WARNING - Overdue FW batch auto-transferred",
                    "ERROR - Split batch FW part NOT placed",
                    "WARNING - Split batch at PR close (FW part not modelled)"]


def test_hybrid_guide_line_is_an_info_hybrid_guide_line():
    from forecast import excel_io
    from forecast.split_batch import hybrid_guide_line
    ln = hybrid_guide_line("B49", 242_257.0, "2026-W38")
    assert ln.startswith("HYBRID GUIDE - ") and "B49" in ln and "242,257" in ln
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(wb, invariant_warnings=[ln])
    (cat,) = [r[1] for r in wb["ValidationLog"].iter_rows(values_only=True)
              if r and isinstance(r[0], int)]
    assert cat == "INFO - Hybrid guide (L1) decision"


# ---- detection: an FW part that was modelled but never placed -------------------

def test_split_check_negative_control():
    """An auto-modelled FW part that did not reach seawater is flagged -- by
    name in the ValidationLog and as a DROP in InputConservationAudit -- and a
    manually handled split (no auto expectation) adds nothing."""
    from forecast import excel_io
    from forecast.split_batch import unplaced_split_parts
    ev = TranOGEntry(batch_id="B49", event_date=date(2026, 9, 14), destinations=[])
    ev.count_placed = 242_257.0
    assert unplaced_split_parts({"B49": 242_257.0}, [ev]) == {}
    assert unplaced_split_parts({"B49": 242_257.0}, []) == {"B49": pytest.approx(242_257.0)}
    short = TranOGEntry(batch_id="B49", event_date=date(2026, 9, 14), destinations=[])
    short.count_placed = 200_000.0
    assert unplaced_split_parts({"B49": 242_257.0}, [short]) == {
        "B49": pytest.approx(42_257.0)}
    assert unplaced_split_parts({}, [short]) == {}              # manual split: no line

    ctrl = SimpleNamespace(forecast_start=date(2026, 9, 1), horizon_weeks=10)
    bt = SimpleNamespace(batch_id="B49", input_count=550_000, tran_og_date=None,
                         tran_og_count=290_000)
    from forecast.placement import BatchLocationRow
    locs = [BatchLocationRow(week_label="2026-W36", week_start=date(2026, 9, 1),
                             batch_id="B49", tank_id=14, location_id="OG1S-14",
                             system_id="OG1S", count=47_743, avg_wt_g=440.0,
                             biomass_kg=21_000.0, density_kg_m3=12.0)]

    def audit(**kw):
        wb = openpyxl.Workbook()
        excel_io.write_input_conservation_audit(wb, [bt], locs, [], ctrl, **kw)
        rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
        return dict(zip(rows[hi], rows[hi + 1])), [str(r[0]) for r in rows[:hi] if r and r[0]]
    row, head = audit()
    assert row["Status"] == "PLACED" and head[2].startswith("OK")
    row, head = audit(unplaced_split_fw={"B49": 242_257.0})
    assert row["Status"] == "*** FW PART DROPPED ***"
    assert row["Fish_At_Risk (fish)"] == 242_257
    assert any("DROPPED" in h and "242,257" in h for h in head)
    assert not head[2].startswith("OK")


def test_the_audit_judges_a_split_against_its_remaining_target():
    """Realized 242,257 against the WHOLE-batch 290,000 read 'FW UNDER plan
    -16%'; the SW part was already in seawater. An auto split is judged against
    its remaining target."""
    from forecast import excel_io
    from forecast.placement import BatchLocationRow
    ctrl = SimpleNamespace(forecast_start=date(2026, 9, 1), horizon_weeks=10)
    bt = SimpleNamespace(batch_id="B49", input_count=550_000,
                         tran_og_date=date(2026, 9, 14), tran_og_count=290_000)
    locs = [BatchLocationRow(week_label="2026-W38", week_start=date(2026, 9, 14),
                             batch_id="B49", tank_id=14, location_id="OG1S-14",
                             system_id="OG1S", count=289_000, avg_wt_g=450.0,
                             biomass_kg=130_000.0, density_kg_m3=50.0)]
    ev = TranOGEntry(batch_id="B49", event_date=date(2026, 9, 14),
                     destinations=[SimpleNamespace(count=242_257.0, tank_id=14)])

    def flag(**kw):
        wb = openpyxl.Workbook()
        excel_io.write_input_conservation_audit(wb, [bt], locs, [], ctrl,
                                                tranog_events=[ev], **kw)
        rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
        return dict(zip(rows[hi], rows[hi + 1]))["FW_Flag"]
    assert flag() == "FW UNDER plan"                              # the old reading
    assert flag(split_remaining={"B49": 242_257.0}) == "auto split (remaining target)"
    assert "target met by the SW part" in flag(split_remaining={"B49": 0.0})


# ---- the ledger follows an auto-modelled split --------------------------------------

W1, W2 = "2026-W41", "2026-W42"
D1, D2 = date(2026, 10, 5), date(2026, 10, 12)
SW0, FW0 = 47_743.0, 250_225.0


def _auto_split_ledger_inputs():
    """B1 split at the close: SW part in tank 7 (4,000 g), FW part on its FW
    track. W1 is the last FW week (FW mort 232, FW feed 900 kg); W2 is the
    entry week: the crossing cull (7,736, TranOG_Date a week start) rides on the
    merged SW row and 242,257 fish top up tank 7."""
    from forecast.placement import BatchLocationRow

    def loc(wk, ws, c):
        return BatchLocationRow(week_label=wk, week_start=ws, batch_id="B1", tank_id=7,
                                location_id="OG1S-7", system_id="OG1S", count=c,
                                avg_wt_g=4_000.0, biomass_kg=c * 4.0, density_kg_m3=50.0)
    fw_close = FW0 - 232.0
    placed = fw_close - 7_736.0
    sw1_close = SW0 - 50.0
    mort2 = (sw1_close + placed) * 0.1 / 100.0
    og1 = _row("B1", W1, D1, "SW", open_c=SW0, close_c=sw1_close, wt=4_000.0)
    og1.mortality_pct_weekly = 50.0 / SW0 * 100.0
    fw1 = _row("B1", W1, D1, "FW", open_c=FW0, close_c=fw_close, wt=369.0, close_wt=380.0,
               mort=232.0, feed_w=900.0)
    merged2 = _row("B1", W2, D2, "SW", open_c=sw1_close + fw_close,
                   close_c=sw1_close + placed - mort2, wt=1_000.0, cull=7_736.0)
    merged2.mortality_pct_weekly = 0.1
    ev = SimpleNamespace(event_date=D2, batch_id="B1",
                         destinations=[SimpleNamespace(count=placed, avg_wt_g=380.0,
                                                       tank_id=7)])
    return dict(batch_locations=[loc(W1, D1, sw1_close), loc(W2, D2, sw1_close + placed - mort2)],
                harvest_events=[], batch_week_states=[og1, merged2], tranog_events=[ev],
                realized_biology={(7, W1, "B1"): [0.0, 50.0], (7, W2, "B1"): [0.0, mort2]},
                fw_openings={"B1": (FW0, FW0 * 0.369)}, fw_projected={"B1"},
                split_fw={"B1": [fw1]}), placed


def test_the_ledger_follows_an_auto_split():
    from forecast.excel_io import _build_batch_week_ledger, _ledger_value_cells
    kw, placed = _auto_split_ledger_inputs()
    r = {x["week"]: x for x in _build_batch_week_ledger(**kw) if x["batch"] == "B1"}
    # the opening holds every fish the PR holds, FW and SW
    assert r[W1]["open_count"] == pytest.approx(SW0 + FW0)
    # FW mortality and FW feed come from the FW track
    assert r[W1]["mort_count"] == pytest.approx(50.0 + 232.0)
    assert r[W1]["feed"] == pytest.approx(900.0)
    assert r[W1]["close_count"] == pytest.approx(SW0 - 50.0 + FW0 - 232.0)
    # chain, and the crossing is a move valued at the FW track's own weight
    assert r[W2]["open_count"] == pytest.approx(r[W1]["close_count"])
    assert r[W2]["open_bio"] == pytest.approx(r[W1]["close_bio"])
    assert r[W2]["xfer_in"] == r[W2]["xfer_out"] == pytest.approx(placed)
    assert r[W2]["input_count"] == 0.0
    assert r[W2]["cull_count"] == pytest.approx(7_736.0)       # booked to the FW phase, once
    for w in (W1, W2):
        assert r[w]["count_check"] == pytest.approx(0.0, abs=1.0), (w, r[w]["count_check"])
        assert not r[w]["rates_unknown"], w                    # real FW biology: rates print
    assert _ledger_value_cells(r[W1])[7] is not None           # SGR prints


def test_the_ledger_without_the_fw_track_loses_the_fw_part():
    """NEGATIVE CONTROL: the same run with the FW track withheld (fw_projected
    still names the batch, so nothing is held) opens on the SW part alone and
    the arrival appears from nowhere in Count_Check."""
    from forecast.excel_io import _build_batch_week_ledger
    kw, placed = _auto_split_ledger_inputs()
    kw.pop("split_fw")
    r = {x["week"]: x for x in _build_batch_week_ledger(**kw) if x["batch"] == "B1"}
    assert r[W1]["open_count"] == pytest.approx(SW0)
    assert abs(r[W2]["count_check"]) > 200_000
