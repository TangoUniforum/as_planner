"""Split FW/SW batch: behaviours an independent mutation proof found unguarded.

2026-09-12. Each test below failed on a one-line mutation of the split-batch
build (forecast/split_batch.py and its wiring) that every test in
tests/test_split_batch.py, tests/test_split_batch_integration.py and
tests/test_ledger_input_is_eggs.py let through, and passes on the real code.
The mutation each one kills is named in its docstring.

Unit tests first (no workbook); the end-to-end ones run the 8/31 PR on TEMP
copies of config/ (costs.yaml never copied) and scenario/.
"""
from __future__ import annotations

import contextlib
import dataclasses
import io
import shutil
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast.models import BatchWeekState, SizeClassSplit
from forecast.state import FacilityState, TankState

ROOT = Path(__file__).resolve().parent.parent
FS = date(2026, 9, 1)
PR_CLOSE = date(2026, 8, 31)


def _tank(tank_id, system_id, batch=None, count=0.0, wt=0.0, stage="SW", vol=1720.0):
    t = TankState(location_id=f"{system_id}-{tank_id}", tank_id=tank_id,
                  system_id=system_id, volume_m3=vol, max_density_kg_m3=85.0,
                  max_feed_kg_day_cap=3000.0, type="OG")
    if batch:
        t.assign(batch_id=batch, count=count, avg_wt_g=wt, cv_pct=16.0, stage=stage)
    return t


def _row(batch, week, ws, stage, *, open_c, close_c, wt, close_wt=None, mort=0.0,
         cull=0.0, feed_w=0.0, feed_d=0.0, sgr=1.0, fcr=1.0):
    cw = wt if close_wt is None else close_wt
    n = (open_c + close_c) / 2.0
    return BatchWeekState(
        batch_id=batch, week_label=week, week_start=datetime.combine(ws, datetime.min.time()),
        days_since_input=300, week_from_input=43, count=n, avg_weight_g=(wt + cw) / 2,
        biomass_kg=n * (wt + cw) / 2000.0, feed_kg_day=feed_d, feed_kg_week=feed_w,
        sgr_pct_day=sgr, fcr=fcr, stage=stage, feed_type="F", mortality_pct_weekly=0.1,
        cull_count_week=cull, cull_biomass_kg_week=cull * wt / 1000.0,
        open_count=open_c, open_avg_weight_g=wt, open_biomass_kg=open_c * wt / 1000.0,
        close_count=close_c, close_avg_weight_g=cw, close_biomass_kg=close_c * cw / 1000.0,
        mort_count_week=mort)


# ---- where the arrival lands: the force-empty floor ---------------------------------

def _one_class(n, wt):
    return SizeClassSplit(batch_id="B49", tran_og_date=datetime(2026, 9, 14),
                          post_cull_count=n, post_cull_avg_wt_g=wt, post_cull_cv_pct=16.0,
                          big_class_count=n, big_class_avg_wt_g=wt,
                          small_class_count=0.0, small_class_avg_wt_g=wt)


def _small_own_tank():
    """One own 500 m3 tank: 36,125 kg at the 85% density target, 10,000 kg in
    it. A 53,000-fish class at 500 g fits 52,250; 750 fish are left over."""
    return FacilityState(today=FS, tanks=[
        _tank(14, "OG1S", "B49", 25_000, 400.0, vol=500.0), _tank(17, "OG1N")])


def test_a_spill_below_min_tank_control_stays_home():
    """Mutation killed: plan_topup's `rest >= min_tank_control` spill guard
    dropped. A 750-fish remainder must NOT open empty tank 17 under the
    force-empty floor (5,000): it stays in the own tank, past the target,
    and the warning names the floor."""
    from forecast.split_batch import own_entry_tanks, plan_topup
    st = _small_own_tank()
    tp = plan_topup(_one_class(53_000.0, 500.0), own_entry_tanks(st, "B49"),
                    [st.tanks_by_id[17]], density_target_pct=0.85, min_tank_control=5_000)
    assert tp.empty_tanks_used == []
    assert {a.tank_id for a in tp.allocations} == {14}
    assert sum(a.count for a in tp.allocations) == pytest.approx(53_000.0)
    assert tp.warnings and "min_tank_control" in tp.warnings[0]
    # the same remainder with no floor does spill (the guard is what stops it)
    tp0 = plan_topup(_one_class(53_000.0, 500.0), own_entry_tanks(st, "B49"),
                     [st.tanks_by_id[17]], density_target_pct=0.85, min_tank_control=0.0)
    assert tp0.empty_tanks_used == [17]


def test_the_move_credit_counts_the_force_empty_floor():
    """Mutation killed: spill_tanks_needed stops passing min_tank_control to
    plan_topup. The arrival make-room / pacing would then free an entry tank
    (a move against the hard weekly cap) the top-up never uses."""
    from forecast.split_batch import own_entry_tanks, spill_tanks_needed
    own = own_entry_tanks(_small_own_tank(), "B49")
    kw = dict(density_target_pct=0.85, empty_cap_kg=1720 * 85 * 0.85)
    assert spill_tanks_needed(_one_class(53_000.0, 500.0), own, **kw) == 1
    assert spill_tanks_needed(_one_class(53_000.0, 500.0), own,
                              min_tank_control=5_000, **kw) == 0


# ---- one SW stream per batch: the rates -------------------------------------------

def test_merged_sgr_and_fcr_are_biomass_weighted():
    """Mutations killed: the merged entry-week row takes SGR from the OG track
    alone, or FCR from the FW track alone. (test_merge_split_states gives
    every row SGR 1.0 and FCR 1.0, so any weighting passed.)"""
    from forecast.split_batch import merge_states
    wk, ws = "2026-W38", date(2026, 9, 14)
    og = [_row("B49", wk, ws, "SW", open_c=47_000, close_c=46_950, wt=520.0,
               sgr=0.60, fcr=1.30)]
    fw = [_row("B49", wk, ws, "SW", open_c=250_000, close_c=242_000, wt=380.0,
               sgr=1.40, fcr=0.90)]
    (m,), _fwp = merge_states(og, fw)
    wo, wf = og[0].biomass_kg, fw[0].biomass_kg
    assert m.sgr_pct_day == pytest.approx((0.60 * wo + 1.40 * wf) / (wo + wf))
    assert m.fcr == pytest.approx((1.30 * wo + 0.90 * wf) / (wo + wf))
    assert 0.60 < m.sgr_pct_day < 1.40 and 0.90 < m.fcr < 1.30


# ---- determinism: natural batch order ------------------------------------------------

def test_split_ids_come_in_number_order_whatever_the_pr_order():
    """Mutations killed: classify walking the PR's own row order, and
    sorted_ids sorting batch ids as strings (B100 before B9). run.py walks
    sorted_ids for its lines, the report states and the audit."""
    from forecast.models import BatchInput
    from forecast.split_batch import classify, sorted_ids

    def meta(b, tog):
        return BatchInput(batch_id=b, input_date=datetime(2025, 9, 11), input_count=550_000,
                          tran_sf_date=datetime(2025, 12, 8), tran_og_date=tog,
                          tran_og_count=290_000, tran_og_avg_wt_g=370.0, tran_og_cv=16.0,
                          fcr_model="FCR_116_Quick", fw_correction=1.0, sgr_correction=1.0)
    ids = ["B100", "B9", "B41", "B10", "B2"]
    split, overdue = {"B100", "B9", "B41"}, {"B10", "B2"}
    by = {b: meta(b, datetime(2026, 9, 14) if b in split else datetime(2026, 8, 20))
          for b in ids}
    og = [SimpleNamespace(batch_id=b, tank_id=i, closing_count=1_000.0,
                          closing_biomass_kg=400.0) for i, b in enumerate(ids) if b in split]
    fw = [SimpleNamespace(batch_id=b, unit_label="PostS.01", fw_system="PostS",
                          closing_count=10_000.0, closing_biomass_kg=3_000.0) for b in ids]
    c = classify(og, fw, transferred_fw=set(), batch_by_id=by, forecast_start=FS)
    assert c.split_ids == ["B9", "B41", "B100"]
    assert c.overdue_ids == ["B2", "B10"]
    assert c.auto_ids == ["B2", "B9", "B10", "B41", "B100"]
    assert sorted_ids({"B100", "B9", "B41", "B10"}) == ["B9", "B10", "B41", "B100"]
    assert sorted_ids({"B100": 1, "B9": 2}) == ["B9", "B100"]


def test_the_fw_part_dropped_headline_lists_batches_by_number():
    """Mutation killed: the FW PART DROPPED headline joining sorted() over the
    ids (a string sort: B10, B100, B9). T7's token scan cannot see it -- the
    expression names no 'batch'."""
    from forecast import excel_io
    ctrl = SimpleNamespace(forecast_start=FS, horizon_weeks=10)
    wb = openpyxl.Workbook()
    excel_io.write_input_conservation_audit(
        wb, [], [], [], ctrl, unplaced_split_fw={"B100": 1.0, "B10": 1.0, "B9": 1.0})
    head = [str(r[0]) for r in wb["InputConservationAudit"].iter_rows(values_only=True)
            if r and r[0] and "FW PART DROPPED" in str(r[0])]
    assert head and "B9, B10, B100" in head[0], head


# ---- the audit: a split off its remaining target ------------------------------------

def _audit_flag(realized, remaining):
    from forecast import excel_io
    from forecast.events import TranOGEntry
    from forecast.placement import BatchLocationRow
    ctrl = SimpleNamespace(forecast_start=FS, horizon_weeks=10)
    bt = SimpleNamespace(batch_id="B49", input_count=550_000,
                         tran_og_date=date(2026, 9, 14), tran_og_count=290_000)
    locs = [BatchLocationRow(week_label="2026-W38", week_start=date(2026, 9, 14),
                             batch_id="B49", tank_id=14, location_id="OG1S-14",
                             system_id="OG1S", count=realized + 47_000, avg_wt_g=450.0,
                             biomass_kg=130_000.0, density_kg_m3=50.0)]
    ev = TranOGEntry(batch_id="B49", event_date=date(2026, 9, 14),
                     destinations=[SimpleNamespace(count=realized, tank_id=14)])
    wb = openpyxl.Workbook()
    excel_io.write_input_conservation_audit(wb, [bt], locs, [], ctrl, tranog_events=[ev],
                                            split_remaining={"B49": remaining})
    rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
    head = [str(r[0]) for r in rows[:hi] if r and r[0]]
    return dict(zip(rows[hi], rows[hi + 1]))["FW_Flag"], head


@pytest.mark.parametrize("remaining,flag", [(300_000.0, "auto split: UNDER remaining target"),
                                            (200_000.0, "auto split: OVER remaining target")])
def test_a_split_off_its_remaining_target_is_flagged(remaining, flag):
    """Mutations killed: an auto split UNDER (or OVER) its remaining target
    reading 'auto split (remaining target)' -- the in-band label -- and
    missing from the >5% divergence NOTE. Only the in-band and target-met
    labels were tested."""
    got, head = _audit_flag(242_257.0, remaining)
    assert got == flag
    assert any(h.startswith("NOTE:") and "B49" in h for h in head), head


# ---- a second manual fw_to_og: the validator mirrors the run -------------------------

def test_the_validator_does_not_consume_a_batch_whose_event_placed_nothing(monkeypatch):
    """Mutation killed: validate_manual_events consuming the batch on an
    fw_to_og that placed no fish. The run already has this rule
    (test_a_refused_first_event_does_not_consume_the_batch); the validator's
    faithful path must agree, or the editor marks a valid event invalid."""
    from forecast import manual_events as me
    from forecast.events import TranOGEntry
    calls = []

    def fake_apply(state, ev, i, count, wt, cv, handling, event_date=None,
                   out_tranog=None, **kw):
        calls.append(i)
        placed = 0.0 if i == 1 else count           # the FIRST event places nothing
        e = TranOGEntry(batch_id=ev.batch, event_date=event_date, destinations=[])
        e.count_placed = placed
        if out_tranog is not None:
            out_tranog.append(e)
        return ([] if placed else ["❌ destination occupied"]), 0.0
    monkeypatch.setattr(me, "_apply_fw_to_og", fake_apply)
    import forecast.manual_window as mw
    monkeypatch.setattr(mw, "_build_fw_lookup",
                        lambda *a, **k: {("B49", "2026-W36"): (250_000.0, 370.0, 16.0)})
    evs = [me.ManualEvent(type="fw_to_og", week=1, batch="B49", count=240_000,
                          destinations=[me.ManualDest(tank=1)]),
           me.ManualEvent(type="fw_to_og", week=1, batch="B49", count=240_000,
                          destinations=[me.ManualDest(tank=3)])]
    st = FacilityState(today=FS, tanks=[_tank(i, "OG1N") for i in (1, 2, 3, 4)])
    from forecast.config_io import load_biology_tables, load_control
    ctl = load_control(ROOT / "config")
    ctl.forecast_start = datetime(2026, 9, 1)
    res = me.validate_manual_events(
        st, evs, batch_by_id={}, tables=load_biology_tables(ROOT / "config"),
        forecast_start=FS, control=ctl, pr_closing=PR_CLOSE, fw_records=[])
    (_i1, _ok1, _m1), (_i2, ok2, m2) = res
    assert calls == [1, 2], calls
    assert ok2, m2
    assert not any("already transferred" in m for m in m2), m2


# ---- app: the Configure knob and the intake editor ---------------------------------

def test_split_batch_fw_is_not_tunable():
    """Mutation killed: split_batch_fw dropped from UNTUNABLE_KNOBS. A search
    free to switch it off would score a lighter facility by forgetting
    ~240,000 fish the PR says are there."""
    from forecast.methods import UNTUNABLE_KNOBS
    assert "split_batch_fw" in UNTUNABLE_KNOBS


def test_split_batch_fw_sits_in_the_horizon_and_scenario_group():
    """Mutation killed: the knob dropped from its Configure group (it would
    fall to 'Everything else', away from the scenario/6N settings it belongs
    with)."""
    pytest.importorskip("streamlit")
    import app
    (keys,) = [k for title, _b, k in app._CONTROL_GROUPS if "sixn_transition_weeks" in k]
    assert "split_batch_fw" in keys


def test_the_intake_editor_hides_already_scripted_batches(monkeypatch):
    """Mutation killed: _mw_fw_intake calling _mw_fw_avail WITHOUT the
    scripted set (test_the_editor_does_not_offer_an_already_scripted_batch
    tests _mw_fw_avail alone, not that the editor passes it)."""
    pytest.importorskip("streamlit")
    import app
    seen = {}

    def spy(ctx, labels, scripted=None):
        seen["scripted"] = scripted
        return {}
    captions = []
    monkeypatch.setattr(app, "_mw_fw_avail", spy)
    monkeypatch.setattr(app, "_mw_events", lambda: [
        SimpleNamespace(type="fw_to_og", batch="B100"),
        SimpleNamespace(type="fw_to_og", batch="B49"),
        SimpleNamespace(type="harvest", batch="B41")])
    monkeypatch.setattr(app, "st", SimpleNamespace(caption=lambda s, *a, **k: captions.append(s)))
    app._mw_fw_intake(None, {}, [], ["2026-W36"], None)
    assert list(seen["scripted"] or ()) == ["B49", "B100"], seen
    assert any("B49, B100" in c for c in captions), captions


# ---- ledger: a mid-week crossing cull rides on the FW track ------------------------

W1, W2 = "2026-W41", "2026-W42"
D1, D2 = date(2026, 10, 5), date(2026, 10, 12)
SW0, FW0 = 47_743.0, 250_225.0


def test_the_ledger_books_a_mid_week_crossing_cull_from_the_fw_track():
    """Mutation killed: the ledger ignoring cull_count_week on the split's FW
    track. With a MID-WEEK tran_og_date the reconcile cull falls in the last
    freshwater week (on the FW track's own row), not on the merged SW row;
    tests/test_split_batch.py covers only the week-start case (the 8/31 B49
    date is a Monday)."""
    from forecast.excel_io import _build_batch_week_ledger
    from forecast.placement import BatchLocationRow

    def loc(wk, ws, c):
        return BatchLocationRow(week_label=wk, week_start=ws, batch_id="B1", tank_id=7,
                                location_id="OG1S-7", system_id="OG1S", count=c,
                                avg_wt_g=4_000.0, biomass_kg=c * 4.0, density_kg_m3=50.0)
    cull = 7_736.0
    fw_close = FW0 - 232.0 - cull                    # culled inside the last FW week
    placed = fw_close
    sw1_close = SW0 - 50.0
    mort2 = (sw1_close + placed) * 0.1 / 100.0
    og1 = _row("B1", W1, D1, "SW", open_c=SW0, close_c=sw1_close, wt=4_000.0)
    og1.mortality_pct_weekly = 50.0 / SW0 * 100.0
    fw1 = _row("B1", W1, D1, "FW", open_c=FW0, close_c=fw_close, wt=369.0, close_wt=380.0,
               mort=232.0, cull=cull, feed_w=900.0)
    merged2 = _row("B1", W2, D2, "SW", open_c=sw1_close + fw_close,
                   close_c=sw1_close + placed - mort2, wt=1_000.0)
    merged2.mortality_pct_weekly = 0.1
    ev = SimpleNamespace(event_date=D2, batch_id="B1",
                         destinations=[SimpleNamespace(count=placed, avg_wt_g=380.0,
                                                       tank_id=7)])
    rows = _build_batch_week_ledger(
        batch_locations=[loc(W1, D1, sw1_close), loc(W2, D2, sw1_close + placed - mort2)],
        harvest_events=[], batch_week_states=[og1, merged2], tranog_events=[ev],
        realized_biology={(7, W1, "B1"): [0.0, 50.0], (7, W2, "B1"): [0.0, mort2]},
        fw_openings={"B1": (FW0, FW0 * 0.369)}, fw_projected={"B1"},
        split_fw={"B1": [fw1]})
    r = {x["week"]: x for x in rows if x["batch"] == "B1"}
    assert r[W1]["cull_count"] == pytest.approx(cull)
    assert r[W2]["cull_count"] == pytest.approx(0.0)
    for w in (W1, W2):
        assert r[w]["count_check"] == pytest.approx(0.0, abs=1.0), (w, r[w]["count_check"])


# ---- end to end: the wiring the unit tests cannot see ------------------------------

PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")
_needs_pr = pytest.mark.skipif(not PR_0831.exists(),
                               reason="the 2026-08-31 ProductionReport is not on this machine")


def _rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _read(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tp = _rows(wb["TransferPlan"])
    th = next(i for i, r in enumerate(tp) if r and r[0] == "Week")
    col = {str(c): j for j, c in enumerate(tp[th]) if c is not None}
    tog: dict = {}
    for r in tp[th + 1:]:
        if r and r[col["Type"]] == "TranOG":
            k = (str(r[col["Batch"]]), str(r[col["Week"]]))
            tog[k] = tog.get(k, 0.0) + float(r[col["Count (fish)"]] or 0)
    vlog = [(str(r[1]), str(r[2])) for r in _rows(wb["ValidationLog"])
            if r and isinstance(r[0], int)]
    ica = _rows(wb["InputConservationAudit"])
    hi = next(i for i, r in enumerate(ica) if r and r[0] == "Batch")
    out = SimpleNamespace(
        tog=tog, vlog=vlog,
        ica={r[0]: dict(zip(ica[hi], r)) for r in ica[hi + 1:] if r and r[0]},
        ica_head=[str(r[0]) for r in ica[:hi] if r and r[0]])
    mr = _rows(wb["MonthlyReport"])
    mh = next(i for i, r in enumerate(mr) if r and r[0] == "Scenario")
    hdr = [str(c) if c is not None else "" for c in mr[mh]]
    out.monthly = [dict(zip(hdr, r)) for r in mr[mh + 1:]
                   if r and any(c is not None for c in r)]
    wb.close()
    return out


def _run_831(tmp, copy_config, *, control_edit=None, scenario_edit=None, patches=()):
    """The 8/31 PR with an EMPTY events file (B49's FW part is then moved by the
    automatic path), on TEMP copies of config/ (no costs.yaml) and scenario/."""
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    if control_edit is not None:
        p = cfg / "control.yaml"
        p.write_text(control_edit(p.read_text(encoding="utf-8")), encoding="utf-8")
    ev = scn / "manual_events" / "2026-08-31.yaml"
    ev.parent.mkdir(exist_ok=True)
    ev.write_text("events: []\n", encoding="utf-8")
    if scenario_edit is not None:
        scenario_edit(scn)
    inp, out = tmp / "pr.xlsm", tmp / "out.xlsm"
    shutil.copy(PR_0831, inp)
    from forecast.run import main
    with pytest.MonkeyPatch.context() as mp:
        for obj, name, val in patches:
            mp.setattr(obj, name, val)
        with contextlib.redirect_stdout(io.StringIO()):
            rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                      calib_log_path="")
    assert rc == 0
    if not out.exists():
        out = out.with_suffix(".xlsx")
    return _read(out)


def _set_b50(scn, **fields):
    p = scn / "batches.yaml"
    txt = p.read_text(encoding="utf-8")
    i = txt.find("- batch_id: B50")
    if i < 0:
        pytest.skip("B50 is not in the live batches.yaml")
    for k, v in fields.items():
        j = txt.index(f"{k}:", i)
        e = txt.index("\n", j)
        txt = txt[:j] + f"{k}: {v}" + txt[e:]
    p.write_text(txt, encoding="utf-8")


@pytest.fixture(scope="module")
def overdue_low(tmp_path_factory, copy_config):
    """B50 overdue (tran_og_date 2026-08-27, before the close) with its
    tran_og_count lowered to 200,000 -- below the ~254,000 FW fish, so the
    reconcile cull has work to do (at the live 290,000 it culls nothing, and
    dropping the cull changed no test). A spy records what the placement's
    FW-inclusive caps are handed."""
    if not PR_0831.exists():
        pytest.skip("the 2026-08-31 ProductionReport is not on this machine")
    import sys
    from forecast import placement
    real = placement.fw_addends_by_week
    calls = []

    def spy(states, extra_fw_states=None):
        calls.append((sys._getframe(1).f_code.co_name, states, extra_fw_states))
        return real(states, extra_fw_states=extra_fw_states)
    r = _run_831(tmp_path_factory.mktemp("gaps_overdue"), copy_config,
                 scenario_edit=lambda s: _set_b50(s, tran_og_date="'2026-08-27'",
                                                  tran_og_count=200000),
                 patches=[(placement, "fw_addends_by_week", spy)])
    r.fw_calls = calls
    return r


@_needs_pr
def test_an_overdue_batch_is_culled_to_its_tran_og_count(overdue_low):
    """Mutation killed: run.py projecting the overdue batch with NO reconcile
    target (target=0). Decision 3 moves it by the scheduled path's own rule,
    which culls to tran_og_count."""
    b50 = {k: v for k, v in overdue_low.tog.items() if k[0] == "B50"}
    assert list(b50) == [("B50", "2026-W36")], b50
    assert b50[("B50", "2026-W36")] == pytest.approx(200_000, abs=1)
    (line,) = [d for c, d in overdue_low.vlog
               if c == "WARNING - Overdue FW batch auto-transferred"]
    assert "reconcile cull to tran_og_count 200,000" in line, line


@_needs_pr
def test_the_placement_caps_carry_the_split_fw_part(overdue_low):
    """Mutations killed: run.py not passing the split's FW track to
    run_placement, and run_placement not passing it on to fw_addends_by_week.
    Either way the FW-inclusive system caps plan 2026-W36/W37 ~90 t lighter
    than the facility is (B49's FW part is kept OUT of states_by_batch, so it
    reaches the caps only through extra_fw_states)."""
    from forecast.placement import fw_addends_by_week
    pc = [(s, x) for who, s, x in overdue_low.fw_calls if who == "run_placement"]
    assert pc, "run_placement never measured the FW load"
    states, extra = pc[-1]
    assert extra and "B49" in extra, "the split's FW track never reached the caps"
    with_x, _f = fw_addends_by_week(states, extra_fw_states=extra)
    without, _f = fw_addends_by_week(states)
    assert with_x["2026-W36"] - without.get("2026-W36", 0.0) > 50_000


def _n(d, k):
    v = d.get(k)
    return float(v) if isinstance(v, (int, float)) else 0.0


@_needs_pr
def test_the_monthly_ledger_follows_the_split(overdue_low):
    """Mutation killed: run.py not handing the split's FW track to the MONTHLY
    ledger (write_monthly_report). The weekly ledger is tested; the monthly
    one was checked only for its printed identity, which a wrong opening
    satisfies. Without the track 2026-09 opens on the seawater part alone
    (47,743) and the arrival lands in Count_Check (-242,136)."""
    from forecast.production_report import read_production_report
    pwb = openpyxl.load_workbook(PR_0831, read_only=True, data_only=True)
    _c, og, fw = read_production_report(pwb)
    pwb.close()
    pr_b49 = (sum(r.closing_count for r in og if r.batch_id == "B49")
              + sum(r.closing_count for r in fw if r.batch_id == "B49"))
    rows = sorted((d for d in overdue_low.monthly if d.get("Batch") == "B49"),
                  key=lambda d: str(d.get("Month")))
    assert rows and str(rows[0]["Month"]) == "2026-09", [d.get("Month") for d in rows[:3]]
    assert _n(rows[0], "Open_Count (fish)") == pytest.approx(pr_b49, abs=1)
    # The entry month and the one after it (the weekly test's rows[:3] rule);
    # later months carry small residuals unrelated to the split.
    bad = [(d["Month"], d["Count_Check (fish)"]) for d in rows[:2]
           if abs(_n(d, "Count_Check (fish)")) > 30]
    assert not bad, bad


@pytest.fixture(scope="module")
def half_placed(tmp_path_factory, copy_config):
    """The automatic top-up forced to place only HALF its fish (plan_topup
    wrapped) -- the loss the detection exists for."""
    if not PR_0831.exists():
        pytest.skip("the 2026-08-31 ProductionReport is not on this machine")
    from forecast import split_batch
    real = split_batch.plan_topup

    def halved(*a, **k):
        tp = real(*a, **k)
        tp.allocations = [dataclasses.replace(x, count=x.count / 2.0) for x in tp.allocations]
        return tp
    return _run_831(tmp_path_factory.mktemp("gaps_half"), copy_config,
                    patches=[(split_batch, "plan_topup", halved)])


@_needs_pr
def test_an_automatic_transfer_that_loses_fish_is_an_error_and_a_drop(half_placed):
    """Mutations killed: run.py not handing the unplaced fish to the audit, and
    not writing the FW PART NOT PLACED line. Detect, don't coerce: ~121,000
    fish leave the plan, so the ValidationLog must say so as an ERROR and the
    audit must read FW PART DROPPED, not OK."""
    errs = [d for c, d in half_placed.vlog if c == "ERROR - Split batch FW part NOT placed"]
    assert len(errs) == 1 and errs[0].startswith("SPLIT BATCH B49: FW PART NOT PLACED"), errs
    assert half_placed.ica["B49"]["Status"] == "*** FW PART DROPPED ***"
    assert not half_placed.ica_head[2].startswith("OK"), half_placed.ica_head[:4]
    assert any("FW PART DROPPED" in h and "B49" in h for h in half_placed.ica_head)


@pytest.fixture(scope="module")
def hybrid_auto(tmp_path_factory, copy_config):
    if not PR_0831.exists():
        pytest.skip("the 2026-08-31 ProductionReport is not on this machine")
    import re

    def hybrid(txt):
        txt, n = re.subn(r"(?m)^hybrid_follow:.*$", "hybrid_follow: 'full'", txt)
        return txt if n else txt.rstrip("\n") + "\nhybrid_follow: 'full'\n"
    return _run_831(tmp_path_factory.mktemp("gaps_hybrid"), copy_config, control_edit=hybrid)


@_needs_pr
def test_a_hybrid_run_says_the_guide_does_not_seed_the_split(hybrid_auto):
    """Mutation killed: run.py not writing the HYBRID GUIDE line for an auto
    split (decision 6: phase 2 is not built, so the gap is written where the
    guide's decisions are read). Only the line's text had a test."""
    hits = [(c, d) for c, d in hybrid_auto.vlog if d.startswith("HYBRID GUIDE - split batch")]
    assert len(hits) == 1, hits
    c, d = hits[0]
    assert c == "INFO - Hybrid guide (L1) decision"
    assert "B49" in d and "242,257" in d and "2026-W38" in d
