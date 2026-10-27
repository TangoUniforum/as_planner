"""Arrival make-room: merge a grow-out batch into one fewer tank before aborting.

MEASURED 2026-10-27: four corpus runs aborted with "TranOG arrival ... CANNOT be
placed" (2025-02-28, 2025-03-31, 2025-05-31 and the long-standing 2025-07-31).
At the abort every entry and grow-out tank was held, 6N had empty tanks but
nothing was ripe to purge -- yet grow-out held half-empty tanks of one batch
(e.g. 3 tanks at ~53 kg/m3) that fit in one fewer. The new last rung of the
make-room ladder frees a grow-out tank that way so an entry tank can move
forward. NEGATIVE CONTROL: the 2025-07-31 corpus PR aborts on the code before
this change (measured on the base worktree), so the first test below cannot
pass on old code.
"""
from __future__ import annotations

import contextlib
import io
import shutil
from pathlib import Path

import openpyxl
import pytest

from conftest import copy_config_without_costs

ROOT = Path(__file__).resolve().parent.parent
PR = ROOT / "pr_corpus" / "2025-07-31.xlsx"
needs_pr = pytest.mark.skipif(not PR.exists(), reason="corpus PR missing")
OG12 = {"OG1N", "OG1S", "OG2N", "OG2S"}


def _rows(wb, sheet):
    hdr = None
    for r in wb[sheet].iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == "Week":
                hdr = list(r)
            continue
        if r and r[0] and str(r[0]).startswith("20"):
            yield dict(zip(hdr, r))


@needs_pr
def test_a_saturated_arrival_is_placed_by_merging_a_growout_batch(tmp_path):
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config_without_costs(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    inp, out = tmp_path / "pr.xlsx", tmp_path / "out.xlsm"
    shutil.copy(PR, inp)
    from forecast.run import main
    with contextlib.redirect_stdout(io.StringIO()):
        rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                  calib_log_path="", growth_history_path="")
    assert rc == 0                      # used to raise "CANNOT be placed"
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    merges = [d for d in _rows(wb, "TransferPlan")
              if d.get("Channel") == "_consolidate_growout_for_arrival"]
    assert merges, "the run was rescued without the new rung -- this test no longer pins it"
    tank_sys = {}
    tank_batch = {}
    for d in _rows(wb, "BatchLocations"):
        tank_sys[int(d["Tank"])] = d["System"]
        tank_batch[(d["Week"], int(d["Tank"]))] = d["Batch"]
    for d in merges:
        src, dst = int(d["From_Tank"]), int(d["To_Tank"])
        # grow-out to grow-out only (never the entry tier, never 6N)
        assert tank_sys.get(src) not in OG12 and tank_sys.get(dst) not in OG12
        assert tank_sys.get(src) != "OG6N" and tank_sys.get(dst) != "OG6N"
        # one batch per tank: the destination holds the moved batch that week
        assert tank_batch.get((d["Week"], dst)) == d["Batch"]
        # the source is emptied: it holds no fish of that batch at the week's end
        assert tank_batch.get((d["Week"], src)) != d["Batch"]
    # no fish lost: the conservation audit's verdict line
    verdict = [str(r[0]) for r in wb["InputConservationAudit"].iter_rows(max_row=4, values_only=True)
               if r and r[0]]
    assert any(v.startswith("OK") and "0 dropped fish" in v for v in verdict), verdict
    log = [str(c) for row in wb["ValidationLog"].iter_rows(values_only=True)
           for c in row if isinstance(c, str)]
    assert any("FREED grow-out tank" in s for s in log)
    assert not any("CANNOT be placed" in s for s in log)
    wb.close()
