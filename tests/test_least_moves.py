"""Least-moves layout (Global v2): solver, independent checker, transfer pairing, companion workbook.

A tiny synthetic facility (2 entry + 3 grow tanks of 120 m3, cap 85 kg/m3 = 10.2 t) and one batch of 24k fish that
outgrows its entry tanks (22.8 t at 0.95 kg): the solver must move it to grow-out before the density cap bites, stage
8k fish into 6N in the last week,
and the checker must find the result clean. Negative controls: the same checker must flag a layout that breaks a
rule (density, two batches in one tank, a wrong batch count).
"""
import copy

import pytest

pytest.importorskip('ortools')
from forecast import least_moves as lm                                        # noqa: E402

WEEKS = [f'2030-W0{i}' for i in range(1, 7)]
WT = [0.5, 0.65, 0.8, 0.95, 1.1, 1.25]


def _endpoint():
    tanks = [(11, 'OG1N', 120, 85), (13, 'OG1N', 120, 85), (31, 'OG3N', 120, 85), (33, 'OG3N', 120, 85),
             (35, 'OG3N', 120, 85)]
    rec = {}
    for i, w in enumerate(WEEKS):
        n = 24000.0 if i < 5 else 16000.0
        rec[w] = {'n': n, 'kg': n * WT[i], 'wt': WT[i], 'arr': 24000.0 if i == 0 else 0.0,
                  'stage': 8000.0 if i == 5 else 0.0, 'sfr': 0.01}
    return {'source': 'synthetic', 'weeks': WEEKS, 'start_from': '2029-W52', 'tanks': tanks,
            'sys_cap': {'OG1N': (10 ** 6, 10 ** 5), 'OG3N': (10 ** 6, 10 ** 5)}, 'buffer': 0.1,
            'min_tank': 7000.0, 'min_move': 7000.0, 'move_rules': True, 'max_moves': 15,
            'dens_cap': {w: None for w in WEEKS}, 'start_layout': {}, 'batch': {'B1': rec},
            'harvested': {w: (8000.0 if w == WEEKS[-1] else 0.0) for w in WEEKS},
            'engine_moved': {w: 0.0 for w in WEEKS}, 'engine_layout': {}, 'engine_plan': []}


@pytest.fixture(scope='module')
def solved():
    D = _endpoint()
    R = lm.solve(D, horizon=6, commit=6, secs=20, workers=4, retries=1, log=lambda s: None)
    return D, R


def test_solver_layout_passes_the_independent_check(solved):
    D, R = solved
    c = lm.check(D, R['plan'])
    assert c['ok'], c['issues']
    assert [p['week'] for p in R['plan']] == WEEKS
    # 22.8 t at 0.95 kg does not fit the 2 entry tanks (2 x 10.2 t): some fish must be in grow-out by week 4
    w4 = R['plan'][3]['rows']
    assert sum(r['n'] for r in w4 if r['t'] in (31, 33, 35)) >= 7000
    assert lm.transfers_per_fish(D, R['plan'])['all'] > 0


def test_checker_flags_broken_layouts(solved):
    D, R = solved
    held = next(r for r in R['plan'][2]['rows'] if r['n'] > 0.5)
    bad = copy.deepcopy(R['plan'])
    row = next(r for r in bad[2]['rows'] if r['t'] == held['t'] and r['b'] == held['b'])
    row['dens'] = 120.0                                                           # over the 85 cap
    assert lm.check(D, bad)['issues'].get('density over cap (tank-weeks)') == 1
    bad = copy.deepcopy(R['plan'])
    D2 = copy.deepcopy(D); D2['batch']['B2'] = copy.deepcopy(D['batch']['B1'])
    for w in WEEKS:
        D2['batch']['B2'][w]['n'] = 0.0
    bad[2]['rows'].append(dict(b='B2', t=held['t'], n=8000.0, inn=0.0, out=0.0, st=0.0, dens=40.0))
    iss = lm.check(D2, bad)['issues']
    assert iss.get('two batches in a tank') == 1 and iss.get('batch count differs from the plan') == 1


def test_pairing_conserves_fish(solved):
    D, R = solved
    legs = lm.pair_moves(R['plan'])
    for p in R['plan']:
        tank_legs = [g for g in legs if g['week'] == p['week'] and g['frm'] not in ('FW',) and g['to'] != '6N']
        assert sum(g['fish'] for g in tank_legs) == pytest.approx(sum(r['inn'] for r in p['rows']), abs=1)
        assert sum(g['fish'] for g in legs if g['week'] == p['week'] and g['to'] == '6N') == \
            pytest.approx(sum(r['st'] for r in p['rows']), abs=1)
    assert sum(g['fish'] for g in legs if g['frm'] == 'FW') == pytest.approx(24000, abs=1)


def test_companion_workbook(solved, tmp_path):
    import openpyxl
    D, R = solved
    D = dict(D, engine_plan=R['plan'])               # the comparison side just needs a layout to judge
    out = tmp_path / 'lm.xlsx'
    lm.write_workbook(D, R, str(out))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ['Summary', 'Transfers', 'Layout', 'Batches']
    assert wb['Transfers'].max_row > 1 and wb['Layout'].max_row > 1


def test_staged_fish_match_the_plan_with_mortality():
    # 50 fish die in the staging week: the staged piece is still the plan's 8,000, not 8,000 x the survivors' factor
    D = _endpoint()
    D['batch']['B1'][WEEKS[-1]]['n'] = 15950.0
    D['batch']['B1'][WEEKS[-1]]['kg'] = 15950.0 * WT[-1]
    R = lm.solve(D, horizon=6, commit=6, secs=20, workers=4, retries=1, log=lambda s: None)
    staged = sum(r['st'] for r in R['plan'][-1]['rows'])
    assert staged == pytest.approx(8000, abs=1)
    assert lm.check(D, R['plan'])['ok']


# --------------------------------------------------------------------------- #
# A real run: the endpoint read back from the workbook matches the engine's own layout
# --------------------------------------------------------------------------- #
from pathlib import Path                                                       # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FIXTURE_PR = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"


@pytest.mark.skipif(not FIXTURE_PR.exists(), reason="fixture PR missing")
def test_endpoint_from_a_real_run(tmp_path):
    import contextlib
    import io
    import shutil
    from conftest import copy_config_without_costs
    from forecast import run as _run_mod
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config_without_costs(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    inp, out = tmp_path / "pr.xlsx", tmp_path / "out.xlsm"
    shutil.copy(FIXTURE_PR, inp)
    with contextlib.redirect_stdout(io.StringIO()):
        rc = _run_mod.main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                           calib_log_path="", growth_history_path="")
    assert rc == 0
    D = lm.read_endpoint(str(out), str(cfg), str(scn))
    assert D['weeks'] and D['start_layout']
    # the engine's own layout, read back, holds exactly the endpoint's fish per batch every week
    c = lm.check(D, D['engine_plan'])
    assert 'batch count differs from the plan' not in c['issues'], c['issues']
    # every re-planned week is after the manual starting events
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    n = next((int(r[1] or 0) for r in wb['RunConfig'].iter_rows(values_only=True)
              if r and r[0] == 'Manual override window (weeks)'), 0)        # no row = no manual events
    first = sorted({d['Week'] for d in lm._sheet_rows(wb, 'BatchLocations')})
    assert D['weeks'][0] == first[min(n, len(first) - 1)]
