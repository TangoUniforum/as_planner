"""BatchInputs: the per-batch scenario inputs beside what the run did with them.

Operator request 2026-10-02 ("an output tab ... the input dates for each
batch, the transfer dates, the counts input and transferred to OG, and the
models used per batch"). REPORT ONLY, so the tests pin two things: the sheet
tells the truth (it reconciles with the TransferPlan TranOG rows and the
InputConservationAudit, and shows the fw_correction the run USED, not just the
input), and it touches nothing else (appended last).
"""
from __future__ import annotations

import contextlib
import io
import shutil
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast.batch_inputs_report import (COLUMNS, SHEET_NAME, batch_input_rows,
                                          write_batch_inputs_sheet)
from forecast.events import TankAllocation, TranOGEntry
from forecast.models import BatchInput

ROOT = Path(__file__).resolve().parent.parent
PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")
COL = {name: i for i, (name, _m) in enumerate(COLUMNS)}


def _b(bid, tog, fw=0.82, **kw):
    d = dict(batch_id=bid, input_date=datetime(2025, 1, 1), input_count=550000,
             tran_sf_date=datetime(2025, 3, 1), tran_og_date=tog,
             tran_og_count=290000, tran_og_avg_wt_g=370.0, tran_og_cv=16.0,
             fcr_model="FCR_116_Quick", fw_correction=fw, sgr_correction=0.95)
    d.update(kw)
    return BatchInput(**d)


def _tranog(bid, when, *dests):
    return TranOGEntry(batch_id=bid, event_date=when, destinations=[
        TankAllocation(tank_id=t, count=n, avg_wt_g=w, cv_pct=16.0)
        for t, n, w in dests])


_CTL = SimpleNamespace(auto_calibrate_fw=True, auto_calibrate_fw_min=0.5,
                       auto_calibrate_fw_max=1.5)
_TABLES = SimpleNamespace(fcr_by_model={"1.16": [1.0]})


def _rows(**kw):
    base = dict(control=_CTL, tables=_TABLES, forecast_start=date(2026, 10, 1))
    base.update(kw)
    return {r[0]: r for r in batch_input_rows(**base)}


# --------------------------------------------------------------------------- #
# The table body
# --------------------------------------------------------------------------- #
def test_rows_are_in_batch_number_order_one_per_batch():
    bs = [_b("B10", datetime(2027, 1, 1)), _b("B9", datetime(2027, 1, 1)),
          _b("B100", datetime(2027, 1, 1))]
    rows = batch_input_rows(batches_input=bs, batch_by_id_used={},
                            tranog_events=[], control=_CTL, tables=_TABLES)
    assert [r[0] for r in rows] == ["B9", "B10", "B100"]
    assert all(len(r) == len(COLUMNS) for r in rows)


def test_run_columns_are_the_tranog_rows_summed():
    b = _b("B51", datetime(2026, 12, 3))
    ev = _tranog("B51", date(2026, 12, 7), (12, 144959, 460.0), (16, 144958, 356.0))
    r = _rows(batches_input=[b], batch_by_id_used={}, tranog_events=[ev],
              pr_fw={"B51"})["B51"]
    assert r[COL["Count into OG (run, fish)"]] == 289917
    assert r[COL["Avg wt at OG entry (run, g)"]] == pytest.approx(
        (144959 * 460 + 144958 * 356) / 289917, abs=0.05)
    assert r[COL["OG entry week (run)"]] == "2026-W50"
    assert r[COL["OG entry date (run)"]] == date(2026, 12, 7)
    assert r[COL["Entry tanks (run)"]] == "12, 16"
    assert r[COL["Source"]] == "In-flight: in freshwater at PR close"


def test_fw_correction_used_comes_from_the_run_not_the_input():
    inp = _b("B52", datetime(2027, 1, 21), fw=0.867)
    used = _b("B52", datetime(2027, 1, 21), fw=0.9402)
    rec = {"batch": "B52", "configured": 0.867, "applied": 0.9402,
           "solved": 0.9402, "target_wt_g": 370.0, "clamped": False,
           "converged": True}
    r = _rows(batches_input=[inp], batch_by_id_used={"B52": used},
              tranog_events=[], calib_records=[rec])["B52"]
    assert r[COL["FW correction (input)"]] == 0.867
    assert r[COL["FW correction (used)"]] == 0.9402
    assert r[COL["FW calibration"]].startswith("rewritten")


def test_clamped_and_off_calibration_are_named():
    inp = _b("B53", datetime(2027, 3, 1), fw=0.82)
    rec = {"batch": "B53", "configured": 0.82, "applied": 1.5, "solved": 1.9,
           "target_wt_g": 370.0, "clamped": True, "converged": True,
           "clamp_lo": 0.5, "clamp_hi": 1.5}
    r = _rows(batches_input=[inp], batch_by_id_used={}, tranog_events=[],
              calib_records=[rec])["B53"]
    assert r[COL["FW calibration"]].startswith("clamped to 1.5")
    off = SimpleNamespace(auto_calibrate_fw=False)
    r = _rows(batches_input=[inp], batch_by_id_used={}, tranog_events=[],
              control=off)["B53"]
    assert r[COL["FW calibration"]].startswith("off")


def test_source_is_judged_on_the_pr_not_the_post_window_state():
    """A window that harvests an in-OG batch out, or scripts an in-FW batch's
    fw_to_og, must not change where the batch was at the PR close."""
    bs = [_b("B42", datetime(2025, 10, 4)), _b("B49", datetime(2026, 9, 14)),
          _b("B50", datetime(2026, 10, 15)), _b("B40", datetime(2025, 7, 4)),
          _b("B60", datetime(2028, 2, 17)), _b("B48", datetime(2026, 9, 1))]
    rows = _rows(batches_input=bs, batch_by_id_used={}, tranog_events=[],
                 pr_og={"B42", "B49"}, pr_fw={"B49", "B50", "B48"},
                 manual_fw_weeks={"B50": (3, "2026-W42")})
    assert rows["B42"][COL["Source"]] == "In-flight: in OG at PR close"
    assert rows["B42"][COL["OG entry week (run)"]] == "in OG before the forecast"
    assert rows["B49"][COL["Source"]].startswith("In-flight: split")
    assert rows["B50"][COL["Source"]] == "In-flight: in freshwater at PR close"
    assert rows["B50"][COL["Manual fw_to_og"]] == "Yes - window week 3 (2026-W42)"
    assert rows["B50"][COL["FW calibration"]].startswith("not applied (transfer "
                                                         "scripted")
    assert rows["B48"][COL["Source"]].endswith("(transfer date already passed)")
    assert rows["B40"][COL["Source"]].startswith("Pre-start")
    assert rows["B60"][COL["Source"]].startswith("Scheduled")


def test_unknown_fcr_model_is_flagged():
    r = _rows(batches_input=[_b("B1", datetime(2027, 1, 1),
                                fcr_model="FCR_199_X")],
              batch_by_id_used={}, tranog_events=[])["B1"]
    assert "NOT in biology.yaml" in r[COL["FCR model"]]


# --------------------------------------------------------------------------- #
# The sheet: appended last, no other sheet touched, idempotent
# --------------------------------------------------------------------------- #
def _snapshot(wb):
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()]
            for ws in wb.worksheets}


def test_sheet_is_appended_last_and_touches_nothing_else():
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb["A"].append(["x", 1])
    wb.create_sheet("B").append(["y", 2])
    before = _snapshot(wb)
    kw = dict(batches_input=[_b("B1", datetime(2027, 1, 1))],
              batch_by_id_used={}, tranog_events=[], control=_CTL,
              tables=_TABLES, pr_closing=date(2026, 9, 30),
              forecast_start=date(2026, 10, 1), horizon_weeks=85)
    assert write_batch_inputs_sheet(wb, **kw) == 1
    assert wb.sheetnames == ["A", "B", SHEET_NAME]
    after = _snapshot(wb)
    assert {k: v for k, v in after.items() if k != SHEET_NAME} == before
    # Re-running replaces it rather than stacking a second copy.
    write_batch_inputs_sheet(wb, **kw)
    assert wb.sheetnames == ["A", "B", SHEET_NAME]
    ws = wb[SHEET_NAME]
    header = next(r for r in ws.iter_rows(values_only=True) if r[0] == "Batch")
    assert list(header[:len(COLUMNS)]) == [c for c, _m in COLUMNS]
    assert ws.auto_filter.ref.startswith("A4:")


# --------------------------------------------------------------------------- #
# End to end: reconciles with the TransferPlan and the input audit
# --------------------------------------------------------------------------- #
@pytest.mark.skipif(not PR_0831.exists(),
                    reason="the 2026-08-31 ProductionReport is not on this machine")
def test_real_run_reconciles_with_transferplan_and_input_audit(tmp_path, copy_config):
    import yaml
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    inp, out = tmp_path / "pr.xlsm", tmp_path / "out.xlsm"
    shutil.copy(PR_0831, inp)
    from forecast.run import main
    with contextlib.redirect_stdout(io.StringIO()):
        assert main(str(inp), str(out), config_dir=str(cfg),
                    scenario_dir=str(scn), calib_log_path="") == 0
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    # Last tab but the RunConfig YAML record, which always goes after it.
    assert [n for n in wb.sheetnames if n != "RunConfig YAML"][-1] == SHEET_NAME

    bi = [tuple(r) for r in wb[SHEET_NAME].iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(bi) if r and r[0] == "Batch")
    body = []
    for r in bi[hi + 1:]:
        if not r or r[0] is None:
            break
        body.append(r)
    scen = yaml.safe_load((scn / "batches.yaml").read_text(encoding="utf-8"))
    from forecast.batch_order import batch_sort_key
    want = sorted((b["batch_id"] for b in scen["batches"]), key=batch_sort_key)
    assert [r[0] for r in body] == want                 # one row per batch, in order
    by = {r[0]: r for r in body}
    # Inputs are batches.yaml verbatim.
    for b in scen["batches"]:
        r = by[b["batch_id"]]
        assert r[COL["Input count (eggs)"]] == b["input_count"]
        assert r[COL["TranOG count (input, fish)"]] == b["tran_og_count"]
        assert r[COL["FW correction (input)"]] == pytest.approx(b["fw_correction"])
        assert r[COL["FCR model"]].startswith(b["fcr_model"])

    # TransferPlan TranOG rows, summed per batch.
    tp = [tuple(r) for r in wb["TransferPlan"].iter_rows(values_only=True)]
    th = next(i for i, r in enumerate(tp) if r and r[0] == "Week")
    c = {str(v): j for j, v in enumerate(tp[th]) if v is not None}
    tog = defaultdict(float)
    for r in tp[th + 1:]:
        if r and r[c["Type"]] == "TranOG":
            tog[str(r[c["Batch"]])] += float(r[c["Count (fish)"]] or 0)
    assert tog, "the 8/31 run has TranOG arrivals"
    for bid, r in by.items():
        got = r[COL["Count into OG (run, fish)"]]
        if bid in tog:
            assert got == pytest.approx(tog[bid], abs=1), bid
        else:
            assert got is None, bid

    # InputConservationAudit Realized_TranOG, wherever the audit states one.
    ica = [tuple(r) for r in wb["InputConservationAudit"].iter_rows(values_only=True)]
    ih = next(i for i, r in enumerate(ica) if r and r[0] == "Batch")
    checked = 0
    for r in ica[ih + 1:]:
        if not r or r[0] not in by:
            continue
        row = dict(zip(ica[ih], r))
        real = row.get("Realized_TranOG (fish)")
        if real is not None:
            assert by[r[0]][COL["Count into OG (run, fish)"]] == pytest.approx(
                real, abs=1), r[0]
            checked += 1
    assert checked > 0

    # Auto-calibration rewrote at least one batch, and the sheet shows it.
    rewritten = [r for r in body if str(r[COL["FW calibration"]]).startswith(
        ("rewritten", "clamped"))]
    assert rewritten
    assert any(abs(r[COL["FW correction (used)"]] - r[COL["FW correction (input)"]])
               > 1e-6 for r in rewritten)
