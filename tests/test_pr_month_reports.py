"""A mid-month PR's month is WHOLE on every monthly harvest and feed report.

Operator 2026-10-10: "the harvest plan report should also show the full
totals, are there any other reports that should include the full month
including everything supplied from the PR and everything from the forecast?"
HarvestPlan Report, MonthlyReport and YearlySummary's harvest already added the
PR's elapsed days (the 1st to its closing date). HarvestReport, FeedForecast-
Monthly and YearlySummary's feed did not: on the 10.8.26 PR October read
159,494 fish on HarvestReport against 214,879 on HarvestPlan Report, and
644,005 kg of feed against MonthlyReport's 830,935. CostsAndProfit stays
forecast-only (operator).

Also: eggs stocked on the PR closing day (not yet in the PR) were booked net of
their first day's mortality -- B57 566,686 of 570,000 -- with the gap in
neither Input nor Mortality.
"""
from __future__ import annotations

import contextlib
import io
import shutil
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from conftest import copy_config_without_costs
import yaml

from forecast import excel_io
from forecast.caps import FacilityLimits
from forecast.events import Harvest
from forecast.production_report import read_pr_period

ROOT = Path(__file__).resolve().parent.parent
CLOSE = date(2026, 10, 8)
START = date(2026, 10, 9)


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


def _table(ws, pred):
    rows = _rows(ws)
    i = next(k for k, r in enumerate(rows) if r and pred(r))
    return [dict(zip(rows[i], r)) for r in rows[i + 1:] if r and any(c is not None for c in r)]


def _hv(d, batch, tank, count, wt=4000.0):
    return Harvest(batch_id=batch, event_date=d, source_tank_id=tank,
                   count=count, avg_wt_g=wt)


def _pb(harv=0.0, gross=0.0, feed=0.0, units=None):
    return SimpleNamespace(harv_count=harv, harv_gross_kg=gross, feed_kg=feed,
                           harv_units=units or {})


def _period(mid=True, **batches):
    return SimpleNamespace(is_mid_month=mid, closing_date=CLOSE, month_label="2026-10",
                           batches=batches)


def _ctl():
    return SimpleNamespace(max_biomass_kg=3_800_000.0, max_feed_per_day_kg=30_000.0,
                           max_harvest_per_week=55_000.0, min_harvest_per_week=26_000.0,
                           default_hog_yield=0.81, max_transfers_per_week=0)


def _hr(wb):
    return _table(wb["HarvestReport"], lambda r: r[0] == "Year" and "Batch" in r)


# ---- HarvestReport -------------------------------------------------------------

def test_harvestreport_month_includes_the_prs_harvest_per_tank():
    pp = _period(B43=_pb(15_927, 62_477, units={63: (5_516, 22_029), 65: (10_411, 40_448)}),
                 B44=_pb(39_458, 144_791, units={71: (39_458, 144_791)}))
    wb = openpyxl.Workbook()
    excel_io.write_harvest_report(wb, [_hv(date(2026, 10, 12), "B44", 61, 1_000)], 0.81, {},
                                  report_start=START, pr_period=pp)
    hr = _hr(wb)
    pr = [r for r in hr if r["Week"] == "PR"]
    assert [(r["Batch"], r["Tank"], r["Count (fish)"]) for r in pr] == [
        ("B43", 63, 5_516), ("B43", 65, 10_411), ("B44", 71, 39_458)]
    assert sum(r["Count (fish)"] for r in hr) == 15_927 + 39_458 + 1_000
    assert pr[2]["HOG Biomass (kg)"] == round(144_791 * 0.81)
    assert all(r["Month"] == date(2026, 10, 1) and r["Year"] == 2026 for r in pr)


def test_harvestreport_keeps_a_batch_whole_when_its_units_fall_short():
    pp = _period(B43=_pb(15_927, 62_477, units={63: (5_516, 22_029)}))
    wb = openpyxl.Workbook()
    excel_io.write_harvest_report(wb, [], 0.81, {}, report_start=START, pr_period=pp)
    pr = [r for r in _hr(wb) if r["Week"] == "PR"]
    assert sum(r["Count (fish)"] for r in pr) == 15_927
    assert sum(r["Gross Biomass (kg)"] for r in pr) == 62_477


def test_harvestreport_month_end_pr_adds_nothing():
    pp = _period(mid=False, B43=_pb(15_927, 62_477, units={63: (15_927, 62_477)}))
    wb = openpyxl.Workbook()
    excel_io.write_harvest_report(wb, [_hv(date(2026, 10, 12), "B44", 61, 1_000)], 0.81, {},
                                  report_start=START, pr_period=pp)
    assert [r["Count (fish)"] for r in _hr(wb)] == [1_000]


def test_yearlysummary_harvest_equals_harvestreport_on_a_mid_month_pr():
    pp = _period(B43=_pb(15_927, 62_477, units={63: (15_927, 62_477)}))
    evs = [_hv(date(2026, 10, 12), "B44", 61, 1_000), _hv(date(2027, 3, 1), "B45", 62, 500)]
    wb = openpyxl.Workbook()
    excel_io.write_harvest_report(wb, evs, 0.81, {}, report_start=START, pr_period=pp)
    excel_io.write_yearly_summary(wb, [], evs, FacilityLimits(), _ctl(),
                                  default_hog_yield=0.81, pr_period=pp, report_start=START)
    by_y = {}
    for r in _hr(wb):
        by_y[r["Year"]] = by_y.get(r["Year"], 0) + r["Count (fish)"]
    ys = {r["Year"]: r["Harvest_Count (fish)"]
          for r in _table(wb["YearlySummary"], lambda r: r[0] == "Year")}
    assert ys == by_y == {2026: 16_927, 2027: 500}


# ---- FeedForecastMonthly / YearlySummary feed -----------------------------------

def test_feed_monthly_grand_total_and_yearly_feed_include_the_prs_feed():
    pp = _period(B43=_pb(feed=4_000.0), B44=_pb(feed=6_000.0))
    wb = openpyxl.Workbook()
    excel_io.write_feed_forecast_monthly(wb, [], {}, START, report_start=START, pr_period=pp)
    rows = _rows(wb["FeedForecastMonthly"])
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Feed Type")
    j = next(k for k, d in enumerate(rows[h]) if d == date(2026, 10, 1))
    grand = next(r for r in rows if r and r[0] == "Grand Total")
    assert grand[j] == 10_000
    b = next(i for i, r in enumerate(rows) if r and r[0] == "Batch" and r[1] == "Feed Type")
    by_batch = {r[0]: r[j] for r in rows[b + 1:] if r and r[0]}
    assert by_batch == {"B43": 4_000, "B44": 6_000}
    excel_io.write_yearly_summary(wb, [], [], FacilityLimits(), _ctl(),
                                  default_hog_yield=0.81, pr_period=pp, report_start=START)
    ys = {r["Year"]: r["Feed (t)"] for r in _table(wb["YearlySummary"], lambda r: r[0] == "Year")}
    assert ys[2026] == 10


def test_feed_monthly_month_end_pr_adds_nothing():
    pp = _period(mid=False, B43=_pb(feed=4_000.0))
    wb = openpyxl.Workbook()
    excel_io.write_feed_forecast_monthly(wb, [], {}, START, report_start=START, pr_period=pp)
    assert not any(r and r[0] and "ProductionReport" in str(r[0])
                   for r in _rows(wb["FeedForecastMonthly"]))


# ---- read_pr_period: per-tank harvest from the Unit rows -------------------------

def test_read_pr_period_reads_each_units_harvest():
    hdr = ["", "", "", "", "", "Opening Count", "Closing Count", "Opening Biomass [kg]",
           "Opening Avg weight", "Gross growth in period", "Feed amount in period",
           "Harvested count (incl discards) in period",
           "Gross harvested biomass, incl. discards [kg] in period",
           "Mortality biomass in period", "Mortality count in period",
           "Culling biomass in period", "Culling count in period", "Deviation count in period"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(hdr)
    z = [0] * 6
    ws.append([None, None, "Fish group name: B43 ", None, None, 39_309, 27_026, 158_803, 4040,
               0, 1_000, 15_927, 62_477] + z)
    ws.append([None, None, None, "Unit: 63 ", None, 0, 0, 0, 0, 0, 0, 5_516, 22_029] + z)
    ws.append([None, None, None, "Unit: 65 ", None, 0, 0, 0, 0, 0, 0, 10_411, 40_448] + z)
    ws.append([None, None, None, "Unit: 67 ", None, 0, 0, 0, 0, 0, 0, 0, 0] + z)
    ws.append([None, None, "Fish group name: B44 ", None, None, 265_382, 226_826, 922_214, 3475,
               0, 2_000, 39_458, 144_791] + z)
    ws.append([None, None, None, "Unit: 71 ", None, 0, 0, 0, 0, 0, 0, 39_458, 144_791] + z)
    p = read_pr_period(ws, CLOSE)
    assert p.batches["B43"].harv_units == {63: (5_516, 22_029), 65: (10_411, 40_448)}
    assert p.batches["B44"].harv_units == {71: (39_458, 144_791)}


# ---- eggs stocked on the PR closing day -----------------------------------------

PR = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"


@pytest.mark.skipif(not (PR.exists() and (ROOT / "config" / "control.yaml").exists()
                         and (ROOT / "scenario" / "batches.yaml").exists()),
                    reason="needs the 9/30 PR fixture + committed config/scenario")
def test_eggs_stocked_on_the_closing_day_are_booked_whole(tmp_path):
    copy_config_without_costs(ROOT / "config", tmp_path / "config")   # never the live costs
    shutil.copytree(ROOT / "scenario", tmp_path / "scenario")
    shutil.rmtree(tmp_path / "scenario" / "manual_events", ignore_errors=True)
    (tmp_path / "scenario" / "manual_events.yaml").unlink(missing_ok=True)
    bf = tmp_path / "scenario" / "batches.yaml"
    doc = yaml.safe_load(bf.read_text(encoding="utf-8"))
    b57 = next(b for b in doc["batches"] if b["batch_id"] == "B57")
    b57["input_date"] = "2026-09-30"          # the PR's closing day; the PR does not list it
    bf.write_text(yaml.safe_dump(doc, sort_keys=False), encoding="utf-8")
    from forecast.run import main
    out = tmp_path / "out.xlsm"
    with contextlib.redirect_stdout(io.StringIO()):
        main(str(PR), str(out), config_dir=str(tmp_path / "config"),
             scenario_dir=str(tmp_path / "scenario"), calib_log_path="")
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    rows = _rows(wb["WeeklyReport"])
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Scenario")
    c = {k: i for i, k in enumerate(rows[h])}
    first = next(r for r in rows[h + 1:] if r and r[c["Batch"]] == "B57")
    assert abs(first[c["Input_Count (fish)"]] - b57["input_count"]) <= 1
    assert abs(first[c["Count_Check (fish)"]]) <= 1
    assert (first[c["Input_Count (fish)"]] - first[c["Mort_Count (fish)"]]
            - first[c["Close_Count (fish)"]]) == pytest.approx(0, abs=1)
