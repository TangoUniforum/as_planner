"""The Costs UI in app.py (2026-09-11): Configure → Targets & prices → Costs,
the Profit objective in both Ideal optimizers, and the Run page's 8th tab,
"Costs & profit".

EVERY NUMBER HERE IS MADE UP. Each case builds a tmp project — a copy of
config/ WITHOUT costs.yaml, plus scenario/ — and writes its own costs.yaml;
the live config/costs.yaml is never read or written (tests/conftest.py
guards it). Feed-type NAMES come from biology.yaml; names are not prices.

The real function bodies are lifted out of app.py by name and run under
AppTest in a fresh interpreter (why a subprocess: see test_app_renders), the
same pattern test_app_ideal_click uses:
  * Configure: renders with no costs.yaml (inputs EMPTY, never 0), with a
    complete one (inputs from the file), with orphan / missing prices (both
    warned), and with a malformed one (named, section hidden); Save refuses
    anything incomplete, and a good Save writes ONLY costs.yaml — every other
    config/ and scenario/ file hashed before and after — without moving
    _config_fingerprint (a negative control proves the fingerprint is live);
  * Ideal step 2: Profit is always an option; without costs, or with a feed
    type unpriced, it warns and disables the Find button (nothing runs),
    while the other objectives still run; with costs the page passes the
    one validated snapshot and shows the cost line and Cost / Profit
    columns;
  * Ideal step 3 (operator, 2026-09-11: show profit, don't rank by it): the
    objectives are exactly Revenue / HOG / Biomass gain, a caption says why
    and points to step 2; Cost / Profit columns and the cost lines are shown
    (blank, never 0, when nothing is priced); a costs save never stales its
    search; a remembered "profit" (widget key, _keep_ shadow, or a kept
    search ranked by it) is reset without an exception, with a one-time
    note;
  * Run tab: the predates / absent / error / not-computed cases, an OK
    sheet written and read back by forecast.costs_report, and the stale
    warnings when costs.yaml changed or was removed since the run;
  * help=: every widget _edit_costs, the Ideal optimizers and the Run costs
    tab create passes a non-empty help (statically over the source, with a
    negative control, and at run time on what the pages draw).
"""
import ast
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from forecast import costs as C
from forecast import costs_report as R
from forecast.config_io import load_biology_tables

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app.py"
NAMES = [n for _s, n in sorted(load_biology_tables(ROOT / "config").feed_types,
                               key=lambda x: x[0])]

pytestmark = pytest.mark.skipif(
    not (APP.exists() and (ROOT / "config" / "control.yaml").exists()
         and (ROOT / "scenario" / "batches.yaml").exists()),
    reason="needs app.py + a seeded config/scenario")


def _costs(drop=(), orphan=False):
    """Made-up costs pricing every model feed type except `drop`."""
    fp = {n: {"item": f"Made-up item {i}", "price_per_kg": 1.0 + i}
          for i, n in enumerate(NAMES) if n not in drop}
    if orphan:
        fp["Retired 0.1"] = {"item": "", "price_per_kg": 9.0}
    return {"schema": 1, "fixed_monthly": 1000.0, "oxygen_per_kg_feed": 0.5,
            "chemicals_per_kg_feed": 0.25, "feed_shipping_per_kg": 0.125,
            "egg_price": 0.01, "feed_prices": fp}


def _project(tmp_path, costs=None):
    """A tmp project: config/ (every file but costs.yaml) + scenario/, and
    `costs` written as config/costs.yaml — a dict through save_costs, a str
    as raw text."""
    d = tmp_path / "proj"
    shutil.copytree(ROOT / "config", d / "config",
                    ignore=shutil.ignore_patterns("costs.yaml*"))
    shutil.copytree(ROOT / "scenario", d / "scenario")
    if isinstance(costs, dict):
        C.save_costs(d / "config", costs)
    elif isinstance(costs, str):
        (d / "config" / "costs.yaml").write_text(costs, encoding="utf-8")
    return d


def _hashes(proj):
    return {str(p.relative_to(proj)).replace("\\", "/"):
            hashlib.sha256(p.read_bytes()).hexdigest()
            for sub in ("config", "scenario")
            for p in sorted((proj / sub).rglob("*")) if p.is_file()}


def _app_functions(*names, **ns):
    """The named top-level functions and assignments lifted out of app.py
    (no Streamlit run), executed in `ns`."""
    tree = ast.parse(APP.read_text(encoding="utf-8"))
    body = [n for n in tree.body
            if (isinstance(n, ast.FunctionDef) and n.name in names)
            or (isinstance(n, ast.Assign)
                and any(getattr(t, "id", None) in names for t in n.targets))]
    assert len(body) == len(names), sorted(names)
    exec(compile(ast.Module(body=body, type_ignores=[]), str(APP), "exec"), ns)
    return ns


def _fingerprint(proj):
    ns = _app_functions("_config_fingerprint", "_NON_ENGINE_CONFIG",
                        CONFIG_DIR=proj / "config",
                        SCENARIO_DIR=proj / "scenario")
    return ns["_config_fingerprint"]()


def _drive(tmp_path, driver, **params):
    params["root"] = str(ROOT)
    pfile = tmp_path / "params.json"
    pfile.write_text(json.dumps(params), encoding="utf-8")
    p = subprocess.run([sys.executable, "-c", driver, str(pfile)],
                       capture_output=True, text=True, timeout=900,
                       cwd=str(ROOT))
    tail = [ln for ln in (p.stdout or "").splitlines()
            if ln.startswith(("OK", "FAIL", "SKIP"))]
    msg = tail[-1] if tail else (p.stderr or "")[-800:]
    if msg.startswith("SKIP"):
        pytest.skip(msg)
    assert p.returncode == 0, msg
    assert msg.startswith("OK"), msg
    return msg


# Shared by every driver: params, AppTest, and `lift` (the app.py functions
# by name) for the page scripts. Each page script is built with ascii(P):
# AppTest.from_string saves the script without an encoding cookie, so a
# non-ASCII character (the sheet note's dash) would make it unreadable.
_PRELUDE = r'''
import json, os, sys
P = json.load(open(sys.argv[1], encoding="utf-8"))
sys.path.insert(0, P["root"])
try:
    from streamlit.testing.v1 import AppTest
except Exception as e:
    print("SKIP no AppTest: %s" % e)
    raise SystemExit(0)

def fail(msg):
    # ASCII only: a Windows console codec cannot print the page's arrows.
    print(ascii("FAIL " + msg)[1:-1])
    raise SystemExit(1)

def boom(at, what):
    if at.exception:
        fail(what + ": " + "; ".join(str(e)[:400] for e in at.exception))

def no_help(at, prefixes):
    """Keys of drawn widgets (key starting with a prefix) with an empty help."""
    bad = []
    for kind in ("number_input", "text_input", "text_area", "radio", "button",
                 "selectbox", "multiselect", "slider", "checkbox", "toggle",
                 "date_input"):
        try:
            ws = list(getattr(at, kind))
        except Exception:
            continue
        for w in ws:
            k = getattr(w, "key", None) or ""
            if k.startswith(tuple(prefixes)) and not (w.help or "").strip():
                bad.append(k)
    return bad

LIFT = r"""
import ast
def lift(root, names=(), prefixes=(), assigns=(), assign_prefixes=()):
    tree = ast.parse(open(root + "/app.py", encoding="utf-8").read())
    def keep(n):
        if isinstance(n, ast.FunctionDef):
            return n.name in names or n.name.startswith(tuple(prefixes))
        if isinstance(n, ast.Assign):
            ids = [getattr(t, "id", "") for t in n.targets]
            return any(i in assigns or i.startswith(tuple(assign_prefixes))
                       for i in ids)
        return False
    return ast.Module(body=[n for n in tree.body if keep(n)], type_ignores=[])
"""
'''


# --------------------------------------------------------------------------- #
# Configure → Targets & prices → Costs
# --------------------------------------------------------------------------- #
_CONFIGURE = _PRELUDE + r'''
SCRIPT = "P = " + ascii(P) + "\n" + LIFT + r"""
import sys
from pathlib import Path
import pandas as pd
import streamlit as st
sys.path.insert(0, P["root"])
ns = {"st": st, "pd": pd, "CONFIG_DIR": Path(P["cfg"])}
exec(compile(lift(P["root"], names=("_edit_costs", "_read_or_explain",
                                    "_records"),
                  assigns=("_READ_FIX_HINT",)), "app.py", "exec"), ns)
ns["_edit_costs"]()
"""
KEYS = ["cost_fixed", "cost_egg", "cost_ship", "cost_o2", "cost_chem"]

def num(at, k):
    """The costs input `k`. Its key carries the file's signature after the
    name, so a costs.yaml changed on disk draws inputs from the new file."""
    ws = [w for w in at.number_input
          if (w.key or "") == k or (w.key or "").startswith(k + "_")]
    if len(ws) != 1:
        fail("want one %s input, got %r" % (k, [w.key for w in at.number_input]))
    return ws[0]

at = AppTest.from_string(SCRIPT, default_timeout=120)
at.run()
boom(at, "render")
case = P["case"]
if case == "bad":
    if not any("config/costs.yaml" in e.value and "could not be read" in e.value
               for e in at.error):
        fail("a malformed costs.yaml is not named: %r"
             % [e.value[:200] for e in at.error])
    if any(b.key == "cost_save" for b in at.button):
        fail("the Save button shows over a file that could not be read")
    print("OK a malformed file is named and the section hidden")
    raise SystemExit(0)

def editor(at):
    eds = [d for d in at.dataframe if "Feed type" in list(d.value.columns)]
    if len(eds) != 1:
        fail("want one feed price editor, got %d" % len(eds))
    return eds[0]

ed = editor(at)
df = ed.value
if list(df["Feed type"]) != P["names"]:
    fail("the feed rows are not biology's feed types in size order: %r"
         % list(df["Feed type"]))
cols = json.loads(ed.proto.columns)
for name in ("Feed type", "Up to size (g)", "Item", "Price / kg"):
    if not (cols.get(name, {}).get("help") or "").strip():
        fail("feed column %r has no help" % name)
if not (cols["Feed type"].get("disabled") and cols["Up to size (g)"].get("disabled")):
    fail("the model feed type / size columns are editable")
if cols["Item"].get("disabled") or cols["Price / kg"].get("disabled"):
    fail("the Item / Price columns are read-only")
bad = no_help(at, ("cost_",))
if bad:
    fail("widgets with no help: %r" % bad)

if case == "empty":
    if not any("Costs not set" in i.value for i in at.info):
        fail("no 'Costs not set' notice without a costs.yaml")
    for k in KEYS:
        if num(at, k).value is not None:
            fail("%s starts at %r — must start EMPTY, never 0"
                 % (k, num(at, k).value))
    if not df["Price / kg"].isna().all() or any(df["Item"]):
        fail("feed prices / items do not start empty")
    at.button(key="cost_save").click().run()
    boom(at, "save empty")
    errs = [e.value for e in at.error]
    if not any("Fixed cost per month is empty" in e for e in errs):
        fail("an empty save did not name the empty fixed cost: %r" % errs)
    if not any("has no price" in e for e in errs):
        fail("an empty save did not name an unpriced feed type: %r" % errs)
    for k, v in zip(KEYS, (1000.0, 0.01, 0.125, 0.5, 0.25)):
        num(at, k).set_value(v)
    at.button(key="cost_save").click().run()
    boom(at, "save with no prices")
    errs = [e.value for e in at.error]
    if any("is empty" in e for e in errs):
        fail("filled costs still said empty: %r" % errs)
    if not any("has no price" in e for e in errs):
        fail("saved with no feed prices: %r" % errs)
    if any("Saved costs" in s.value for s in at.success):
        fail("claimed a save without feed prices")
    print("OK empty inputs, and an incomplete save is refused")
elif case == "saved":
    if any("Costs not set" in i.value for i in at.info):
        fail("'Costs not set' with a costs.yaml present")
    for k, v in P["want"].items():
        if num(at, k).value != v:
            fail("%s shows %r, the file says %r" % (k, num(at, k).value, v))
    if [round(float(x), 6) for x in df["Price / kg"]] != P["prices"]:
        fail("prices are not the file's: %r" % list(df["Price / kg"]))
    if list(df["Item"]) != P["items"]:
        fail("items are not the file's: %r" % list(df["Item"]))
    # The price column shows the stored number in full ("plain") and takes
    # any number of decimals (no step): a rounded display, retyped, would
    # save a different price.
    tc = cols["Price / kg"].get("type_config") or {}
    if tc.get("format") != "plain" or tc.get("step") is not None:
        fail("the feed price column rounds what it shows or accepts: %r" % tc)
    if at.warning:
        fail("warnings on a complete file: %r" % [w.value for w in at.warning])
    num(at, "cost_fixed").set_value(1234.5)
    at.button(key="cost_save").click().run()
    boom(at, "save")
    if not any("Saved costs" in s.value for s in at.success):
        fail("no save: %r" % [e.value[:200] for e in at.error])
    print("OK inputs from the file, and a good save")
elif case == "warn":
    ws = [w.value for w in at.warning]
    if not any("No price yet for 1 feed type(s)" in w and P["dropped"] in w
               for w in ws):
        fail("no missing-price warning naming %r: %r" % (P["dropped"], ws))
    if not any("Retired 0.1" in w and "removed when you save" in w for w in ws):
        fail("no orphan warning: %r" % ws)
    if not df["Price / kg"].isna().iloc[-1]:
        fail("the unpriced feed type is not shown blank")
    at.button(key="cost_save").click().run()
    boom(at, "save with a missing price")
    if not any(("feed type %r has no price" % P["dropped"]) in e.value
               for e in at.error):
        fail("a missing price did not block the save: %r"
             % [e.value[:200] for e in at.error])
    if any("Saved costs" in s.value for s in at.success):
        fail("saved with a missing price")
    print("OK orphan and missing prices warned, save refused")
elif case == "disk":
    if num(at, "cost_fixed").value != P["before"]:
        fail("cost_fixed shows %r, the file says %r"
             % (num(at, "cost_fixed").value, P["before"]))
    # Another session (or a hand edit) changes the file while Configure is
    # open: the next render must show the file, never the numbers it drew
    # first (Save would write those back).
    with open(os.path.join(P["cfg"], "costs.yaml"), "w", encoding="utf-8") as fh:
        fh.write(P["other_costs"])
    at.run()
    boom(at, "rerun after a change on disk")
    if num(at, "cost_fixed").value != P["after"]:
        fail("the inputs kept the old numbers after costs.yaml changed on "
             "disk: cost_fixed shows %r, the file says %r"
             % (num(at, "cost_fixed").value, P["after"]))
    if float(editor(at).value["Price / kg"].iloc[0]) != P["after_price"]:
        fail("the feed prices are not the changed file's: %r"
             % list(editor(at).value["Price / kg"]))
    print("OK the inputs follow a change on disk")
'''


def test_configure_costs_start_empty_and_refuse_an_incomplete_save(tmp_path):
    proj = _project(tmp_path)                      # no costs.yaml at all
    before = _hashes(proj)
    _drive(tmp_path, _CONFIGURE, case="empty", cfg=str(proj / "config"),
           names=NAMES)
    assert _hashes(proj) == before                 # nothing written
    assert not (proj / "config" / "costs.yaml").exists()


def test_configure_costs_save_writes_only_costs_yaml(tmp_path):
    made_up = _costs()
    proj = _project(tmp_path, made_up)
    before = _hashes(proj)
    fp0 = _fingerprint(proj)
    _drive(tmp_path, _CONFIGURE, case="saved", cfg=str(proj / "config"),
           names=NAMES,
           want={"cost_fixed": 1000.0, "cost_egg": 0.01, "cost_ship": 0.125,
                 "cost_o2": 0.5, "cost_chem": 0.25},
           prices=[1.0 + i for i in range(len(NAMES))],
           items=[f"Made-up item {i}" for i in range(len(NAMES))])
    after = _hashes(proj)
    changed = {k for k in set(before) | set(after)
               if before.get(k) != after.get(k)}
    assert changed == {"config/costs.yaml"}        # ONLY the costs file
    got = C.load_costs(proj / "config")
    assert got == C.validate_costs(dict(made_up, fixed_monthly=1234.5))
    assert not list((proj / "config").glob("costs.yaml.tmp-*"))
    # A costs save does not move the engine fingerprint...
    assert _fingerprint(proj) == fp0
    # ...which is live: an engine input edit does move it.
    with open(proj / "config" / "control.yaml", "a", encoding="utf-8") as fh:
        fh.write("\n# touched by the negative control\n")
    assert _fingerprint(proj) != fp0


def test_configure_costs_warn_on_orphan_and_missing_prices(tmp_path):
    proj = _project(tmp_path, _costs(drop=(NAMES[-1],), orphan=True))
    before = _hashes(proj)
    _drive(tmp_path, _CONFIGURE, case="warn", cfg=str(proj / "config"),
           names=NAMES, dropped=NAMES[-1])
    assert _hashes(proj) == before


def test_configure_costs_follow_a_change_on_disk(tmp_path):
    made_up = _costs()
    proj = _project(tmp_path, made_up)
    other = tmp_path / "other"
    other.mkdir()
    C.save_costs(other, dict(
        made_up, fixed_monthly=4321.0,
        feed_prices=dict(made_up["feed_prices"],
                         **{NAMES[0]: {"item": "Made-up item 0",
                                       "price_per_kg": 7.5}})))
    text = (other / "costs.yaml").read_text(encoding="utf-8")
    _drive(tmp_path, _CONFIGURE, case="disk", cfg=str(proj / "config"),
           names=NAMES, other_costs=text, before=1000.0, after=4321.0,
           after_price=7.5)
    # The page wrote nothing: the file is the one written on disk.
    assert (proj / "config" / "costs.yaml").read_text(encoding="utf-8") == text


def test_configure_costs_name_a_malformed_file(tmp_path):
    proj = _project(tmp_path, "schema: 1\nfixed_monthly: -5\n")
    before = _hashes(proj)
    _drive(tmp_path, _CONFIGURE, case="bad", cfg=str(proj / "config"))
    assert _hashes(proj) == before


def test_costs_yaml_never_enters_the_run_config_sheet():
    """RunConfig stays byte-identical: costs.yaml is neither snapshotted into
    a workbook (so an import never restores it) nor named in the sheet's
    'excluded' header (whose text is printed into every run's RunConfig)."""
    from forecast import config_snapshot as cs
    assert all(name != "costs.yaml" for name, _d in cs._FILES)
    assert all("costs.yaml" not in path for path, _why in cs._EXCLUDED)


def test_costs_yaml_is_not_an_engine_input_and_an_import_never_moves_it():
    ns = _app_functions("_NON_ENGINE_CONFIG")
    assert "costs.yaml" in ns["_NON_ENGINE_CONFIG"]
    src = APP.read_text(encoding="utf-8")
    tree = ast.parse(src)
    io_sec = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                  and n.name == "_config_io_section")
    assert "`costs.yaml`" in ast.get_source_segment(src, io_sec)


# --------------------------------------------------------------------------- #
# Ideal: the Profit objective in both optimizers
# --------------------------------------------------------------------------- #
_IDEAL = _PRELUDE + r'''
import datetime as dt
from types import SimpleNamespace
from forecast import ideal_optimize as io
from forecast import transition_optimize as to

NAMES = P["names"]
YEARS = (2026, 2027, 2028, 2029, 2030)
CALLS, TR_CALLS = [], []
# "undriven": every read lacks its cost drivers, so under Profit every cell
# RAN but none can be priced (require_price makes each an error cell).
DRIVERS = P["state"] != "undriven"

def read(size, year=2029):
    d = {a: 0 for _k, a in to.YEAR_COUNTS}
    # Whole ISO years (2026 has 53 weeks), so no first_week is needed.
    return SimpleNamespace(year=year,
                           weeks=dt.date(year, 12, 28).isocalendar()[1],
                           revenue=size * 300.0, hog_t=4000.0, gain_t=4000.0,
                           feed_kg_by_type=({NAMES[0]: 1000.0,
                                             NAMES[-1]: size / 4.0}
                                            if DRIVERS else None),
                           eggs=size * 1.5 if DRIVERS else None, **d)

def snapshot(costs, objective, project_dir):
    # The optimizers' own refusal and snapshot, on what the page passed.
    return io.check_costs(costs, objective,
                          io.feed_type_names(project_dir)
                          if objective == "profit" and costs is not None
                          else None)

def mk(c, s, p, snap, objective):
    y = read(s)
    return io.require_price(io.Cell(
        cadence_days=c, batch_size=s, fish_per_week=s * 7 / c,
        cap_kg=float(p), year=2029, revenue=y.revenue, hog_t=y.hog_t,
        gain_t=y.gain_t, avg_gross_kg=4.2, peak_pct_of_cap=0.8,
        breaches={k: 0 for k in io.BREACH_KEYS}, total=0, within_limits=True,
        **io.cost_fields([y], y.revenue, snap)), objective)

def fake_optimize(cadences, sizes, cap_kg, project_dir, *, caps=None,
                  objective="revenue", costs=None, **kw):
    CALLS.append(dict(objective=objective, costs=costs))
    snap = snapshot(costs, objective, project_dir)
    cells = tuple(mk(c, s, p, snap, objective)
                  for c in cadences for s in sizes for p in caps)
    b = io.best(cells, objective)
    stab = ()
    if b is not None:
        stab = ((b, tuple(mk(b.cadence_days, b.batch_size + d, b.cap_kg, snap,
                             objective) for d in (-5000, 5000)), True),)
    return io.Result(cells=cells, objective=objective, best=b,
                     closest=None if b else io.closest(cells, objective),
                     unconstrained=io.ignoring_limits(cells, objective),
                     stability=stab, best_stable=b)

def fake_tr(live, fs, cutoff, sizes, caps, pr_path, *, ceiling_kg,
            project_dir, objective="revenue", costs=None, **kw):
    TR_CALLS.append(dict(objective=objective, costs=costs))
    if objective not in to.TR_OBJECTIVES:        # the real refusal, first
        raise ValueError(to.PROFIT_NOT_RANKED)
    snap = io.check_costs(costs, objective)
    reads = lambda s: {y: read(s, year=y) for y in YEARS}
    gates = {y: () for y in YEARS}
    today = to.judge_today(to.TrCell(batch_size=0, cap_kg=3.8e6, today=True,
                                     reads=reads(340_000),
                                     gates_by_year=gates), costs=snap)
    cells = tuple(to.judge_cell(to.TrCell(
        batch_size=s, cap_kg=float(p), reads=reads(s), gates_by_year=gates,
        n_changed=3, first_tran_og=dt.date(2027, 9, 23)), today, costs=snap)
        for s in sizes for p in caps)
    b = io.best(cells, objective)
    return to.TrResult(today=today, cells=cells, objective=objective, best=b,
                       closest=None if b else io.closest(cells, objective),
                       unconstrained=io.ignoring_limits(cells, objective),
                       ceiling_kg=float(ceiling_kg))

io.optimize, to.optimize_transition = fake_optimize, fake_tr   # app imports these

SCRIPT = "P = " + ascii(P) + "\n" + LIFT + r"""
import os, sys, datetime as dt
from pathlib import Path
from types import SimpleNamespace
import streamlit as st
sys.path.insert(0, P["root"])
from forecast import ideal as im
from forecast.scenario_io import load_batches
ns = {"st": st, "os": os, "_ROOT": Path(P["proj"]),
      "_config_fingerprint": lambda: "fp", "_cpu_workers": lambda: 1}
exec(compile(lift(P["root"], prefixes=("_ideal",),
                  assign_prefixes=("_IDEAL", "_TR_", "_REF_")),
             "app.py", "exec"), ns)
ctx = im.load_context(P["proj"])
c = ctx["control"]
ns["_ideal_restore"]()
if st.session_state.get("which", "step2") == "step2":
    ov = dict(max_biomass_kg=3.8e6,
              max_harvest_per_week=float(c.max_harvest_per_week),
              min_harvest_per_week=float(c.min_harvest_per_week),
              min_harvest_weight_g=float(c.min_harvest_weight_g),
              max_feed_per_day_kg=float(c.max_feed_per_day_kg))
    ns["_ideal_optimizer"](ctx, ov, {}, {}, "controller", {}, 3800)
else:
    live = load_batches(os.path.join(P["proj"], "scenario"))
    fs = dt.date(2026, 9, 1)
    ns["_ideal_tr_optimizer"](
        ctx, live, fs, fs,
        SimpleNamespace(name="pr.xlsm", getvalue=lambda: b"PR bytes"),
        {}, {}, {}, "controller", {}, 3800, ns["_ideal_tr_seeds"](c))
ns["_ideal_save"]()
"""

state = P["state"]
# Step 2 ranks by Profit; step 3 only shows it (its own block below).
STEPS = (("step2", "ideal_opt_obj", "ideal_opt_run", "ideal_opt",
          "Best within every limit", "Profit $M/yr", CALLS, "ideal_opt_use",
          "Within every limit"),)
COSTS_YAML = os.path.join(P["proj"], "config", "costs.yaml")
for which, rkey, bkey, pref, winword, pcol, calls, ukey, lcol in STEPS:
    at = AppTest.from_string(SCRIPT, default_timeout=120)
    at.session_state["which"] = which
    at.run()
    boom(at, which + " render")
    if "Profit" not in list(at.radio(key=rkey).options):
        fail(which + ": no Profit option: %r" % list(at.radio(key=rkey).options))
    at.radio(key=rkey).set_value("profit").run()
    boom(at, which + " profit")
    warn = [w.value for w in at.warning if "Profit is unavailable" in w.value]
    n0 = len(calls)
    if state == "full":
        if warn or at.button(key=bkey).disabled:
            fail(which + ": Profit blocked with complete costs: %r" % warn)
        at.button(key=bkey).click().run()
        boom(at, which + " find profit")
        if len(calls) != n0 + 1 or calls[-1]["objective"] != "profit":
            fail(which + ": the optimizer was not asked for profit: %r"
                 % [x["objective"] for x in calls])
        if calls[-1]["costs"] != P["costs"]:
            fail(which + ": the page did not pass the saved costs")
        if not any(winword in s.value and "profit $" in s.value
                   for s in at.success):
            fail(which + ": no winner line with its profit: %r"
                 % [s.value[:200] for s in at.success])
        lines = [x.value for x in at.caption
                 if x.value.startswith("Cost $") and "→ profit" in x.value]
        if not lines:
            fail(which + ": no cost breakdown under the winner")
        if which == "step2" and "/kg HOG" not in lines[0]:
            fail("step2: the steady year's line has no cost per kg HOG")
        tbl = [d for d in at.dataframe if pcol in list(d.value.columns)]
        if not tbl or tbl[0].value[pcol].isna().any():
            fail(which + ": no filled %r column" % pcol)
        # The costs are in a PROFIT search's signature: a costs edit makes
        # it stale (Use greyed out). A search on another objective does not
        # depend on them: it stays current and notes the change.
        orig = open(COSTS_YAML, "rb").read()
        try:
            with open(COSTS_YAML, "wb") as fh:
                fh.write(orig.replace(b"fixed_monthly: 1000.0",
                                      b"fixed_monthly: 2000.0"))
            at.run()
            boom(at, which + " costs edited")
            if not any("Showing the last" in w.value for w in at.warning):
                fail(which + ": a costs edit left a Profit search current")
            if not at.button(key=ukey).disabled:
                fail(which + ": Use is live on a stale Profit search")
            at.radio(key=rkey).set_value("revenue").run()
            at.button(key=bkey).click().run()
            boom(at, which + " find revenue")
            with open(COSTS_YAML, "wb") as fh:
                fh.write(orig)
            at.run()
            boom(at, which + " costs edited back")
            if any("Showing the last" in w.value for w in at.warning):
                fail(which + ": a costs edit made a Revenue search stale")
            if at.button(key=ukey).disabled:
                fail(which + ": Use greyed out on a current Revenue search")
            if not any("Your costs have changed since this search" in x.value
                       for x in at.caption):
                fail(which + ": no note that the cost columns used old costs")
        finally:
            with open(COSTS_YAML, "wb") as fh:
                fh.write(orig)
    elif state == "undriven":
        # Complete costs, but no read can be priced: every cell RAN and is
        # within the limits; Profit just cannot rank them.
        if warn or at.button(key=bkey).disabled:
            fail(which + ": Profit blocked with complete costs: %r" % warn)
        at.button(key=bkey).click().run()
        boom(at, which + " find undriven")
        errs = [e.value for e in at.error]
        if not any("Profit cannot rank this grid" in e for e in errs):
            fail(which + ": no 'cannot rank' headline: %r"
                 % [e[:200] for e in errs])
        if any("ran at all" in e for e in errs):
            fail(which + ": says nothing ran, though every cell ran")
        tbl = [d for d in at.dataframe if lcol in list(d.value.columns)]
        if not tbl:
            fail(which + ": no table")
        df = tbl[0].value
        if list(df[lcol]) != ["✓"] * len(df):
            fail(which + ": the limits column hides the verdict the cells "
                 "ran with: %r" % list(df[lcol]))
        if "Priced for Profit" not in df.columns or not all(
                str(x).startswith("✗ ") and "without cost drivers" in str(x)
                for x in df["Priced for Profit"]):
            fail(which + ": no Priced for Profit column naming why: %r"
                 % (list(df.get("Priced for Profit", [])),))
        if "Error" in df.columns:
            fail(which + ": a pricing failure is shown as an engine error")
    else:
        if state == "none":
            want = "costs are not set"
        else:
            # The page's Profit gate IS the optimizers' refusal: the same
            # rule (io.check_costs), word for word.
            try:
                io.check_costs(P["costs"], "profit",
                               io.feed_type_names(P["proj"]))
                want = None
            except ValueError as e:
                want = str(e)
            if not want or P["dropped"] not in want:
                fail(which + ": check_costs does not refuse the partial "
                     "costs: %r" % want)
        if not any(want in w for w in warn):
            fail(which + ": no warning naming %r: %r" % (want, warn))
        if not at.button(key=bkey).disabled:
            fail(which + ": Find is live for Profit without complete costs")
        at.button(key=bkey).click().run()
        if len(calls) != n0:
            fail(which + ": the optimizer ran on a disabled Profit")
        at.radio(key=rkey).set_value("revenue").run()
        boom(at, which + " back to revenue")
        if at.button(key=bkey).disabled:
            fail(which + ": revenue is disabled too")
        at.button(key=bkey).click().run()
        boom(at, which + " find revenue")
        if len(calls) != n0 + 1 or calls[-1]["objective"] != "revenue":
            fail(which + ": revenue did not run")
        if state == "none" and calls[-1]["costs"] is not None:
            fail(which + ": costs passed although none are set")
        if state == "partial":
            if calls[-1]["costs"] != P["costs"]:
                fail(which + ": the page did not pass the saved costs")
            if not any("leave out feed with no price" in w.value
                       and P["dropped"] in w.value for w in at.warning):
                fail(which + ": the unpriced feed is not warned")
    bad = no_help(at, (pref,))
    if bad:
        fail(which + ": widgets with no help: %r" % bad)

# ---- Step 3: show profit, don't rank by it (operator, 2026-09-11) ----
TR_LABELS = ["Revenue", "Harvest tonnage (HOG)", "Biomass gain"]
RK3, BK3, UK3, LC3 = ("ideal_tr_opt_obj", "ideal_tr_opt_run",
                      "ideal_tr_opt_use", "Within the limits")

def tr_table(at):
    tbl = [d for d in at.dataframe if LC3 in list(d.value.columns)]
    if not tbl:
        fail("step3: no table")
    return tbl[0].value

def has(at, k):
    try:
        at.session_state[k]
    except KeyError:
        return False
    return True

at = AppTest.from_string(SCRIPT, default_timeout=120)
at.session_state["which"] = "step3"
at.run()
boom(at, "step3 render")
opts = list(at.radio(key=RK3).options)
if opts != TR_LABELS:
    fail("step3: the objectives are not exactly Revenue / HOG / Biomass "
         "gain: %r" % opts)
if not any("does not rank by profit" in c.value and "step 2" in c.value
           and "smaller future batches" in c.value for c in at.caption):
    fail("step3: no caption saying why it does not rank by profit")
if any("Profit is unavailable" in w.value for w in at.warning):
    fail("step3: a Profit gate is drawn")
if at.button(key=BK3).disabled:
    fail("step3: Find is disabled")
n0 = len(TR_CALLS)
at.button(key=BK3).click().run()
boom(at, "step3 find")
if len(TR_CALLS) != n0 + 1 or TR_CALLS[-1]["objective"] != "revenue":
    fail("step3: the optimizer was not asked for revenue: %r"
         % [x["objective"] for x in TR_CALLS])
if TR_CALLS[-1]["costs"] != P["costs"]:
    fail("step3: the page did not pass the saved costs (None when unset)")
if not any("Best within the limits" in s.value for s in at.success):
    fail("step3: no winner line: %r" % [s.value[:200] for s in at.success])
if any("Showing the last" in w.value for w in at.warning):
    fail("step3: a fresh search is marked stale")
df = tr_table(at)
for col in ("Priced for Profit", "Error"):
    if col in df.columns:
        fail("step3: a %r column: %r" % (col, list(df[col])))
if list(df[LC3]) != ["✓"] * len(df):
    fail("step3: the limits column: %r" % list(df[LC3]))
lines = [x.value for x in at.caption
         if x.value.startswith("Cost $") and "→ profit" in x.value]
today_cost = any("; cost $" in x.value for x in at.caption)
if state in ("full", "partial"):
    if df["Cost $M"].isna().any():
        fail("step3: the Cost column is not filled: %r"
             % list(df["Cost $M"]))
    if state == "full" and df["Profit $M"].isna().any():
        fail("step3: the Profit column is not filled: %r"
             % list(df["Profit $M"]))
    if state == "partial":
        # Feed with no price left out of the cost: step 3's profit is
        # display only, so it is withheld (it would be overstated), as the
        # Run page withholds it — in the table, the winner's line and
        # today's line.
        if not df["Profit $M"].isna().all():
            fail("step3: an overstated profit is shown with unpriced feed: "
                 "%r" % list(df["Profit $M"]))
        if not lines or "profit not shown" not in lines[0]:
            fail("step3: the winner's line shows a profit with unpriced "
                 "feed: %r" % lines[:1])
        if not any("; cost $" in x.value and "profit not shown" in x.value
                   for x in at.caption):
            fail("step3: today's line shows a profit with unpriced feed")
    elif lines and "profit not shown" in lines[0]:
        fail("step3: profit withheld with every feed priced")
    if not lines:
        fail("step3: no cost breakdown under the winner")
    # Over a run window the spend includes fish not yet sold: no per-kg
    # figure, and the cash-window caveat instead.
    if "/kg HOG" in lines[0]:
        fail("step3: a cost per kg HOG over the run window")
    if "Cash over the run window" not in lines[0]:
        fail("step3: the winner's cost line has no cash caveat")
    if not today_cost:
        fail("step3: today's plan line has no cost")
else:
    # Nothing priced: blank, never 0.
    for col in ("Cost $M", "Profit $M"):
        if not df[col].isna().all():
            fail("step3: %r shows a number for an unpriced plan: %r"
                 % (col, list(df[col])))
    if lines or today_cost:
        fail("step3: a cost is shown for an unpriced plan")
    if state == "undriven" and not any(
            x.value.startswith("Cost not priced")
            and "without cost drivers" in x.value for x in at.caption):
        fail("step3: the unpriced winner does not say why")
if state == "partial" and not any("leaves out feed with no price" in w.value
                                  and "Profit is not shown" in w.value
                                  and P["dropped"] in w.value
                                  for w in at.warning):
    fail("step3: the unpriced feed (and the withheld profit) is not warned")
if state == "full":
    # No step-3 search ranks by profit, so a costs save never stales one:
    # the pick stays current (Use live), and a note says its Cost / Profit
    # columns used the costs as of the search.
    orig = open(COSTS_YAML, "rb").read()
    try:
        with open(COSTS_YAML, "wb") as fh:
            fh.write(orig.replace(b"fixed_monthly: 1000.0",
                                  b"fixed_monthly: 2000.0"))
        at.run()
        boom(at, "step3 costs edited")
        if any("Showing the last" in w.value for w in at.warning):
            fail("step3: a costs edit made the search stale")
        if at.button(key=UK3).disabled:
            fail("step3: Use greyed out after a costs edit")
        if not any("Your costs have changed since this search" in x.value
                   for x in at.caption):
            fail("step3: no note that the cost columns used old costs")
    finally:
        with open(COSTS_YAML, "wb") as fh:
            fh.write(orig)
bad = no_help(at, ("ideal_tr_opt",))
if bad:
    fail("step3: widgets with no help: %r" % bad)

# A "profit" remembered from before the ruling — in the widget key, in its
# _keep_ shadow (a mode round trip), or as a kept search ranked by it — is
# reset to revenue before the radio draws: no exception, a one-time note.
if state in ("none", "full"):
    kept = {"sig": "old", "res": SimpleNamespace(objective="profit"),
            "caps_kg": (), "costs_sig": None}
    for case, seed in (("widget key", {RK3: "profit"}),
                       ("_keep_ shadow", {"_keep_" + RK3: "profit"}),
                       ("kept search", {"_keep_" + RK3: "profit",
                                        "_ideal_tr_opt": kept})):
        at = AppTest.from_string(SCRIPT, default_timeout=120)
        at.session_state["which"] = "step3"
        for k, v in seed.items():
            at.session_state[k] = v
        at.run()
        boom(at, "step3 remembered profit (%s)" % case)
        if at.radio(key=RK3).value != "revenue":
            fail("step3 (%s): the remembered profit was not reset: %r"
                 % (case, at.radio(key=RK3).value))
        notes = [i.value for i in at.info if "no longer ranks by" in i.value]
        if len(notes) != 1:
            fail("step3 (%s): want one reset note, got %r" % (case, notes))
        if case == "kept search" and ("cleared" not in notes[0]
                                      or has(at, "_ideal_tr_opt")):
            fail("step3: the kept profit-ranked search was not cleared")
        if at.session_state["_keep_" + RK3] != "revenue":
            fail("step3 (%s): the _keep_ shadow still says %r"
                 % (case, at.session_state["_keep_" + RK3]))
        at.run()
        boom(at, "step3 rerun after the reset (%s)" % case)
        if any("no longer ranks by" in i.value for i in at.info):
            fail("step3 (%s): the reset note is not one-time" % case)
print("OK ideal Profit (%s)" % state)
'''


@pytest.mark.parametrize("state", ["none", "partial", "full", "undriven"])
def test_ideal_profit_ranks_in_step2_and_is_only_shown_in_step3(tmp_path,
                                                                 state):
    made_up = {"none": None, "partial": _costs(drop=(NAMES[-1],)),
               "full": _costs(), "undriven": _costs()}[state]
    proj = _project(tmp_path, made_up)
    before = _hashes(proj)
    _drive(tmp_path, _IDEAL, state=state, proj=str(proj), names=NAMES,
           dropped=NAMES[-1],
           costs=None if made_up is None else C.validate_costs(made_up))
    assert _hashes(proj) == before                 # the page writes nothing


# --------------------------------------------------------------------------- #
# Run forecast: the Costs & profit tab
# --------------------------------------------------------------------------- #
def _ok_sheet(tmp_path, odd=False):
    """A CostsAndProfit sheet written and read back by forecast.costs_report
    on test_costs_report's made-up inputs: (workbook, parsed, config dir).
    With `odd`, one feed type (test_costs_report.ODD) is fed but has no
    price, so its kg are left out of the cost."""
    import test_costs_report as tcr
    inp = tcr._inputs(odd=odd)
    wb = tcr._workbook(inp)
    cfg = tcr._cfg(tmp_path)                    # made-up costs + economics
    assert R.write_costs_sheet(wb, config_dir=cfg, **inp)["status"] == "ok"
    return wb, R.read_costs_sheet(wb[R.COSTS_SHEET]), cfg


_RUNTAB = _PRELUDE + r'''
SCRIPT = "P = " + ascii(P) + "\n" + LIFT + r"""
import sys
from pathlib import Path
import pandas as pd
import plotly.express as px
import streamlit as st
sys.path.insert(0, P["root"])
ns = {"st": st, "pd": pd, "px": px}
exec(compile(lift(P["root"], names=("_run_costs_tab", "_run_costs_caption",
                                    "_run_costs_unpriced", "_run_money")),
             "app.py", "exec"), ns)
r = st.session_state["r"]
line = ns["_run_costs_caption"](r)
if line:
    st.caption(line)
ns["_run_costs_tab"](r, Path(P["cfg"]))
"""

def page(r):
    at = AppTest.from_string(SCRIPT, default_timeout=120)
    at.session_state["r"] = r
    at.run()
    boom(at, "render %r" % (r.get("costs") or {}).get("status"))
    return at

at = page({})
if not any("predates costs" in i.value for i in at.info) or at.metric:
    fail("a result from before costs does not say so")
at = page({"costs": {"status": "absent"}})
if not any("Costs not set" in i.value for i in at.info) or at.metric:
    fail("costs not set is not said")
at = page({"costs": {"status": "error", "error": "KeyError: boom"}})
if not any("could not be read" in e.value and "boom" in e.value
           for e in at.error) or at.metric:
    fail("an unreadable sheet is not named")
at = page({"costs": {"status": "not_computed",
                     "error": "config/costs.yaml: egg_price is missing"}})
if not any("not computed" in e.value and "egg_price is missing" in e.value
           for e in at.error) or at.metric:
    fail("a NOT COMPUTED sheet is not named")

ok = {"costs": P["ok"]}
at = page(ok)
labels = [m.label for m in at.metric]
if labels != ["Revenue", "Total cost", "Profit", "Spend per kg HOG sold"]:
    fail("the four metrics: %r" % labels)
if [m.value for m in at.metric if m.label == "Profit"] == ["—"]:
    fail("Profit withheld on a fully priced sheet")
if any(not (m.help or "").strip() for m in at.metric):
    fail("a metric has no help")
if at.warning:
    fail("stale warning on a current sheet: %r" % [w.value for w in at.warning])
if not any(c.value.startswith("Profit over the horizon") for c in at.caption):
    fail("no one-line summary")
if not any("Cash view:" in c.value and "shows a loss" in c.value
           for c in at.caption):
    fail("no cash-view caption")
if not any("economics.yaml price bands" in c.value and "config/costs.yaml" in c.value
           for c in at.caption):
    fail("the tab does not say which pricing it uses")
months = [d for d in at.dataframe if "Period" in list(d.value.columns)]
if not months or len(months[0].value) != len(P["ok"]["months"]):
    fail("no monthly table with one row per month")
feed = [d for d in at.dataframe if "Item" in list(d.value.columns)]
if not feed or not any(P["item"] == x for x in feed[0].value["Item"]):
    fail("the feed-by-type table has no item names")
if not at.get("plotly_chart"):
    fail("no monthly chart")

# Feed of a type with no price is left out of the cost: Profit would be
# overstated, so it is withheld (caption, metric, rows) and warned.
at = page({"costs": P["unpriced"]})
if [m.value for m in at.metric if m.label == "Profit"] != ["—"]:
    fail("Profit shown although feed is unpriced: %r"
         % [(m.label, m.value) for m in at.metric])
if not any("has no price" in w.value and P["odd"] in w.value
           and "Profit is not shown" in w.value for w in at.warning):
    fail("no warning naming the unpriced feed: %r"
         % [w.value[:200] for w in at.warning])
kpi = [c.value for c in at.caption if "over the horizon" in c.value]
if not (kpi and kpi[0].startswith("Cost over the horizon")
        and "profit not shown" in kpi[0] and P["odd"] in kpi[0]):
    fail("the KPI caption shows profit with unpriced feed: %r" % kpi)
mt = [d for d in at.dataframe if "Period" in list(d.value.columns)][0].value
pcol = [c for c in mt.columns if c.startswith("Profit (")][0]
unp = mt["Unpriced feed (kg)"].fillna(0) > 0
if not unp.any() or not mt.loc[unp, pcol].isna().all():
    fail("a month with unpriced feed shows a profit")
if mt.loc[~unp, pcol].isna().all():
    fail("fully priced months lost their profit too")
if not any("has no price" in c.value and "profit overstated" in c.value
           for c in at.caption):
    fail("the sheet's note does not say feed was unpriced")

with open(os.path.join(P["cfg"], "costs.yaml"), "w", encoding="utf-8") as fh:
    fh.write(P["other_costs"])
at = page(ok)
if not any("Costs changed since this run" in w.value for w in at.warning):
    fail("no stale warning after costs.yaml changed")
os.remove(os.path.join(P["cfg"], "costs.yaml"))
at = page(ok)
if not any("removed since this run" in w.value for w in at.warning):
    fail("no warning after costs.yaml was removed")
if any("price bands" in w.value for w in at.warning):
    fail("a price-band warning with the bands unchanged")
# Revenue and profit are priced with the bands as they were: a later edit
# to economics.yaml is warned too.
with open(os.path.join(P["cfg"], "economics.yaml"), "a",
          encoding="utf-8") as fh:
    fh.write("\n# edited after the run\n")
at = page(ok)
if not any("price bands" in w.value and "changed since this run" in w.value
           for w in at.warning):
    fail("no warning after economics.yaml changed")
print("OK run costs tab: predates / absent / error / not computed / ok / "
      "unpriced / stale costs / stale price bands")
'''


def test_run_costs_tab_every_case(tmp_path):
    import test_costs_report as tcr
    _wb, parsed, cfg = _ok_sheet(tmp_path)
    assert parsed["status"] == "ok"
    (tmp_path / "u").mkdir()
    _wb2, unpriced, _cfg2 = _ok_sheet(tmp_path / "u", odd=True)
    assert unpriced["total"]["unpriced_kg"] > 0
    other = cfg.parent / "other"
    other.mkdir()
    C.save_costs(other, dict(C.load_costs(cfg), fixed_monthly=4321.0))
    _drive(tmp_path, _RUNTAB, ok=json.loads(json.dumps(parsed)),
           unpriced=json.loads(json.dumps(unpriced)), odd=tcr.ODD,
           cfg=str(cfg), item="Acme Crumble No.1",
           other_costs=(other / "costs.yaml").read_text(encoding="utf-8"))


def test_parse_costs_sheet_absent_ok_and_error(tmp_path, monkeypatch):
    parse = _app_functions("_parse_costs_sheet")["_parse_costs_sheet"]
    assert parse(openpyxl.Workbook()) == {"status": "absent"}
    wb, parsed, _cfg = _ok_sheet(tmp_path)
    assert parse(wb) == parsed and parsed["status"] == "ok"
    foreign = openpyxl.Workbook()
    foreign.create_sheet(R.COSTS_SHEET).append(["hello"])
    assert parse(foreign)["status"] == "error"

    def raising(ws):
        raise ValueError("boom")
    monkeypatch.setattr(R, "read_costs_sheet", raising)
    assert parse(wb) == {"status": "error", "error": "ValueError: boom"}


def _func(tree, name):
    return next(n for n in tree.body
                if isinstance(n, ast.FunctionDef) and n.name == name)


def test_the_run_page_wires_the_sheet_and_an_eighth_tab_last():
    src = APP.read_text(encoding="utf-8")
    tree = ast.parse(src)
    parse = _func(tree, "_parse_output_workbook")
    ret = [n for n in ast.walk(parse) if isinstance(n, ast.Return)
           and isinstance(n.value, ast.Dict)][-1].value
    keys = [k.value for k in ret.keys]
    call = ret.values[keys.index("costs")]
    assert isinstance(call, ast.Call) and call.func.id == "_parse_costs_sheet"
    tabs = [n for n in ast.walk(tree) if isinstance(n, ast.Assign)
            and isinstance(n.value, ast.Call)
            and getattr(n.value.func, "attr", "") == "tabs"
            and any(isinstance(e, ast.Constant) and e.value == "Costs & profit"
                    for e in n.value.args[0].elts)]
    assert len(tabs) == 1
    names = [e.value for e in tabs[0].value.args[0].elts]
    targets = [e.id for e in tabs[0].targets[0].elts]
    assert names == ["Overview", "Per-Batch", "Period Summary", "Harvest",
                     "Feed", "Yearly", "Plan", "Costs & profit"]
    assert targets[-1] == "tab_costs" and len(targets) == 8
    assert "with tab_costs:\r\n        _run_costs_tab(r, CONFIG_DIR)" in src \
        or "with tab_costs:\n        _run_costs_tab(r, CONFIG_DIR)" in src


# --------------------------------------------------------------------------- #
# help= on every new input (the operator asked for help popups throughout)
# --------------------------------------------------------------------------- #
_WIDGETS = {"button", "download_button", "number_input", "text_input",
            "text_area", "radio", "selectbox", "multiselect", "slider",
            "select_slider", "checkbox", "toggle", "date_input", "time_input",
            "file_uploader", "color_picker", "metric",
            "Column", "TextColumn", "NumberColumn", "SelectboxColumn",
            "CheckboxColumn", "DateColumn", "ListColumn", "LinkColumn",
            "ProgressColumn"}
_HELP_FUNCS = ("_edit_costs", "_ideal_optimizer", "_ideal_tr_optimizer",
               "_ideal_opt_stability", "_ideal_tr_opt_stability",
               "_ideal_ref_cost_metrics", "_run_costs_tab")


def _empty_help(v) -> bool:
    if isinstance(v, ast.Constant):
        return not (isinstance(v.value, str) and v.value.strip())
    if isinstance(v, ast.JoinedStr):         # an f-string with no words
        return not any(isinstance(p, ast.Constant) and str(p.value).strip()
                       for p in v.values)
    return False    # a name or expression: the page tests judge it drawn


def _widgets_without_help(fn):
    """(widget, line) for every widget call in `fn` with no or empty help=,
    and the number of widget calls seen."""
    bad, n = [], 0
    for node in ast.walk(fn):
        if (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
                and node.func.attr in _WIDGETS):
            n += 1
            kw = {k.arg: k.value for k in node.keywords}
            if "help" not in kw or _empty_help(kw["help"]):
                bad.append((node.func.attr, node.lineno))
    return bad, n


def test_every_new_widget_has_help():
    tree = ast.parse(APP.read_text(encoding="utf-8"))
    for name in _HELP_FUNCS:
        bad, n = _widgets_without_help(_func(tree, name))
        assert n, f"{name}: no widget seen — the scan would pass vacuously"
        assert not bad, f"{name}: widgets with no help= {bad}"
    # The data editor's every column is configured (with help, above).
    ec = _func(tree, "_edit_costs")
    de = next(n for n in ast.walk(ec) if isinstance(n, ast.Call)
              and getattr(n.func, "attr", "") == "data_editor")
    cc = {k.arg: k.value for k in de.keywords}["column_config"]
    assert sorted(k.value for k in cc.keys) == sorted(
        ["Feed type", "Up to size (g)", "Item", "Price / kg"])


def test_the_help_scan_catches_a_missing_help():
    fn = ast.parse(
        "def f(c1, x):\n"
        "    st.number_input('a', key='k')\n"
        "    c1.button('b', help='')\n"
        "    st.metric('m', 1, help=f'{x}')\n"
        "    st.column_config.NumberColumn('n')\n"
        "    st.button('ok', help='Says what it does.')\n").body[0]
    bad, n = _widgets_without_help(fn)
    assert n == 5
    assert [b[0] for b in bad] == ["number_input", "button", "metric",
                                   "NumberColumn"]


# --------------------------------------------------------------------------- #
# The Ideal texts for Profit (pure)
# --------------------------------------------------------------------------- #
def test_profit_texts_never_show_an_unpriced_zero():
    from types import SimpleNamespace as NS
    ns = _app_functions("_ideal_opt_value_text", "_ideal_opt_diff_text",
                        "_ideal_money")
    val, diff = ns["_ideal_opt_value_text"], ns["_ideal_opt_diff_text"]
    priced = NS(costs_applied=True, profit=-1.5e6, revenue=2e6, hog_t=1.0,
                gain_t=1.0)
    other = NS(costs_applied=True, profit=0.5e6, revenue=1e6, hog_t=1.0,
               gain_t=1.0)
    bare = NS(costs_applied=False, profit=0.0, revenue=2e6, hog_t=1.0,
              gain_t=1.0)
    assert val(priced, "profit") == "profit −$1.5M"
    assert val(bare, "profit") == "profit — (not priced)"
    assert val(bare, "revenue") == "revenue $2.0M"      # unchanged
    assert diff(other, priced, "profit") == "+$2.0M"
    assert diff(bare, priced, "profit") == "—"
    assert diff(priced, bare, "profit") == "—"
    assert diff(priced, other, "revenue") == "+$1.0M"   # unchanged
    assert ns["_ideal_money"](-4e5) == "−$0.4M"
