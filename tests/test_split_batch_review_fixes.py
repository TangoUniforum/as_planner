"""Split FW/SW batch: the fixes from the 2026-09-12 review.

Each test names the review finding it guards and was proven by a negative
control (the fix reverted on a copy outside the sandbox -> the test fails;
restored -> it passes). Unit tests first (no workbook); the end-to-end ones run
the 8/31 PR on TEMP copies of config/ (costs.yaml never copied) and scenario/.
"""
from __future__ import annotations

import contextlib
import dataclasses
import io
import re
import shutil
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast.models import BatchInput, BatchWeekState

ROOT = Path(__file__).resolve().parent.parent
FS = date(2026, 9, 1)
PR_CLOSE = date(2026, 8, 31)


def _row(batch, week, ws, stage, *, open_c, close_c, wt=370.0, mort=0.0, cull=0.0):
    n = (open_c + close_c) / 2.0
    return BatchWeekState(
        batch_id=batch, week_label=week, week_start=datetime.combine(ws, datetime.min.time()),
        days_since_input=300, week_from_input=43, count=n, avg_weight_g=wt,
        biomass_kg=n * wt / 1000.0, feed_kg_day=0.0, feed_kg_week=0.0,
        sgr_pct_day=1.0, fcr=1.0, stage=stage, feed_type="F", mortality_pct_weekly=0.1,
        cull_count_week=cull, cull_biomass_kg_week=cull * wt / 1000.0,
        open_count=open_c, open_avg_weight_g=wt, open_biomass_kg=open_c * wt / 1000.0,
        close_count=close_c, close_avg_weight_g=wt, close_biomass_kg=close_c * wt / 1000.0,
        mort_count_week=mort)


def _loc(batch, wk, ws, count, tank=11):
    from forecast.placement import BatchLocationRow
    return BatchLocationRow(week_label=wk, week_start=ws, batch_id=batch, tank_id=tank,
                            location_id=f"OG1S-{tank}", system_id="OG1S", count=count,
                            avg_wt_g=400.0, biomass_kg=count * 0.4, density_kg_m3=20.0)


def _audit(bt, locs, states, placed, ctrl, **kw):
    from forecast import excel_io
    from forecast.events import TranOGEntry
    ev = TranOGEntry(batch_id=bt.batch_id, event_date=ctrl.forecast_start,
                     destinations=[SimpleNamespace(count=placed, tank_id=11)])
    wb = openpyxl.Workbook()
    excel_io.write_input_conservation_audit(
        wb, [bt], locs, [], ctrl, tranog_events=[ev],
        biology_states_by_batch={bt.batch_id: states}, **kw)
    rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
    return (dict(zip(rows[hi], rows[hi + 1])),
            [str(r[0]) for r in rows[:hi] if r and r[0]])


# ---- MAJOR 1: no false FW mass-balance breach on an auto split ----------------------

def _b40_split_track():
    """2025-06-30 B40, as measured: a one-week FW track whose mid-week
    tran_og_date puts the reconcile cull INSIDE its first (only) freshwater
    week. 97,133 FW fish = 91,783 placed + 40 mortality + 5,310 cull."""
    ws0, ws1 = date(2025, 6, 30), date(2025, 7, 7)
    fw = _row("B40", "2025-W27", ws0, "FW", open_c=97_133.0, close_c=91_783.0,
              mort=40.0, cull=5_310.0)
    sw = _row("B40", "2025-W28", ws1, "SW", open_c=91_783.0 + 88_000.0,
              close_c=91_700.0 + 87_900.0)
    bt = SimpleNamespace(batch_id="B40", input_count=400_000,
                         tran_og_date=date(2025, 7, 3), tran_og_count=185_000)
    ctrl = SimpleNamespace(forecast_start=date(2025, 7, 1), horizon_weeks=10)
    locs = [_loc("B40", "2025-W28", ws1, 179_600.0)]
    return bt, locs, [fw, sw], ctrl


def test_an_auto_split_that_conserves_reads_no_fw_breach():
    """Review MAJOR 1. The balance started from the first FW week's MEAN
    count, which for a one-week track already has half the reconcile cull
    taken out: B40 read -2,675 (-2.8%) and the headline said FISH LEAK. A
    split's balance starts from that week's OPENING (the PR's own count)."""
    bt, locs, states, ctrl = _b40_split_track()
    row, head = _audit(bt, locs, states, 91_783.0, ctrl,
                       split_remaining={"B40": 97_000.0})
    assert not any("BREACH" in h for h in head), head
    assert row["FW_Bal_Residual (fish)"] == 0, row["FW_Bal_Residual (fish)"]
    assert row["FW_Cull (fish)"] == 5_310 and row["FW_Mort (fish)"] == 40


def test_a_batch_that_is_not_an_auto_split_keeps_the_mean_basis():
    """Only auto splits changed: the same rows for an ordinary batch keep the
    long-standing mean basis (V1 byte for byte) -- here that basis reads the
    breach the split no longer does."""
    bt, locs, states, ctrl = _b40_split_track()
    row, head = _audit(bt, locs, states, 91_783.0, ctrl)
    assert any("BREACH" in h and "B40" in h for h in head), head
    assert row["FW_Bal_Residual (fish)"] == -2_675


def test_a_past_date_split_balances_on_its_fw_track_entry_row():
    """2024-11-30 B36: its date had passed at the close, so its FW part
    crosses on day 1 -- no freshwater week -- and its seawater rows are the
    merged stream (SW part + arrival). Its FW balance was left blank; it is
    now taken from the FW track's own entry row (201,807 in, the 21-fish
    handling loss and 201,786 placed out)."""
    ws = date(2024, 12, 1)
    merged = _row("B36", "2024-W48", ws, "SW", open_c=106_207.0 + 201_807.0,
                  close_c=106_100.0 + 201_700.0, cull=21.0)
    entry = _row("B36", "2024-W48", ws, "SW", open_c=201_807.0, close_c=201_700.0,
                 cull=21.0)
    bt = SimpleNamespace(batch_id="B36", input_count=320_000,
                         tran_og_date=date(2024, 11, 30), tran_og_count=106_207)
    ctrl = SimpleNamespace(forecast_start=ws, horizon_weeks=10)
    locs = [_loc("B36", "2024-W48", ws, 307_800.0)]
    row, head = _audit(bt, locs, [merged], 201_786.0, ctrl,
                       split_remaining={"B36": 0.0}, split_entry_rows={"B36": entry})
    assert row["FW_Bal_Residual (fish)"] == 0 and row["FW_Cull (fish)"] == 21
    assert not any("BREACH" in h for h in head), head
    row0, _h = _audit(bt, locs, [merged], 201_786.0, ctrl, split_remaining={"B36": 0.0})
    assert row0["FW_Bal_Residual (fish)"] in ("", None)          # unmeasured before


# ---- overdue batch: In_Horizon and its FW balance ----------------------------------

def _b50_overdue():
    """B50 on 8/31 with tran_og_date 2026-08-27: every fish in freshwater,
    moved on day 1 of the first forecast week -- no FW week at all; the entry
    row opens on the PR's 254,135 and carries the 25-fish handling loss."""
    st = _row("B50", "2026-W36", FS, "SW", open_c=254_135.0, close_c=253_983.0,
              wt=200.0, cull=25.0)
    bt = SimpleNamespace(batch_id="B50", input_count=550_000,
                         tran_og_date=date(2026, 8, 27), tran_og_count=290_000)
    ctrl = SimpleNamespace(forecast_start=FS, horizon_weeks=10)
    return bt, [_loc("B50", "2026-W36", FS, 253_983.0)], [st], ctrl


def test_an_overdue_batch_reads_in_horizon_and_balances_on_its_entry_row():
    """Numbers-check C15: the overdue batch the automatic path placed in week
    1 read In_Horizon N (its scenario date is before the start) and FW_Cull 0
    while the ledger booked the 25-fish handling loss."""
    bt, locs, states, ctrl = _b50_overdue()
    row, _head = _audit(bt, locs, states, 254_110.0, ctrl,
                        overdue_effective={"B50": FS})
    assert row["In_Horizon"] == "Y"
    assert row["FW_Cull (fish)"] == 25
    assert row["FW_Bal_Residual (fish)"] == 0
    # without the name the audit reads it as before
    row0, _h = _audit(bt, locs, states, 254_110.0, ctrl)
    assert row0["In_Horizon"] == "N" and row0["FW_Cull (fish)"] == 0
    assert row0["FW_Bal_Residual (fish)"] in ("", None)


# ---- MINOR 8: a refused manual fw_to_og does not "move" the part ---------------------

def test_a_refused_fw_to_og_leaves_the_split_flagged_not_placed():
    """Review MINOR 8. The manual fw_to_og always wins, even when every one
    of its destinations is refused -- but then it moved nothing, and the
    split's freshwater fish must be named, not read PLACED."""
    from forecast.events import TranOGEntry
    from forecast.excel_io import unmodelled_fw_warnings
    held = {"B49": (250_225.0, 92_350.0)}
    refused = TranOGEntry(batch_id="B49", event_date=FS, destinations=[])
    refused.count_placed, refused.count_refused = 0.0, 240_000.0
    lines, fish = unmodelled_fw_warnings(held, [refused], {"B49"})
    assert fish == {"B49": 250_225.0}
    assert len(lines) == 1 and lines[0].startswith("SPLIT BATCH AT PR CLOSE - B49")
    ok = TranOGEntry(batch_id="B49", event_date=FS, destinations=[])
    ok.count_placed = 240_000.0
    assert unmodelled_fw_warnings(held, [ok], {"B49"}) == ([], {})
    # an event never applied is judged by its presence, as before
    assert unmodelled_fw_warnings(held, [SimpleNamespace(batch_id="B49")], {"B49"}) == ([], {})


# ---- MINOR 5: the line never claims a cull that was not made -------------------------

def test_the_auto_line_claims_no_cull_that_was_not_made():
    """Review MINOR 5. B50 overdue on 8/31: 254,135 FW fish against
    tran_og_count 290,000 -- only the 25-fish handling loss, yet the line said
    'after ... the reconcile cull to tran_og_count 290,000'."""
    from forecast.split_batch import auto_transfer_line
    kw = dict(fw_count=254_135.0, week_label="2026-W36", placed=254_110.0,
              tran_og_date=date(2026, 8, 27), pr_closing=PR_CLOSE, sw_count=0.0,
              overdue=True)
    none = auto_transfer_line("B50", tran_og_count=290_000, arrival=254_110.0, **kw)
    assert "NO reconcile cull" in none and "cull to tran_og_count" not in none
    culled = auto_transfer_line("B50", tran_og_count=200_000, arrival=200_000.0,
                                **{**kw, "placed": 200_000.0})
    assert "reconcile cull to tran_og_count 200,000" in culled
    sk = dict(fw_count=250_225.0, week_label="2026-W38", tran_og_date=date(2026, 9, 14),
              pr_closing=PR_CLOSE, tran_og_count=290_000, sw_count=47_743.0)
    short = auto_transfer_line("B49", placed=230_000.0, arrival=230_000.0, **sk)
    assert "NO reconcile cull: fewer fish than the remaining target 242,257" in short
    full = auto_transfer_line("B49", placed=242_257.0, arrival=242_257.0, **sk)
    assert "reconcile cull to the remaining target 242,257" in full


# ---- MINOR 6: the top-up past the density target has its own category ----------------

def test_the_top_up_past_the_density_target_lands_in_its_own_category():
    """Review MINOR 6: the warning fell into the generic Phase-D bin."""
    from forecast import excel_io
    from forecast.split_batch import own_entry_tanks, plan_topup
    from forecast.models import SizeClassSplit
    from forecast.state import FacilityState, TankState

    def tank(tid, sysid, count, wt):
        t = TankState(location_id=f"{sysid}-{tid}", tank_id=tid, system_id=sysid,
                      volume_m3=500.0, max_density_kg_m3=85.0,
                      max_feed_kg_day_cap=3000.0, type="OG")
        t.assign(batch_id="B49", count=count, avg_wt_g=wt, cv_pct=16.0, stage="SW")
        return t
    st = FacilityState(today=FS, tanks=[tank(14, "OG1S", 25_000, 400.0),
                                        tank(24, "OG2S", 22_743, 480.0)])
    sp = SizeClassSplit(batch_id="B49", tran_og_date=datetime(2026, 9, 14),
                        post_cull_count=242_257.0, post_cull_avg_wt_g=464.0,
                        post_cull_cv_pct=16.0, big_class_count=121_128.5,
                        big_class_avg_wt_g=523.0, small_class_count=121_128.5,
                        small_class_avg_wt_g=405.0)
    tp = plan_topup(sp, own_entry_tanks(st, "B49"), [], density_target_pct=0.85)
    assert tp.warnings
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(
        wb, placement_warnings=[f"[D] 2026-W38: {w}" for w in tp.warnings])
    cats = {r[1] for r in wb["ValidationLog"].iter_rows(values_only=True)
            if r and isinstance(r[0], int)}
    assert cats == {"WARNING - Split batch top-up past the density target"}, cats


# ---- MINOR 11: exactly the two documented words --------------------------------------

def test_split_batch_fw_accepts_only_auto_off_or_a_yaml_boolean():
    """Review MINOR 11: the docs say auto / off and anything else is refused,
    but 'on', 'true', 'yes', '1', 'false', 'no', '0' were silently accepted."""
    from forecast.split_batch import mode
    for bad in ("on", "true", "yes", "1", "false", "no", "0", "of"):
        with pytest.raises(ValueError, match="split_batch_fw"):
            mode(SimpleNamespace(split_batch_fw=bad))
    assert mode(SimpleNamespace(split_batch_fw="Auto")) == "auto"
    assert mode(SimpleNamespace(split_batch_fw=" OFF ")) == "off"
    assert mode(SimpleNamespace(split_batch_fw=True)) == "auto"     # YAML on / yes
    assert mode(SimpleNamespace(split_batch_fw=False)) == "off"     # YAML off / no


def test_the_configure_list_shows_what_the_run_would_use():
    pytest.importorskip("streamlit")
    import app
    assert app._CONTROL_CHOICES["split_batch_fw"] == ("auto", "off")
    idx = app._control_choice_index
    assert idx("split_batch_fw", "auto") == 0
    assert idx("split_batch_fw", None) == 0                  # blank = the default
    assert idx("split_batch_fw", False) == 1                 # YAML off
    assert idx("split_batch_fw", "maybe") is None            # a typo is shown as one


# ---- MINOR 9: the hybrid guide line for an overdue batch -----------------------------

def test_the_overdue_hybrid_guide_line_is_an_info_hybrid_guide_line():
    from forecast import excel_io
    from forecast.split_batch import hybrid_guide_overdue_line
    ln = hybrid_guide_overdue_line("B50", 254_110.0, "2026-W36")
    assert ln.startswith("HYBRID GUIDE - overdue FW batch B50:")
    assert "254,110" in ln and "2026-W36" in ln and "whole horizon" in ln
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(wb, invariant_warnings=[ln])
    (cat,) = [r[1] for r in wb["ValidationLog"].iter_rows(values_only=True)
              if r and isinstance(r[0], int)]
    assert cat == "INFO - Hybrid guide (L1) decision"


# ---- MAJOR 2: the editor offers what the run accepts ---------------------------------

def _meta(batch_id, tog):
    return BatchInput(
        batch_id=batch_id, input_date=datetime(2025, 9, 11), input_count=550_000,
        tran_sf_date=datetime(2025, 12, 8), tran_og_date=tog,
        tran_og_count=290_000, tran_og_avg_wt_g=370.0, tran_og_cv=16.0,
        fcr_model="FCR_116_Quick", fw_correction=0.9517, sgr_correction=0.95)


def _fwrec(batch, count, kg):
    return SimpleNamespace(batch_id=batch, unit_label="PostS.01", fw_system="PostS",
                           closing_count=count, closing_biomass_kg=kg)


@pytest.fixture(scope="module")
def tables():
    from forecast.config_io import load_biology_tables
    return load_biology_tables(ROOT / "config")


@pytest.fixture(scope="module")
def control():
    from forecast.config_io import load_control
    c = load_control(ROOT / "config")
    c.forecast_start = datetime(2026, 9, 1)
    return c


def test_the_editor_offers_a_past_date_batch_in_week_1_like_the_run(tables, control):
    """Review MAJOR 2. The run's lookup got the week-1 fallback, the editor's
    list did not: for a past-date split (B49 dated 2026-08-20) and an overdue
    wholly-FW batch (B50 dated 2026-08-27) _mw_fw_avail returned {} while the
    run demanded an fw_to_og for them inside a window -- advice the app could
    not follow. Both must now offer week 1 at the PR's own count and weight,
    exactly as the run's _build_fw_lookup does."""
    pytest.importorskip("streamlit")
    import app
    from forecast.manual_events import ManualDest, ManualEvent
    from forecast.manual_window import _build_fw_lookup
    meta = {"B49": _meta("B49", datetime(2026, 8, 20)),
            "B50": _meta("B50", datetime(2026, 8, 27)),
            "B51": _meta("B51", datetime(2026, 10, 5))}         # future: FW weeks
    recs = [_fwrec("B49", 250_225.0, 92_333.0), _fwrec("B50", 254_135.0, 50_987.0),
            _fwrec("B51", 200_000.0, 40_000.0)]
    labels = ["2026-W36", "2026-W37"]
    ctx = dict(fw_records=recs, control=control, batch_by_id=meta, tables=tables,
               pr_closing=PR_CLOSE, forecast_start=FS)
    avail = app._mw_fw_avail(ctx, labels)
    evs = [ManualEvent(type="fw_to_og", week=1, batch=b, destinations=[ManualDest(tank=1)])
           for b in ("B49", "B50", "B51")]
    run = _build_fw_lookup(evs, recs, control, PR_CLOSE, tables, meta)
    for b, n in (("B49", 250_225.0), ("B50", 254_135.0)):
        assert b in avail, (b, sorted(avail))
        assert list(avail[b]) == ["2026-W36"], avail[b]           # week 1 only
        assert avail[b]["2026-W36"] == pytest.approx(run[(b, "2026-W36")])
        assert avail[b]["2026-W36"][0] == pytest.approx(n)        # the PR's count
    # a batch still in freshwater keeps its projected weeks (no fallback):
    # week 2 is offered too, and each week offers the fish at its START
    # (manual_window.fw_week_start_states) -- week 1 the PR's own count,
    # week 2 less (a week of freshwater losses). Until 2026-09-15 week 1
    # offered the week's CLOSE, grown a week the fish then grew again.
    assert set(avail["B51"]) == {"2026-W36", "2026-W37"}
    assert avail["B51"]["2026-W36"][0] == pytest.approx(200_000.0)
    assert avail["B51"]["2026-W37"][0] < 200_000.0
    assert avail["B51"]["2026-W36"] == pytest.approx(run[("B51", "2026-W36")])


# ---- end to end: the 8/31 PR ------------------------------------------------------

PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")
_needs_pr = pytest.mark.skipif(not PR_0831.exists(),
                               reason="the 2026-08-31 ProductionReport is not on this machine")


def _rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _table(wb, name, first):
    rows = _rows(wb[name])
    h = next(i for i, r in enumerate(rows) if r and r[0] == first)
    hdr = [str(c) if c is not None else "" for c in rows[h]]
    return [dict(zip(hdr, r)) for r in rows[h + 1:] if r and any(c is not None for c in r)]


def _read(path, stdout):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tp = _rows(wb["TransferPlan"])
    th = next(i for i, r in enumerate(tp) if r and r[0] == "Week")
    col = {str(c): j for j, c in enumerate(tp[th]) if c is not None}
    tog: dict = {}
    for r in tp[th + 1:]:
        if r and r[col["Type"]] == "TranOG":
            k = (str(r[col["Batch"]]), str(r[col["Week"]]))
            tog[k] = tog.get(k, 0.0) + float(r[col["Count (fish)"]] or 0)
    vlog = [(str(r[1]), str(r[2])) for r in _rows(wb["ValidationLog"])
            if r and isinstance(r[0], int)]
    ica = _rows(wb["InputConservationAudit"])
    hi = next(i for i, r in enumerate(ica) if r and r[0] == "Batch")
    out = SimpleNamespace(
        tog=tog, vlog=vlog, stdout=stdout,
        ica={r[0]: dict(zip(ica[hi], r)) for r in ica[hi + 1:] if r and r[0]},
        ica_head=[str(r[0]) for r in ica[:hi] if r and r[0]],
        bio=_table(wb, "BiologyProjection", "Batch"),
        weekly=[d for d in _table(wb, "WeeklyReport", "Scenario") if d.get("Week")])
    wb.close()
    return out


def _edit_batch(scn, batch, **fields):
    p = scn / "batches.yaml"
    txt = p.read_text(encoding="utf-8")
    i = txt.find(f"- batch_id: {batch}\n")
    if i < 0:
        i = txt.find(f"- batch_id: {batch}\r\n")
    if i < 0:
        pytest.skip(f"{batch} is not in the live batches.yaml")
    for k, v in fields.items():
        j = txt.index(f"{k}:", i)
        e = txt.index("\n", j)
        if txt[e - 1] == "\r":
            e -= 1                                   # keep a CRLF line ending
        txt = txt[:j] + f"{k}: {v}" + txt[e:]
    p.write_text(txt, encoding="utf-8")


def _run(tmp, copy_config, *, events="shipped", scenario_edit=None, control_edit=None):
    """The 8/31 PR on TEMP copies of config/ (no costs.yaml) and scenario/."""
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    if control_edit is not None:
        p = cfg / "control.yaml"
        p.write_text(control_edit(p.read_text(encoding="utf-8")), encoding="utf-8")
    ev = scn / "manual_events" / "2026-08-31.yaml"
    if events == "empty":
        ev.parent.mkdir(exist_ok=True)
        ev.write_text("events: []\n", encoding="utf-8")
    elif events == "no_fw_to_og":
        import yaml
        d = yaml.safe_load(ev.read_text(encoding="utf-8")) or {}
        d["events"] = [e for e in d.get("events", []) if e.get("type") != "fw_to_og"]
        ev.write_text(yaml.safe_dump(d, sort_keys=False), encoding="utf-8")
    if scenario_edit is not None:
        scenario_edit(scn)
    inp, out = tmp / "pr.xlsm", tmp / "out.xlsm"
    shutil.copy(PR_0831, inp)
    from forecast.run import main
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                  calib_log_path="")
    assert rc == 0
    if not out.exists():
        out = out.with_suffix(".xlsx")
    return _read(out, buf.getvalue())


def _past_b49(scn):
    _edit_batch(scn, "B49", tran_og_date="'2026-08-20'")


def _overdue_b50(scn):
    _edit_batch(scn, "B50", tran_og_date="'2026-08-27'")


@_needs_pr
@pytest.mark.parametrize("batch,events,edit", [
    ("B49", "no_fw_to_og", _past_b49),       # a past-date split, its fw_to_og removed
    ("B50", "shipped", _overdue_b50),         # an overdue wholly-FW batch
])
def test_the_window_guard_sends_a_past_date_batch_to_week_1(tmp_path, copy_config,
                                                           batch, events, edit):
    """Review MAJOR 2. With the shipped two-week window, a batch whose date
    had already passed at the PR close crosses in week 1. The refusal said
    'shorten the window so it ends before those dates' -- impossible for a
    week-1 date. It now names the one route: an fw_to_og in week 1."""
    with pytest.raises(ValueError) as ei:
        _run(tmp_path, copy_config, events=events, scenario_edit=edit)
    msg = str(ei.value)
    assert batch in msg and "PAST the automatic FW->OG entry" in msg
    assert f"{batch}: its transfer date had already passed at the PR close" in msg
    assert "script its fw_to_og in week 1" in msg


@_needs_pr
def test_a_past_date_split_scripted_in_week_1_is_planned(tmp_path, copy_config):
    """The route the refusal names works end to end: B49 dated 2026-08-20
    (past at the close) with the shipped week-1 fw_to_og runs, the event
    taking B49's PR freshwater count (the week-1 fallback) -- one arrival in
    2026-W36 at the event's 240,000 target, no automatic transfer."""
    r = _run(tmp_path, copy_config, events="shipped", scenario_edit=_past_b49)
    b49 = {k: v for k, v in r.tog.items() if k[0] == "B49"}
    assert list(b49) == [("B49", "2026-W36")], b49
    assert b49[("B49", "2026-W36")] == pytest.approx(240_000, abs=1)
    assert not any("auto-transferred" in d for _c, d in r.vlog)
    assert r.ica["B49"]["Status"] == "PLACED"
    assert r.ica_head[2].startswith("OK"), r.ica_head[:4]


@pytest.fixture(scope="module")
def hyb_both(tmp_path_factory, copy_config):
    """8/31, EMPTY events, hybrid: B49 split (auto, 2026-W38) and B50 made
    overdue (2026-08-27; auto, week 1) in one run."""
    if not PR_0831.exists():
        pytest.skip("the 2026-08-31 ProductionReport is not on this machine")

    def hybrid(txt):
        txt, n = re.subn(r"(?m)^hybrid_follow:.*$", "hybrid_follow: 'full'", txt)
        return txt if n else txt.rstrip("\n") + "\nhybrid_follow: 'full'\n"
    return _run(tmp_path_factory.mktemp("rf_hyb"), copy_config, events="empty",
                scenario_edit=_overdue_b50, control_edit=hybrid)


@_needs_pr
def test_biology_projection_lists_the_split_fw_part(hyb_both):
    """Review MINOR 7: BiologyProjection showed B49's 47,743 seawater fish
    alone for 2026-W36/W37 while the ledger opened on 297,968."""
    fw = {d["Week"]: d for d in hyb_both.bio if d["Batch"] == "B49" and d["Stage"] == "FW"}
    assert set(fw) == {"2026-W36", "2026-W37"}, sorted(fw)
    for d in fw.values():
        assert 240_000 < float(d["Count"]) < 251_000, d
    sw36 = [d for d in hyb_both.bio if d["Batch"] == "B49" and d["Week"] == "2026-W36"
            and d["Stage"] == "SW"]
    assert len(sw36) == 1 and float(sw36[0]["Count"]) < 50_000


@_needs_pr
def test_the_hybrid_guide_names_the_overdue_batch_too(hyb_both):
    """Review MINOR 9: the L1 guide never sees an overdue batch's fish, but
    the HYBRID GUIDE line was written for splits only."""
    ov = [(c, d) for c, d in hyb_both.vlog if d.startswith("HYBRID GUIDE - overdue FW batch")]
    assert len(ov) == 1, ov
    assert ov[0][0] == "INFO - Hybrid guide (L1) decision"
    assert "B50" in ov[0][1] and "2026-W36" in ov[0][1]
    sp = [d for _c, d in hyb_both.vlog if d.startswith("HYBRID GUIDE - split batch")]
    assert len(sp) == 1 and "B49" in sp[0]


@_needs_pr
def test_the_overdue_line_names_no_cull_that_did_not_happen(hyb_both):
    """Review MINOR 5, wired: B50's 254,135 fish are below tran_og_count
    290,000, so only the handling loss applies."""
    (line,) = [d for c, d in hyb_both.vlog if c == "WARNING - Overdue FW batch auto-transferred"]
    assert "NO reconcile cull" in line and "cull to tran_og_count" not in line, line


@_needs_pr
def test_the_audit_balances_the_split_and_the_overdue_batch(hyb_both):
    """Review MAJOR 1 + numbers-check C11/C15, wired: B49's residual read -63
    (its first FW week's mean count) and B50 read In_Horizon N with FW_Cull 0.
    Both balance exactly now."""
    b49, b50 = hyb_both.ica["B49"], hyb_both.ica["B50"]
    assert abs(b49["FW_Bal_Residual (fish)"]) <= 2, b49["FW_Bal_Residual (fish)"]
    assert b50["In_Horizon"] == "Y"
    assert b50["FW_Cull (fish)"] > 0
    assert abs(b50["FW_Bal_Residual (fish)"]) <= 2, b50["FW_Bal_Residual (fish)"]
    assert not any("BREACH" in h for h in hyb_both.ica_head), hyb_both.ica_head


@_needs_pr
def test_the_ledger_count_check_holds_for_the_whole_horizon(hyb_both):
    """Review MINOR 10: Count_Check ~0 was asserted on B49's first three weeks
    only. Every weekly row of both auto-moved batches, whole horizon: the
    failures this guards are +250,225 / +254,135 (fish held, never moved);
    the tolerance covers the small residuals measured on this run."""
    bad = [(d["Batch"], d["Week"], d["Count_Check (fish)"]) for d in hyb_both.weekly
           if d.get("Batch") in ("B49", "B50")
           and isinstance(d.get("Count_Check (fish)"), (int, float))
           and abs(d["Count_Check (fish)"]) > 100]
    assert not bad, bad[:5]
    assert sum(1 for d in hyb_both.weekly if d.get("Batch") == "B50") > 40
