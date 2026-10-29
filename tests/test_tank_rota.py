"""Tank rota: steps, the one-second tank count, the full design on a small synthetic facility, the independent check.

Synthetic facility: 2 entry tanks (OG1N, OG1S) and 3 grow tanks (OG3N) of 120 m3 at 85 kg/m3 (10.2 t). A batch of
16,000 fish every 2 weeks: weights 0.3..0.8 kg over 6 weeks, 8,000 fish staged into 6N at week 4 and 8,000 at week 5.
Steps 0:1, 2:2 (one entry tank for 2 weeks, then 2 grow tanks) fit: the busiest week needs 3 grow tanks.
Negative controls: 3 grow tanks per batch cannot fit (4 needed); the checker flags a rota whose batches share a tank.
"""
import copy

import pytest

pytest.importorskip('ortools')
from forecast import tank_rota as tr                                            # noqa: E402

NW = 24
WEEKS = [f'2030-W{i + 1:02d}' for i in range(NW)]
WT = [0.3, 0.4, 0.5, 0.6, 0.7, 0.8]
N = [16000, 16000, 16000, 16000, 8000, 0]
STAGE = [0, 0, 0, 0, 8000, 8000]


def _plan():
    batch = {}
    for k, i0 in enumerate(range(0, 16, 2)):
        rec = {}
        for i, w in enumerate(WEEKS):
            a = i - i0
            if 0 <= a < 6:
                rec[w] = {'n': float(N[a]), 'kg': N[a] * WT[a], 'wt': WT[a], 'arr': 16000.0 if a == 0 else 0.0,
                          'stage': float(STAGE[a]), 'sfr': 0.01}
            else:
                rec[w] = {'n': 0.0, 'kg': 0.0, 'wt': 0.0, 'arr': 0.0, 'stage': 0.0, 'sfr': 0.0}
        batch[f'B{k + 1}'] = rec
    tanks = [(11, 'OG1N', 120, 85), (12, 'OG1S', 120, 85), (31, 'OG3N', 120, 85), (32, 'OG3N', 120, 85),
             (33, 'OG3N', 120, 85)]
    return {'source': 'synthetic', 'weeks': WEEKS, 'tanks': tanks,
            'sys_cap': {'OG1N': (10 ** 6, 10 ** 5), 'OG1S': (10 ** 6, 10 ** 5), 'OG3N': (10 ** 6, 10 ** 5)},
            'buffer': 0.1, 'min_tank': 7000.0, 'min_move': 7000.0, 'move_rules': True, 'max_moves': 15,
            'dens_cap': {}, 'batch': batch}


def test_steps_parse_and_validate():
    assert tr.parse_steps('0:4, 21:6, 32:+1') == [{'week': 0, 'tanks': 4}, {'week': 21, 'tanks': 6},
                                                   {'week': 32, 'add': 1}]
    assert tr.validate_steps(tr.parse_steps('0:4, 21:6, 32:+1'), 12) == []
    assert tr.validate_steps(tr.parse_steps('0:4'), 12)                       # no move to grow-out
    assert tr.validate_steps(tr.parse_steps('0:4, 21:6, 32:2'), 12)           # a later step must be +n
    assert tr.validate_steps(tr.parse_steps('0:13, 21:6'), 12)                # more entry tanks than exist
    assert tr.validate_steps(tr.parse_steps('0:4, 21:6, 21:+1'), 12)          # weeks must increase


def test_profile_and_cadence():
    D = _plan()
    p = tr.batch_profile(D)
    assert p['cadence'] == 2 and p['cap'] == 85
    assert [a['n'] for a in p['ages']] == N                                   # the whole life, arrival to empty
    assert tr.batch_profile(D, fish=8000)['ages'][4]['stage'] == pytest.approx(4000)   # scaled


def test_quick_check_passes_and_fails():
    D = _plan(); p = tr.batch_profile(D)
    ok = tr.quick_check(D, p, tr.parse_steps('0:1, 2:2'))
    assert ok['ok'] and ok['peak'] == 3, ok['lines']
    bad = tr.quick_check(D, p, tr.parse_steps('0:1, 2:3'))                   # 3 + 1 grow tanks > 3
    assert not bad['ok'] and bad['peak'] == 4
    p22 = tr.batch_profile(D, fish=22000)                                    # 11 t at week 2 in one 10.2 t tank
    dense = tr.quick_check(D, p22, tr.parse_steps('0:1, 3:2'))
    assert not dense['ok'] and 'too dense' in ' '.join(dense['lines'])
    late = tr.quick_check(D, p, tr.parse_steps('0:1, 4:2'))                  # move after staging starts
    assert not late['ok']


@pytest.fixture(scope='module')
def done():
    D = _plan()
    return D, tr.run(D, tr.parse_steps('0:1, 2:2'), secs_batch=20, secs_tanks=30, workers=4, log=lambda s: None)


def test_full_run_is_clean(done):
    D, res = done
    assert res['stage'] == 'done' and res['ok'], res.get('lines') or res['eval']['check']['issues']
    ev = res['eval']
    assert ev['tpf'] == pytest.approx(1.0, abs=0.01)                          # every fish moves exactly once
    assert ev['max_moves'] <= 15


def test_checker_flags_a_shared_tank(done):
    D, res = done
    bad = copy.deepcopy(res['tanks'])
    for k in bad['grow']:
        bad['grow'][k] = 31                                                    # every grow role in one tank
    ev = tr.evaluate(D, res['profile'], res['design'], bad)
    assert not ev['check']['ok'] and ev['check']['issues'].get('two batches in a tank', 0) > 0


def test_workbook(done, tmp_path):
    import openpyxl
    D, res = done
    out = tmp_path / 'rota.xlsx'
    tr.write_workbook(res, str(out), 'synthetic')
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ['Summary', 'Batch recipe', 'Tank calendar', 'Transfers']
    assert wb['Transfers'].max_row > 1
