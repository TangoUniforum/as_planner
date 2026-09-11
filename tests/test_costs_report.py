"""forecast/costs_report.py: the CostsAndProfit sheet (openpyxl in memory, no
engine run).

EVERY NUMBER AND ITEM NAME HERE IS MADE UP. The operator's real costs never
appear in a committed file: each test writes its own costs.yaml under
tmp_path and never reads config/costs.yaml.
"""
import datetime as dt
import re
from collections import defaultdict
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
import yaml

from forecast import costs as C
from forecast import costs_report as R
from forecast import excel_format
from forecast.analysis import harvest_rows_from_ws, revenue_for
from forecast.excel_io import (write_feed_forecast_monthly,
                               write_feed_forecast_weekly,
                               write_harvest_plan_output)
from forecast.models import BatchInput, BiologyTables
from forecast.placement import BatchLocationRow

ROOT = Path(__file__).resolve().parents[1]
REPORT_START = dt.date(2026, 9, 1)      # the day after a month-end PR closing
STARTER, ODD, GROWER = "Starter 0.5", "Mystery 3.0", "Grower 9.0"
WEEKS = [f"2026-W{w:02d}" for w in range(35, 54)]   # W35 wholly pre-start,
                                                    # W53 ends 2027-01-03
ECON = {"currency": "XYZ", "basis": "hog", "model_cv_pct": 18.0,
        "price_bands": [
            {"min_kg": 0.0, "max_kg": 3.0, "price_per_kg": 5.0},
            {"min_kg": 3.0, "max_kg": 100.0, "price_per_kg": 7.0}]}


def _costs():
    return {"schema": 1, "fixed_monthly": 3100.0, "oxygen_per_kg_feed": 0.5,
            "chemicals_per_kg_feed": 0.25, "feed_shipping_per_kg": 0.1,
            "egg_price": 0.01,
            "feed_prices": {
                STARTER: {"item": "Acme Crumble No.1", "price_per_kg": 4.0},
                GROWER: {"item": "", "price_per_kg": 2.0}}}


def _tables(odd=False):
    ft = [(500.0, STARTER), (1e9, GROWER)] + ([(3000.0, ODD)] if odd else [])
    return BiologyTables(sgr_size_g=[0.0, 1e9], sgr_fw_pct_day=[1.0, 1.0],
                         sgr_sw_pct_day=[1.0, 1.0], fcr_size_g=[0.0, 1e9],
                         fcr_by_model={"1.18": [1.2, 1.2]}, feed_types=ft)


def _monday(label):
    return dt.date.fromisocalendar(int(label[:4]), int(label[6:8]), 1)


def _batch(bid, input_date, count):
    return BatchInput(batch_id=bid, input_date=dt.datetime.combine(
        input_date, dt.time()), input_count=count, tran_sf_date=None,
        tran_og_date=None, tran_og_count=None, tran_og_avg_wt_g=None,
        tran_og_cv=10.0, fcr_model="FCR_118", fw_correction=1.0,
        sgr_correction=1.0)


def _inputs(odd=False):
    locs = []
    for i, w in enumerate(WEEKS):
        for bid, tank, stage in (("B1", 1, ""), ("B2", 2, "STARVE")):
            locs.append(BatchLocationRow(
                week_label=w, week_start=_monday(w), batch_id=bid,
                tank_id=tank, location_id=f"OG1-{tank}", system_id="OG1",
                count=4000.0, avg_wt_g=300.0 + 150.0 * i,
                biomass_kg=10_000.0, density_kg_m3=50.0, stage=stage))
    fw = [SimpleNamespace(stage="FW", feed_kg_week=70.0 + k,
                          feed_type=STARTER, week_label=w,
                          week_start=_monday(w))
          for k, w in enumerate(("2026-W40", "2026-W41", "2026-W44"))]
    fw.append(SimpleNamespace(stage="SW", feed_kg_week=999.0,
                              feed_type=GROWER, week_label="2026-W41",
                              week_start=_monday("2026-W41")))
    return dict(batch_locations=locs, states_by_batch={"B3": fw},
                tables=_tables(odd), batch_by_id={},
                batches=[_batch("B0", dt.date(2026, 8, 15), 50_000),
                         _batch("B3", dt.date(2026, 10, 7), 100_000),
                         _batch("B9", dt.date(2027, 6, 1), 70_000)],
                sixn_move_in_feed={("B1", "2026-W44", GROWER): 50.0},
                report_start=REPORT_START)


_HARVESTS = [("2026-W36", 2, 5000, 2800.0),   # Monday 08-31 -> September
             ("2026-W40", 3, 4000, 3100.0),
             ("2026-W41", 1, 4500, 3300.0),
             ("2026-W45", 2, 3000, 2500.0),
             ("2026-W49", 4, 5200, 3600.0)]   # Monday 11-30 -> November


def _workbook(inp):
    wb = openpyxl.Workbook()
    wb.active.title = "Control"
    evs = [SimpleNamespace(
        event_date=dt.datetime.combine(_monday(w) + dt.timedelta(days=d),
                               dt.time()),
        source_tank_id=1, batch_id="B1", count=n, avg_wt_g=g)
        for w, d, n, g in _HARVESTS]
    write_harvest_plan_output(wb, evs, default_hog_yield=0.9,
                              facility_limits_hog={})
    args = (wb, inp["batch_locations"], inp["states_by_batch"],
            REPORT_START, inp["tables"], inp["batch_by_id"])
    write_feed_forecast_weekly(
        *args, sixn_move_in_feed=inp["sixn_move_in_feed"])
    write_feed_forecast_monthly(
        *args, sixn_move_in_feed=inp["sixn_move_in_feed"],
        report_start=REPORT_START)
    return wb


def _cfg(tmp_path, costs=True, economics=True):
    d = tmp_path / "cfg"
    d.mkdir()
    if costs:
        C.save_costs(d, _costs())
    if economics:
        (d / "economics.yaml").write_text(yaml.safe_dump(ECON),
                                          encoding="utf-8")
    return d


def _write(wb, cfg, inp):
    return R.write_costs_sheet(wb, config_dir=cfg, **inp)


def _snapshot(wb):
    """Every cell that EXISTS (keys, values, style arrays) and the sheet
    settings, per sheet. Reading `_cells` never creates one; `style_id`
    would register styles, so the raw style array is compared instead."""
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = dict(
            cells=sorted((k, repr(c.value),
                          None if c._style is None else tuple(c._style))
                         for k, c in ws._cells.items()),
            freeze=ws.freeze_panes, filt=ws.auto_filter.ref,
            tab=repr(ws.sheet_properties.tabColor),
            widths={k: v.width for k, v in ws.column_dimensions.items()},
            rows={k: v.height for k, v in ws.row_dimensions.items()})
    return out


def _rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- #
# Neutrality: gated, appended last, nothing else touched
# --------------------------------------------------------------------------- #
def test_no_costs_file_leaves_the_workbook_untouched(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    names, before = list(wb.sheetnames), _snapshot(wb)
    got = _write(wb, _cfg(tmp_path, costs=False), inp)
    assert got == {"status": "absent"}
    assert wb.sheetnames == names
    assert _snapshot(wb) == before


def test_costs_append_one_sheet_last_and_touch_no_other(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    names, before = list(wb.sheetnames), _snapshot(wb)
    got = _write(wb, _cfg(tmp_path), inp)
    assert got["status"] == "ok"
    assert wb.sheetnames == names + [R.COSTS_SHEET]
    after = _snapshot(wb)
    assert {k: v for k, v in after.items() if k != R.COSTS_SHEET} == before


def test_plain_iter_rows_would_have_created_cells(tmp_path):
    """Negative control for the test above: openpyxl's own iter_rows on a
    writable HarvestPlan creates the empty cells it passes (row 3 is blank),
    which would add a <row> to the saved XML. The writer's reader does not."""
    wb = _workbook(_inputs())
    ws = wb["HarvestPlan"]
    n = len(ws._cells)
    via_reader = harvest_rows_from_ws(R._Values(ws))
    assert len(ws._cells) == n
    plain = harvest_rows_from_ws(ws)
    assert len(ws._cells) > n                   # it did create cells
    assert via_reader == plain and len(plain) == len(_HARVESTS)


def test_a_sheet_without_cells_map_is_not_computed_never_read_plainly(
        tmp_path, monkeypatch):
    """An openpyxl whose Worksheet lost `_cells` must not fall back to
    iter_rows (which creates cells, see the test above): the sheet goes NOT
    COMPUTED instead, and HarvestPlan is never read the plain way."""
    inp = _inputs()
    wb = _workbook(inp)
    ws = wb["HarvestPlan"]
    monkeypatch.delattr(ws, "_cells")

    def _plain(*a, **k):
        raise AssertionError("fell back to the plain iter_rows")
    monkeypatch.setattr(ws, "iter_rows", _plain)
    got = _write(wb, _cfg(tmp_path), inp)
    assert got["status"] == "not_computed" and "_cells" in got["error"]
    assert wb.sheetnames[-1] == R.COSTS_SHEET


def test_drop_stale_costs_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.create_sheet(R.COSTS_SHEET)
    wb.create_sheet("B")
    assert R.drop_stale_costs_sheet(wb) is True
    assert wb.sheetnames == ["A", "B"]
    assert R.drop_stale_costs_sheet(wb) is False
    assert wb.sheetnames == ["A", "B"]


def test_rewriting_keeps_exactly_one_sheet_last(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    wb.move_sheet(wb.create_sheet(R.COSTS_SHEET), offset=-2)  # a stale copy
    cfg = _cfg(tmp_path)
    _write(wb, cfg, inp)
    _write(wb, cfg, inp)
    assert wb.sheetnames.count(R.COSTS_SHEET) == 1
    assert wb.sheetnames[-1] == R.COSTS_SHEET


def test_styled_alone_never_through_the_workbook_pass(tmp_path, monkeypatch):
    def _boom(*a, **k):
        raise AssertionError("apply_workbook_formatting restyles every sheet")
    monkeypatch.setattr(excel_format, "apply_workbook_formatting", _boom)
    inp = _inputs()
    wb = _workbook(inp)
    got = _write(wb, _cfg(tmp_path), inp)
    ws = wb[R.COSTS_SHEET]
    assert ws.sheet_properties.tabColor.rgb.endswith(excel_format.TAB_REPORT)
    assert ws.freeze_panes == f"A{R.HEADER_ROW + 1}"
    assert ws.auto_filter.ref == f"A{R.HEADER_ROW}:P{got['body_last']}"
    assert ws.cell(R.HEADER_ROW, 1).font.bold


# --------------------------------------------------------------------------- #
# Not computed: loud, never a raise
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("text, words", [
    ("schema: 1\nfixed_monthly: 1\n", "missing"),
    ("", "costs not set"),
    ("fixed_monthly: [1,\n", "not valid YAML"),
])
def test_invalid_costs_write_not_computed_and_never_raise(tmp_path, capsys,
                                                          text, words):
    inp = _inputs()
    wb = _workbook(inp)
    names, before = list(wb.sheetnames), _snapshot(wb)
    cfg = _cfg(tmp_path, costs=False)
    (cfg / C.COSTS_FILE).write_text(text, encoding="utf-8")
    got = _write(wb, cfg, inp)
    assert got["status"] == "not_computed" and words in got["error"]
    assert wb.sheetnames == names + [R.COSTS_SHEET]
    rows = _rows(wb[R.COSTS_SHEET])
    assert len(rows) == 2 and rows[0][0] == R.TITLE
    assert rows[1][0].startswith("NOT COMPUTED — config/costs.yaml: ")
    assert "costs.yaml: costs.yaml" not in rows[1][0]
    assert {k: v for k, v in _snapshot(wb).items()
            if k != R.COSTS_SHEET} == before
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert back["status"] == "not_computed" and words in back["error"]
    assert "NOTE: CostsAndProfit NOT COMPUTED" in capsys.readouterr().out


def test_a_failure_while_pricing_is_not_computed_too(tmp_path):
    inp = _inputs()
    inp["batches"].append(_batch("BX", dt.date(2026, 9, 9), None))
    wb = _workbook(inp)
    got = _write(wb, _cfg(tmp_path), inp)
    assert got["status"] == "not_computed"
    assert "BX" in got["error"] and "input_count" in got["error"]
    # No HarvestPlan: revenue cannot be read, so nothing is priced as 0.
    del wb["HarvestPlan"]
    got = _write(wb, tmp_path / "cfg", _inputs())
    assert got["status"] == "not_computed" and "HarvestPlan" in got["error"]
    assert wb.sheetnames.count(R.COSTS_SHEET) == 1


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), -1, True])
def test_both_egg_counters_reject_the_same_bad_count(tmp_path, bad):
    """The Run sheet's counter and ideal_engine.eggs_by_year refuse the same
    input_count with the same words, at count time (never priced as 0)."""
    from forecast import ideal_engine as ie
    bx = [_batch("BX", dt.date(2026, 9, 9), bad)]
    with pytest.raises(ValueError) as by_month:
        R.eggs_by_month(bx, REPORT_START, dt.date(2027, 1, 3))
    with pytest.raises(ValueError) as by_year:
        ie.eggs_by_year(bx, REPORT_START, [2026])
    assert str(by_month.value) == str(by_year.value)
    assert "BX" in str(by_month.value) and "input_count" in str(by_month.value)
    inp = _inputs()
    inp["batches"] += bx
    wb = _workbook(inp)
    got = _write(wb, _cfg(tmp_path), inp)
    assert got["status"] == "not_computed" and "BX" in got["error"]


# --------------------------------------------------------------------------- #
# The figures
# --------------------------------------------------------------------------- #
def _ffm(wb):
    """FeedForecastMonthly's by-type matrix: ({(type, 'YYYY-MM'): kg}, months,
    {month: grand total})."""
    rows = _rows(wb["FeedForecastMonthly"])
    hdr = rows[2]
    months = [f"{d.year}-{d.month:02d}" for d in hdr[2:] if d is not None]
    cells, grand = {}, {}
    for r in rows[3:]:
        if r[0] == "Grand Total":
            grand = dict(zip(months, r[2:]))
            break
        for m, v in zip(months, r[2:]):
            cells[(r[0], m)] = v
    return cells, months, grand


def test_feed_ties_to_feed_forecast_monthly(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    cells, months, grand = _ffm(wb)
    feed, _ws = R.feed_by_type_month(
        inp["batch_locations"], inp["states_by_batch"], inp["tables"],
        inp["batch_by_id"], inp["sixn_move_in_feed"], REPORT_START)
    assert sorted(feed) == months
    for (name, m), v in cells.items():
        assert round(feed[m].get(name, 0.0), 0) == v, (name, m)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert [r["period"] for r in back["months"]] == months
    for r in back["months"]:
        assert abs(r["feed_kg"] - grand[r["period"]]) <= 0.5
    assert back["feed_total"]["feed_kg"] == pytest.approx(
        back["total"]["feed_kg"], rel=1e-12)
    by_type = {r["feed_type"]: r["feed_kg"] for r in back["feed_by_type"]}
    for name in (STARTER, GROWER):
        assert round(by_type[name], 0) == pytest.approx(
            sum(v for (n, _m), v in cells.items() if n == name), abs=1.0)


def test_one_month_roll_up_behind_both_sheets():
    """One source of truth: FeedForecastMonthly's by-type matrix and this
    sheet both take their months from excel_io._feed_by_type_month; neither
    carries its own copy of the calendar-day split loop."""
    import inspect
    from forecast import excel_io
    ffm = inspect.getsource(excel_io.write_feed_forecast_monthly)
    ours = inspect.getsource(R.feed_by_type_month)
    for src in (ffm, ours):
        assert "_feed_by_type_month(ftw, wk_start, report_start)" in src
    assert "ftm[(name, mo)] +=" not in ffm
    assert "calendar_day_month_split(" not in inspect.getsource(R)


def test_month_rows_wire_every_driver(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    mo = {r["period"]: r for r in back["months"]}
    # W35 is wholly before the report opens: its feed is kept (the sheet ties
    # to FeedForecastMonthly) and labelled, and no fixed cost falls there.
    assert mo["2026-08"]["kind"] == R.KIND_PRE_START
    assert mo["2026-08"]["fixed"] == 0 and mo["2026-08"]["feed_kg"] > 0
    assert mo["2026-09"]["kind"] == "Month"
    assert mo["2026-09"]["fixed"] == pytest.approx(3100.0)         # 30 / 30
    assert mo["2027-01"]["fixed"] == pytest.approx(3100.0 * 3 / 31)  # to 01-03
    # Eggs: only B3 (stocked inside the horizon), in its input month.
    assert {m: r["eggs_n"] for m, r in mo.items() if r["eggs_n"]} == {
        "2026-10": 100_000}
    assert mo["2026-10"]["eggs"] == pytest.approx(1000.0)
    # Revenue and HOG: HarvestPlan rows booked by Monday, clipped.
    hrows = harvest_rows_from_ws(wb["HarvestPlan"])
    want = defaultdict(list)
    for r, m in zip(hrows, ("2026-09", "2026-09", "2026-10", "2026-11",
                            "2026-11")):
        want[m].append(r)
    econ = {**ECON, "price_bands": [dict(b, monthly={})
                                    for b in ECON["price_bands"]]}
    for m, r in mo.items():
        rows = want.get(m, [])
        assert r["revenue"] == pytest.approx(
            revenue_for(rows, econ)["total"] if rows else 0.0)
        assert r["hog_kg"] == pytest.approx(sum(x["hog_kg"] for x in rows))
        assert r["profit"] == pytest.approx(r["revenue"] - r["total"])
        assert r["total"] == pytest.approx(
            r["feed"] + r["shipping"] + r["oxygen"] + r["chemicals"]
            + r["eggs"] + r["fixed"])
        assert r["shipping"] == pytest.approx(0.1 * r["feed_kg"])
        if r["hog_kg"] > 0:
            assert r["cost_per_kg_hog"] == pytest.approx(
                r["total"] / r["hog_kg"])
        else:
            assert r["cost_per_kg_hog"] is None       # "—", never 0
    # Years and the total are sums of the months.
    assert [y["period"] for y in back["years"]] == ["2026", "2027"]
    for f in ("feed", "total", "revenue", "profit", "fixed"):
        assert back["total"][f] == pytest.approx(
            sum(r[f] for r in back["months"]))
        assert back["total"][f] == pytest.approx(
            sum(y[f] for y in back["years"]))


def test_revenue_sums_to_one_call_over_all_harvest(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    from forecast.analysis import load_economics
    one = revenue_for(harvest_rows_from_ws(wb["HarvestPlan"]),
                      load_economics(tmp_path / "cfg"))["total"]
    assert back["total"]["revenue"] == pytest.approx(one, rel=1e-12)


def test_unpriced_feed_is_reported_never_priced_at_zero(tmp_path):
    inp = _inputs(odd=True)
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    odd = next(r for r in back["feed_by_type"] if r["feed_type"] == ODD)
    assert odd["price_per_kg"] == R.UNPRICED_TEXT and odd["feed"] is None
    assert odd["feed_kg"] > 0
    assert back["total"]["unpriced_kg"] == pytest.approx(odd["feed_kg"])
    # the priced types alone make the feed cost; shipping etc. on ALL kg
    priced = sum(r["feed"] for r in back["feed_by_type"]
                 if r["feed"] is not None)
    assert back["total"]["feed"] == pytest.approx(priced)
    assert back["total"]["shipping"] == pytest.approx(
        0.1 * back["total"]["feed_kg"])
    assert [r["feed_type"] for r in back["feed_by_type"]] == [STARTER, ODD,
                                                             GROWER]
    items = {r["feed_type"]: r["item"] for r in back["feed_by_type"]}
    assert items == {STARTER: "Acme Crumble No.1", ODD: "", GROWER: ""}


def test_the_note_says_when_feed_was_left_unpriced(tmp_path):
    """Profit is overstated when feed of a type with no price is left out of
    the cost: row 2 says so, naming the kg and the type (the Run page reads
    the same figures and withholds Profit)."""
    inp = _inputs(odd=True)
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    kg = back["total"]["unpriced_kg"]
    assert kg > 0
    assert (f"{kg:,.0f} kg of feed ({ODD}) has no price in "
            f"config/costs.yaml") in back["note"]
    assert "understated and profit overstated" in back["note"]


def test_a_fully_priced_sheet_says_nothing_about_unpriced_feed(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert back["total"]["unpriced_kg"] == 0
    assert "has no price" not in back["note"]


def test_the_note_carries_both_pricing_signatures(tmp_path):
    """costs.yaml's and economics.yaml's first 8 hex, each read back under
    its own name (the costs sha is no longer the first "(sha" in the note:
    the price-band sha sits beside it), "none" when economics.yaml is
    missing — so the Run page can say either changed since the run."""
    inp = _inputs()
    wb = _workbook(inp)
    cfg = _cfg(tmp_path)
    _write(wb, cfg, inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert back["sha"] == C.costs_sig(cfg)[:8]
    assert back["economics_sha"] == C.economics_sig(cfg)[:8] != "none"
    assert back["sha"] != back["economics_sha"]
    other = tmp_path / "other"
    other.mkdir()
    wb2 = _workbook(inp)
    _write(wb2, _cfg(other, economics=False), inp)
    back2 = R.read_costs_sheet(wb2[R.COSTS_SHEET])
    assert back2["economics_sha"] == "none"
    assert back2["sha"] == C.costs_sig(other / "cfg")[:8]
    # A note that carries no price-band sha (none was given) reads None.
    rows, _lay = R.build_costs_rows(
        R.cost_drivers(harvest_rows=harvest_rows_from_ws(wb["HarvestPlan"]),
                       **inp), C.load_costs(cfg), None,
        feed_types=[STARTER, GROWER], sha="0123abcd", revenue_note="")
    assert "economics.yaml (sha" not in rows[1][0]


def test_spend_per_kg_is_labelled_a_cash_figure_not_a_unit_cost():
    heads = R._headers(R.COLUMNS, "XYZ")
    assert "Spend per kg HOG sold (XYZ)" in heads
    assert not any("Cost per kg" in h for h in heads)


def test_without_economics_revenue_is_unpriced_not_zero(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path, economics=False), inp)
    back = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert back["status"] == "ok" and back["currency"] == R.NO_CURRENCY
    assert "Revenue: NOT PRICED" in back["note"]
    assert all(r["revenue"] is None and r["profit"] is None
               for r in back["months"] + back["years"] + [back["total"]])
    assert back["total"]["total"] > 0


def _same(a, b):
    """Equal, numbers to 1e-12 relative (and an int equal to its float)."""
    if isinstance(a, dict) and isinstance(b, dict):
        assert a.keys() == b.keys()
        for k in a:
            _same(a[k], b[k])
    elif isinstance(a, list) and isinstance(b, list):
        assert len(a) == len(b)
        for x, y in zip(a, b):
            _same(x, y)
    elif isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        assert a == pytest.approx(b, rel=1e-12, abs=1e-9)
    else:
        assert a == b


def test_read_round_trips_in_memory_and_from_a_saved_file(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    cfg = _cfg(tmp_path)
    _write(wb, cfg, inp)
    mem = R.read_costs_sheet(wb[R.COSTS_SHEET])
    assert mem["status"] == "ok" and mem["currency"] == "XYZ"
    assert mem["sha"] == C.costs_sig(cfg)[:8]
    # A saved file keeps ~15 significant digits and reads a whole float back
    # as an int, so the file round trip is compared to that precision.
    p = tmp_path / "out.xlsx"
    wb.save(p)
    for ro in (True, False):
        w2 = openpyxl.load_workbook(p, read_only=ro, data_only=True)
        try:
            _same(R.read_costs_sheet(w2[R.COSTS_SHEET]), mem)
        finally:
            w2.close()
    drivers = R.cost_drivers(harvest_rows=harvest_rows_from_ws(
        wb["HarvestPlan"]), **inp)
    pl = C.monthly_pl(drivers["feed"], drivers["eggs"], None,
                      {}, drivers["month_days"], C.load_costs(cfg))
    assert mem["total"]["total"] == pytest.approx(pl["total"]["total"],
                                                  rel=1e-12)
    assert len(mem["months"]) == len(pl["months"])


def test_a_foreign_sheet_reads_as_an_error():
    wb = openpyxl.Workbook()
    wb.active.append(["something else"])
    assert R.read_costs_sheet(wb.active)["status"] == "error"


def test_the_sheet_never_looks_like_a_production_report(tmp_path):
    inp = _inputs()
    wb = _workbook(inp)
    _write(wb, _cfg(tmp_path), inp)
    name = R.COSTS_SHEET.lower()
    assert len(R.COSTS_SHEET) <= 31 and " " not in R.COSTS_SHEET
    assert not any(w in name for w in ("month", "batch", "summary"))
    text = [str(v) for r in _rows(wb[R.COSTS_SHEET]) for v in r
            if isinstance(v, str)]
    assert not any("Closing Month" in t or "Unit:" in t for t in text)


# --------------------------------------------------------------------------- #
# run.py: exactly the two hooks
# --------------------------------------------------------------------------- #
def test_run_py_has_exactly_the_two_hooks():
    lines = (ROOT / "forecast" / "run.py").read_text(
        encoding="utf-8").splitlines()
    load = next(i for i, s in enumerate(lines)
                if s.strip() == "wb = load_workbook(in_path)")
    nxt = [s.strip() for s in lines[load + 1:load + 8]
           if s.strip() and not s.strip().startswith("#")]
    assert nxt[:2] == ["from .costs_report import drop_stale_costs_sheet",
                       "drop_stale_costs_sheet(wb)"]
    text = "\n".join(lines)
    assert text.count("write_costs_sheet(") == 1
    assert text.count("drop_stale_costs_sheet(") == 1
    ann = text.index("annotate_batch_plan_handling(wb)\n")
    gate = text.index('if (Path(config_dir) / "costs.yaml").is_file():')
    call = text.index("write_costs_sheet(wb")
    save = text.index("wb.save(out_path)")
    assert ann < gate < call < save
    # nothing but the call itself between it and the save
    tail = "report_start=report_start)"
    between = text[call:save]
    assert between[between.index(tail) + len(tail):].strip() == ""
