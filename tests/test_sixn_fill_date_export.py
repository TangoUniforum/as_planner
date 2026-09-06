"""BatchLocations exports each 6N tank's purge clock.

Without it there is no way to tell, from an output workbook, how long a cohort
has been in depuration: occupancy alone cannot separate "drained fully and
refilled the same week" from "topped up on top of fish that never left". Both
show a tank that never reads zero, and reading residency off occupancy spans
therefore reports phantom over-stays -- it did, and produced a table of five
that measurement by fill_date later showed to be none.

The exported value is the tank's CURRENT clock (state.sixn_fill_date), NOT the
arrival date of the fish in the row. A same-batch top-up overwrites it (see
_freeze_6n_dest), so a fill_date that moves forward while the tank stays
occupied IS the clock reset -- that is the signal, and the last test pins it.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook

from forecast.excel_io import write_batch_locations
from forecast.placement import BatchLocationRow

FILLED = date(2026, 11, 11)


def _row(tank_id, system_id, fill_date):
    return BatchLocationRow(
        week_label="2026-W46", week_start=date(2026, 11, 9), batch_id="B45",
        tank_id=tank_id, location_id=f"{system_id}-{tank_id}", system_id=system_id,
        count=46042.0, avg_wt_g=3100.0, biomass_kg=142730.0,
        density_kg_m3=88.0, stage="STARVE", purge_fill_date=fill_date,
    )


def _sheet(rows):
    wb = Workbook()
    write_batch_locations(wb, rows)
    ws = wb["BatchLocations"]
    header = [c.value for c in ws[4]]
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
    return header, body


def test_column_is_present_and_last():
    header, _ = _sheet([_row(61, "OG6N", FILLED)])
    assert header[-1] == "Purge_Fill_Date"
    # Appended at the end so the by-name column mapping in accuracy.py keeps
    # working against workbooks written before this column existed.
    assert header[:10] == [
        "Week", "Week_Start", "Batch", "Tank", "System",
        "Count (fish)", "AvgWt (kg)", "Biomass (kg)", "Density (kg/m3)", "Stage",
    ]


def test_sixn_tank_carries_its_clock():
    _, body = _sheet([_row(61, "OG6N", FILLED)])
    assert body[0][-1] == FILLED


def test_non_sixn_tank_has_no_clock():
    # sixn_fill_date only ever holds 6N keys, so every other tank exports None
    # rather than an invented date.
    _, body = _sheet([_row(31, "OG3N", None)])
    assert body[0][-1] is None


def test_default_is_none_so_other_producers_need_no_change():
    # manual_window.py builds rows without the field; it must stay constructible.
    r = BatchLocationRow(
        week_label="2026-W46", week_start=date(2026, 11, 9), batch_id="B45",
        tank_id=61, location_id="OG6N-61", system_id="OG6N",
        count=1.0, avg_wt_g=1.0, biomass_kg=1.0, density_kg_m3=1.0,
    )
    assert r.purge_fill_date is None


def test_a_moving_clock_on_an_occupied_tank_is_visible():
    """The top-up signature: occupancy never drops, the clock jumps forward.

    This is what occupancy-only measurement cannot see, and the reason the
    column exists.
    """
    later = date(2026, 11, 18)
    _, body = _sheet([_row(61, "OG6N", FILLED), _row(61, "OG6N", later)])
    counts = [r[5] for r in body]
    assert counts[0] == counts[1] != 0        # tank never read empty
    assert body[0][-1] != body[1][-1]         # but the clock moved
