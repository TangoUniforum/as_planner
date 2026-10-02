"""The readable RunConfig tab and its `RunConfig YAML` companion.

Operator, 2026-10-02: "make the RunConfig output more visually pleasing".
RunConfig became a formatted view; the lossless YAML moved, in the SAME block
format, to `RunConfig YAML`. What must not move:

  * import: a new workbook round-trips the config files byte-for-byte, and a
    workbook written before the split (YAML in RunConfig itself) still imports;
  * run_config_kind: snapshot / stamp / unknown / none are still told apart;
  * analysis.convergence_review, which scrapes RunConfig CELLS for the harvest
    inputs, still reads the CONTROL values -- never a per-week limits row.
"""
from __future__ import annotations

import contextlib
import io
import shutil
from pathlib import Path

import openpyxl
import pytest

from forecast import config_snapshot as cs

ROOT = Path(__file__).resolve().parent.parent
CRLF, LF = bytes([13, 10]), bytes([10])
PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")


def _lf_copy(tmp_path, copy_config) -> tuple[Path, Path]:
    """The repo's real config + scenario, LF-normalised (the snapshot stores
    lines, so an LF source is the one that can round-trip byte-for-byte)."""
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    for p in list(cfg.glob("*.yaml")) + list(scn.glob("*.yaml")):
        p.write_bytes(p.read_bytes().replace(b"\r\n", b"\n"))
    return cfg, scn


def _legacy_wb(blocks: dict) -> openpyxl.Workbook:
    """A workbook exactly as the pre-split writer (fdc297c) laid it out: the
    YAML blocks in RunConfig itself, no companion tab."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cs.SNAPSHOT_SHEET
    ws.append([f"{cs.KIND_SNAPSHOT_MARK} — the exact config + scenario this "
               "forecast ran from."])
    ws.append(["Re-importable: app Configure mode > 'Import from workbook', "
               "or forecast.config_snapshot.import_config_snapshot()."])
    ws.append([cs._excluded_line()])
    ws.append([])
    for key, text in blocks.items():
        ws.append([f"# ===== {key} ====="])
        for line in text.split("\n"):
            ws.append([line])
        ws.append([])
    return wb


def _saved(wb, tmp_path, name="wb.xlsx"):
    p = tmp_path / name
    wb.save(p)
    return openpyxl.load_workbook(p)


def _write(tmp_path, copy_config, **kw):
    cfg, scn = _lf_copy(tmp_path, copy_config)
    wb = openpyxl.Workbook()
    cs.write_config_snapshot(wb, config_dir=cfg, scenario_dir=scn, **kw)
    return _saved(wb, tmp_path), cfg, scn


# --------------------------------------------------------------------------- #
# Import
# --------------------------------------------------------------------------- #
def test_new_workbook_round_trips_the_files_exactly(tmp_path, copy_config):
    wb, cfg, scn = _write(tmp_path, copy_config)
    assert cs.YAML_SHEET in wb.sheetnames
    out_c, out_s = tmp_path / "c2", tmp_path / "s2"
    restored = cs.import_config_snapshot(wb, out_c, out_s)
    assert set(restored) == {f"{k}/{n}" for n, k in cs._FILES}
    for name, kind in cs._FILES:
        src = (cfg if kind == "config" else scn) / name
        dst = (out_c if kind == "config" else out_s) / name
        # import_config_snapshot writes in text mode, so on Windows its lines
        # end CRLF -- as they always have. Byte-for-byte modulo that newline.
        got = dst.read_bytes().replace(CRLF, LF)
        assert got == src.read_bytes(), name


def test_new_and_legacy_layouts_import_identically(tmp_path, copy_config):
    wb, _cfg, _scn = _write(tmp_path, copy_config)
    blocks = cs.read_config_snapshot(wb)
    legacy = _saved(_legacy_wb(blocks), tmp_path, "legacy.xlsx")
    assert cs.read_config_snapshot(legacy) == blocks
    a, b = tmp_path / "a", tmp_path / "b"
    cs.import_config_snapshot(wb, a / "c", a / "s")
    cs.import_config_snapshot(legacy, b / "c", b / "s")
    for sub in ("c", "s"):
        for f in sorted((a / sub).iterdir()):
            assert f.read_bytes() == (b / sub / f.name).read_bytes(), f.name


def test_an_old_workbook_still_imports_and_is_a_snapshot(tmp_path):
    legacy = _saved(_legacy_wb({"config/control.yaml": "scenario_name: old",
                                "scenario/batches.yaml": "batches: []"}), tmp_path)
    assert cs.run_config_kind(legacy) == "snapshot"
    assert "re-importable" in cs.describe_run_config_sheet(legacy)
    got = cs.import_config_snapshot(legacy, tmp_path / "c", tmp_path / "s")
    assert got == ["config/control.yaml", "scenario/batches.yaml"]
    assert (tmp_path / "c" / "control.yaml").read_text() == "scenario_name: old\n"


def test_the_readable_view_is_never_parsed_as_yaml(tmp_path, copy_config):
    """The fallback reader treats `# ===== key =====` cells as blocks. The view
    must never hold one, or a workbook missing its YAML tab would 'import'
    garbage instead of saying so."""
    wb, _c, _s = _write(tmp_path, copy_config)
    assert not cs.is_yaml_dump(wb[cs.SNAPSHOT_SHEET])
    assert cs._read_blocks(wb, cs.SNAPSHOT_SHEET) == {}


def test_view_without_its_yaml_tab_refuses_loudly(tmp_path, copy_config):
    wb, _c, _s = _write(tmp_path, copy_config)
    del wb[cs.YAML_SHEET]
    assert cs.run_config_kind(wb) == "snapshot"
    msg = cs.describe_run_config_sheet(wb)
    assert "missing" in msg and cs.YAML_SHEET in msg
    with pytest.raises(ValueError):
        cs.import_config_snapshot(wb, tmp_path / "c", tmp_path / "s")


# --------------------------------------------------------------------------- #
# Kind classification still holds (stamp / unknown / none)
# --------------------------------------------------------------------------- #
def test_kinds_are_still_told_apart(tmp_path, copy_config):
    wb, _c, _s = _write(tmp_path, copy_config)
    assert cs.run_config_kind(wb) == "snapshot"
    stamp = openpyxl.Workbook()
    stamp.active.title = cs.SNAPSHOT_SHEET
    stamp.active["A1"] = cs.KIND_STAMP_MARK
    stamp.active["A2"], stamp.active["B2"] = "planning_method", "global-lp"
    assert cs.run_config_kind(stamp) == "stamp"
    odd = openpyxl.Workbook()
    odd.active.title = cs.SNAPSHOT_SHEET
    odd.active["A1"] = "something else entirely"
    assert cs.run_config_kind(odd) == "unknown"
    assert cs.run_config_kind(openpyxl.Workbook()) is None
    # A stamp is never rescued by a stray YAML companion's absence/presence.
    with pytest.raises(ValueError):
        cs.import_config_snapshot(stamp, tmp_path / "c", tmp_path / "s")


# --------------------------------------------------------------------------- #
# The view itself
# --------------------------------------------------------------------------- #
def test_view_shows_control_as_key_value_and_pivots_weekly_limits(tmp_path, copy_config):
    wb, cfg, _s = _write(tmp_path, copy_config)
    ws = wb[cs.SNAPSHOT_SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert str(rows[0][0]).startswith(cs.KIND_SNAPSHOT_MARK)
    import yaml
    ctl = yaml.safe_load((cfg / "control.yaml").read_text())
    kv = {r[0]: r[1] for r in rows if r and isinstance(r[0], str)}
    assert kv["min_harvest_per_week"] == ctl["min_harvest_per_week"]
    assert kv["min_harvest_weight_g"] == ctl["min_harvest_weight_g"]
    # No cell that names a harvest input is followed by a NUMBER except the
    # control row itself -- that is what convergence_review scrapes.
    for key in ("min_harvest_weight_g", "min_harvest_per_week",
                "max_harvest_per_week"):
        hits = [r for r in rows for i, c in enumerate(r)
                if c == key and i + 1 < len(r) and isinstance(r[i + 1], (int, float))]
        assert len(hits) == 1, key
    assert ws.freeze_panes == "A5"
    assert any(d.outlineLevel for d in ws.row_dimensions.values())


def test_view_points_batches_to_batchinputs(tmp_path, copy_config):
    wb, _c, _s = _write(tmp_path, copy_config)
    text = " ".join(str(c) for r in wb[cs.SNAPSHOT_SHEET].iter_rows(values_only=True)
                    for c in r if c is not None)
    assert "BatchInputs" in text
    assert "batch_id" not in text      # batches are not repeated here


# --------------------------------------------------------------------------- #
# End to end on a real run: sheet places, import, convergence_review
# --------------------------------------------------------------------------- #
@pytest.mark.skipif(not PR_0831.exists(),
                    reason="the 2026-08-31 ProductionReport is not on this machine")
def test_real_run_sheet_places_and_analysis_reads_the_same(tmp_path, copy_config):
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    inp, out = tmp_path / "pr.xlsm", tmp_path / "out.xlsm"
    shutil.copy(PR_0831, inp)
    from forecast.run import main
    with contextlib.redirect_stdout(io.StringIO()):
        assert main(str(inp), str(out), config_dir=str(cfg),
                    scenario_dir=str(scn), calib_log_path="") == 0
    wb = openpyxl.load_workbook(out, keep_vba=True)
    names = wb.sheetnames
    assert names[-1] == cs.YAML_SHEET
    assert names[-2] == "BatchInputs"
    assert names.index(cs.SNAPSHOT_SHEET) == names.index("FacilityMap") + 1
    blocks = cs.read_config_snapshot(wb)
    for name, kind in cs._FILES:
        src = (cfg if kind == "config" else scn) / name
        assert blocks[f"{kind}/{name}"] == src.read_text(
            encoding="utf-8").strip("\n").replace("\r\n", "\n"), name

    # The same workbook with the pre-split RunConfig must judge identically:
    # the readable view may not change what convergence_review reads.
    from forecast.analysis import convergence_review
    new = convergence_review(out)
    assert new is not None and new["inputs_from"] == "workbook"
    legacy = openpyxl.load_workbook(out, keep_vba=True)
    idx = legacy.sheetnames.index(cs.SNAPSHOT_SHEET)
    del legacy[cs.SNAPSHOT_SHEET]
    del legacy[cs.YAML_SHEET]
    ws = legacy.create_sheet(cs.SNAPSHOT_SHEET, idx)
    for r in _legacy_wb(blocks)[cs.SNAPSHOT_SHEET].iter_rows(values_only=True):
        ws.append(list(r))
    lp = tmp_path / "legacy.xlsm"
    legacy.save(lp)
    assert convergence_review(lp) == new
