"""Ledger INPUT = eggs stocked, and nothing else.

THE OPERATOR (2026-09-11): "input is only egg inputs; moving FW -> OG/SW is
not an input; the opening at the forecast start holds every fish the PR
holds, FW and SW, including the FW part of a batch that straddles both."

What the ledgers did instead: on the week a batch crossed from freshwater to
seawater they RESET the opening to 0 and booked the TranOG as Input. The chain
close[w] -> open[w+1] broke on every arrival week; Sum(Input) over the horizon
counted every smolt twice (eggs, then again at TranOG: 10.6 M against 6.84 M
eggs on the 9.10 PR); a month holding an arrival printed an identity off by the
whole arrival (B50 2026-10: open 253,547 - ... + input 253,196 - close) while
its Count_Check read 0; and the PR's freshwater part of a split batch was not
in the opening at all.

Now: the FW->SW week opens on the freshwater fish about to cross and shows the
move in Xfer_In/Xfer_Out (they cancel); the crossing cull and a manual
fw_to_og's cull and pre-transfer FW loss are booked; a split batch's FW part is
held in the opening; eggs are booked whole in the month of their input date;
a split whose FW part nothing models is left visible in Count_Check and named
in the ValidationLog (detect, don't coerce).
"""
from __future__ import annotations

import contextlib
import io
import shutil
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast import excel_io
from forecast.excel_io import (_build_batch_week_ledger, _ledger_value_cells,
                               held_fw_openings, iso_week_label,
                               unmodelled_fw_warnings)
from forecast.models import BatchWeekState
from forecast.placement import BatchLocationRow
from forecast.production_report import PRBatchPeriod, PRPeriod

ROOT = Path(__file__).resolve().parent.parent
PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")

W1, W2, W3 = "2026-W41", "2026-W42", "2026-W43"
D1, D2, D3 = date(2026, 10, 5), date(2026, 10, 12), date(2026, 10, 19)


# ---- synthetic fixtures -----------------------------------------------------

def _loc(week, ws, batch, tank, count, wt_g):
    return BatchLocationRow(
        week_label=week, week_start=ws, batch_id=batch, tank_id=tank,
        location_id=f"OG1-{tank}", system_id="OG1N", count=count,
        avg_wt_g=wt_g, biomass_kg=count * wt_g / 1000.0, density_kg_m3=50.0)


def _state(week, ws, batch, stage, *, open_c=0.0, close_c=0.0, wt=0.0,
           close_wt=None, mort=0.0, cull=0.0, wfi=30, mort_pct=0.0):
    cw = wt if close_wt is None else close_wt
    n = close_c or open_c
    return BatchWeekState(
        batch_id=batch, week_label=week, week_start=ws,
        days_since_input=wfi * 7, week_from_input=wfi, count=n,
        avg_weight_g=cw, biomass_kg=n * cw / 1000.0, feed_kg_day=0.0,
        feed_kg_week=0.0, sgr_pct_day=0.0, fcr=0.0, stage=stage, feed_type="",
        mortality_pct_weekly=mort_pct, cull_count_week=cull,
        cull_biomass_kg_week=cull * wt / 1000.0,
        open_count=open_c, open_avg_weight_g=wt, open_biomass_kg=open_c * wt / 1000.0,
        close_count=close_c, close_avg_weight_g=cw,
        close_biomass_kg=close_c * cw / 1000.0, mort_count_week=mort)


def _tog(ws, batch, count, wt=380.0):
    return SimpleNamespace(event_date=ws, batch_id=batch, destinations=[
        SimpleNamespace(count=count, avg_wt_g=wt, tank_id=5)])


def _auto(cull_at_crossing=0.0):
    """Automatic TranOG. W1 is the freshwater projection's last week (opens
    100,000, loses 1,000, closes 99,000 @ 380 g); W2 is the FW->SW week: the
    TranOG takes that close (less a crossing cull when TranOG_Date is itself a
    week start, which then sits in the SW-labelled state)."""
    arrive = 99_000.0 - cull_at_crossing
    fw = _state(W1, D1, "B1", "FW", open_c=100_000, close_c=99_000, wt=350.0,
                close_wt=380.0, mort=1_000)
    sw = _state(W2, D2, "B1", "SW", close_c=arrive - 100, wt=380.0, close_wt=420.0,
                cull=cull_at_crossing, mort_pct=0.1)
    return dict(batch_locations=[_loc(W2, D2, "B1", 5, arrive - 100, 420.0)],
                harvest_events=[], batch_week_states=[fw, sw],
                tranog_events=[_tog(D2, "B1", arrive)],
                realized_biology={(5, W2, "B1"): [0.0, 100.0]})


H, H_KG = 250_225.0, 250_225.0 * 0.35        # held PR freshwater part
SW0, SW_WT = 47_743.0, 4_000.0              # the seawater part at the PR close
P, C, F = 240_000.0, 8_000.0, 248_000.0      # placed, culled, fw_count (F = P + C)


def _split(week_of_transfer=None, manual=True):
    """A batch split across FW and SW at the PR close (the B49 shape), in a
    manual window: the SW part in tank 7, the FW part held. With a
    week_of_transfer, a scripted fw_to_og places P and culls C then."""
    locs, opens, rb = [], [], {}
    sw = SW0
    for wk, ws in ((W1, D1), (W2, D2)):
        opens.append(_loc(wk, ws, "B1", 7, sw, SW_WT))
        placed = P if wk == week_of_transfer else 0.0
        rb[(7, wk, "B1")] = [0.0, 50.0]
        close = sw + placed - 50.0
        locs.append(_loc(wk, ws, "B1", 7, close, SW_WT))
        sw = close
    kw = dict(batch_locations=locs, harvest_events=[], batch_week_states=[],
              realized_biology=rb, window_openings=opens,
              fw_openings={"B1": (H, H_KG)})
    if week_of_transfer is not None:
        d = D1 if week_of_transfer == W1 else D2
        kw.update(tranog_events=[_tog(d, "B1", P)],
                  window_culls={("B1", week_of_transfer): [C, C * 0.35]},
                  fw_transfer_basis={"B1": [F, C]} if manual else None)
    return kw


def _by_week(rows, batch="B1"):
    return {r["week"]: r for r in rows if r["batch"] == batch}


def _check(r):
    """The identity from the row's own columns (L1)."""
    return (r["open_count"] - r["mort_count"] - r["harv_count"] - r["cull_count"]
            + r["input_count"] + r["xfer_in"] - r["xfer_out"] - r["close_count"])


# ---- (a)-(g): the weekly rule -------------------------------------------------

def test_a_fw_to_sw_week_opens_on_the_fw_close_and_shows_a_move():
    r = _by_week(_build_batch_week_ledger(**_auto()))
    assert r[W2]["open_count"] == pytest.approx(r[W1]["close_count"])  # chain
    assert r[W2]["input_count"] == 0.0
    assert r[W2]["xfer_in"] == r[W2]["xfer_out"] == pytest.approx(99_000.0)
    assert r[W2]["count_check"] == pytest.approx(0.0)
    assert _check(r[W2]) == pytest.approx(r[W2]["count_check"])
    # valued at its own weight: the FW close's biomass, no arrival crutch
    assert r[W2]["open_bio"] == pytest.approx(r[W1]["close_bio"])


def test_b_the_crossing_cull_is_booked_when_tranog_date_is_a_week_start():
    r = _by_week(_build_batch_week_ledger(**_auto(cull_at_crossing=1_000.0)))
    assert r[W2]["cull_count"] == pytest.approx(1_000.0)
    assert r[W2]["count_check"] == pytest.approx(0.0)


def test_c_egg_stocking_is_the_only_input():
    egg = _state(W1, D1, "B1", "EGG", close_c=500_000, wfi=0)
    r = _by_week(_build_batch_week_ledger([], [], [egg]))
    assert r[W1]["open_count"] == 0.0
    assert r[W1]["input_count"] == pytest.approx(500_000)
    assert r[W1]["count_check"] == pytest.approx(0.0)


@pytest.mark.parametrize("wk", [W1, W2])
def test_d_a_split_batch_opens_on_every_fish_and_a_manual_transfer_balances(wk):
    rows = _by_week(_build_batch_week_ledger(**_split(wk)))
    assert rows[W1]["open_count"] == pytest.approx(SW0 + H)    # 297,968 on B49
    t = rows[wk]
    assert t["input_count"] == 0.0
    assert t["xfer_in"] == t["xfer_out"] == pytest.approx(P)
    assert t["cull_count"] == pytest.approx(C)                 # booked, once
    assert t["mort_count"] == pytest.approx(50.0 + (H - F))    # FW loss to transfer
    for w in (W1, W2):
        assert rows[w]["count_check"] == pytest.approx(0.0), w
    assert rows[W2]["open_count"] == pytest.approx(rows[W1]["close_count"])


def test_e_a_split_nothing_models_stays_visible_in_count_check():
    rows = _by_week(_build_batch_week_ledger(**_split(None)))
    assert rows[W1]["open_count"] == pytest.approx(SW0 + H)
    assert rows[W1]["count_check"] == pytest.approx(H)         # detected
    assert rows[W2]["open_count"] == pytest.approx(rows[W1]["close_count"])
    assert rows[W2]["count_check"] == pytest.approx(0.0)


def test_f_first_ledger_week_is_the_crossing_no_double_count():
    """The state branch used to open on the FW balance AND credit the TranOG
    as input: the same 99,000 fish twice."""
    sw = _state(W1, D1, "B1", "SW", open_c=100_000, close_c=98_900, wt=380.0,
                close_wt=420.0, cull=1_000, mort_pct=0.1)
    kw = dict(batch_locations=[_loc(W1, D1, "B1", 5, 98_900, 420.0)],
              harvest_events=[], batch_week_states=[sw],
              tranog_events=[_tog(D1, "B1", 99_000)],
              realized_biology={(5, W1, "B1"): [0.0, 100.0]})
    r = _by_week(_build_batch_week_ledger(**kw))[W1]
    assert r["open_count"] == pytest.approx(100_000)
    assert r["input_count"] == 0.0
    assert r["xfer_in"] == pytest.approx(99_000)
    assert r["count_check"] == pytest.approx(0.0)


def test_g_a_fw_to_sw_week_keeps_its_sgr():
    """Same fish, open > 0, no input: the growth rate is meaningful again."""
    r = _by_week(_build_batch_week_ledger(**_auto()))[W2]
    assert _ledger_value_cells(r)[7] is not None


def test_a_projected_fw_part_is_never_counted_twice():
    """A later planning fix that projects a split batch's FW part: the ledger
    takes the FW part from the projection and ignores a held copy."""
    base = _build_batch_week_ledger(**_auto())
    held = _build_batch_week_ledger(**_auto(), fw_openings={"B1": (5_000.0, 100.0)})
    assert held == base


def test_a_batch_the_projector_ran_for_is_never_held_even_without_an_fw_row():
    """A wholly-FW batch whose tran_og_date is already past is projected from
    the PR's FW count starting in SEAWATER: no FW/EGG row, so a stage test
    cannot see it. The projector's own list (fw_projected) must keep a held
    copy out -- it opened the same fish twice (8/31 PR, B50: 508,270)."""
    kw = _split(None)
    held = _by_week(_build_batch_week_ledger(**kw))
    assert held[W1]["open_count"] == pytest.approx(SW0 + H)
    kw["fw_projected"] = {"B1"}
    r = _by_week(_build_batch_week_ledger(**kw))
    assert r[W1]["open_count"] == pytest.approx(SW0)
    assert r[W1]["count_check"] == pytest.approx(0.0)
    # and the run.py side: the held map is built from that same list
    agg = {"B50": {"count": 254_135.0, "biomass_kg": 51_730.0}}
    assert held_fw_openings(agg, {"B50"}) == {}


# ---- D6 only from a balance that describes ONE applied fw_to_og --------------------

@pytest.mark.parametrize("shape", ["second_refused", "applied_twice", "loss_negative"])
def test_d6_books_no_fw_loss_from_a_balance_it_cannot_trust(shape):
    """manual_fw_balance SUMS every fw_to_og for a batch, a refused one included.
    8/31 PR plus a second (refused) B49 fw_to_og printed Mort_Count -249,742 and
    a TOTAL of -225,968. The loss is booked only when the balance describes one
    applied event; otherwise nothing is booked and the gap stays visible."""
    kw = _split(W1)
    if shape == "second_refused":            # fw_count of both events summed
        kw["fw_transfer_basis"] = {"B1": [2 * F, C]}
    elif shape == "applied_twice":           # two TranOGs, balance consistent
        kw["tranog_events"] = kw["tranog_events"] + [_tog(D1, "B1", P)]
        kw["fw_transfer_basis"] = {"B1": [2 * F, 2 * C]}
    else:                                    # consistent, but more than was held
        kw["fw_transfer_basis"] = {"B1": [H + 1_000, H + 1_000 - P]}
    rows = _build_batch_week_ledger(**kw)
    r = _by_week(rows)[W1]
    assert r["mort_count"] == pytest.approx(50.0)        # seawater only, no D6
    assert all(x["mort_count"] >= 0 for x in rows)
    if shape == "second_refused":
        assert r["count_check"] == pytest.approx(H - F)  # the FW loss, visible


def test_d6_books_the_fw_loss_from_one_applied_event():
    """Positive control for the guard above: the 8/31 shape books H - F."""
    r = _by_week(_build_batch_week_ledger(**_split(W1)))[W1]
    assert r["mort_count"] == pytest.approx(50.0 + (H - F))
    assert r["count_check"] == pytest.approx(0.0)


@pytest.mark.parametrize("shape", ["two_applied_in_range", "not_placed_plus_culled"])
def test_d6_each_guard_holds_on_its_own(shape):
    """Each of D6's guards must hold where the OTHERS would let the loss
    through (independent mutation proof, 2026-09-11: the shapes above are all
    stopped by the sign/range check, so removing the one-event guard or the
    fw_count == placed + culled leg alone left every test green).

    two_applied_in_range: two applied fw_to_og (100,000 placed + 4,000 culled
    each) whose summed balance is consistent and whose "loss" (H - 208,000)
    is positive and below H -- only the one-event guard stops it.
    not_placed_plus_culled: one TranOG, but fw_count (249,000) is not placed +
    culled (248,000), e.g. a refused second event's count summed in -- the
    loss (1,225) is in range, so only the conservation leg stops it."""
    kw = _split(W1)
    if shape == "two_applied_in_range":
        kw["tranog_events"] = [_tog(D1, "B1", 100_000.0), _tog(D1, "B1", 100_000.0)]
        kw["fw_transfer_basis"] = {"B1": [208_000.0, 8_000.0]}
    else:
        kw["fw_transfer_basis"] = {"B1": [F + 1_000.0, C]}
    rows = _build_batch_week_ledger(**kw)
    r = _by_week(rows)[W1]
    assert r["mort_count"] == pytest.approx(50.0)        # seawater only, no D6
    if shape == "not_placed_plus_culled":
        assert r["count_check"] == pytest.approx(H - F)  # the FW loss, visible


# ---- rates over a held part, and its biomass when nothing moves it -----------------

def _sheet_rows(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Scenario")
    hdr = [str(c) if c is not None else "" for c in rows[hi]]
    return [dict(zip(hdr, r)) for r in rows[hi + 1:] if r and any(c is not None for c in r)]


RATES = ("SGR (%/day)", "SFR (%/day)", "Bio_FCR (ratio)", "Econ_FCR (ratio)")


def test_rates_are_blank_on_rows_carrying_a_held_part_but_not_on_totals():
    """The held part has no FW biology or feed; on its transfer week its whole FW
    growth lands at once (8/31 PR: B49 2026-W36 Bio_FCR 0.52, 2026-09 0.86), so
    its OWN rows are blank. The period TOTAL is the facility's and prints by the
    ordinary rules (operator, 2026-09-11: "Show them anyway") -- here SFR 0.0,
    since this synthetic batch has no feed; the FCRs then blank by the SFR
    rule, not the held one."""
    rows = _by_week(_build_batch_week_ledger(**_split(W2)))
    for w in (W1, W2):
        cells = _ledger_value_cells(rows[w])
        assert rows[w]["rates_unknown"], w
        assert [cells[i] for i in (7, 11, 12, 13)] == [None] * 4, w
    for writer, sheet, per in ((excel_io.write_weekly_report, "WeeklyReport", "Week"),
                               (excel_io.write_monthly_report, "MonthlyReport", "Month")):
        wb = openpyxl.Workbook()
        writer(wb, **_split(W2))
        got = _sheet_rows(wb, sheet)
        tots = [d for d in got if d["Batch"] == "TOTAL"]
        assert tots and len(tots) < len(got), sheet
        for d in got:
            if d["Batch"] == "TOTAL":
                assert d["SFR (%/day)"] == 0.0, (sheet, d[per])
            else:
                assert all(d[k] is None for k in RATES), (sheet, d[per], d["Batch"])


def test_rates_print_again_where_no_held_part_is_left():
    """Negative control: an ordinary FW->SW week (no held part) keeps its rates,
    and so does its TOTAL."""
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, **_auto())
    tot = [d for d in _sheet_rows(wb, "WeeklyReport") if d["Batch"] == "TOTAL"
           and d["Week"] == W2]
    assert tot and tot[0]["SFR (%/day)"] is not None


def test_an_unmodelled_held_part_goes_to_bio_check_not_growth():
    """Nothing moves it: its fish show in Count_Check and its biomass in
    Bio_Check -- not as -H_KG of "growth" (B49 read -90,854 kg)."""
    r = _by_week(_build_batch_week_ledger(**_split(None)))[W1]
    assert r["count_check"] == pytest.approx(H)
    assert r["bio_check"] == pytest.approx(H_KG)
    assert abs(r["gross_growth"]) < 0.01 * H_KG
    (m,) = [d for d in _sheet_rows(_wb_monthly(**_split(None)), "MonthlyReport")
            if d["Batch"] == "B1"]
    ident_kg = (_n(m, "Open_Bio (kg)") + _n(m, "Gross_Growth (kg)") - _n(m, "Mort_Bio (kg)")
                - _n(m, "Harv_Gross (kg)") - _n(m, "Cull_Bio (kg)") - _n(m, "Close_Bio (kg)"))
    assert abs(ident_kg - _n(m, "Bio_Check (kg)")) <= 2
    assert _n(m, "Bio_Check (kg)") == pytest.approx(H_KG, abs=1)


def _wb_monthly(**kw):
    wb = openpyxl.Workbook()
    excel_io.write_monthly_report(wb, **kw)
    return wb


# ---- (h) monthly, D4 eggs, D5, the PR merge ---------------------------------------

def _monthly(**kw):
    wb = openpyxl.Workbook()
    excel_io.write_monthly_report(wb, **kw)
    rows = list(wb["MonthlyReport"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Scenario")
    hdr = [str(c) if c is not None else "" for c in rows[hi]]
    return [d for d in (dict(zip(hdr, r)) for r in rows[hi + 1:])
            if d.get("Batch") not in (None, "TOTAL")]


def _n(d, k):
    v = d.get(k)
    return float(v) if isinstance(v, (int, float)) else 0.0


def _printed_identity(d):
    return (_n(d, "Open_Count (fish)") + _n(d, "Input_Count (fish)")
            - _n(d, "Mort_Count (fish)") - _n(d, "Harv_Count (fish)")
            - _n(d, "Cull_Count (fish)") + _n(d, "Xfer_In (fish)")
            - _n(d, "Xfer_Out (fish)") - _n(d, "Close_Count (fish)"))


def test_h_a_month_holding_an_arrival_closes_from_its_printed_columns():
    (m,) = _monthly(**_auto())
    assert _n(m, "Input_Count (fish)") == 0
    assert abs(_printed_identity(m) - _n(m, "Count_Check (fish)")) <= 2


def _eggs_straddling(with_date=True):
    """Stocked Thu 2026-10-01, in the week starting Mon 2026-09-28."""
    egg = _state("2026-W40", date(2026, 9, 28), "B1", "EGG", close_c=500_000, wfi=0)
    nxt = _state(W1, D1, "B1", "EGG", open_c=500_000, close_c=499_000, mort=1_000, wfi=1)
    kw = dict(batch_locations=[], harvest_events=[], batch_week_states=[egg, nxt])
    if with_date:
        kw["batches"] = {"B1": SimpleNamespace(input_date=date(2026, 10, 1))}
    return kw


def test_eggs_are_booked_whole_in_the_month_of_the_input_date():
    rows = {d["Month"]: d for d in _monthly(**_eggs_straddling())}
    assert _n(rows["2026-10"], "Input_Count (fish)") == 500_000
    assert _n(rows.get("2026-09", {}), "Input_Count (fish)") == 0
    for d in rows.values():
        assert abs(_printed_identity(d) - _n(d, "Count_Check (fish)")) <= 2


def test_without_an_input_date_eggs_keep_the_calendar_day_split():
    """Negative control: the date is what moves them (3 of 7 days in Sep)."""
    rows = {d["Month"]: d for d in _monthly(**_eggs_straddling(with_date=False))}
    assert _n(rows["2026-09"], "Input_Count (fish)") == pytest.approx(500_000 * 3 / 7, abs=1)


def test_d5_a_batch_the_pr_harvested_out_carries_its_deviation():
    pr = PRPeriod(closing_date=date(2026, 10, 7), batches={"B_GONE": PRBatchPeriod(
        batch_id="B_GONE", open_count=9_000.0, open_bio_kg=36_000.0,
        open_avg_wt_g=4000.0, growth_kg=0.0, feed_kg=1_000.0, harv_count=8_800.0,
        harv_gross_kg=35_000.0, mort_count=50.0, mort_bio_kg=200.0,
        cull_count=0.0, cull_bio_kg=0.0)})
    (g,) = [d for d in _monthly(batch_locations=[], harvest_events=[], pr_period=pr)
            if d["Batch"] == "B_GONE"]
    assert _n(g, "Count_Check (fish)") == 150          # never a forced 0
    assert _printed_identity(g) == _n(g, "Count_Check (fish)")


def test_a_mid_month_pr_with_a_split_batch_books_no_input():
    """The month opens on the PR's facility-wide balance (FW + SW), so crediting
    the fw_to_og as input would count the same fish twice."""
    kw = _split(W2)
    kw["batch_week_states"] = [_state(W1, D1, "B1", "SW", open_c=SW0, close_c=SW0 - 50,
                                      wt=SW_WT, wfi=90)]
    kw.pop("window_openings")
    pr = PRPeriod(closing_date=date(2026, 10, 7), batches={"B1": PRBatchPeriod(
        batch_id="B1", open_count=300_000.0, open_bio_kg=280_000.0,
        open_avg_wt_g=933.0, growth_kg=5_000.0, feed_kg=6_000.0, harv_count=0.0,
        harv_gross_kg=0.0, mort_count=2_000.0, mort_bio_kg=1_800.0,
        cull_count=0.0, cull_bio_kg=0.0, close_count=SW0 + H)})
    (m,) = _monthly(**kw, pr_period=pr, report_start=date(2026, 10, 8))
    assert _n(m, "Input_Count (fish)") == 0
    assert _n(m, "Open_Count (fish)") == 300_000
    assert abs(_printed_identity(m) - _n(m, "Count_Check (fish)")) <= 2


# ---- (i) negative control: the old reset must fail every check ----------------------

def _old_reset(real):
    """Re-create the retired rule: FW->SW week opens at 0 and books the TranOG
    as input (the move no longer shown as a transfer)."""
    def wrapped(*a, **kw):
        rows = real(*a, **kw)
        tog = defaultdict(float)
        for ev in kw.get("tranog_events") or ():
            tog[(ev.batch_id, iso_week_label(ev.event_date))] += sum(
                d.count for d in ev.destinations)
        for r in rows:
            t = tog.get((r["batch"], r["week"]))
            if t:
                r["open_count"] = r["open_bio"] = r["open_wt"] = 0.0
                r["input_count"] += t
                r["xfer_in"] -= t
                r["xfer_out"] -= t
                r["count_check"] = _check(r)
        return rows
    return wrapped


def _chain_breaks(rows):
    by = defaultdict(list)
    for r in rows:
        by[r["batch"]].append(r)
    return sum(1 for rs in by.values() for a, b in zip(rs, rs[1:])
               if abs(b["open_count"] - a["close_count"]) > 1)


def test_i_negative_control_the_old_reset_fails_chain_input_and_identity(monkeypatch):
    kw = _auto()
    good = _build_batch_week_ledger(**kw)
    old = _old_reset(_build_batch_week_ledger)(**kw)
    eggs = 0.0                                             # none stocked here
    assert _chain_breaks(good) == 0 and _chain_breaks(old) > 0          # L3
    assert sum(r["input_count"] for r in good) == eggs                  # L7
    assert sum(r["input_count"] for r in old) > eggs
    assert all(r["input_count"] == 0 for r in good if r["xfer_in"] > 0)  # L2
    # and the month's printed identity -- the check MonthlyReport's summed
    # Count_Check hid on the live workbook (B50 2026-10: 253,196 off)
    (m,) = _monthly(**kw)
    assert abs(_printed_identity(m) - _n(m, "Count_Check (fish)")) <= 2
    monkeypatch.setattr(excel_io, "_build_batch_week_ledger",
                        _old_reset(excel_io._build_batch_week_ledger))
    (m_old,) = _monthly(**_auto())
    assert abs(_printed_identity(m_old) - _n(m_old, "Count_Check (fish)")) > 1_000


# ---- the split-batch warning (D3 interim) ---------------------------------------------

def test_held_parts_come_from_the_pr_only_when_no_projection_carries_them():
    agg = {"B49": {"count": H, "biomass_kg": H_KG}, "B50": {"count": 1.0, "biomass_kg": 0.1},
           "B51": {"count": 0.0, "biomass_kg": 0.0}}
    assert held_fw_openings(agg, {"B50"}) == {"B49": (H, H_KG)}
    assert held_fw_openings(agg, {"B49", "B50"}) == {}


def test_the_split_warning_fires_only_when_nothing_moves_the_fw_part():
    held = {"B49": (H, H_KG)}
    lines, fish = unmodelled_fw_warnings(held, [], {"B49"})
    assert fish == {"B49": H}
    assert len(lines) == 1 and lines[0].startswith("SPLIT BATCH AT PR CLOSE - B49")
    assert "250,225" in lines[0]
    # a scripted fw_to_og models it: silent
    assert unmodelled_fw_warnings(held, [_tog(D1, "B49", P)], {"B49"}) == ([], {})
    # nothing held: silent
    assert unmodelled_fw_warnings({}, [], {"B49"}) == ([], {})
    # a wholly-FW batch nothing projects is named too, as what it is
    lines, _ = unmodelled_fw_warnings(held, [], set())
    assert lines[0].startswith("FW BATCH AT PR CLOSE NOT MODELLED - B49")


def test_the_warning_is_loud_in_the_validation_log():
    lines, _ = unmodelled_fw_warnings({"B49": (H, H_KG)}, [], {"B49"})
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(wb, invariant_warnings=lines)
    body = [r for r in wb["ValidationLog"].iter_rows(values_only=True) if r and r[0] == 1]
    assert body and body[0][1] == "WARNING - Split batch at PR close (FW part not modelled)"
    assert "B49" in body[0][2] and "250,225" in body[0][2]


def test_reconciliation_report_names_the_arrival_tranog_in_not_input():
    """Only eggs are input: the seawater-only ReconciliationReport's arrival
    columns are TranOG_In and TranOG_In_kg. The count header was pinned only
    through a PR run (test_sixn_mortality_accounting reads TranOG_In) and the
    kg header not at all (independent mutation proof, 2026-09-11: renaming it
    back to Input_kg left every test green)."""
    wb = openpyxl.Workbook()
    excel_io.write_reconciliation_report(wb, [], [], [], [], None)
    hdr = next(r for r in wb["ReconciliationReport"].iter_rows(values_only=True)
               if r and r[0] == "Week")
    assert "TranOG_In" in hdr and "TranOG_In_kg" in hdr
    assert not [c for c in hdr if c and str(c).startswith("Input")], hdr


def test_a_wholly_fw_batch_nothing_models_is_logged_under_its_own_category():
    """The second warning shape (no seawater part) gets its own WARNING
    category, not the 'WARNING - Hydration' catch-all, which tells an operator
    the PR read badly (independent mutation proof, 2026-09-11: deleting that
    branch left every test green -- only the split-batch category was pinned)."""
    lines, _ = unmodelled_fw_warnings({"B49": (H, H_KG)}, [], set())
    wb = openpyxl.Workbook()
    excel_io.write_validation_log(wb, invariant_warnings=lines)
    body = [r for r in wb["ValidationLog"].iter_rows(values_only=True) if r and r[0] == 1]
    assert body and body[0][1] == "WARNING - FW batch at PR close (not modelled)"
    assert "B49" in body[0][2] and "250,225" in body[0][2]


def test_input_conservation_does_not_claim_placed_for_unmodelled_fish():
    ctrl = SimpleNamespace(forecast_start=date(2026, 10, 5), horizon_weeks=10)
    bt = SimpleNamespace(batch_id="B49", input_count=520_000, tran_og_date=None,
                         tran_og_count=0)
    locs = [_loc(W1, D1, "B49", 7, SW0, SW_WT)]

    def status(unm):
        wb = openpyxl.Workbook()
        excel_io.write_input_conservation_audit(wb, [bt], locs, [], ctrl,
                                                unmodelled_fw=unm)
        rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
        hdr = list(rows[hi])
        row = dict(zip(hdr, rows[hi + 1]))
        return row, [r[0] for r in rows[:hi] if r and r[0]]
    row, head = status(None)
    assert row["Status"] == "PLACED"
    row, head = status({"B49": H})
    assert row["Status"] != "PLACED" and "NOT MODELLED" in row["Status"]
    assert row["Fish_At_Risk (fish)"] == round(H)
    assert any("250,225" in str(h) for h in head)
    # the conservation gate scans for "DROP"; this is not a dropped batch
    assert not any("DROP" in str(h).upper() for h in head if "NOT MODELLED" in str(h)
                   or "nothing models" in str(h))


def test_a_dropped_batch_keeps_its_dropped_verdict():
    """ideal_engine and the gates count the exact '*** DROPPED ***' string; the
    not-modelled override must never replace it."""
    ctrl = SimpleNamespace(forecast_start=date(2026, 10, 5), horizon_weeks=10)
    bt = SimpleNamespace(batch_id="B49", input_count=520_000,
                         tran_og_date=date(2026, 10, 20), tran_og_count=290_000)
    wb = openpyxl.Workbook()
    excel_io.write_input_conservation_audit(wb, [bt], [], [], ctrl,
                                            unmodelled_fw={"B49": H})
    rows = list(wb["InputConservationAudit"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
    row = dict(zip(rows[hi], rows[hi + 1]))
    assert row["Status"] == "*** DROPPED ***"
    assert row["Fish_At_Risk (fish)"] == 520_000


# ---- a real run ------------------------------------------------------------------------

def _run_pr(tmp, copy_config, edit_scenario=None):
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(ROOT / "scenario", scn)
    if edit_scenario is not None:
        edit_scenario(scn)             # a TEMP copy; never the live scenario
    inp, out = tmp / "pr.xlsm", tmp / "out.xlsm"
    shutil.copy(PR_0831, inp)
    from forecast.run import main
    with contextlib.redirect_stdout(io.StringIO()):
        rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                  calib_log_path="")
    assert rc == 0
    return out, scn


@pytest.fixture(scope="module", params=["reference", "pr_2026_08_31"])
def run(request, tmp_path_factory, copy_config):
    tmp = tmp_path_factory.mktemp("eggs_" + request.param)
    if request.param == "reference":
        from tests.fixtures.freeze_golden import REF, run_reference
        with contextlib.redirect_stdout(io.StringIO()):
            path = run_reference(tmp)
        scn = REF / "scenario"
    else:
        if not PR_0831.exists():
            pytest.skip("the 2026-08-31 ProductionReport is not on this machine")
        path, scn = _run_pr(tmp, copy_config)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for name in ("WeeklyReport", "MonthlyReport"):
        rows = list(wb[name].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "Scenario")
        hdr = [str(c) if c is not None else "" for c in rows[hi]]
        sheets[name] = [dict(zip(hdr, r)) for r in rows[hi + 1:]
                        if r and r[3 if name == "WeeklyReport" else 2] is not None]
    tp = list(wb["TransferPlan"].iter_rows(values_only=True))
    th = next(i for i, r in enumerate(tp) if r and r[0] == "Week")
    col = {str(c): j for j, c in enumerate(tp[th]) if c is not None}
    tog = defaultdict(float)
    for r in tp[th + 1:]:
        if r and r[col["Type"]] == "TranOG":
            tog[(str(r[col["Batch"]]), str(r[col["Week"]]))] += float(r[col["Count (fish)"]] or 0)
    vlog = [str(r[1]) for r in wb["ValidationLog"].iter_rows(values_only=True)
            if r and isinstance(r[0], int)]
    wb.close()
    return SimpleNamespace(name=request.param, path=path, scenario=scn, tog=dict(tog),
                           vlog=vlog, **{k.lower(): v for k, v in sheets.items()})


def _monday(label):
    return date.fromisocalendar(int(label[:4]), int(label[6:8]), 1)


def _batch_rows(rows):
    return [d for d in rows if d.get("Batch") != "TOTAL"]


def test_real_run_weekly_chain_close_equals_next_open(run):
    by = defaultdict(list)
    for d in _batch_rows(run.weeklyreport):
        by[d["Batch"]].append(d)
    breaks = []
    for b, rs in by.items():
        rs.sort(key=lambda d: d["Week"])
        for a, n in zip(rs, rs[1:]):
            if _monday(n["Week"]) - _monday(a["Week"]) != timedelta(days=7):
                continue
            if abs(_n(n, "Open_Count (fish)") - _n(a, "Close_Count (fish)")) > 1:
                breaks.append((b, n["Week"], _n(a, "Close_Count (fish)"),
                               _n(n, "Open_Count (fish)")))
    assert not breaks, f"{len(breaks)} close->open breaks, e.g. {breaks[:5]}"


def test_real_run_total_input_equals_the_eggs_stocked(run):
    from forecast.scenario_io import load_batches
    weeks = sorted({d["Week"] for d in run.weeklyreport})
    first, last = _monday(weeks[0]), _monday(weeks[-1]) + timedelta(days=6)
    eggs = []
    for b in load_batches(run.scenario):
        d = b.input_date.date() if hasattr(b.input_date, "date") else b.input_date
        if d and first <= d <= last:
            eggs.append(float(b.input_count or 0))
    total = sum(_n(d, "Input_Count (fish)") for d in _batch_rows(run.weeklyreport))
    if run.name == "pr_2026_08_31":
        assert eggs, "the 2026-08-31 horizon stocks 12 batches -- none found"
    # No eggs in the horizon (the reference fixture) means NO input at all:
    # every smolt arrival is a move, never an input.
    assert abs(total - sum(eggs)) <= len(eggs) + 1, (total, sum(eggs))


def test_real_run_fw_to_sw_weeks_are_moves_not_input(run):
    rows = {(d["Batch"], d["Week"]): d for d in _batch_rows(run.weeklyreport)}
    assert run.tog, "no TranOG in this plan"
    for k, n in run.tog.items():
        d = rows[k]
        assert _n(d, "Input_Count (fish)") == 0, k
        assert _n(d, "Open_Count (fish)") > 0, k
        assert _n(d, "Xfer_In (fish)") + 1 >= n and _n(d, "Xfer_Out (fish)") + 1 >= n, k
    for d in _batch_rows(run.weeklyreport):          # L2: input only on stocking
        if _n(d, "Input_Count (fish)") > 0:
            assert _n(d, "Open_Count (fish)") == 0, (d["Batch"], d["Week"])


@pytest.mark.parametrize("sheet", ["weeklyreport", "monthlyreport"])
def test_real_run_identity_holds_from_the_printed_columns(run, sheet):
    bad = [(d["Batch"], d.get("Week") or d.get("Month"), round(_printed_identity(d)),
            _n(d, "Count_Check (fish)"))
           for d in _batch_rows(getattr(run, sheet))
           if abs(_printed_identity(d) - _n(d, "Count_Check (fish)")) > 5]
    assert not bad, f"{len(bad)} rows whose printed columns do not give Count_Check: {bad[:5]}"


def test_real_run_eggs_land_whole_in_their_input_month(run):
    from forecast.scenario_io import load_batches
    months = sorted({d["Month"] for d in run.monthlyreport})
    meta = {b.batch_id: b for b in load_batches(run.scenario)}
    got = defaultdict(dict)
    for d in _batch_rows(run.monthlyreport):
        if _n(d, "Input_Count (fish)") > 0:
            got[d["Batch"]][d["Month"]] = _n(d, "Input_Count (fish)")
    if not got:
        if run.name == "pr_2026_08_31":
            pytest.fail("the 2026-08-31 horizon stocks eggs -- no Input found")
        pytest.skip("no egg stocking in this horizon")
    for b, per in got.items():
        dd = meta[b].input_date
        dd = dd.date() if hasattr(dd, "date") else dd
        want = max(f"{dd.year}-{dd.month:02d}", months[0])
        assert list(per) == [want], (b, per, want)
        assert abs(per[want] - float(meta[b].input_count)) <= 1, (b, per)


def test_real_run_week_totals_chain(run):
    tot = {d["Week"]: d for d in run.weeklyreport if d.get("Batch") == "TOTAL"}
    n = defaultdict(int)
    for d in _batch_rows(run.weeklyreport):
        n[d["Week"]] += 1
    weeks = sorted(tot)
    bad = [(w1, _n(tot[w0], "Close_Count (fish)"), _n(tot[w1], "Open_Count (fish)"))
           for w0, w1 in zip(weeks, weeks[1:])
           if abs(_n(tot[w1], "Open_Count (fish)") - _n(tot[w0], "Close_Count (fish)"))
           > n[w0] + n[w1]]
    assert not bad, bad[:5]


def test_real_run_the_split_warning_is_silent_when_the_split_is_scripted(run):
    """On the 2026-08-31 PR, B49 is split 47,743 SW + 250,225 FW and the FW part
    is moved by a scripted fw_to_og: modelled, so no warning -- and the batch
    opens on all of it."""
    assert not any("Split batch at PR close" in c or "not modelled" in c
                   for c in run.vlog)
    if run.name != "pr_2026_08_31":
        return
    from forecast.production_report import read_production_report
    pwb = openpyxl.load_workbook(PR_0831, read_only=True, data_only=True)
    _closing, og, fw = read_production_report(pwb)
    pwb.close()
    b49 = (sum(r.closing_count for r in og if r.batch_id == "B49")
           + sum(r.closing_count for r in fw if r.batch_id == "B49"))
    first = min(d["Week"] for d in run.weeklyreport)
    (row,) = [d for d in _batch_rows(run.weeklyreport)
              if d["Batch"] == "B49" and d["Week"] == first]
    assert abs(_n(row, "Open_Count (fish)") - b49) <= 1
    # its held FW part crosses this week with no FW feed behind it: no rate
    assert all(row[k] is None for k in RATES)


@pytest.mark.parametrize("sheet,per", [("weeklyreport", "Week"),
                                       ("monthlyreport", "Month")])
def test_real_run_the_total_over_a_held_part_prints_the_facility_rates(run, sheet, per):
    """Operator, 2026-09-11 ("Show them anyway"): B49's own first week and
    month are blank (above), but the facility TOTAL of that same period
    prints SFR and both FCRs -- and a salmon's FCR, not the 0.52 B49 alone
    would read."""
    if run.name != "pr_2026_08_31":
        pytest.skip("the reference fixture has no split batch")
    rows = getattr(run, sheet)
    first = min(d[per] for d in rows)
    (b49,) = [d for d in _batch_rows(rows) if d["Batch"] == "B49" and d[per] == first]
    assert all(b49[k] is None for k in RATES), b49
    (tot,) = [d for d in rows if d["Batch"] == "TOTAL" and d[per] == first]
    for k in ("SFR (%/day)", "Bio_FCR (ratio)", "Econ_FCR (ratio)"):
        assert isinstance(tot[k], (int, float)), (k, tot[k])
    assert tot["Bio_FCR (ratio)"] >= 0.6, tot["Bio_FCR (ratio)"]


def _pr_fish():
    """(total fish, {batch: FW count}, {batch: FW kg}) at the 2026-08-31 close."""
    from forecast.production_report import read_production_report
    pwb = openpyxl.load_workbook(PR_0831, read_only=True, data_only=True)
    _closing, og, fw = read_production_report(pwb)
    pwb.close()
    fw_n, fw_kg = defaultdict(float), defaultdict(float)
    for r in fw:
        fw_n[r.batch_id] += r.closing_count
        fw_kg[r.batch_id] += r.closing_biomass_kg
    total = sum(r.closing_count for r in og) + sum(fw_n.values())
    return total, fw_n, fw_kg


def _first_total_open(rows):
    first = min(d["Week"] for d in rows)
    (tot,) = [d for d in rows if d["Batch"] == "TOTAL" and d["Week"] == first]
    n_rows = sum(1 for d in _batch_rows(rows) if d["Week"] == first)
    return _n(tot, "Open_Count (fish)"), n_rows


def test_real_run_the_opening_holds_every_fish_the_pr_holds(run):
    """L6: the first week's TOTAL Open is the PR's own fish count, FW + SW."""
    if run.name != "pr_2026_08_31":
        pytest.skip("the reference fixture has no ProductionReport to tie to")
    total, _n_fw, _kg = _pr_fish()
    got, n_rows = _first_total_open(run.weeklyreport)
    assert abs(got - total) <= n_rows, (got, total)


def _impossible_fcr(rows):
    """test_coordinator_regression's bar, on every printed row (TOTAL too):
    feed and growth both over a tonne and Bio_FCR under 0.6."""
    out = []
    for d in rows:
        x = d.get("Bio_FCR (ratio)")
        if (_n(d, "Feed (kg)") > 1000 and _n(d, "Gross_Growth (kg)") > 1000
                and isinstance(x, (int, float)) and 0 < x < 0.6):
            out.append((d.get("Batch"), d.get("Week") or d.get("Month"), x))
    return out


@pytest.mark.parametrize("sheet", ["weeklyreport", "monthlyreport"])
def test_real_run_no_ledger_row_reports_an_impossible_fcr(run, sheet):
    """The held FW part of B49 (8/31 PR) crossed with no FW feed and printed
    Bio_FCR 0.52 (W36) / 0.86 (2026-09). This guard only ran on Forecast.xlsm,
    which has no fw_to_og, so nothing saw it."""
    bad = _impossible_fcr(getattr(run, sheet))
    assert not bad, bad[:5]


# ---- a real split nothing models, and a projection that starts in seawater ------

def _unscripted_split_and_late_b50(scn):
    """The 8/31 PR with NO events file (B49's FW part is then modelled by
    nothing) and B50's tran_og_date moved before the close (its projection
    then starts in seawater from the PR's FW count)."""
    ev = scn / "manual_events" / "2026-08-31.yaml"
    if ev.exists():
        ev.unlink()
    p = scn / "batches.yaml"
    txt = p.read_text(encoding="utf-8")
    i = txt.find("- batch_id: B50")
    if i < 0:
        pytest.skip("B50 is not in the live batches.yaml")
    j = txt.index("tran_og_date:", i)
    k = txt.index("\n", j)
    p.write_text(txt[:j] + "tran_og_date: '2026-08-27'" + txt[k:], encoding="utf-8")


@pytest.fixture(scope="module")
def unscripted(tmp_path_factory, copy_config):
    if not PR_0831.exists():
        pytest.skip("the 2026-08-31 ProductionReport is not on this machine")
    tmp = tmp_path_factory.mktemp("eggs_unscripted")
    path, _scn = _run_pr(tmp, copy_config, _unscripted_split_and_late_b50)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = SimpleNamespace(
        weekly=_sheet_rows(wb, "WeeklyReport"), monthly=_sheet_rows(wb, "MonthlyReport"),
        vlog=[(str(r[1]), str(r[2])) for r in wb["ValidationLog"].iter_rows(values_only=True)
              if r and isinstance(r[0], int)])
    ica = list(wb["InputConservationAudit"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(ica) if r and r[0] == "Batch")
    out.ica = {r[0]: dict(zip(ica[hi], r)) for r in ica[hi + 1:] if r and r[0]}
    wb.close()
    return out


def test_real_run_an_unscripted_split_is_warned_by_name(unscripted):
    """End to end through run.main: the ValidationLog names B49 and its count,
    the audit does not say PLACED, and the first week shows the fish in
    Count_Check and their biomass in Bio_Check (detect, don't coerce)."""
    _total, fw_n, fw_kg = _pr_fish()
    h = fw_n["B49"]
    assert h > 0, "B49 is no longer split on the 8/31 PR"
    hits = [d for c, d in unscripted.vlog
            if c == "WARNING - Split batch at PR close (FW part not modelled)"]
    assert len(hits) == 1 and "B49" in hits[0] and f"{h:,.0f}" in hits[0], hits
    a = unscripted.ica["B49"]
    assert a["Status"] == "*** FW PART NOT MODELLED ***"
    assert a["Fish_At_Risk (fish)"] == round(h)
    first = min(d["Week"] for d in unscripted.weekly)
    (r,) = [d for d in unscripted.weekly if d["Batch"] == "B49" and d["Week"] == first]
    assert abs(_n(r, "Count_Check (fish)") - h) <= 30, r["Count_Check (fish)"]
    assert abs(_n(r, "Bio_Check (kg)") - fw_kg["B49"]) <= 1, r["Bio_Check (kg)"]
    assert _n(r, "Gross_Growth (kg)") > -0.05 * fw_kg["B49"]       # not in growth
    assert all(r[k] is None for k in RATES)


def test_real_run_a_projection_starting_in_seawater_is_not_held_again(unscripted):
    """B50's FW projection opens on the PR's FW count in SEAWATER (no FW row).
    It was held too: 2 x 254,135 in the opening, a +508,143 Count_Check, and a
    'not modelled' warning for a batch that IS in the Batches sheet."""
    total, fw_n, _kg = _pr_fish()
    first = min(d["Week"] for d in unscripted.weekly)
    (r,) = [d for d in unscripted.weekly if d["Batch"] == "B50" and d["Week"] == first]
    assert abs(_n(r, "Open_Count (fish)") - fw_n["B50"]) <= 1, r["Open_Count (fish)"]
    assert not any("B50" in d for c, d in unscripted.vlog
                   if "PR close" in c or "PR CLOSE" in d)
    got, n_rows = _first_total_open(unscripted.weekly)
    assert abs(got - total) <= n_rows, (got, total)


@pytest.mark.parametrize("sheet", ["weekly", "monthly"])
def test_real_run_unscripted_split_no_impossible_fcr_or_negative_mort(unscripted, sheet):
    rows = getattr(unscripted, sheet)
    assert not _impossible_fcr(rows)
    assert not [d["Batch"] for d in rows if _n(d, "Mort_Count (fish)") < 0]
