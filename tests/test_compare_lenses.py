"""The RunComparison PRODUCTION block: what the plan ships, and handling in jobs.

Added 2026-09-18 because the sheet scored placement, caps, steadiness and
balance across 30 rows and never stated the tonnage, and counted handling as
moves per FISH -- which ranked the arm with the MOST transfer rows (686 vs 534
on the 9.16 PR) as the best on handling.

Everything here is read from a workbook, so the fixture is a workbook: no
engine run, no config.
"""
from __future__ import annotations

import openpyxl
import pytest

from tools.run_compare import _growout_m3, _production


def _wb(tmp_path, *, years, transfers, harvest_weeks):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("YearlySummary")
    ws.append(["YEARLY SUMMARY"])
    ws.append([])
    ws.append(["Year", "Harvest_Count (fish)", "Harvest_HOG (t)",
               "Harvest_Gross (t)", "Avg_HOG_Wt (kg)", "Feed (t)",
               "Peak_Biomass (t)", "Mean_Biomass (t)", "Mean_Utilisation (%)"])
    for y, hog, gross, feed, mean_b in years:
        ws.append([y, 0, hog, gross, 0, feed, 0, mean_b, 0])
    ws = wb.create_sheet("TransferPlan")
    ws.append(["TRANSFER PLAN"])
    ws.append([])
    ws.append(["Week", "Batch", "Type", "From_Tank", "To_Tank", "Count (fish)",
               "Avg_Weight (kg)", "Grade", "CV (%)", "Channel"])
    for wk, kind, channel in transfers:
        ws.append([wk, "B01", kind, 1, 2, 100, 4.0, "", 16, channel])
    ws = wb.create_sheet("HarvestPlan")
    ws.append(["HARVEST PLAN"])
    ws.append([])
    ws.append(["Week", "Batch", "Tank", "Count (fish)", "Gross_AvgWt (kg)",
               "Gross_Biomass (kg)", "HOG_Yield (ratio)", "HOG_AvgWt (kg)",
               "HOG_Biomass (kg)"])
    for wk in harvest_weeks:
        ws.append([wk, "B01", 1, 1000, 4.0, 4000, 0.81, 3.24, 3240])
    p = tmp_path / "plan.xlsm"
    wb.save(p)
    return p


def test_it_reports_the_tonnage_the_sheet_never_had(tmp_path):
    p = _wb(tmp_path,
            years=[(2026, 1000.0, 1200.0, 1500.0, 500.0),
                   (2027, 2000.0, 2400.0, 3000.0, 700.0)],
            transfers=[], harvest_weeks=[f"2026-W{w:02d}" for w in range(1, 53)])
    got = _production(p, og_m3=1000.0)
    assert got["hog_t"] == 3000.0
    assert got["feed_t"] == 4500.0
    # 52 harvest weeks ~ one year, so per-year == the total
    assert got["hog_t_per_yr"] == pytest.approx(3000.0 * 52.1775 / 52, rel=1e-9)


def test_tank_year_productivity_is_per_m3_of_growout(tmp_path):
    """The rotation lens: tonnage per m3 of grow-out water per year."""
    p = _wb(tmp_path, years=[(2026, 1000.0, 1200.0, 100.0, 500.0)],
            transfers=[], harvest_weeks=[f"2026-W{w:02d}" for w in range(1, 53)])
    got = _production(p, og_m3=2000.0)
    assert got["kg_hog_per_m3_yr"] == pytest.approx(
        got["hog_t_per_yr"] * 1000.0 / 2000.0, rel=1e-9)
    # no grow-out volume known -> the lens is BLANK, never a wrong number
    assert _production(p, og_m3=0.0)["kg_hog_per_m3_yr"] is None


def test_production_to_capacity_is_annual_output_over_standing_stock(tmp_path):
    """The RAS literature's 2-3 benchmark. Mean standing biomass is the mean of
    the per-year means, so a two-year plan cannot double-count it."""
    p = _wb(tmp_path,
            years=[(2026, 1000.0, 1200.0, 0.0, 400.0),
                   (2027, 1000.0, 1200.0, 0.0, 600.0)],
            transfers=[], harvest_weeks=[f"2026-W{w:02d}" for w in range(1, 53)])
    got = _production(p, og_m3=1000.0)
    gross_per_yr = 2400.0 / (52 / 52.1775)
    assert got["prod_to_capacity"] == pytest.approx(gross_per_yr / 500.0, rel=1e-9)


def test_handling_is_counted_in_jobs_not_fish(tmp_path):
    """One TransferPlan row is one job the crew rigs for. The older
    transfers-per-fish metric divides by the fish moved, which is why it
    preferred the highest-handling arm."""
    p = _wb(tmp_path, years=[(2026, 1000.0, 1200.0, 0.0, 500.0)],
            transfers=([("2026-W01", "Transfer", "rotation_fill")] * 5
                       + [("2026-W02", "Grade", "grade_event")] * 3
                       + [("2026-W02", "TranOG", "TranOG")] * 2),
            harvest_weeks=[f"2026-W{w:02d}" for w in range(1, 53)])
    got = _production(p, og_m3=1000.0)
    assert got["move_events"] == 5        # TranOG is not a planner move
    assert got["grade_events"] == 3
    assert got["moves_per_1000t_hog"] == pytest.approx(5 / 1.0, rel=1e-9)


def test_an_empty_plan_gives_blanks_not_zeros_or_crashes(tmp_path):
    p = _wb(tmp_path, years=[], transfers=[], harvest_weeks=[])
    got = _production(p, og_m3=1000.0)
    assert got["hog_t"] == 0.0
    assert got["hog_t_per_yr"] is None and got["prod_to_capacity"] is None
    assert got["moves_per_1000t_hog"] is None


def test_the_growout_volume_comes_from_the_runs_own_facility(tmp_path):
    """Never hardcoded: the denominator is read from the config the run used,
    and only OG tanks count (FW is not grow-out water)."""
    import yaml
    cfg = tmp_path / "config"
    cfg.mkdir()
    (cfg / "facility.yaml").write_text(yaml.safe_dump({"tanks": [
        {"tank_id": 1, "type": "OG", "volume_m3": 100},
        {"tank_id": 2, "type": "OG", "volume_m3": 50},
        {"tank_id": 3, "type": "FW", "volume_m3": 999},
    ]}), encoding="utf-8")
    assert _growout_m3(cfg) == 150.0
    assert _growout_m3(tmp_path / "nope") == 0.0     # unreadable -> blank lens
