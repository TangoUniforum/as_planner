"""Density metrics must exclude OFF-FEED PURGE fish, not the OG6N system.

DEFECT (2026-07-13 audit): _density_quality/_density_peak/_density_overshoot
excluded every System == OG6N row. From 2028 (production mode) the 6N MAINS are
grow-out production tanks — their fish are reared and fed, and excluding the
whole system blinds the welfare/compliance metrics to real crowding there.
The exclusion now keys on Stage == STARVE (off-feed purge, wherever it sits),
falling back to the legacy System rule only when the workbook has no Stage.
"""
from __future__ import annotations

from datetime import date

import openpyxl

from forecast import optimize as opt

WK = "2028-W10"
WS = date(2028, 3, 6)


def _wb(rows, with_stage=True):
    """A workbook whose BatchLocations matches excel_io.write_batch_locations."""
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("BatchLocations")
    ws.append(["BATCH LOCATIONS"])
    ws.append(["blurb"])
    ws.append([])
    hdr = ["Week", "Week_Start", "Batch", "Tank", "System",
           "Count (fish)", "AvgWt (kg)", "Biomass (kg)", "Density (kg/m3)"]
    if with_stage:
        hdr.append("Stage")
    ws.append(hdr)
    for r in rows:
        ws.append(r if with_stage else r[:9])
    return wb


ROWS = [
    # A 6N MAIN rearing production fish (2028+): SW stage — must COUNT.
    [WK, WS, "B1", 61, "OG6N", 5000, 4.0, 1000.0, 90.0, "SW"],
    # A 6N tank purging: STARVE — expected high density, must NOT count.
    [WK, WS, "B2", 63, "OG6N", 2000, 5.0, 500.0, 120.0, "STARVE"],
    # An ordinary grow-out tank under the welfare line.
    [WK, WS, "B3", 31, "OG3N", 5000, 3.0, 1000.0, 50.0, "SW"],
    # An OG tank starving IN PLACE pre-harvest — purge too, must NOT count.
    [WK, WS, "B4", 32, "OG3N", 3000, 5.0, 800.0, 110.0, "STARVE"],
]


def test_production_mode_6n_mains_count_in_the_quality_metric():
    wb = _wb(ROWS)
    mean_d, crowded_fw, frac = opt._density_quality(wb, welfare=80.0)
    # Only B1 (90 > 80) is crowded; B2/B4 are purge; B3 is under the line.
    assert frac == 1000.0 / 2000.0, "the 6N main's crowding must be visible"
    assert crowded_fw == 5000
    assert mean_d == (90.0 * 1000 + 50.0 * 1000) / 2000.0


def test_peak_and_overshoot_see_the_6n_main_but_not_purge():
    wb = _wb(ROWS)
    assert opt._density_peak(wb) == 90.0, "STARVE rows (110/120) are purge"
    # Overshoot vs the hard 95 cap: neither counted row exceeds it.
    assert opt._density_overshoot(wb) == 0.0
    # Push the 6N main over the hard cap -> it must register.
    hot = [r[:] for r in ROWS]
    hot[0][8] = 97.0
    assert opt._density_overshoot(_wb(hot)) == 0.5   # 1 of 2 counted rows


def test_workbook_without_stage_falls_back_to_the_system_rule():
    wb = _wb(ROWS, with_stage=False)
    mean_d, _fw, frac = opt._density_quality(wb, welfare=80.0)
    # Legacy behaviour: every OG6N row excluded; OG3N STARVE has no stage info
    # so it counts (110 > 80) — exactly what the old code did.
    assert frac == 800.0 / 1800.0
    assert opt._density_peak(wb) == 110.0


# ---- Where the caps come from (2026-09-08) ----
# The 2026-08-26 fix read per-tank caps from a `FacilityConfig` SHEET. This
# engine never writes one -- a controller output has 27 sheets and none is that
# -- so the lookup returned {} on every real workbook and every caller fell back
# to the hardcoded 95.0 the fix existed to remove. Measured on the operator's
# live plan, the metric was scoring 0.076% over-cap against a true 2.87%: a 38x
# understatement, invisible because 95 is a plausible number. Caps now fall back
# to the RunConfig snapshot, which every run stamps.

_FACILITY_YAML = """tanks:
- {location_id: Hatchery1, system_id: HA1, tank_id: 1, volume_m3: 5.5,
   max_density_kg_m3: 30, max_feed_kg_day: 10, type: FW}
- {location_id: L, system_id: OG3N, tank_id: 31, volume_m3: 100,
   max_density_kg_m3: 85, max_feed_kg_day: 100, type: OG}
- {location_id: L, system_id: OG6N, tank_id: 61, volume_m3: 100,
   max_density_kg_m3: 120, max_feed_kg_day: 100, type: OG}
"""


def _with_snapshot(wb, yaml_text=_FACILITY_YAML):
    """Attach a RunConfig snapshot the way write_config_snapshot lays it out."""
    ws = wb.create_sheet("RunConfig")
    ws.append(["RUN CONFIG SNAPSHOT"])
    ws.append([])
    ws.append(["# ===== config/facility.yaml ====="])
    for line in yaml_text.splitlines():
        ws.append([line])
    return wb


def test_caps_come_from_the_run_config_snapshot_when_there_is_no_sheet():
    caps = opt._tank_density_caps(_with_snapshot(_wb(ROWS)))
    assert caps == {"31": 85.0, "61": 120.0}, caps
    # FW is skipped: facility.yaml repeats tank_id 1 across every FW system, so
    # keeping it would collide with a real OG tank id.
    assert "1" not in caps


def test_without_either_source_the_lookup_is_empty_not_wrong():
    """The 95.0 default lives at the call site; this must not invent caps."""
    assert opt._tank_density_caps(_wb(ROWS)) == {}


def test_the_snapshot_caps_actually_reach_the_overshoot_metric():
    """The regression that hid for two weeks: caps resolved but never applied.

    Tank 31 sits at 90 kg/m3. Against its real 85 cap that is over; against the
    hardcoded 95 fallback it is compliant. One row, opposite verdicts.
    """
    rows = [[WK, WS, "B3", 31, "OG3N", 5000, 3.0, 1000.0, 90.0, "SW"]]
    assert opt._density_overshoot(_wb(rows)) == 0.0            # 90 < 95 fallback
    assert opt._density_overshoot(_with_snapshot(_wb(rows))) == 1.0   # 90 > 85 cap


def test_a_6n_production_tank_is_judged_against_its_own_higher_cap():
    """The error ran BOTH ways: 6N production rows were judged 25 too strictly."""
    rows = [[WK, WS, "B1", 61, "OG6N", 5000, 4.0, 1000.0, 100.0, "SW"]]
    assert opt._density_overshoot(_wb(rows)) == 1.0            # 100 > 95 fallback
    assert opt._density_overshoot(_with_snapshot(_wb(rows))) == 0.0   # 100 < 120 cap
