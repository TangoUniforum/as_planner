"""Two presentation defects the operator hit reading the workbook (2026-09-07).

1. WEEK_START SHOWED A TIME. The number-format pass touches only int/float, so
   a date cell kept whatever openpyxl defaulted to when it was written -- and
   that depends on the TYPE. A `date` gets yyyy-mm-dd; a `datetime` gets
   yyyy-mm-dd h:mm:ss. Ledger rows carry both (PR-derived week starts arrive
   from Excel as datetimes, computed ones are dates), so WeeklyReport rendered
   619 of 1,266 Week_Start cells with a meaningless 00:00:00 beside the rest.

2. AV WEIGHT LOST ITS DECIMALS. HarvestPlan Report is laid out by ROW -- each
   batch has Units / Av Weight / Biomass lines across month COLUMNS -- so a
   per-column format cannot tell a count from a weight, and every month cell
   took the integer "#,##0". An average weight of 2.63 kg displayed as "3",
   while the TOTAL column kept its decimals only because it is General.
"""
from __future__ import annotations

from datetime import date, datetime

from openpyxl import Workbook

from forecast.excel_format import _format_table


def _sheet(rows, headers):
    wb = Workbook()
    ws = wb.active
    ws.append(["Title"])
    ws.append([])
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _format_table(ws, header_row=3)
    return ws


def test_a_date_and_a_datetime_render_the_same_way():
    """The whole defect: two types, one column, two different renderings."""
    ws = _sheet([["2026-W36", date(2026, 9, 1), 10],
                 ["2026-W37", datetime(2026, 9, 8, 0, 0), 20]],
                ["Week", "Week_Start", "Count (fish)"])
    fmts = {ws.cell(r, 2).number_format for r in (4, 5)}
    assert fmts == {"yyyy-mm-dd"}, f"mixed date rendering: {fmts}"


def test_no_body_date_carries_a_time():
    ws = _sheet([["w", datetime(2026, 9, 1, 0, 0), 1]],
                ["Week", "Week_Start", "Count (fish)"])
    assert "h:mm" not in ws.cell(4, 2).number_format


def test_numbers_beside_a_date_are_untouched():
    ws = _sheet([["w", date(2026, 9, 1), 1234]],
                ["Week", "Week_Start", "Count (fish)"])
    assert ws.cell(4, 3).number_format == "#,##0"


# The real sheet's month headers are DATES, which is what makes those columns
# take the integer "#,##0" in the first place. A string header would not, and a
# test using one would pass while proving nothing about the sheet that broke.
_MONTHS = [datetime(2026, 9, 1), datetime(2026, 10, 1)]


def test_an_av_weight_row_keeps_two_decimals():
    """Row-labelled, not column-labelled -- 2.63 must not display as 3."""
    ws = _sheet([["B41", "Units", 1000, 2000],
                 ["", "Av Weight - Kg", 2.63, 3.21],
                 ["", "Biomass - Tons", 5.0, 6.0]],
                ["Batch", "Line"] + _MONTHS)
    for c in (3, 4):
        assert ws.cell(5, c).number_format == "#,##0.00"


def test_the_rows_around_it_stay_integers():
    """Only the weight line gains decimals; Units stays a count."""
    ws = _sheet([["B41", "Units", 1000, 2000],
                 ["", "Av Weight - Kg", 2.63, 3.21]],
                ["Batch", "Line"] + _MONTHS)
    assert ws.cell(4, 3).number_format == "#,##0"
    assert ws.cell(5, 3).number_format == "#,##0.00"
