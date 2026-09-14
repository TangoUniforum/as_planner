"""One number per quantity, end to end through run.main on real PRs.

The unit tests in test_numbers_consistency.py pin each writer. These run the
whole pipeline on two corpus ProductionReports (with their era registries) and
assert the CROSS-SHEET invariants the 2026-09-12 numbers sandbox restored --
invariants, not pinned numbers, so a config change that moves the plan cannot
break them, only a writer that stops agreeing with another:

  2026-02-28 (a Sunday report start; batch roll-ups with no Unit rows; the
             realized plan used to stop 11 weeks before the horizon --
             it now walks every horizon week, calendar-09)
  2025-07-31 (a 3-day first week; the planner leaves empty harvest weeks)

Every mutable input is a TEMP copy; costs.yaml is never copied.
"""
from __future__ import annotations

import collections
import contextlib
import datetime as dt
import io
import re
import shutil
from pathlib import Path

import openpyxl
import pytest

from conftest import copy_config_without_costs

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "pr_corpus"


def _run(tmp: Path, close: str, horizon=None):
    pr, reg = CORPUS / f"{close}.xlsx", CORPUS / "registries" / close
    if not (pr.exists() and reg.is_dir()):
        pytest.skip(f"corpus PR {close} not on this machine")
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config_without_costs(ROOT / "config", cfg)
    shutil.copytree(reg, scn)
    if horizon is not None:
        ctl = cfg / "control.yaml"
        ctl.write_text(re.sub(r"(?m)^horizon_weeks:.*$", f"horizon_weeks: {horizon}",
                              ctl.read_text(encoding="utf-8")), encoding="utf-8")
    inp, out = tmp / ("pr" + pr.suffix), tmp / "out.xlsm"
    shutil.copy(pr, inp)
    from forecast.run import main
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                  calib_log_path="")
    assert rc == 0
    if not out.exists():
        out = out.with_suffix(".xlsx")
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    sheets = {n: [tuple(r) for r in wb[n].iter_rows(values_only=True)]
              for n in wb.sheetnames}
    wb.close()
    return {"sheets": sheets, "pr": pr, "horizon": horizon}


@pytest.fixture(scope="module")
def runs(tmp_path_factory):
    return {"2026-02-28": _run(tmp_path_factory.mktemp("n0228"), "2026-02-28"),
            "2025-07-31": _run(tmp_path_factory.mktemp("n0731"), "2025-07-31",
                               horizon=30)}


CLOSES = ("2026-02-28", "2025-07-31")


def _table(rows, pred):
    i = next(k for k, r in enumerate(rows) if r and pred(r))
    hdr = [str(c) if c is not None else "" for c in rows[i]]
    out = []
    for r in rows[i + 1:]:
        if not r or all(c is None for c in r):
            break
        out.append(dict(zip(hdr, r)))
    return out


def _n(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def _report_start(r):
    from forecast.production_report import find_pr_sheet, parse_pr_worksheet
    wb = openpyxl.load_workbook(r["pr"], read_only=True, data_only=True)
    try:
        closing, _og, _fw = parse_pr_worksheet(find_pr_sheet(wb), quiet=True)
    finally:
        wb.close()
    return closing + dt.timedelta(days=1)


def _hp(s):
    return _table(s["HarvestPlan"], lambda r: r[0] == "Week" and "Batch" in r)


def _vlog(s):
    return [(str(r[1]), str(r[2])) for r in s["ValidationLog"] if r and isinstance(r[0], int)]


@pytest.mark.parametrize("close", CLOSES)
def test_runs_tca_labels_and_staging_tanks(runs, close):
    s = runs[close]["sheets"]
    tca = _table(s["TankContinuityAudit"], lambda r: r[0] == "Week" and "Tank" in r)
    lab = {(str(t["Week"]), t["Tank"]): str(t["Batch"] or "") for t in tca}
    for h in _hp(s):                                        # counts-01
        assert h["Batch"] in lab[(str(h["Week"]), h["Tank"])].split("->"), h
    # counts-08: the tank-week's Harvest_Out is the whole fish HarvestPlan
    # shows for that tank-week (it rounded the fractional sum before).
    hp_tw = collections.defaultdict(float)
    for h in _hp(s):
        hp_tw[(str(h["Week"]), h["Tank"])] += _n(h["Count (fish)"])
    out_tw = {(str(t["Week"]), t["Tank"]): _n(t["Harvest_Out"]) for t in tca}
    for k, v in hp_tw.items():
        assert out_tw.get(k, 0.0) == v, (k, out_tw.get(k), v)
    tp = _table(s["TransferPlan"], lambda r: r[0] == "Week" and "Type" in r)
    for t in tp:                                            # counts-02
        if t["Type"] in ("Transfer", "Grade") and _n(t["Count (fish)"]) > 0:
            for tk in (t["From_Tank"], t["To_Tank"]):
                if isinstance(tk, (int, float)):
                    assert (str(t["Week"]), int(tk)) in lab, (t["Week"], tk)


@pytest.mark.parametrize("close", CLOSES)
def test_runs_per_batch_sheets_agree(runs, close):
    s = runs[close]["sheets"]
    hp = _hp(s)
    bp = _table(s["Batch Plan"], lambda r: r[0] == "Batch" and "Peak_tanks" in r)
    tt = _table(s["TransferTemplate"],
                lambda r: r[0] == "Batch" and "SW_Entry_Week" in r)
    assert {h["Batch"] for h in hp} <= {b["Batch"] for b in bp}          # counts-05
    assert {h["Batch"] for h in hp} <= {b["Batch"] for b in tt}
    sw = {t["Batch"]: str(t["SW_Entry_Week"]) for t in tt}
    for b in bp:                                                          # calendar-07
        if b["Batch"] in sw:
            assert str(b["SW_entry"]) == sw[b["Batch"]], b["Batch"]
    tp = _table(s["TransferPlan"], lambda r: r[0] == "Week" and "Type" in r)
    tog = collections.defaultdict(lambda: collections.defaultdict(float))
    for t in tp:
        if t["Type"] == "TranOG":
            tog[t["Batch"]][str(t["Week"])] += _n(t["Count (fish)"])
    for t in tt:                                                          # counts-04
        if t["Batch"] in tog:
            assert _n(t["Entry_Count (fish)"]) == tog[t["Batch"]][min(tog[t["Batch"]])]


@pytest.mark.parametrize("close", CLOSES)
def test_runs_period_totals_are_the_rows_under_them(runs, close):
    s = runs[close]["sheets"]
    hr = _table(s["HarvestReport"], lambda r: r[0] == "Year" and "Batch" in r)
    by_y = collections.defaultdict(float)
    for r in hr:
        by_y[int(r["Year"])] += _n(r["Count (fish)"])
    ys = _table(s["YearlySummary"], lambda r: r[0] == "Year")
    for y in ys:                                               # counts-08 / biomass-06
        assert _n(y["Harvest_Count (fish)"]) == by_y.get(int(y["Year"]), 0.0), y["Year"]
    units = collections.defaultdict(float)
    hpr_month = {}
    year = None
    for row in s["HarvestPlan Report"]:
        if row and isinstance(row[0], str) and re.match(r".* \d{4}$", row[0]) \
                and row[1] is None:
            year = int(row[0][-4:])
        if row and row[1] == "Units":
            key = "TOTAL" if row[0] == "TOTAL" else "rows"
            if key == "rows":
                for i in range(12):
                    units[i] += _n(row[2 + i])
            else:
                for i in range(12):
                    assert _n(row[2 + i]) == units[i], i
                    if isinstance(row[2 + i], (int, float)):
                        hpr_month[f"{year}-{i + 1:02d}"] = row[2 + i]
                units.clear()
    # counts-08: the MonthlyReport TOTAL harvest count is HarvestPlan Report's.
    mr = _table(s["MonthlyReport"], lambda r: r[0] == "Scenario" and "Month" in r)
    for r in mr:
        if r["Batch"] == "TOTAL" and str(r["Month"]) in hpr_month:
            assert _n(r["Harv_Count (fish)"]) == hpr_month[str(r["Month"])], r["Month"]


@pytest.mark.parametrize("close", CLOSES)
def test_runs_named_weeks_and_days(runs, close):
    s = runs[close]["sheets"]
    rs = _report_start(runs[close])
    bl = _table(s["BatchLocations"], lambda r: r[0] == "Week" and "Batch" in r)
    zero = {str(r["Week"]) for r in bl} - {str(h["Week"]) for h in _hp(s)}
    named = {re.search(r"\d{4}-W\d{2}", d).group(0) for _c, d in _vlog(s)
             if "NO HARVEST" in d}
    assert zero <= named                                      # harvest zero weeks
    rows = s["Daily Harvest Schedule"]                         # harvest daily / calendar-05
    for r in rows:
        if r and isinstance(r[0], int) and isinstance(r[2], (dt.date, dt.datetime)):
            d = r[2].date() if isinstance(r[2], dt.datetime) else r[2]
            assert d >= rs, r


@pytest.mark.parametrize("close", CLOSES)
def test_runs_facility_feed_is_one_number(runs, close):
    s = runs[close]["sheets"]
    adv = {str(r["Week"]): _n(r["Total_Feed (kg/day)"])
           for r in _table(s["Advisory"], lambda r: r[0] == "Week" and "Week_Start" in r)}
    fm = s["FacilityMap"]
    i = next(k for k, r in enumerate(fm) if r and str(r[0]).startswith("TOTAL PLANNED FEED"))
    hdr = fm[i + 1]
    fac = next(r for r in fm[i + 1:] if r and r[0] == "FACILITY")
    for j in range(2, len(hdr)):                               # feed-04
        if hdr[j]:
            assert abs(_n(fac[j]) - adv.get(str(hdr[j]), 0.0)) <= 1.0, hdr[j]


@pytest.mark.parametrize("close", CLOSES)
def test_runs_detection_lines_match_the_facts(runs, close):
    s = runs[close]["sheets"]
    vl = _vlog(s)
    # calendar-09: PLAN ENDS EARLY iff the realized plan stops before the
    # horizon with fish still in the tanks.
    bl = _table(s["BatchLocations"], lambda r: r[0] == "Week" and "Batch" in r)
    wks = sorted({str(r["Week"]) for r in bl})
    early = [d for c, d in vl if c == "WARNING - Plan ends before the horizon"]
    last_fish = sum(_n(r["Count (fish)"]) for r in bl if str(r["Week"]) == wks[-1])
    if early:
        assert wks[-1] in early[0] and last_fish > 0
    # counts-03: PR NOT HYDRATED lines iff the PR's facility close is not the
    # sum of its Unit rows under Bnn batches.
    from forecast.production_report import pr_structure_warnings, find_pr_sheet
    wb = openpyxl.load_workbook(runs[close]["pr"], read_only=True, data_only=True)
    try:
        want = pr_structure_warnings(find_pr_sheet(wb))
    finally:
        wb.close()
    got = [d for c, d in vl if c == "WARNING - PR fish not hydrated"]
    assert got == want


def test_runs_the_2026_02_28_plan_walks_the_whole_horizon(runs):
    """calendar-09 (engine change 2 of 7, taken one at a time).

    This corpus PR's projection runs out of batches 11 weeks before the end of
    the horizon (an era registry with no future stocking). The realized plan
    used to walk only the weeks the projection had load for, so it STOPPED at
    2027-W29 with 7,128 fish / 37,762 kg left in tank 63 -- never harvested.
    It now walks every horizon week and harvests them (2027-W30).

    Measured on the tanks themselves, not on the week list: BatchLocations
    lists only non-empty tanks, so an emptied facility has no rows after its
    last full week either way. The proof is that every fish still in a tank
    in that last week is harvested after it -- and that on this PR there are
    such fish (the case the fix is for). Without calendar-09 none are."""
    s = runs["2026-02-28"]["sheets"]
    early = [d for c, d in _vlog(s) if c == "WARNING - Plan ends before the horizon"]
    assert not early, early
    bl = _table(s["BatchLocations"], lambda r: r[0] == "Week" and "Batch" in r)
    last = max(str(r["Week"]) for r in bl)
    left = sum(_n(r["Count (fish)"]) for r in bl if str(r["Week"]) == last)
    hp = _table(s["HarvestPlan"], lambda r: r[0] == "Week" and "Batch" in r)
    later = sum(_n(r["Count (fish)"]) for r in hp if str(r["Week"]) > last)
    assert left > 0 and later > 0, (
        f"no fish harvested after the last tank week {last} "
        f"({left:,.0f} fish still in tanks then): the plan stopped early")
    assert abs(later - left) <= 1, (last, left, later)
