"""The split FW/SW batch, end to end through run.main on real ProductionReports.

The 2026-08-31 PR holds B49 split: 47,743 fish in seawater (tanks 14 + 24) and
250,225 in freshwater. With its shipped manual events the FW part is moved by a
scripted fw_to_og; with an EMPTY events file nothing moved it and the fish sat
in the opening forever (Count_Check +250,225, "FW PART NOT MODELLED").

Operator decisions 2026-09-11 (see tests/test_split_batch.py for the rules):
auto-model the FW part (split_batch_fw: auto, the default); `off` keeps the V1
engine and its detection; a manual fw_to_og always wins; a second one for the
same batch is refused. Every mutable input here is a TEMP copy -- the live
config/ and scenario/ are never written, and costs.yaml is never copied.
"""
from __future__ import annotations

import contextlib
import io
import re
import shutil
from collections import defaultdict
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
PR_0831 = Path(r"C:\Users\julian.f\Downloads"
               r"\8 31 2026 AS Monthly Production Report_planned (31).xlsm")
EV_0831 = "2026-08-31.yaml"

pytestmark = pytest.mark.skipif(not PR_0831.exists(),
                                reason="the 2026-08-31 ProductionReport is not on this machine")

SECOND_FW = """- type: fw_to_og
  week: 2
  from_tank: null
  destinations:
  - tank: 24
    count: null
    avg_wt_g: null
    size_class: big
  count: 240000.0
  batch: B49
  mode: transfer
  notes: ''
"""


def _events_without_fw_to_og(txt: str) -> str:
    """The shipped 8/31 events with the fw_to_og entry removed (text edit on a
    TEMP copy: the events are a YAML list under `events:`)."""
    import yaml
    d = yaml.safe_load(txt) or {}
    d["events"] = [e for e in d.get("events", []) if e.get("type") != "fw_to_og"]
    return yaml.safe_dump(d, sort_keys=False)


def run_pr(tmp: Path, copy_config, *, events="shipped", mode=None, pr=PR_0831,
           scenario_src=None, horizon=None, edit_scenario=None):
    cfg, scn = tmp / "config", tmp / "scenario"
    copy_config(ROOT / "config", cfg)
    shutil.copytree(scenario_src or (ROOT / "scenario"), scn)
    ctl = cfg / "control.yaml"
    txt = ctl.read_text(encoding="utf-8")
    if horizon is not None:
        txt = re.sub(r"(?m)^horizon_weeks:.*$", f"horizon_weeks: {horizon}", txt)
    if mode is not None:
        txt = txt.rstrip("\n") + f"\nsplit_batch_fw: '{mode}'\n"
    ctl.write_text(txt, encoding="utf-8")
    ev = scn / "manual_events" / EV_0831
    if events == "empty":
        ev.parent.mkdir(exist_ok=True)
        ev.write_text("events: []\n", encoding="utf-8")
    elif events == "second_fw":
        ev.write_text(ev.read_text(encoding="utf-8").rstrip("\n") + "\n" + SECOND_FW,
                      encoding="utf-8")
    elif events == "no_fw_to_og":
        ev.write_text(_events_without_fw_to_og(ev.read_text(encoding="utf-8")),
                      encoding="utf-8")
    elif events in ("week1_no_fw_to_og", "week1_only"):
        import yaml
        d = yaml.safe_load(ev.read_text(encoding="utf-8")) or {}
        d["events"] = [e for e in d.get("events", [])
                       if int(e.get("week") or 1) == 1
                       and (events == "week1_only" or e.get("type") != "fw_to_og")]
        ev.write_text(yaml.safe_dump(d, sort_keys=False), encoding="utf-8")
    if edit_scenario is not None:
        edit_scenario(scn)
    inp, out = tmp / ("pr" + pr.suffix), tmp / "out.xlsm"
    shutil.copy(pr, inp)
    from forecast.run import main
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = main(str(inp), str(out), config_dir=str(cfg), scenario_dir=str(scn),
                  calib_log_path="")
    assert rc == 0
    if not out.exists():
        out = out.with_suffix(".xlsx")
    return read_out(out, buf.getvalue())


def _rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def read_out(path, stdout):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tp = _rows(wb["TransferPlan"])
    th = next(i for i, r in enumerate(tp) if r and r[0] == "Week")
    col = {str(c): j for j, c in enumerate(tp[th]) if c is not None}
    # TransferPlan writes one TranOG row PER DESTINATION tank: one arrival is
    # the (batch, week) sum.
    _tog = defaultdict(float)
    tog_rows = []                      # (batch, week, to_tank, count, avg_wt_kg)
    for r in tp[th + 1:]:
        if r and r[col["Type"]] == "TranOG":
            _tog[(str(r[col["Batch"]]), str(r[col["Week"]]))] += float(
                r[col["Count (fish)"]] or 0)
            tog_rows.append((str(r[col["Batch"]]), str(r[col["Week"]]),
                             r[col["To_Tank"]], float(r[col["Count (fish)"]] or 0),
                             float(r[col["Avg_Weight (kg)"]] or 0)))
    tog = [(b, w, n) for (b, w), n in sorted(_tog.items())]
    vlog = [(str(r[1]), str(r[2])) for r in _rows(wb["ValidationLog"])
            if r and isinstance(r[0], int)]
    ica = _rows(wb["InputConservationAudit"])
    hi = next(i for i, r in enumerate(ica) if r and r[0] == "Batch")
    ica_rows = {r[0]: dict(zip(ica[hi], r)) for r in ica[hi + 1:] if r and r[0]}
    ica_head = [str(r[0]) for r in ica[:hi] if r and r[0]]
    bl = _rows(wb["BatchLocations"])
    bh = next(i for i, r in enumerate(bl) if r and r[0] == "Week")
    bc = {str(c): j for j, c in enumerate(bl[bh]) if c is not None}
    locs = defaultdict(dict)           # (batch, week) -> {tank: (count, avg_wt_kg)}
    for r in bl[bh + 1:]:
        if r and r[0]:
            locs[(str(r[bc["Batch"]]), str(r[bc["Week"]]))][r[bc["Tank"]]] = (
                float(r[bc["Count (fish)"]] or 0), float(r[bc["AvgWt (kg)"]] or 0))
    tca = _rows(wb["TankContinuityAudit"])
    drift = sum(1 for i, r in enumerate(tca) if i >= 5 and len(r) > 14 and r[14] == "TANK_DRIFT")
    weekly = []
    wr = _rows(wb["WeeklyReport"])
    wh = next(i for i, r in enumerate(wr) if r and r[0] == "Scenario")
    hdr = [str(c) if c is not None else "" for c in wr[wh]]
    weekly = [dict(zip(hdr, r)) for r in wr[wh + 1:] if r and r[3] is not None]
    wb.close()
    m = re.search(r"biomass over-cap (\d+).*?feed over-cap (\d+)", stdout)
    return SimpleNamespace(path=path, tog=tog, tog_rows=tog_rows, locs=dict(locs),
                           vlog=vlog, ica=ica_rows, ica_head=ica_head,
                           drift=drift, weekly=weekly, stdout=stdout,
                           sys_over=(int(m.group(1)), int(m.group(2))) if m else None)


def _n(d, k):
    v = d.get(k)
    return float(v) if isinstance(v, (int, float)) else 0.0


def _pr_b49():
    from forecast.production_report import read_production_report
    pwb = openpyxl.load_workbook(PR_0831, read_only=True, data_only=True)
    _c, og, fw = read_production_report(pwb)
    pwb.close()
    sw = sum(r.closing_count for r in og if r.batch_id == "B49")
    f = sum(r.closing_count for r in fw if r.batch_id == "B49")
    return sw, f


def _lands_on_own_tanks(r, batch):
    """The arrival TOPS UP the batch's own tanks -- every destination closes
    the entry week holding MORE of the batch than the arrival put there (its
    seawater part is in it too), which an empty tank never does -- and the
    bigger class goes to the tank of heavier fish. (Own tanks are judged at
    entry, not at the PR close: the planner may move the seawater part within
    the entry tier first -- B42 2025-09-30 moved from tank 23 to 13 on a
    30-week horizon.)"""
    rows = [x for x in r.tog_rows if x[0] == batch]
    assert rows, batch
    wk = rows[0][1]
    held = r.locs.get((batch, wk), {})
    for _b, _w, tank, n, _wt in rows:
        assert tank in held and held[tank][0] > n, (tank, n, held.get(tank))
    if len(rows) >= 2:
        big = max(rows, key=lambda x: x[4])
        small = min(rows, key=lambda x: x[4])
        assert held[big[2]][1] >= held[small[2]][1], (big, small, held)


@pytest.fixture(scope="module")
def auto_empty(tmp_path_factory, copy_config):
    return run_pr(tmp_path_factory.mktemp("split_auto"), copy_config, events="empty")


@pytest.fixture(scope="module")
def off_empty(tmp_path_factory, copy_config):
    return run_pr(tmp_path_factory.mktemp("split_off"), copy_config, events="empty",
                  mode="off")


# ---- auto: B49's FW part is placed once, in its own week ---------------------------

def test_split_8_31_auto_places_the_fw_part_once(auto_empty):
    sw, fw = _pr_b49()
    assert sw > 0 and fw > 0, "B49 is no longer split on the 8/31 PR"
    b49 = [t for t in auto_empty.tog if t[0] == "B49"]
    assert len(b49) == 1, b49
    assert b49[0][1] == "2026-W38"                      # scenario tran_og_date 2026-09-14
    assert b49[0][2] == pytest.approx(290_000 - sw, abs=1)   # remaining target 242,257
    assert not any("SHORT-STOCKED" in d and "B49" in d for _c, d in auto_empty.vlog)
    _lands_on_own_tanks(auto_empty, "B49")                  # tanks 24 + 14


@pytest.mark.parametrize("case", ["auto_empty", "w1_auto"])
def test_split_8_31_auto_spends_no_move_freeing_tanks_for_the_top_up(case, request):
    """The top-up needs no empty tank, so neither the anticipatory pacing nor
    the arrival make-room may free one for it. Measured with the credits
    removed: with both, a 27,205-fish pacing move in 2026-W37 and an
    entry-tank vacate in 2026-W38 (empty events); with the pacing credit alone,
    a 7,140-fish move into 6N in 2026-W37 "holding a tank for TranOG arrival
    in 2026-W38" (one-week window) -- moves against the hard 15-move weekly
    cap, for nothing."""
    r = request.getfixturevalue(case)
    freed = [d for _c, d in r.vlog
             if "arrival in 2026-W38" in d
             or ("VACATED entry tank" in d and ("2026-W37" in d[:16] or "2026-W38" in d[:16]))]
    assert not freed, freed[:3]


def test_split_8_31_auto_warning_line_once(auto_empty):
    hits = [d for c, d in auto_empty.vlog if "auto-transferred" in d and "B49" in d]
    assert len(hits) == 1, hits
    assert hits[0].startswith("SPLIT BATCH B49: 250,225 FW fish auto-transferred 2026-W38")
    assert not any("Split batch at PR close" in c for c, _d in auto_empty.vlog)


def test_split_8_31_auto_conserves(auto_empty):
    a = auto_empty
    assert a.ica_head[2].startswith("OK"), a.ica_head[:4]
    assert a.ica["B49"]["Status"] == "PLACED"
    assert not any("NOT MODELLED" in h or "BREACH" in h for h in a.ica_head), a.ica_head
    assert a.ica["B49"]["FW_Flag"] == "auto split (remaining target)"
    res = a.ica["B49"]["FW_Bal_Residual (fish)"]
    assert isinstance(res, (int, float)) and abs(res) <= 0.02 * 250_225
    assert a.drift == 0


def test_split_8_31_auto_ledger_opens_on_every_fish_and_chains(auto_empty):
    sw, fw = _pr_b49()
    rows = [d for d in auto_empty.weekly if d.get("Batch") == "B49"]
    rows.sort(key=lambda d: d["Week"])
    assert rows[0]["Week"] == "2026-W36"
    assert _n(rows[0], "Open_Count (fish)") == pytest.approx(sw + fw, abs=1)
    for a, b in zip(rows, rows[1:]):
        assert abs(_n(b, "Open_Count (fish)") - _n(a, "Close_Count (fish)")) <= 1, b["Week"]
    for d in rows[:3]:
        assert abs(_n(d, "Count_Check (fish)")) <= 30, (d["Week"], d["Count_Check (fish)"])
    w38 = next(d for d in rows if d["Week"] == "2026-W38")
    assert _n(w38, "Xfer_In (fish)") == pytest.approx(290_000 - sw, abs=1)
    assert _n(w38, "Input_Count (fish)") == 0
    # FW feed now reaches the ledger for the FW weeks, and rates print
    w36 = rows[0]
    assert w36["SGR (%/day)"] is not None


def test_split_8_31_off_keeps_v1_and_its_detection(off_empty):
    o = off_empty
    assert not [t for t in o.tog if t[0] == "B49"]
    hits = [d for c, d in o.vlog
            if c == "WARNING - Split batch at PR close (FW part not modelled)"]
    assert len(hits) == 1 and "B49" in hits[0]
    assert o.ica["B49"]["Status"] == "*** FW PART NOT MODELLED ***"
    assert not any("auto-transferred" in d for _c, d in o.vlog)


def _breaches(r):
    dens = sum(1 for c, _d in r.vlog if c == "WARNING - Density")
    ceil = sum(1 for c, _d in r.vlog if c.startswith("WARNING - Harvest ceiling"))
    hb = sum(1 for c, _d in r.vlog if c.startswith("WARNING - Handling budget"))
    zero = sum(1 for c, d in r.vlog if "zero-harvest" in d.lower())
    return dict(density=dens, ceiling=ceil, handling=hb, zero_weeks=zero,
                sys_bio=r.sys_over[0], sys_feed=r.sys_over[1])


@pytest.fixture(scope="module")
def w1_auto(tmp_path_factory, copy_config):
    """The shipped week-1 events without their fw_to_og: B49's FW part is
    moved by the automatic path (2026-W38)."""
    return run_pr(tmp_path_factory.mktemp("split_w1_auto"), copy_config,
                  events="week1_no_fw_to_og")


@pytest.fixture(scope="module")
def w1_manual(tmp_path_factory, copy_config):
    """The same week-1 events WITH their fw_to_og: the operator's own route
    for the same fish (2026-W36, tanks 24 + 14)."""
    return run_pr(tmp_path_factory.mktemp("split_w1_manual"), copy_config,
                  events="week1_only")


def test_split_auto_carries_the_fish_no_worse_than_the_manual_route(w1_auto, w1_manual):
    """The hard-limit comparison that means something. `off` is NOT the
    baseline for it: `off` never moves ~242,000 of the PR's own fish out of
    freshwater, so it plans a lighter facility that does not exist. Measured
    2026-09-12 on the empty-events synthetic (a plan already far outside its
    limits in V1 -- the operator's week-1/2 script is load-bearing):
      off  1,395 density tank-weeks / 482 system-biomass / 410 system-feed
      auto 1,748 / 601 / 613 (and a 10-fish nudge of `off` alone moved it to
      1,479 / 534 / 434 -- the planner is mode-discontinuous).
    Carrying the same fish the operator's way (a week-1 fw_to_og into tanks
    24 + 14) measured 1,511 / 518 / 544 against auto's 1,550 / 529 / 538. So
    the automatic path must be no worse than the manual route, within the
    measured discontinuity band.

    The band is 40% (+3 for the small counts), set by the operator
    (2026-09-14) from a measurement: numbers-audit finding E1 (a short first
    week walked for its own days) makes B43 0.4% lighter in 2026-W37 -- one
    day less growth -- and the planner then picks different moves on both
    routes. Density warnings / system-feed breaches (manual, auto): before
    85 / 47 and 91 / 40; after 77 / 38 and 101 / 51 (auto 1.31x and 1.34x the
    manual route). The sums hardly moved (176 / 87 -> 178 / 89), the extra
    breaches sit a year and more out in both directions, no tank is filled
    into a full one on either route, and the automatic route's worst tank
    fell 106 -> 98 kg/m3. The old 12% band was narrower than the planner's
    own swing."""
    a, m = _breaches(w1_auto), _breaches(w1_manual)
    worse = {k: (m[k], a[k]) for k in a if a[k] > m[k] * 1.40 + 3}
    assert not worse, f"auto breaches more than the manual route on {worse} (manual, auto)"


def _feed_total(path, week):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = _rows(wb["FeedForecastWeekly"])
    wb.close()
    hdr = next(r for r in rows if r and r[0] == "Feed Type")
    tot = next(r for r in rows if r and r[0] == "Total (kg)")
    return float(tot[list(hdr).index(week)] or 0)


def test_split_fw_feed_reaches_the_feed_report_as_in_the_ledger(auto_empty, off_empty):
    """The split's freshwater feed is real feed: FeedForecastWeekly carries it
    exactly as the ledger does (B49 2026-W36: ~9,900 kg of FW feed that `off`
    never plans). The report writers read FW rows only, so the FW track joins
    them -- never the planner's one-row-per-week states."""
    def led(r, b, w):
        return sum(_n(d, "Feed (kg)") for d in r.weekly if d["Batch"] == b and d["Week"] == w)
    for w in ("2026-W36", "2026-W37"):
        b49 = led(auto_empty, "B49", w) - led(off_empty, "B49", w)
        tot = led(auto_empty, "TOTAL", w) - led(off_empty, "TOTAL", w)
        ffw = _feed_total(auto_empty.path, w) - _feed_total(off_empty.path, w)
        assert b49 > 5_000, (w, b49)                      # B49's FW part eats
        assert abs(ffw - tot) <= 2, (w, ffw, tot)          # report == ledger


def test_a_past_date_split_on_a_full_entry_tier_tops_up_without_moves(tmp_path, copy_config):
    """B49 with its tran_og_date moved before the close (2026-08-20): the FW
    part enters in the FIRST forecast week, on the entry tier the PR left 12/12
    full. The top-up needs no empty tank, so neither the arrival make-room nor
    the anticipatory pacing may spend a move on it (15 is a hard weekly cap)."""
    def past(scn):
        p = scn / "batches.yaml"
        t = p.read_text(encoding="utf-8")
        i = t.index("- batch_id: B49")
        j = t.index("tran_og_date:", i)
        k = t.index("\n", j)
        p.write_text(t[:j] + "tran_og_date: '2026-08-20'" + t[k:], encoding="utf-8")
    sw, _fw = _pr_b49()
    r = run_pr(tmp_path, copy_config, events="empty", edit_scenario=past)
    b49 = [t for t in r.tog if t[0] == "B49"]
    assert len(b49) == 1 and b49[0][1] == "2026-W36", b49
    assert b49[0][2] == pytest.approx(290_000 - sw, abs=1)
    line = next(d for _c, d in r.vlog if d.startswith("SPLIT BATCH B49:"))
    assert "already past at the PR close" in line and "first forecast week" in line
    freed = [d for _c, d in r.vlog if "2026-W36" in d[:16]
             and ("VACATED entry tank" in d or "arrival in 2026-W36" in d)]
    assert not freed, freed[:3]
    assert r.drift == 0


# ---- a manual fw_to_og wins; a second one is refused ------------------------------

@pytest.fixture(scope="module")
def shipped(tmp_path_factory, copy_config):
    return run_pr(tmp_path_factory.mktemp("split_shipped"), copy_config)


def test_split_8_31_manual_wins(shipped):
    b49 = [t for t in shipped.tog if t[0] == "B49"]
    assert len(b49) == 1 and b49[0][1] == "2026-W36"
    assert b49[0][2] == pytest.approx(240_000, abs=1)
    assert not any("auto-transferred" in d for _c, d in shipped.vlog)
    assert not any("Split batch at PR close" in c for c, _d in shipped.vlog)


def test_second_fw_to_og_refused_end_to_end(tmp_path, copy_config):
    r = run_pr(tmp_path, copy_config, events="second_fw")
    b49 = [t for t in r.tog if t[0] == "B49"]
    assert len(b49) == 1, b49                         # was 2: 240,000 fish created
    ref = [d for c, d in r.vlog if c == "ERROR - Manual window (REFUSED)" and "fw_to_og" in d]
    assert len(ref) == 1 and "B49" in ref[0] and "already" in ref[0]
    assert r.ica_head[2].startswith("OK")


def test_split_window_crossing_guard(tmp_path, copy_config):
    """The shipped week-1+2 window minus its fw_to_og opens the planner at
    2026-W38 -- on B49's own entry week. The window does not auto-transfer, so
    the existing guard must refuse the run, naming B49 (the auto path never
    handles fish inside a manual window)."""
    with pytest.raises(ValueError, match="B49"):
        run_pr(tmp_path, copy_config, events="no_fw_to_og")


def test_split_after_a_one_week_window_is_placed_once(w1_auto):
    """A one-week manual window shifts the planner to a TUESDAY start, and the
    ragged first planner week's day loop ran to week_start + 7 -- onto the
    Monday B49 enters (since numbers-audit finding E1 it stops at that Monday). An ordinary arrival has no Phase-C row the week before
    it enters seawater; a split batch does, so it was topped up TWICE (484,514
    fish placed for 242,257, a FW mass-balance breach and 4 TANK_DRIFT rows)."""
    sw, _fw = _pr_b49()
    r = w1_auto
    b49 = [t for t in r.tog if t[0] == "B49"]
    assert len(b49) == 1 and b49[0][1] == "2026-W38", b49
    assert b49[0][2] == pytest.approx(290_000 - sw, abs=1)
    assert r.drift == 0
    assert not any("BREACH" in h for h in r.ica_head), r.ica_head


# ---- the split corpus: every historical split places its FW part -----------------

CORPUS = ROOT / "pr_corpus"


@pytest.mark.parametrize("close,batch", [("2024-11-30", "B36"), ("2025-06-30", "B40"),
                                         ("2025-07-31", "B41"), ("2025-09-30", "B42")])
def test_split_corpus(tmp_path, copy_config, close, batch):
    pr = CORPUS / f"{close}.xlsx"
    reg = CORPUS / "registries" / close
    if not (pr.exists() and reg.is_dir()):
        pytest.skip(f"corpus PR {close} not on this machine")
    r = run_pr(tmp_path, copy_config, pr=pr, scenario_src=reg, events=None, horizon=30)
    got = [t for t in r.tog if t[0] == batch]
    assert len(got) == 1, got                            # placed once, no abort
    assert got[0][2] > 50_000
    _lands_on_own_tanks(r, batch)
    assert any(d.startswith(f"SPLIT BATCH {batch}:") and "auto-transferred" in d
               for _c, d in r.vlog)
    assert r.ica[batch]["Status"] == "PLACED"
    assert r.drift == 0
    # No FALSE FW mass-balance breach for the split (2026-09-12 review:
    # B40 read -2.8% and B42 -2.5%, every fish accounted for, because the
    # reconcile cull sat in the one-week FW track's MEAN count). A breach of
    # ANOTHER batch is not this test's business -- 2024-11-30 B44 -2.1% is
    # identical in V1.
    assert not any("BREACH" in h and f"{batch} (" in h for h in r.ica_head), r.ica_head
    res = r.ica[batch]["FW_Bal_Residual (fish)"]
    assert isinstance(res, (int, float)) and abs(res) <= 2, res
    if batch == "B36":                                   # target met by the SW part
        assert any("target already met by the SW part" in d for _c, d in r.vlog
                   if d.startswith("SPLIT BATCH B36"))
