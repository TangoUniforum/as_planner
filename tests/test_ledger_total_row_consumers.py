"""A TOTAL row must not be double-counted by anything that reads the ledger.

WeeklyReport gained a per-week TOTAL row (2026-09-07). It repeats that week's
figures in the SAME columns as the batch rows, so any consumer that sums a
column over all rows now counts every week twice. That is not hypothetical:
app.py's feed-vs-cap chart summed `Feed (kg)` per week and would have shown
double the feed against the cap on the very first run.

These pin the shape a consumer must expect, so the next reader of this sheet
hits a failing test rather than a chart that is quietly 2x.
"""
from __future__ import annotations

from datetime import date

import openpyxl

from forecast import excel_io
from forecast.placement import BatchLocationRow

_MON = date(2026, 8, 3)          # 2026-W32


def _loc(batch, tank, bio, dens=70.0):
    return BatchLocationRow(
        week_label="2026-W32", week_start=_MON, batch_id=batch, tank_id=tank,
        location_id=f"OG3-{tank}", system_id="OG3", count=10_000,
        avg_wt_g=bio / 10_000 * 1000.0, biomass_kg=bio, density_kg_m3=dens)


def _sheet():
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, [_loc("B1", 40, 40_000.0),
                                      _loc("B2", 41, 30_000.0)], [], [])
    ws = wb["WeeklyReport"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Scenario")
    return rows[hi], [r for r in rows[hi + 1:] if r and r[0] is not None]


def test_the_total_row_is_labelled_TOTAL_in_the_batch_column():
    """The only thing a consumer can filter on. If this label ever changes,
    every guard downstream breaks silently."""
    hdr, data = _sheet()
    bi = hdr.index("Batch")
    labels = [r[bi] for r in data]
    assert "TOTAL" in labels
    assert labels.count("TOTAL") == 1        # one week in this fixture


def test_summing_a_column_without_skipping_TOTAL_double_counts():
    """The failure mode itself, pinned so it stays visible."""
    hdr, data = _sheet()
    bi, ci = hdr.index("Batch"), hdr.index("Close_Bio (kg)")
    naive = sum(r[ci] or 0 for r in data)
    guarded = sum(r[ci] or 0 for r in data
                  if str(r[bi]).strip().upper() != "TOTAL")
    assert guarded == 70_000
    assert naive == 2 * guarded


def test_app_feed_reader_skips_the_total_row():
    """app.py's feed-vs-cap chart sums Feed (kg) per week off this sheet."""
    import re
    src = open("app.py", encoding="utf-8").read()
    i = src.index('w, f = g("Week"), g("Feed (kg)")')
    window = src[i:i + 600]
    assert 'g("Batch")' in window and "TOTAL" in window, (
        "app.py sums Feed (kg) per week without skipping the TOTAL row")


def test_grouped_sheet_carries_blank_rows_and_no_filter():
    """The readable twin. A filter over blanks covers only the first group."""
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, [_loc("B1", 40, 40_000.0)], [], [])
    assert wb["WeeklyReport"].auto_filter.ref is not None
    g = wb["WeeklyReport Grouped"]
    assert g.auto_filter.ref is None
