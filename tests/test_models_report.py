"""Models sheet (report only, operator 2026-10-19): the growth, FCR and other
biology models a run used, spelled out. Each check has a negative control."""
from __future__ import annotations

import shutil
from dataclasses import asdict
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast.biology import BASE_GROWTH_MODEL, _interp
from forecast.config_io import load_config
from forecast.models_report import (adjustment_text, fcr_rows, growth_curve_rows,
                                    growth_model_rows, other_rows)

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"


def _b(bid, gm="", corr=1.0, fcr="FCR_116_Quick", fw=1.0):
    return SimpleNamespace(batch_id=bid, growth_model=gm, sgr_correction=corr,
                           fcr_model=fcr, fw_correction=fw)


@pytest.fixture(scope="module")
def cfg():
    control, tables, _fac = load_config(ROOT / "config")
    return control, tables


def _store(tmp_path, tables, name="Measured X"):
    """A growth history holding one measured fit and one ASSUMED version."""
    import json
    from forecast.growth_history import (STORE_FORMAT, DerivedModel, GrowthFit)
    fit = GrowthFit(name=name, as_of="2026-10-15", window_months=24,
                    window_start="2024-10-15", breakpoint_g=2000.0, a=0.6, b=0.4,
                    cap_g_day=13.9, n=370, n_provisional=23, n_below=161,
                    n_above=209, sizes_g=list(tables.sgr_size_g),
                    sgr_pct_day=list(tables.sgr_sw_pct_day),
                    holdout={"n": 79, "from": "2026-07-15",
                             "fit": {"median_err_pct": 1.03, "median_abs_err_pct": 1.6},
                             "base": {"median_err_pct": 3.12, "median_abs_err_pct": 3.2}})
    der = DerivedModel(name="Assumed Y", source=name, method="pct_per_year",
                       params={"pct": 3}, growing_in="2027-06-01", factor=1.03,
                       explain="ASSUMED: 3% a year better than Measured X.",
                       sizes_g=list(tables.sgr_size_g),
                       sgr_pct_day=list(tables.sgr_sw_pct_day))
    p = tmp_path / "gh.json"
    p.write_text(json.dumps({"format": STORE_FORMAT, "snapshots": [],
                             "fits": [asdict(fit)], "derived": [asdict(der)]}),
                 encoding="utf-8")
    return p


def test_adjustment_is_words_never_a_bare_one():
    assert adjustment_text([1.0]) == "none - the curve as it is"
    assert adjustment_text([0.9]) == "x0.9 - 10% slower than the curve"
    assert adjustment_text([1.25], "FW curve") == "x1.25 - 25% faster than the FW curve"
    assert adjustment_text([1.0, 0.9, 1.0]).count(";") == 1        # distinct values
    # negative control: the old display was the bare number
    assert adjustment_text([1.0]) != "1"


def test_growth_rows_say_how_each_model_was_made(tmp_path, cfg):
    _c, tables = cfg
    p = _store(tmp_path, tables)
    bs = [_b("B44", "Measured X"), _b("B45", "Measured X", 0.9),
          _b("B46", "Assumed Y"), _b("B47", ""), _b("B48", "Missing Z")]
    rows = {r[0]: r for r in growth_model_rows(bs, tables, p)}
    m = rows["Measured X"]
    assert m[1] == "Measured" and m[2] == "B44, B45"
    assert "370 clean tank-months" in m[4] and "13.9 g/day above 2,000 g" in m[4]
    assert "+1.0%" in m[5] and "+3.1%" in m[5]
    assert "x0.9 - 10% slower" in m[3] and "none" in m[3]
    assert rows["Assumed Y"][1] == "ASSUMED improvement"
    assert "ASSUMED: 3% a year" in rows["Assumed Y"][4]
    assert rows[BASE_GROWTH_MODEL][1] == "Company table (assumed)"
    # negative controls: a name the history lacks, and no history at all, are
    # never called measured
    assert rows["Missing Z"][1] == "From biology.yaml"
    assert "not in the growth history" in rows["Missing Z"][4]
    none = {r[0]: r for r in growth_model_rows(bs, tables, tmp_path / "nope.json")}
    assert none["Measured X"][1] == "From biology.yaml"
    assert "not found" in none["Measured X"][4]


def test_growth_curves_are_the_runs_tables(cfg):
    _c, tables = cfg
    tables = SimpleNamespace(**vars(tables))
    tables.sgr_sw_by_model = {"M": [v * 0.8 if v else v for v in tables.sgr_sw_pct_day]}
    head, rows = growth_curve_rows([_b("B1", "M")], tables)
    assert head[1] == "M - SGR (%/day)" and "not used" in head[3]
    for r in rows:
        w = r[0]
        assert r[1] == pytest.approx(_interp(w, tables.sgr_size_g, tables.sgr_sw_by_model["M"]))
        assert r[2] == pytest.approx(w * r[1] / 100.0)
        assert r[3] == pytest.approx(_interp(w, tables.sgr_size_g, tables.sgr_sw_pct_day))
    # negative control: the model is not the base curve
    assert any(abs(r[1] - r[3]) > 1e-6 for r in rows)


def test_fcr_rows_group_batches_by_model(cfg):
    _c, tables = cfg
    head, rows = fcr_rows([_b("B2", fcr="FCR_118_Quick"), _b("B1"), _b("B3")], tables)
    by = {r[0]: r for r in rows}
    assert by["1.16 (FCR_116_Quick)"][2] == "B1, B3"
    assert by["1.18 (FCR_118_Quick)"][2] == "B2"
    i = head.index("FCR at 2,000 g")
    assert by["1.16 (FCR_116_Quick)"][i] == pytest.approx(
        _interp(2000, tables.fcr_size_g, tables.fcr_by_model["1.16"]))
    assert "assumed" in by["1.16 (FCR_116_Quick)"][1]
    # negative control: different models give different values
    assert by["1.16 (FCR_116_Quick)"][i] != by["1.18 (FCR_118_Quick)"][i]


def test_weekly_factors_listed_for_the_forecast_weeks_only(cfg):
    control, tables = cfg
    tables = SimpleNamespace(**vars(tables))
    tables.og_sgr_by_week = {"2026-W42": 0.93, "2026-W44": 0.95, "2026-W45": 1.0}
    rows = {r[0]: r for r in other_rows([_b("B1")], tables, control,
                                        forecast_start=date(2026, 10, 19))}
    assert rows["Seawater growth per-week factors"][1] == "2026-W44 x0.95"
    # negative control: without a start the past week shows too
    allw = {r[0]: r for r in other_rows([_b("B1")], tables, control)}
    assert "2026-W42 x0.93" in allw["Seawater growth per-week factors"][1]
    tables.og_sgr_by_week = {"2026-W42": 0.93}
    rows = {r[0]: r for r in other_rows([_b("B1")], tables, control,
                                        forecast_start=date(2026, 10, 19))}
    assert rows["Seawater growth per-week factors"][1].startswith("none")
    assert "x0.8 - 20% slower" in {r[0]: r for r in other_rows(
        [_b("B1", fw=0.8), _b("B2", fw=1.1)], tables, control)}["Freshwater growth"][1]


@pytest.mark.skipif(not FIXTURE.exists(), reason="fixture PR missing")
def test_fixture_run_writes_the_models_sheet(tmp_path, copy_config):
    from forecast.run import main
    cfg_dir, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config(ROOT / "config", cfg_dir)
    shutil.copytree(ROOT / "scenario", scn)
    main(str(FIXTURE), str(tmp_path / "out.xlsx"), config_dir=str(cfg_dir),
         scenario_dir=str(scn), calib_log_path="", growth_history_path="")
    out = next(tmp_path.glob("out.xls*"))
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    assert "Models" in wb.sheetnames
    cells = [c for r in wb["Models"].iter_rows(values_only=True) for c in r
             if isinstance(c, str)]
    for title in ("SEAWATER GROWTH MODELS (OG) - which curve each batch grows on",
                  "SEAWATER GROWTH BY FISH SIZE - the curves above",
                  "FEED CONVERSION (FCR) MODELS - kg feed per kg growth, by fish size",
                  "OTHER BIOLOGY AND HARVEST ASSUMPTIONS"):
        assert title in cells
    # every batch the run planned is named under a growth model
    planned = {r[0] for r in wb["BatchInputs"].iter_rows(min_row=5, values_only=True)
               if r and isinstance(r[0], str) and r[0].startswith("B") and r[0][1:].isdigit()}
    named = set()
    for r in wb["Models"].iter_rows(values_only=True):
        if r and isinstance(r[2], str) and r[2].startswith("B"):
            named |= {x.strip() for x in r[2].split(",")}
    assert planned and planned <= named
    assert "1" not in [c for c in cells]                 # no bare '1' adjustment
