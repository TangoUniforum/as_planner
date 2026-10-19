"""Lifetime per-batch results when the upload is a FULL production report.

Operator 2026-10-10: "the code should recognize if it's just a PR for
forecasting or a full production report and build the accumulated results if
that's possible given the upload ... we should keep all the other
functionality too". A full report carries an accumulated tab ("Acc": FishTalk
columns, one row per fish group, lifetime to the closing date, openings 0).
With it, the run writes BatchAccumulated + BatchAccumulatedMonthly; without
it, nothing changes but one INFO line.
"""
from __future__ import annotations

import contextlib
import io
import shutil
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from conftest import copy_config_without_costs
from forecast import accumulated as acc_mod
from forecast import excel_io
from forecast.production_report import find_pr_sheet, read_pr_period

ROOT = Path(__file__).resolve().parent.parent
PR = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"
HAVE_RUN = (PR.exists() and (ROOT / "config" / "control.yaml").exists()
            and (ROOT / "scenario" / "batches.yaml").exists())

HDR = ["", "Opening Count", "Closing Count", "Opening Biomass [kg]",
       "Closing Biomass [kg]", "Opening Avg weight", "Closing Avg weight",
       "Transfer in count in period", "Transfer in biomass in period",
       "Transfer out count in period", "Transfer out biomass in period",
       "Deviation count in period", "Closing Density", "SGR in period",
       "Net growth (excl discards) [kg] in period",
       "Net growth (incl discards) [kg] in period", "Gross growth in period",
       "Gross growth % in period", "Closing Growth vs model [%]",
       "Feed amount in period", "SFR in period", "Biological FCR in period",
       "Economical FCR in period",
       "Net harvested biomass, excl. discards [kg] in period",
       "Harvested count (incl discards) in period",
       "Gross harvested biomass, incl. discards [kg] in period",
       "Net harvested biomass, incl. discards [kg] in period",
       "Harvest deviation count in period",
       "Harvest deviation biomass [kg] in period", "Mortality biomass in period",
       "Mortality average weight in period", "Mortality percent in period",
       "Mortality count in period", "Culling biomass in period",
       "Culling count in period", "Temperature for growth calculation in period"]
C = {h: i for i, h in enumerate(HDR)}


def _acc_row(batch, **v):
    r = [f"Fish group name: {batch} "] + [0] * (len(HDR) - 1)
    for k, x in v.items():
        r[C[k]] = x
    return r


# B49 exactly as the Sep'26 report's Acc tab states it.
B49 = dict(**{"Closing Count": 284361, "Closing Biomass [kg]": 162116,
              "Closing Avg weight": 570.1, "Deviation count in period": -43315,
              "SGR in period": 2.21, "Net growth (incl discards) [kg] in period": 162020,
              "Gross growth in period": 170396, "Gross growth % in period": 177794.93,
              "Feed amount in period": 155418.4, "Biological FCR in period": 0.91,
              "Economical FCR in period": 0.96, "Mortality biomass in period": 1937.9,
              "Mortality percent in period": 27.63, "Mortality count in period": 151954,
              "Culling biomass in period": 6441, "Culling count in period": 70370})


def _acc_wb(title="Acc", rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["Stock", "Density", "Growth"])
    ws.append(HDR)
    for r in rows or [_acc_row("B49", **B49)]:
        ws.append(r)
    return wb


# ---- detection -------------------------------------------------------------------

def test_the_tab_is_found_by_content_whatever_its_name():
    wb = _acc_wb("Lifetime to date")
    assert acc_mod.find_acc_sheet(wb).title == "Lifetime to date"


def test_a_pr_sheet_and_a_workbook_without_the_tab_are_not_taken():
    assert acc_mod.find_acc_sheet(openpyxl.load_workbook(PR)) is None if PR.exists() else True
    wb = _acc_wb()
    wb["Acc"].append([None, None, None, "Unit: 63 ", None, 0, 100])   # tank rows = a PR
    assert acc_mod.find_acc_sheet(wb) is None


def test_a_sheet_whose_groups_open_above_zero_is_a_period_not_a_lifetime():
    row = _acc_row("B49", **B49)
    row[C["Opening Count"]] = 290_000
    assert acc_mod.find_acc_sheet(_acc_wb(rows=[row])) is None


def test_eggs_and_stocked_biomass_come_from_the_tabs_own_balance():
    acc, missing = acc_mod.read_acc(_acc_wb()["Acc"])
    assert not missing
    b = acc["B49"]
    assert b.eggs == 550_000                       # batches.yaml: input_count 550000
    assert round(b.egg_bio, 1) == 95.8             # ~0.17 g an egg
    assert round(100 * b.mort_count / b.eggs, 2) == b.mort_pct == 27.63


def test_a_tab_from_another_report_is_refused():
    """Same closing date, different report: the 10.1.26 PR's B49 is not the
    Sep'26 report's -- the guard that stops lifetime figures being built on a
    tab that does not belong to the PR."""
    acc, _ = acc_mod.read_acc(_acc_wb()["Acc"])
    pp = SimpleNamespace(batches={"B49": SimpleNamespace(close_count=284_000.0),
                                  "B50": SimpleNamespace(close_count=252_972.0)})
    bad = acc_mod.check_against_pr(acc, pp)
    assert len(bad) == 2 and "B49" in bad[0] and "B50" in bad[1]
    ok = SimpleNamespace(batches={"B49": SimpleNamespace(close_count=284_361.0)})
    assert acc_mod.check_against_pr(acc, ok) == []


# ---- history feed split ------------------------------------------------------------

def test_history_feed_split_sums_to_one_and_follows_size():
    tables = SimpleNamespace(feed_types=[(1.0, "S"), (100.0, "M"), (10_000.0, "L")],
                             fcr_size_g=[0.1, 10_000.0], fcr_by_model={"1.18": [1.0, 1.0]})
    small = acc_mod.history_feed_shares(0.2, 50.0, tables, "1.18")
    big = acc_mod.history_feed_shares(0.2, 4_000.0, tables, "1.18")
    assert abs(sum(small.values()) - 1) < 1e-9 and abs(sum(big.values()) - 1) < 1e-9
    assert "L" not in small and big["L"] > 0.95       # weight gained drives the split
    assert acc_mod.history_feed_shares(5.0, 5.0, tables, "1.18") == {"M": 1.0}


# ---- the monthly ledger hands over the forecast's OWN months ------------------------

def test_month_rows_out_holds_the_forecast_months_before_the_pr_merge():
    from forecast.events import Harvest
    ev = Harvest(batch_id="B42", event_date=date(2026, 10, 12), source_tank_id=1,
                 count=1_000, avg_wt_g=4_000.0)
    pp = SimpleNamespace(is_mid_month=True, closing_date=date(2026, 10, 8),
                         month_label="2026-10", batches={"B42": SimpleNamespace(
                             open_count=5_000.0, open_bio_kg=20_000.0, input_count=0.0,
                             harv_count=2_000.0, harv_gross_kg=8_000.0, feed_kg=0.0,
                             mort_count=0.0, mort_bio_kg=0.0, cull_count=0.0,
                             cull_bio_kg=0.0, growth_kg=0.0)})
    out: list = []
    wb = openpyxl.Workbook()
    excel_io.write_monthly_report(wb, [], [ev], [], report_start=date(2026, 10, 9),
                                  hog_yield=0.81, pr_period=pp, month_rows_out=out)
    oct_ = [a for mo, b, a in out if mo == "2026-10" and b == "B42"]
    assert oct_ and oct_[0]["harv_count"] == 1_000        # the PR's 2,000 are NOT in it


# ---- whole runs ----------------------------------------------------------------------

def _run(tmp_path, pr_path):
    copy_config_without_costs(ROOT / "config", tmp_path / "config")   # never the live costs
    shutil.copytree(ROOT / "scenario", tmp_path / "scenario")
    shutil.rmtree(tmp_path / "scenario" / "manual_events", ignore_errors=True)
    (tmp_path / "scenario" / "manual_events.yaml").unlink(missing_ok=True)
    from forecast.run import main
    out = tmp_path / "out.xlsm"
    with contextlib.redirect_stdout(io.StringIO()):
        main(str(pr_path), str(out), config_dir=str(tmp_path / "config"),
             scenario_dir=str(tmp_path / "scenario"), calib_log_path="")
    return openpyxl.load_workbook(out, read_only=True, data_only=True)


def _info(wb):
    return [r[2] for r in wb["ValidationLog"].iter_rows(values_only=True)
            if r and len(r) > 2 and r[1] == "INFO - Accumulated"]


@pytest.mark.skipif(not HAVE_RUN, reason="needs the 9/30 PR fixture + committed config/scenario")
def test_a_pr_only_upload_builds_no_lifetime_sheet_and_says_why(tmp_path):
    wb = _run(tmp_path, PR)
    assert "BatchAccumulated" not in wb.sheetnames
    assert any("no accumulated (Acc) tab" in m for m in _info(wb))


@pytest.mark.skipif(not HAVE_RUN, reason="needs the 9/30 PR fixture + committed config/scenario")
def test_a_full_report_builds_history_forecast_and_lifetime(tmp_path):
    src = openpyxl.load_workbook(PR)
    pp = read_pr_period(find_pr_sheet(src), date(2026, 9, 30))
    ws = src.create_sheet("Acc")
    ws.append(["Stock"])
    ws.append(HDR)
    for b, pb in pp.batches.items():
        row = {"Closing Count": pb.close_count, "Feed amount in period": 1_000.0,
               "Gross growth in period": 2_000.0, "Mortality count in period": 100.0,
               "Net growth (incl discards) [kg] in period": 1_900.0}
        if b == "B49":
            row = dict(B49, **{"Closing Count": pb.close_count})
        ws.append(_acc_row(b, **row))
    b49_row = _acc_row("B49", **dict(B49, **{"Closing Count": pp.batches["B49"].close_count}))
    b49_eggs = acc_mod.read_acc(_acc_wb(rows=[b49_row])["Acc"])[0]["B49"].eggs
    full = tmp_path / "full.xlsx"
    src.save(full)
    wb = _run(tmp_path, full)
    assert "BatchAccumulated" in wb.sheetnames and "BatchAccumulatedMonthly" in wb.sheetnames
    rows = list(wb["BatchAccumulated"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Batch")
    t = [dict(zip(rows[h], r)) for r in rows[h + 1:] if r and r[0]]
    b49 = {r["Part"]: r for r in t if r["Batch"] == "B49"}
    assert b49["History"]["Mort_Count (fish)"] == 151_954
    assert b49["History"]["Mortality (% of fish in)"] == 27.63
    for col in ("Mort_Count (fish)", "Feed (kg)", "Harv_Count (fish)"):
        assert abs(b49["History"][col] + b49["Forecast"][col] - b49["Lifetime"][col]) <= 1
    assert b49["Lifetime"]["Eggs_In (fish)"] == round(b49_eggs)
    m = [r for r in wb["BatchAccumulatedMonthly"].iter_rows(values_only=True)
         if r and r[0] == "B49"]
    assert m[0][3] == "history" and m[-1][3] == "lifetime to month end"
    assert any("built from tab 'Acc'" in x for x in _info(wb))
