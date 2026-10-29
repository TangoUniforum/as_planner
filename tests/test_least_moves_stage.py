"""Least-moves layout as a pipeline step (`placement_method: least_moves`, forecast/least_moves_stage.py).

End to end on the 9/30 fixture PR over a short horizon: the step must keep the plan's endpoint exactly (harvest,
fish and biomass per batch per week, everything in 6N), leave the tank continuity audit clean, and move fewer fish.
Negative controls: a rebuilt layout that loses a transfer is refused by the per-tank biology gate (the normal plan
stands); a pinned end layout the rules cannot reach is infeasible; the guard flags a harvest from a grow-out tank.
"""
import collections
import copy
import os
import shutil
from types import SimpleNamespace

import pytest
import yaml

pytest.importorskip('ortools')
from tests.conftest import copy_config_without_costs                          # noqa: E402
from forecast import least_moves as lm                                        # noqa: E402
from forecast import least_moves_stage as st                                  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PR = os.path.join(ROOT, 'tests', 'fixtures', 'pr_2026-09-30.xlsx')
SIXN = {61, 63, 65, 67, 69, 71}
SECS = 30                     # per solver step; much less and the step (rightly) refuses a rougher layout


def _run(tmp, method, monkeypatch=None):
    d = tmp / method
    d.mkdir()
    copy_config_without_costs(os.path.join(ROOT, 'config'), d / 'config')      # never the operator's costs
    shutil.copytree(os.path.join(ROOT, 'scenario'), d / 'scenario')
    cy = d / 'config' / 'control.yaml'
    c = yaml.safe_load(open(cy)) or {}
    c.update(horizon_weeks=20, placement_method=method, least_moves_secs=SECS, least_moves_workers=8)
    yaml.safe_dump(c, open(cy, 'w'))
    from forecast.run import main
    out = d / 'out.xlsm'
    main(PR, str(out), config_dir=str(d / 'config'), scenario_dir=str(d / 'scenario'), calib_log_path='')
    return out


def _sheet(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hdr = None
    for r in wb[name].iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == 'Week':
                hdr = r
            continue
        if r and r[0] is not None and str(r[0]).startswith('20'):
            yield dict(zip(hdr, r))


def _summary(path):
    bw = collections.defaultdict(lambda: [0.0, 0.0]); six = collections.Counter(); moved = 0.0
    for d in _sheet(path, 'BatchLocations'):
        bw[(d['Week'], d['Batch'])][0] += d['Count (fish)']; bw[(d['Week'], d['Batch'])][1] += d['Biomass (kg)']
        if int(d['Tank']) in SIXN:
            six[(d['Week'], d['Tank'], d['Batch'])] += d['Count (fish)']
    for d in _sheet(path, 'TransferPlan'):
        try:
            s, t = int(d['From_Tank']), int(d['To_Tank'])
        except (TypeError, ValueError):
            continue
        if d['Type'] != 'TranOG' and s != t and s not in SIXN and t not in SIXN:
            moved += d['Count (fish)']
    harvest = sorted((d['Week'], d['Batch'], d['Tank'], d['Count (fish)']) for d in _sheet(path, 'HarvestPlan'))
    flags = sum(1 for d in _sheet(path, 'TankContinuityAudit') if d.get('Flag') or d.get('Bio_Flag'))
    import openpyxl
    log = [' '.join(str(c) for c in r if c is not None)
           for r in openpyxl.load_workbook(path, read_only=True)['ValidationLog'].iter_rows(values_only=True)]
    return dict(bw=bw, six=six, moved=moved, harvest=harvest, flags=flags, log=' | '.join(log))


@pytest.fixture(scope='module')
def runs(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('lms')
    return _summary(_run(tmp, 'greedy')), _summary(_run(tmp, 'least_moves'))


def test_same_endpoint_fewer_moves(runs):
    a, b = runs
    assert a['harvest'] == b['harvest']                                       # harvest identical
    assert a['six'].keys() == b['six'].keys()
    assert all(abs(a['six'][k] - b['six'][k]) < 1 for k in a['six'])            # everything in 6N identical
    for k in set(a['bw']) | set(b['bw']):                                     # fish + biomass per batch-week
        assert abs(a['bw'][k][0] - b['bw'][k][0]) <= 10 and abs(a['bw'][k][1] - b['bw'][k][1]) <= 10, k
    assert b['flags'] <= a['flags']                                           # the continuity audit stays clean
    assert b['moved'] < a['moved']                                            # fewer fish moved between tanks
    assert 'Least-moves layout used' in b['log']


def test_broken_rebuild_is_refused(monkeypatch, tmp_path):
    """Negative control: the re-laid layout loses its biggest transfer; the per-tank biology gate must refuse it and
    the normal plan must stand."""
    real = st._layout

    def broken(*a, **k):
        P, close = real(*a, **k)
        evs = [e for e in P.transfer_events if getattr(e, 'channel', None) == st.CHANNEL]
        P.transfer_events.remove(max(evs, key=lambda e: e.count_transferred))
        return P, close

    monkeypatch.setattr(st, '_layout', broken)
    s_ = _summary(_run(tmp_path, 'least_moves'))
    assert 'Least-moves layout NOT used' in s_['log'] and 'biology' in s_['log']
    assert s_['flags'] == 0                                                   # the normal plan stood, intact


# --------------------------------------------------------------------------------------- solver pins (unit)
def _endpoint():
    from tests.test_least_moves import _endpoint as e
    return e()


def test_pinned_end_layout_is_honoured():
    D = _endpoint()
    R = lm.solve(D, horizon=6, commit=6, secs=10, workers=4, retries=0, log=lambda s: None)
    last = {str(r['t']): [r['b'], r['n']] for r in R['plan'][-1]['rows'] if r['n'] > 0.5}
    D2 = copy.deepcopy(D); D2['end_layout'] = last
    R2 = lm.solve(D2, horizon=6, commit=6, secs=10, workers=4, retries=0, log=lambda s: None)
    got = {str(r['t']) for r in R2['plan'][-1]['rows'] if r['n'] > 0.5}
    assert got == set(last)


def test_unreachable_pin_is_infeasible():
    """Negative control: a batch that starts in grow-out cannot be pinned back into an entry tank."""
    D = _endpoint()
    for w in D['weeks']:
        D['batch']['B1'][w]['arr'] = 0.0
    D['start_layout'] = {'31': ['B1', 12000.0], '33': ['B1', 12000.0]}
    D['end_layout'] = {'11': ['B1', 16000.0]}
    with pytest.raises(RuntimeError):
        lm.solve(D, horizon=6, commit=6, secs=5, workers=4, retries=0, log=lambda s: None)


def test_one_kg_rule_uses_the_weight_when_fish_move():
    B = {'B1': {'w1': {'n': 10.0, 'wt': 0.98}, 'w2': {'n': 10.0, 'wt': 1.05}}}
    assert lm.open_wt(B, 'B1', 'w2', ['w1', 'w2']) == 0.98                   # moves happen at the week's open
    assert lm.open_wt(B, 'B1', 'w1', ['w1', 'w2']) == 0.98


def test_guard_flags_harvest_from_a_growout_tank():
    sys_of = {31: 'OG3N', 61: 'OG6N'}
    from datetime import date
    h = SimpleNamespace(source_tank_id=31, event_date=date(2028, 1, 19), count=1000, avg_wt_g=4000)
    P = SimpleNamespace(harvest_events=[h], transfer_events=[], grade_events=[], batch_locations=[])
    why = st.unsupported(P, sys_of)
    assert why and 'harvest straight from a grow-out tank' in next(iter(why.values()))
    h.source_tank_id = 61
    assert not st.unsupported(P, sys_of)
