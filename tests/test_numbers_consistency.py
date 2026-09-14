"""One number per quantity: the 2026-09-12 numbers sandbox.

Operator, 2026-09-12: "sand box every part of the code and make sure all the
numbers are consistant and correct". A 14-workbook corpus was audited and every
place that states the same quantity was compared with every other and with an
independent recomputation. Each test below pins one report/display fix. Every
one was PROVEN by a negative control: the fix reverted on a copy outside the
sandbox, the test run there, and seen to FAIL (see engine_patches/ notes and
the numbers-sandbox report). Report layer only -- nothing here touches what the
planner decides (plan sheets re-proved identical on all 11 corpus re-runs).

Test names carry the finding id (counts-01, feed-04, ...) so a negative control
can select them.
"""
from __future__ import annotations

import ast
import datetime as dt
import hashlib
import math
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from forecast import analysis as A
from forecast import excel_io
from forecast.caps import FacilityLimits
from forecast.events import Harvest, TankAllocation, Transfer, TranOGEntry
from forecast.models import BatchWeekState
from forecast.placement import BatchLocationRow

ROOT = Path(__file__).resolve().parent.parent


# ---- helpers ------------------------------------------------------------------

def _wk(d):
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def _loc(d, batch, tank, count, wt=1000.0, system="OG3N", stage="", dens=50.0):
    return BatchLocationRow(
        week_label=_wk(d), week_start=d, batch_id=batch, tank_id=tank,
        location_id=f"{system}-{tank}", system_id=system, count=count,
        avg_wt_g=wt, biomass_kg=count * wt / 1000.0, density_kg_m3=dens,
        stage=stage)


def _hv(d, batch, tank, count, wt=4000.0):
    return Harvest(batch_id=batch, event_date=d, source_tank_id=tank,
                   count=count, avg_wt_g=wt)


def _xfer(d, batch, src, dests):
    return Transfer(batch_id=batch, event_date=d, source_tank_id=src,
                    destinations=[TankAllocation(t, n, w, 10.0) for t, n, w in dests],
                    count_transferred=sum(n for _t, n, _w in dests))


def _tog(d, batch, dests):
    return TranOGEntry(batch_id=batch, event_date=d,
                       destinations=[TankAllocation(t, n, w, 16.0) for t, n, w in dests])


def _bws(batch, d, stage="SW", **kw):
    base = dict(batch_id=batch, week_label=_wk(d),
                week_start=datetime.combine(d, datetime.min.time()),
                days_since_input=100, week_from_input=14, count=1000.0,
                avg_weight_g=1000.0, biomass_kg=1000.0, feed_kg_day=0.0,
                feed_kg_week=0.0, sgr_pct_day=0.0, fcr=1.1, stage=stage,
                feed_type="T", mortality_pct_weekly=0.0)
    base.update(kw)
    return BatchWeekState(**base)


def _rows(ws):
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _table(ws, pred):
    rows = _rows(ws)
    i = next(k for k, r in enumerate(rows) if r and pred(r))
    hdr = [str(c) if c is not None else "" for c in rows[i]]
    out = []
    for r in rows[i + 1:]:
        if not r or all(c is None for c in r):
            break
        out.append(dict(zip(hdr, r)))
    return out


def _tca(wb):
    return _table(wb["TankContinuityAudit"], lambda r: r[0] == "Week" and "Tank" in r)


def _app():
    pytest.importorskip("streamlit")
    import app
    return app


def _ctl(**kw):
    """The Control fields caps.resolve_facility_cap reads."""
    base = dict(max_biomass_kg=3_800_000.0, max_feed_per_day_kg=30_000.0,
                max_harvest_per_week=55_000.0, min_harvest_per_week=26_000.0,
                default_hog_yield=0.81, max_transfers_per_week=0)
    base.update(kw)
    return SimpleNamespace(**base)


M1, M2, M3 = date(2028, 2, 7), date(2028, 2, 14), date(2028, 2, 21)   # Mondays


# ---- counts-01: TankContinuityAudit names both batches on a changeover ----------

def test_counts_01_tca_changeover_row_names_departing_and_arriving_batch():
    """Tank 31: B54 25,504 at the W06 close; in W07 B54 is harvested out whole
    and 4,704 B56 arrive. The row holds B54's opening and harvest AND B56's
    arrival and close, so its Batch cell must name both, departing first --
    it read 'B56', booking B54's 25,504-fish harvest under B56."""
    locs = [_loc(M1, "B54", 31, 25_504.0), _loc(M2, "B56", 31, 4_701.0),
            _loc(M1, "B56", 40, 4_704.0)]
    wb = openpyxl.Workbook()
    excel_io.write_tank_continuity_audit(
        wb, locs, [], [_hv(M2, "B54", 31, 25_504.0)],
        [_xfer(M2, "B56", 40, [(31, 4_704.0, 1000.0)])], [], [], None,
        realized_biology={})
    row = next(r for r in _tca(wb) if r["Tank"] == 31 and r["Week"] == _wk(M2))
    assert row["Batch"] == "B54->B56"
    assert row["Harvest_Out"] == 25_504 and row["Transfer_In"] == 4_704


# ---- counts-02: a same-week staging tank is audited ---------------------------

def test_counts_02_tca_audits_a_tank_empty_at_both_ends_of_its_week():
    """Tank 62 is empty at the W36 and W37 closes but 14,814 fish pass through
    it in W37 (in from 16, out to 46). A one-fish leak there (in 14,814, out
    14,813, close 0) must reach the facility summary -- the tank-week used to
    be skipped before its events were looked at, so the proof read 0."""
    locs = [_loc(M1, "B52", 16, 14_814.0), _loc(M2, "B52", 46, 14_813.0)]
    ev = [_xfer(M2, "B52", 16, [(62, 14_814.0, 1000.0)]),
          _xfer(M2, "B52", 62, [(46, 14_813.0, 1000.0)])]
    # Tank 16 opens the audit holding its fish, so every other tank-week
    # balances and the only drift left is the staging tank's one fish.
    init = SimpleNamespace(tanks_by_id={16: SimpleNamespace(
        batch_id="B52", count=14_814.0, biomass_kg=14_814.0, is_empty=False)})
    wb = openpyxl.Workbook()
    excel_io.write_tank_continuity_audit(wb, locs, [], [], ev, [], [], init,
                                         realized_biology={})
    row = next(r for r in _tca(wb) if r["Tank"] == 62 and r["Week"] == _wk(M2))
    assert (row["Transfer_In"], row["Transfer_Out"], row["Delta"]) == (14_814, 14_813, -1)
    assert row["Batch"] == "B52"
    summ = next(r for r in _rows(wb["TankContinuityAudit"]) if r and r[0] == "Count (fish)")
    assert summ[1] == -1, "the facility proof must see the staging tank's leak"


def test_counts_02_lns_accept_gate_sees_staging_tanks():
    """Engine change 5 of 7 (counts-02, operator-approved one at a time): the
    LNS placement accept gate (lns_placement.drift_count) audits exactly what
    the shipped TankContinuityAudit does, staging tanks included. Tank 62 here
    leaks 114 fish (in 14,814, out 14,700, close 0): the shipped sheet flags it
    TANK_DRIFT, and the gate now counts it too -- a move that leaves such a
    leak is rejected. Before this change the gate counted 0 (its pre-2026-09-12
    scope ignored a tank empty at both ends of a week)."""
    from forecast import lns_placement
    locs = [_loc(M1, "B52", 16, 14_814.0), _loc(M2, "B52", 46, 14_700.0)]
    ev = [_xfer(M2, "B52", 16, [(62, 14_814.0, 1000.0)]),
          _xfer(M2, "B52", 62, [(46, 14_700.0, 1000.0)])]
    init = SimpleNamespace(tanks_by_id={16: SimpleNamespace(
        batch_id="B52", count=14_814.0, biomass_kg=14_814.0, is_empty=False)})
    wb = openpyxl.Workbook()
    excel_io.write_tank_continuity_audit(wb, locs, [], [], ev, [], [], init,
                                         realized_biology={})
    row = next(r for r in _tca(wb) if r["Tank"] == 62 and r["Week"] == _wk(M2))
    assert row["Flag"] == "TANK_DRIFT", "the shipped sheet must flag the leak"
    assert lns_placement._count_drift_rows(wb["TankContinuityAudit"]) == 1
    placement = SimpleNamespace(batch_locations=locs, harvest_events=[],
                                transfer_events=ev, grade_events=[],
                                tranog_events=[])
    assert lns_placement.drift_count(placement, [], init) == 1, (
        "the LNS accept gate must see the staging-tank leak the shipped sheet "
        "flags (counts-02): a move that leaves one must be rejected")


# ---- counts-04 / calendar-07: TransferTemplate entry = the TranOG fish ---------

def _tt_fixture():
    # B49 split at the PR close: 47,700 already in SW (seen from W36, tank 30);
    # its FW part enters W38 as two TranOG rows of 121,128.
    w36, w37, w38 = date(2026, 8, 31), date(2026, 9, 7), date(2026, 9, 14)
    locs = [_loc(w36, "B49", 30, 47_700.0, 300.0), _loc(w37, "B49", 30, 47_650.0, 310.0),
            _loc(w38, "B49", 30, 47_600.0, 320.0), _loc(w38, "B49", 31, 121_000.0, 300.0),
            _loc(w38, "B49", 32, 121_000.0, 300.0),
            _loc(w36, "B41", 63, 1.0, 4000.0)]          # keeps W36 on the grid
    tog = [_tog(w38, "B49", [(31, 121_128.0, 300.0), (32, 121_128.0, 300.0)])]
    facility = SimpleNamespace(tanks=[SimpleNamespace(
        tank_id=t, system_id="OG1N", max_density_kg_m3=85.0, type="OG")
        for t in (30, 31, 32, 63)])
    control = SimpleNamespace(sixn_growth=False, sixn_production_start=None)
    return locs, tog, facility, control


def test_counts_04_transfer_template_entry_count_is_the_tranog_fish():
    locs, tog, facility, control = _tt_fixture()
    wb = openpyxl.Workbook()
    excel_io.write_transfer_template(wb, locs, [], tog, control, facility)
    b = next(r for r in _table(wb["TransferTemplate"], lambda r: r[0] == "Batch" and "SW_Entry_Week" in r)
             if r["Batch"] == "B49")
    assert b["Entry_Count (fish)"] == 242_256          # 121,128 x 2, not the close
    assert b["Entry_AvgWt (kg)"] == 0.3
    assert b["SW_Entry_Week"] == "2026-W38 (split: part in SW from 2026-W36)"


def test_calendar_07_batch_plan_and_transfer_template_share_one_sw_entry_week():
    locs, tog, facility, control = _tt_fixture()
    wb = openpyxl.Workbook()
    excel_io.write_transfer_template(wb, locs, [], tog, control, facility)
    excel_io.write_batch_plan(wb, locs, [], tranog_events=tog)
    tt = {r["Batch"]: r["SW_Entry_Week"] for r in _table(
        wb["TransferTemplate"], lambda r: r[0] == "Batch" and "SW_Entry_Week" in r)}
    bp = {r["Batch"]: r["SW_entry"] for r in _table(
        wb["Batch Plan"], lambda r: r[0] == "Batch" and "Peak_tanks" in r)}
    assert bp["B49"] == tt["B49"] == "2026-W38 (split: part in SW from 2026-W36)"
    assert bp["B41"] == tt["B41"] == "2026-W36 (in-flight)"


# ---- counts-05 / harvest: a batch harvested out in week 1 keeps its row ---------

def test_counts_05_batch_plan_and_template_carry_a_batch_with_no_tank_week():
    """B41 is harvested out whole before the first closing snapshot, so it has
    no BatchLocations row. Both per-batch sheets dropped it (5,322 fish, 14 t
    HOG on the 8/31 PR) while HarvestPlan carried it."""
    w36 = date(2026, 8, 31)
    locs = [_loc(w36, "B42", 12, 50_000.0, 3500.0)]
    hv = [_hv(w36, "B41", 63, 5_322.0, 3252.0)]
    wb = openpyxl.Workbook()
    excel_io.write_batch_plan(wb, locs, hv, default_hog_yield=0.81, tranog_events=[])
    facility = SimpleNamespace(tanks=[SimpleNamespace(
        tank_id=12, system_id="OG5N", max_density_kg_m3=85.0, type="OG")])
    excel_io.write_transfer_template(wb, locs, hv, [], SimpleNamespace(
        sixn_growth=False, sixn_production_start=None), facility)
    bp = {r["Batch"]: r for r in _table(wb["Batch Plan"], lambda r: r[0] == "Batch" and "Peak_tanks" in r)}
    assert "B41" in bp and bp["B41"]["HOG_tonnes"] == round(5_322 * 3.252 * 0.81 / 1000.0)
    tt = {r["Batch"] for r in _table(wb["TransferTemplate"], lambda r: r[0] == "Batch" and "SW_Entry_Week" in r)}
    assert "B41" in tt


def test_harvest_batch_plan_harvest_weight_is_fish_weighted():
    """1,000 fish at 5.0 kg and 9,000 at 3.0 kg weigh 3.2 kg on average, not
    4.0 -- the Batch Plan milestone averaged the events, TransferTemplate
    weighted them."""
    w = date(2026, 10, 5)
    hv = [_hv(w, "B46", 60, 1_000.0, 5000.0), _hv(w, "B46", 61, 9_000.0, 3000.0)]
    wb = openpyxl.Workbook()
    excel_io.write_batch_plan(wb, [_loc(w, "B46", 60, 1.0)], hv)
    ms = _rows(wb["Batch Plan"])
    got = [r for r in ms if r and r[2] == "Harvest"]
    assert got and got[0][4] == 3.2


# ---- counts-06 / counts-08 / biomass-06: period totals = the rows under them ----

def _pr_period():
    return SimpleNamespace(is_mid_month=True, closing_date=date(2026, 9, 10),
                           month_label="2026-09",
                           batches={"B41": SimpleNamespace(harv_count=6_164.0,
                                                           harv_gross_kg=18_782.0)})


def _year_events():
    # Fractional fish, as the engine carries them. W36/W37 split whole fish per
    # WEEK: two 1000.5 events in one week show as 1001 + 1000.
    d1, d2 = date(2026, 9, 14), date(2026, 9, 21)
    d3, d4 = date(2027, 3, 1), date(2027, 3, 8)
    return [_hv(d1, "B42", 1, 1000.5), _hv(d1, "B43", 2, 1000.5),
            _hv(d2, "B42", 1, 10.6), _hv(d3, "B44", 3, 10.6), _hv(d4, "B44", 4, 10.6)]


def test_counts_08_harvest_plan_report_total_is_the_sum_of_its_rows_and_of_harvestplan():
    evs = _year_events()
    wb = openpyxl.Workbook()
    excel_io.write_harvest_plan_output(wb, evs, 0.81, {})
    excel_io.write_harvest_plan_report(wb, evs, "S", 0.81, {},
                                       report_start=date(2026, 9, 1))
    hp = _table(wb["HarvestPlan"], lambda r: r[0] == "Week" and "Batch" in r)
    shown = sum(r["Count (fish)"] for r in hp if r["Week"] in ("2026-W38", "2026-W39"))
    rows = _rows(wb["HarvestPlan Report"])
    i = next(k for k, r in enumerate(rows) if r and r[0] == "S 2026")
    block = []
    for r in rows[i + 2:]:
        if not r or r[0] is None and r[1] is None:
            break
        block.append(r)
    sep_col = 2 + 8                         # September
    batch_units = sum(r[sep_col] for r in block if r[1] == "Units" and r[0] != "TOTAL"
                      and isinstance(r[sep_col], (int, float)))
    total = next(r for r in block if r[0] == "TOTAL" and r[1] == "Units")[sep_col]
    assert total == batch_units == shown


def test_counts_08_ledger_harvest_totals_are_the_harvest_sheets_whole_fish():
    """Two batches harvest 1,000.5 fish each in one week. HarvestPlan shows
    1,001 + 1,000 (whole fish tied to the week); the ledgers rounded each
    batch's fractional sum (1,000 + 1,000 under round-half-even), so the
    WeeklyReport TOTAL missed the HarvestPlan week and the MonthlyReport TOTAL
    missed HarvestPlan Report's month (up to 3 fish on the corpus)."""
    d = date(2026, 9, 14)
    locs = [_loc(d, "B42", 1, 5_000.0), _loc(d, "B43", 2, 5_000.0)]
    evs = [_hv(d, "B42", 1, 1000.5), _hv(d, "B43", 2, 1000.5)]
    wb = openpyxl.Workbook()
    excel_io.write_harvest_plan_output(wb, evs, 0.81, {})
    excel_io.write_harvest_plan_report(wb, evs, "S", 0.81, {},
                                       report_start=date(2026, 9, 1))
    excel_io.write_weekly_report(wb, locs, evs, [])
    excel_io.write_monthly_report(wb, locs, evs, [], report_start=date(2026, 9, 1))
    hp = sum(r["Count (fish)"] for r in _table(
        wb["HarvestPlan"], lambda r: r[0] == "Week" and "Batch" in r))
    wk_tot = next(x for x in _table(wb["WeeklyReport"], lambda r: r[0] == "Scenario")
                  if x["Batch"] == "TOTAL")
    mo_tot = next(x for x in _table(wb["MonthlyReport"], lambda r: r[0] == "Scenario")
                  if x["Batch"] == "TOTAL")
    hpr_tot = next(r for r in _rows(wb["HarvestPlan Report"])
                   if r and r[0] == "TOTAL" and r[1] == "Units")[2 + 8]
    assert hp == 2_001
    assert wk_tot["Harv_Count (fish)"] == mo_tot["Harv_Count (fish)"] == hpr_tot == 2_001


def test_counts_06_yearly_summary_harvest_ties_to_harvestreport_and_hpr():
    """YearlySummary = HarvestReport rows by year, and on a mid-month PR it adds
    the PR's elapsed harvest to the closing year exactly as HarvestPlan Report
    does (the 9.10 PR's 2026 read 863,457 here against 921,265 there)."""
    evs = _year_events()
    pp = _pr_period()
    wb = openpyxl.Workbook()
    excel_io.write_harvest_report(wb, evs, 0.81, {}, report_start=date(2026, 9, 11))
    excel_io.write_harvest_plan_report(wb, evs, "S", 0.81, {}, pr_period=pp,
                                       report_start=date(2026, 9, 11))
    excel_io.write_yearly_summary(wb, [], evs, FacilityLimits(),
                                  _ctl(), default_hog_yield=0.81, pr_period=pp,
                                  report_start=date(2026, 9, 11))
    hr = _table(wb["HarvestReport"], lambda r: r[0] == "Year" and "Batch" in r)
    by_y = {}
    for r in hr:
        by_y[r["Year"]] = by_y.get(r["Year"], 0) + r["Count (fish)"]
    ys = {r["Year"]: r["Harvest_Count (fish)"] for r in _table(
        wb["YearlySummary"], lambda r: r[0] == "Year")}
    hpr_tot = {int(str(r[0])[-4:]): None for r in _rows(wb["HarvestPlan Report"])
               if r and isinstance(r[0], str) and r[0].startswith("S ")}
    rows = _rows(wb["HarvestPlan Report"])
    y = None
    for r in rows:
        if r and isinstance(r[0], str) and r[0].startswith("S "):
            y = int(r[0][-4:])
        if r and r[0] == "TOTAL" and r[1] == "Units":
            hpr_tot[y] = r[14]
    assert ys[2026] == by_y[2026] + 6_164 == hpr_tot[2026]
    assert ys[2027] == by_y[2027] == hpr_tot[2027] == 22     # 11 + 11, not round(21.2)


# ---- counts-07: RR fallback charges no mortality on pre-biology harvest -------

def test_counts_07_reconciliation_fallback_does_not_kill_harvested_fish():
    w36 = date(2026, 8, 31)
    tank = SimpleNamespace(batch_id="B41", count=5_322.0, biomass_kg=17_300.0,
                           is_empty=False)
    init = SimpleNamespace(tanks_by_id={63: tank})
    wb = openpyxl.Workbook()
    excel_io.write_reconciliation_report(
        wb, [_loc(w36, "B42", 12, 50_000.0)],
        [_bws("B41", w36, mortality_pct_weekly=0.05)],
        [_hv(w36, "B41", 63, 5_322.0, 3252.0)], [], init, realized_biology={})
    rr = _table(wb["ReconciliationReport"], lambda r: r[0] == "Week" and "Batch" in r)
    b41 = next(r for r in rr if r["Batch"] == "B41")
    dcol = next(k for k in b41 if k.startswith("Count_Delta"))
    assert b41[dcol] == 0


# ---- counts-09: a PR batch with no metadata opens on its PR balance -------------

def test_counts_09_ledger_opens_a_no_metadata_batch_on_its_pr_balance():
    w = date(2025, 7, 28)
    locs = [_loc(w, "B34", 12, 112_168.0, 2800.0)]
    hv = [_hv(w, "B34", 63, 17_744.0, 3000.0)]
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, locs, hv, [],
                                 pr_openings={"B34": (129_912.0, 129_912 * 2.8)})
    r = next(x for x in _table(wb["WeeklyReport"], lambda r: r[0] == "Scenario")
             if x["Batch"] == "B34")
    assert r["Open_Count (fish)"] == 129_912
    assert r["Count_Check (fish)"] == 0


# ---- biomass-03: the MonthlyReport density legend says what the column is -----

def test_biomass_03_monthly_density_legend_and_value_agree():
    locs = [_loc(M1, "B1", 1, 1000.0, 1000.0, dens=90.0),
            _loc(M1, "B1", 2, 3000.0, 1000.0, dens=30.0)]
    wb = openpyxl.Workbook()
    excel_io.write_monthly_report(wb, locs, [], [])
    legend = " ".join(str(r[0]) for r in _rows(wb["MonthlyReport"])[:3] if r and r[0])
    assert "max of the month" not in legend
    assert "total biomass over the total water" in legend
    row = next(x for x in _table(wb["MonthlyReport"], lambda r: r[0] == "Scenario")
               if x["Batch"] == "B1")
    water = 1000.0 / 90.0 + 3000.0 / 30.0
    assert row["Avg_Density (kg/m³)"] == round(4000.0 / water, 1)


# ---- biomass-05: FacilityMap / TransferTemplate judge prep tanks at 150 ------

def test_biomass_05_facility_map_and_template_judge_harvest_prep_at_150():
    w = date(2025, 4, 7)
    locs = [_loc(w, "B38", 61, 10_000.0, 4000.0, system="OG6N", stage="STARVE", dens=185.0),
            _loc(w, "B38", 62, 10_000.0, 4000.0, system="OG6N", stage="STARVE", dens=140.0),
            _loc(w, "B39", 40, 10_000.0, 2000.0, system="OG3N", dens=90.0)]
    facility = SimpleNamespace(tanks=[SimpleNamespace(
        tank_id=t, system_id=s, max_density_kg_m3=85.0, type="OG")
        for t, s in ((61, "OG6N"), (62, "OG6N"), (40, "OG3N"))])
    control = SimpleNamespace(sixn_growth=False, sixn_production_start=None)
    wb = openpyxl.Workbook()
    excel_io.write_facility_map(wb, locs, facility, control=control)
    ws = wb["FacilityMap"]
    red = set()
    for row in ws.iter_rows():
        for c in row:
            f = c.font
            if f is not None and f.bold and f.color is not None and str(f.color.rgb).endswith("9C0006"):
                red.add(ws.cell(row=c.row, column=1).value)
    assert red == {61, 40}                  # 185 > 150 and 90 > 85; 140 <= 150 is legal
    excel_io.write_transfer_template(wb, locs, [], [], control, facility)
    tt = {r["Batch"]: r for r in _table(wb["TransferTemplate"],
                                        lambda r: r[0] == "Batch" and "SW_Entry_Week" in r)}
    assert tt["B38"]["Peak_Density (×cap)"] == round(185.0 / 150.0, 2)
    assert tt["B38"]["Density_Status"] == "OVER CAP"


# ---- biomass-07: a TCA row re-adds from its own printed columns ---------------

def test_biomass_07_tca_prints_negative_growth_so_the_row_re_adds():
    locs = [_loc(M1, "B1", 5, 1000.0, 1000.0), _loc(M2, "B1", 5, 1000.0, 950.0)]
    wb = openpyxl.Workbook()
    excel_io.write_tank_continuity_audit(
        wb, locs, [], [], [], [], [], None,
        realized_biology={(5, _wk(M2), "B1"): (-50.0, 0.0)})
    r = next(x for x in _tca(wb) if x["Tank"] == 5 and x["Week"] == _wk(M2))
    n = lambda k: r.get(k) or 0.0
    readd = (n("Open_Bio_kg") + n("Growth_kg") - n("Mort_kg") - n("Harvest_Out_kg")
             - n("Transfer_Out_kg") + n("Transfer_In_kg") - n("Grade_Out_kg")
             + n("Grade_In_kg") + n("TranOG_In_kg"))
    assert r["Growth_kg"] == -50
    assert readd == r["Expected_Close_kg"] == 950


def test_biomass_07_tca_prints_a_negative_mortality_so_the_row_re_adds():
    """The recorded mortality of a tank-week can be NEGATIVE (c 2026-W49 tank
    32: -19 kg). Mort_kg printed only when > 0, so that row stayed 19 kg short
    of re-adding even after Growth_kg printed its negative."""
    locs = [_loc(M1, "B1", 5, 1000.0, 1000.0), _loc(M2, "B1", 5, 1000.0, 950.0)]
    wb = openpyxl.Workbook()
    excel_io.write_tank_continuity_audit(
        wb, locs, [], [], [], [], [], None,
        realized_biology={(5, _wk(M2), "B1"): (-50.0, -20.0)})
    r = next(x for x in _tca(wb) if x["Tank"] == 5 and x["Week"] == _wk(M2))
    n = lambda k: r.get(k) or 0.0
    readd = (n("Open_Bio_kg") + n("Growth_kg") - n("Mort_kg") - n("Harvest_Out_kg")
             - n("Transfer_Out_kg") + n("Transfer_In_kg") - n("Grade_Out_kg")
             + n("Grade_In_kg") + n("TranOG_In_kg"))
    assert r["Mort_kg"] == -20
    assert readd == r["Expected_Close_kg"]


# ---- feed-02: week-0 opening before the first day's growth; FCR denominator ----

def test_feed_02_week0_opens_on_the_weight_before_the_first_days_growth():
    w = date(2026, 9, 11)                                     # a Friday start
    s = _bws("B44", w, open_count=1000.0, open_avg_weight_g=3231.2,
             open_biomass_kg=3231.2, open_avg_weight_pre_g=3216.1,
             close_count=999.0, close_avg_weight_g=3290.0, close_biomass_kg=3286.7)
    locs = [BatchLocationRow(week_label=_wk(w), week_start=w, batch_id="B44",
                             tank_id=5, location_id="OG5N-5", system_id="OG5N",
                             count=999.0, avg_wt_g=3290.0, biomass_kg=3286.7,
                             density_kg_m3=60.0)]
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, locs, [], [s])
    r = next(x for x in _table(wb["WeeklyReport"], lambda r: r[0] == "Scenario")
             if x["Batch"] == "B44")
    assert r["Open_AvgWt (g)"] == 3216.1
    assert r["Open_Bio (kg)"] == round(1000 * 3.2161)


def test_feed_02_split_batch_fw_part_opens_before_its_first_days_growth():
    """An automatically modelled split (B49 on the 8/31 PR without events):
    its FW track's week-0 opening is the PR's fish at their weight BEFORE the
    first day's growth -- it opened +1,564 kg heavy."""
    w = date(2026, 9, 1)                                     # a Tuesday start
    locs = [_loc(w, "B49", 14, 47_000.0, 300.0)]
    fw = _bws("B49", w, stage="FW", open_count=250_000.0, open_avg_weight_g=250.8,
              open_avg_weight_pre_g=250.0, open_biomass_kg=62_700.0,
              close_count=249_900.0, close_avg_weight_g=256.0,
              close_biomass_kg=63_974.0, feed_kg_week=5_000.0)
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, locs, [], [], split_fw={"B49": [fw]},
                                 fw_projected={"B49"})
    r = next(x for x in _table(wb["WeeklyReport"], lambda r: r[0] == "Scenario")
             if x["Batch"] == "B49")
    assert r["Open_Count (fish)"] == 297_000
    assert r["Open_Bio (kg)"] == round(47_000 * 0.3 + 250_000 * 0.25)


def test_feed_02_no_fcr_over_a_growth_that_prints_as_zero():
    d = dict(open_count=1000.0, open_bio=100.0, harv_count=0.0, cull_count=0.0,
             input_count=0.0, sfr=1.0, gross_growth=1e-12, net_prod=1e-12,
             feed=196.0, bio_fcr=196.0 / 1e-12, econ_fcr=196.0 / 1e-12,
             open_wt=100.0, close_count=1000.0, close_wt=100.0, close_bio=100.0,
             sgr=0.0, mort_count=0.0, mort_bio=0.0, harv_gross=0.0, harv_hog=0.0,
             harv_avg_hog=0.0, cull_bio=0.0, xfer_in=0.0, xfer_out=0.0,
             count_check=0.0, bio_check=0.0)
    cells = excel_io._ledger_value_cells(d)
    assert cells[12] is None and cells[13] is None          # Bio_FCR, Econ_FCR


# ---- feed-03 / money-02 / calendar-04: a year's feed by calendar day -----------

def test_feed_03_yearly_feed_splits_a_year_boundary_week_by_calendar_day():
    """7,000 kg of FW feed in the week of Mon 2026-12-28: 4 days in 2026, 3 in
    2027. YearlySummary booked all 7 t to 2026 (the week's Monday); the
    CostsAndProfit / FeedForecastMonthly roll-up books 4 + 3."""
    ws_ = date(2026, 12, 28)
    s = _bws("B60", ws_, stage="FW", feed_kg_week=7_000.0, biomass_kg=10.0)
    wb = openpyxl.Workbook()
    excel_io.write_yearly_summary(wb, [], [], FacilityLimits(),
                                  _ctl(), biology_states_by_batch={"B60": [s]})
    ys = {r["Year"]: r["Feed (t)"] for r in _table(wb["YearlySummary"], lambda r: r[0] == "Year")}
    assert (ys[2026], ys[2027]) == (4, 3)
    ftw, wks = excel_io._feed_by_type_week([], {"B60": [s]}, None)
    ftm, _m = excel_io._feed_by_type_month(ftw, wks)
    cp = {}
    for (_n, mo), kg in ftm.items():
        cp[int(mo[:4])] = cp.get(int(mo[:4]), 0.0) + kg
    assert (round(cp[2026] / 1000.0), round(cp[2027] / 1000.0)) == (ys[2026], ys[2027])


# ---- feed-04: FacilityMap feed FACILITY = Advisory Total_Feed ------------------

def test_feed_04_facility_map_feed_total_includes_hatchery_feed():
    w = date(2026, 9, 7)
    s = _bws("B55", w, stage="FW", feed_kg_day=1_800.0, feed_kg_week=12_600.0,
             biomass_kg=20_000.0)
    locs = [_loc(w, "B44", 5, 1000.0)]
    facility = SimpleNamespace(tanks=[SimpleNamespace(
        tank_id=5, system_id="OG5N", max_density_kg_m3=85.0, type="OG")])
    wb = openpyxl.Workbook()
    excel_io.write_facility_map(wb, locs, facility, biology_states_by_batch={"B55": [s]})
    excel_io.write_advisory(wb, locs, [], FacilityLimits(), _ctl(),
                            biology_states_by_batch={"B55": [s]})
    rows = _rows(wb["FacilityMap"])
    i = next(k for k, r in enumerate(rows) if r and str(r[0]).startswith("TOTAL PLANNED FEED"))
    fac = next(r for r in rows[i:] if r and r[0] == "FACILITY")
    adv = _table(wb["Advisory"], lambda r: r[0] == "Week" and "Week_Start" in r)[0]
    assert fac[2] == round(adv["Total_Feed (kg/day)"]) == 1_800


# ---- feed-06: freshwater feed billed to each day's own size band ---------------

def test_feed_06_fw_feed_by_type_uses_each_days_band():
    w = date(2026, 10, 5)
    s = _bws("B51", w, stage="FW", feed_kg_week=700.0, feed_type="Big",
             feed_kg_by_type={"Small": 300.0, "Big": 400.0})
    ftw, _w = excel_io._feed_by_type_week([], {"B51": [s]}, None)
    assert ftw[("Small", _wk(w))] == 300.0 and ftw[("Big", _wk(w))] == 400.0
    fbtw, _w2 = excel_io._feed_by_batch_type_week([], {"B51": [s]}, None)
    assert fbtw[("B51", "Small", _wk(w))] == 300.0


def test_feed_06_biology_types_each_day_by_its_own_weight():
    from forecast.biology import _feed_by_type_days
    days = [(None,) * 9 + (100.0, "Small"), (None,) * 9 + (120.0, "Small"),
            (None,) * 9 + (150.0, "Big"), (None,) * 9 + (0.0, "")]
    assert _feed_by_type_days(days) == {"Small": 220.0, "Big": 150.0}


# ---- harvest: a week with NO harvest is named in the floor audit ---------------

def test_harvest_zero_week_is_named_by_the_realized_floor_audit():
    evs = [_hv(date(2026, 3, 2), "B1", 1, 20_000.0), _hv(date(2026, 3, 16), "B1", 1, 30_000.0)]
    control = _ctl()
    out = A.realized_plan_audit(evs, [], FacilityLimits(), control,
                                plan_weeks=["2026-W10", "2026-W11", "2026-W12"])
    assert any("2026-W11" in m and "NO HARVEST" in m for m in out)
    assert any("2026-W10" in m for m in out)                 # 20,000 < 26,000


# ---- the first-week month rule on the app side ------------------------------

def test_harvest_month_of_first_week_is_the_reports_first_month():
    rs = date(2026, 9, 1)                  # 8/31 PR: W36's Monday is 2026-08-31
    assert A.week_to_month("2026-W36") == "2026-08"            # plain rule kept
    assert A.week_to_month("2026-W36", rs) == "2026-09"
    rows = [{"week": "2026-W36", "hog_kg": 72_069.0, "gross_kg": 0.0},
            {"week": "2026-W37", "hog_kg": 100.0, "gross_kg": 0.0}]
    mo, _y = A.harvest_by_period(rows, clip_start=rs)
    assert mo == {"2026-09": 72_169.0}
    from forecast import costs_report as R
    cp = R.harvest_by_month(rows, rs)
    assert set(cp) == set(mo)
    full, _fy = A.full_periods([f"2026-W{w}" for w in range(36, 41)], rs)
    assert "2026-09" in full and "2026-08" not in full
    from forecast import harvest_plan as hp
    assert hp.month_of("2026-W36", rs) == "2026-09"


# ---- the daily schedule never schedules a day before the report opens ----------

def test_harvest_daily_schedule_has_no_day_before_a_sunday_report_start():
    from forecast.time_grid import harvest_schedule_days
    rs = date(2024, 12, 1)                                    # a Sunday
    assert harvest_schedule_days(rs, rs) == [rs]
    assert harvest_schedule_days(date(2026, 9, 1), date(2026, 9, 1)) == [
        date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3), date(2026, 9, 4)]
    wb = openpyxl.Workbook()
    excel_io.write_daily_harvest_schedule(wb, [_hv(rs, "B30", 61, 19_433.0)],
                                          rs, 0.81, {}, report_start=rs)
    rows = [r for r in _rows(wb["Daily Harvest Schedule"])
            if r and isinstance(r[0], int)]
    days = [r for r in rows if isinstance(r[2], (date, datetime))]
    assert all((d[2].date() if isinstance(d[2], datetime) else d[2]) >= rs for d in days)
    total = next(r for r in rows if r[2] == "Total")
    assert sum(d[5] for d in days) == total[5] == 19_433


def test_display_daily_table_matches_the_sheet():
    import pandas as pd
    rs = date(2026, 9, 1)
    he = pd.DataFrame({"Week": ["2026-W36"], "Batch": ["B41"], "Tank": [63],
                       "Count": [23_387.0], "Gross_kg": [70_000.0], "HOG_kg": [56_700.0]})
    df, _tot, _blank = _app()._daily_harvest_table(he, rs)
    dates = [d for d in df["Date"] if d and d != "Total"]
    assert dates == ["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"]
    counts = [int(str(c).replace(",", "")) for c, d in zip(df["Count"], df["Date"])
              if d and d != "Total"]
    assert sum(counts) == 23_387


# ---- money-04: the fingerprint is labelled what it is ------------------------

def test_money_04_costs_note_labels_the_md5_and_old_notes_still_parse(tmp_path):
    from forecast import costs as C
    from forecast import costs_report as R
    (tmp_path / "costs.yaml").write_text("schema: 1\n", encoding="utf-8")
    sig = C.costs_sig(tmp_path)[:8]
    assert sig == hashlib.md5((tmp_path / "costs.yaml").read_bytes()).hexdigest()[:8]
    drivers = dict(start=date(2026, 9, 1), end=date(2026, 9, 30), months=["2026-09"],
                   feed={"2026-09": {}}, harvest={}, eggs={},
                   month_days={"2026-09": (30, 30)}, pre_start=set())
    costs = {"schema": 1, "fixed_monthly": 1000.0, "oxygen_per_kg_feed": 0.0,
             "chemicals_per_kg_feed": 0.0, "feed_shipping_per_kg": 0.0,
             "egg_price": 0.0, "feed_prices": {}}
    rows, _l = R.build_costs_rows(drivers, costs, None, feed_types=[], sha=sig,
                                  revenue_note="")
    assert f"(md5 {sig})" in rows[1][0] and "(sha " not in rows[1][0]
    for label in ("md5", "sha"):
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in rows:
            ws.append([c if not (isinstance(c, str) and "(md5 " in c) else
                       c.replace("(md5 ", f"({label} ") for c in r])
        assert R.read_costs_sheet(ws)["sha"] == sig


# ---- money-03 / biomass-04 / calendar-06: the part month is labelled ----------

def test_money_03_costs_first_month_of_a_mid_month_report_is_labelled():
    from forecast import costs_report as R
    costs = {"schema": 1, "fixed_monthly": 1000.0, "oxygen_per_kg_feed": 0.0,
             "chemicals_per_kg_feed": 0.0, "feed_shipping_per_kg": 0.0,
             "egg_price": 0.0, "feed_prices": {}}
    for start, expect in ((date(2026, 9, 11), True), (date(2026, 9, 1), False)):
        drivers = dict(start=start, end=date(2026, 10, 31), months=["2026-09", "2026-10"],
                       feed={}, harvest={}, eggs={},
                       month_days={"2026-09": (30 - start.day + 1, 30), "2026-10": (31, 31)},
                       pre_start=set())
        rows, _l = R.build_costs_rows(drivers, costs, None, feed_types=[],
                                      sha="0123abcd", revenue_note="")
        assert ("forecast days only" in rows[1][0]) is expect
        kind = rows[4][1]
        assert kind.startswith("Month")
        assert (kind != "Month") is expect


# ---- money-06: the Ideal year's fixed cost starts on the forecast's first day --

def test_money_06_ideal_first_year_fixed_cost_starts_at_the_forecast():
    from forecast import costs as C
    from forecast.ideal_engine import _first_day
    y = SimpleNamespace(year=2026, weeks=18, first_week="2026-W36",
                        first_day=_first_day("2026-W36", date(2026, 9, 1)))
    assert y.first_day == date(2026, 9, 1)
    got = C.year_weeks_fixed_months(y)
    assert math.isclose(got, C.fixed_months_between(date(2026, 9, 1), date(2027, 1, 3)))
    assert math.isclose(C.iso_weeks_fixed_months("2026-W36", 18) - got, 1 / 31)
    assert _first_day("2027-W01", date(2027, 1, 4)) is None      # a Monday start


# ---- display: board badge, 6N spells, conservation label, target line ----------

def _advisory_wb(tmp_path, series):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Advisory"
    ws.append(["CAPACITY ADVISORY"])
    ws.append([])
    ws.append(["Week", "Week_Start", "Total_Biomass (kg)", "Biomass_Limit (kg)",
               "Biomass_Excess (kg)", "Total_Feed (kg/day)", "Feed_Limit (kg/day)",
               "Feed_Excess (kg/day)", "Harvest_Count", "Harvest_Biomass (kg)",
               "Advisory"])
    for wk, bio, cap, feed, fcap in series:
        ws.append([wk, None, bio, cap, max(0, bio - cap), feed, fcap,
                   max(0, feed - fcap), 30_000, 100_000, "OK"])
    p = tmp_path / "adv.xlsx"
    wb.save(p)
    return p


def test_display_board_under_cap_judges_each_week_against_its_own_cap(tmp_path):
    """Peak 3.82M sits under 1.02 x the first week's 3.80M cap, but the week at
    3.82M had a 3.65M cap: 104.7%. The badge passed it (2026-02-28 PR)."""
    p = _advisory_wb(tmp_path, [("2026-W09", 3_700_000, 3_800_000, 0, 0),
                                ("2026-W10", 3_820_000, 3_650_000, 0, 0)])
    m = SimpleNamespace(overall_peak_biomass=3_820_000, biomass_cap=3_800_000)
    assert _app()._board_under_cap(str(p), m, 1.02) is False


def test_display_feed_tab_reads_the_advisory(tmp_path):
    p = _advisory_wb(tmp_path, [("2026-W40", 1, 1, 27_000.0, 27_500.0),
                                ("2026-W41", 1, 1, 28_431.0, 27_500.0)])
    df, _note = _app()._feed_weekly(str(p))
    assert df["Feed (kg/day)"].tolist() == [27_000, 28_431]
    assert df["Over"].tolist() == ["", "OVER"]


def test_display_sixn_spell_crossing_a_52_week_year_is_one_spell(tmp_path):
    rows = []
    for d, n in ((date(2027, 12, 20), 5_000.0), (date(2027, 12, 27), 4_990.0),
                 (date(2028, 1, 3), 4_980.0)):
        rows.append(_loc(d, "B43", 65, n, 4000.0, system="OG6N", stage="STARVE"))
    wb = openpyxl.Workbook()
    excel_io.write_batch_locations(wb, rows)
    p = tmp_path / "bl.xlsx"
    wb.save(p)
    r = A.sixn_trapped_review(str(p))
    assert r["spells"] == 1


def test_display_conservation_gate_says_batches_not_fish():
    st, txt = A._gate_conservation({"dropped": 0, "overprod": 1})
    assert st == "FAIL" and "batch(es) over-produced" in txt
    assert "over-produced fish" not in txt
    assert A._gate_conservation({"dropped": 0, "overprod": 0}) == (
        "PASS", "0 dropped / 0 over-produced")


def _app_src():
    return (ROOT / "app.py").read_text(encoding="utf-8")


def test_display_per_system_target_line_reads_the_planners_target():
    src = _app_src()
    assert "85% target" not in src
    assert "add_hline(y=85," not in src


def test_display_app_months_use_the_reports_first_week_rule():
    """No second copy of the plain Monday month in app.py: the Harvest tab and
    the target panels go through analysis.week_to_month(..., clip_start)."""
    src = _app_src()
    assert 'fromisocalendar(y, w, 1).strftime("%Y-%m")' not in src
    tree = ast.parse(src)
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and getattr(n.func, "attr", "") == "harvest_by_period"]
    assert calls and all(any(k.arg == "clip_start" for k in c.keywords) for c in calls)


def test_display_ideal_page_states_its_iso_year_basis():
    src = _app_src()
    assert src.count("st.caption(_IDEAL_YEAR_BASIS_NOTE)") == 2
    note = _app().__dict__["_IDEAL_YEAR_BASIS_NOTE"]
    assert "ISO weeks" in note and "calendar years" in note


def test_display_batch_plans_view_mirrors_the_sheet():
    import pandas as pd
    bl = pd.DataFrame({"Batch": ["B42"], "System": ["OG5N"], "Week": ["2026-W36"],
                       "Tank": [12], "AvgWt_kg": [3.5]})
    he = pd.DataFrame({"Week": ["2026-W36", "2026-W36"], "Batch": ["B41", "B41"],
                       "Tank": [63, 64], "Count": [1_000.0, 9_000.0],
                       "Gross_kg": [5_000.0, 27_000.0], "HOG_kg": [4_050.0, 21_870.0],
                       "Avg_wt_kg": [5.0, 3.0]})
    plans = {p["Batch"]: p for p in _app()._derive_batch_plans(bl, he)}
    assert "B41" in plans
    assert plans["B41"]["milestones"][-1]["AvgWt (kg)"] == 3.2


# ---- counts-03: PR fish that no tank receives are named ----------------------

def test_counts_03_pr_structure_warnings_name_unhydrated_fish():
    from forecast.production_report import pr_structure_warnings
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Closing Month: 11/30/2024 ", None, None, None, None, 0, 1_000_000])
    ws.append([None, None, "Fish group name: 35A ", None, None, 0, 83_698])
    ws.append([None, None, None, "Unit: 26 ", None, 0, 83_698])
    ws.append([None, None, "Fish group name: B40 ", None, None, 0, 413_064])
    ws.append([None, None, None, "Unit: 30 ", None, 0, 413_064])
    ws.append([None, None, "Fish group name: B41 ", None, None, 0, 449_643])
    got = pr_structure_warnings(ws)
    assert any("'35A'" in g and "83,698" in g for g in got)
    assert any(g.startswith("PR NOT HYDRATED - B41") and "449,643" in g for g in got)
    assert any("facility" in g and "413,064" in g for g in got)
    assert not any("B40:" in g for g in got)


# ---- calendar-03: a freshwater row's rates use its own days ------------------

def test_calendar_03_fw_row_in_a_one_day_first_week_uses_one_day():
    """d1 2024-W48 is the single day Sunday 2024-12-01. A freshwater projection
    row grows 1 day there; its weekly and monthly SGR must divide by 1, not 7
    (the month read 37 days in a 31-day December)."""
    w = date(2024, 12, 1)
    s = _bws("B37", w, stage="FW", open_count=100_000.0, open_avg_weight_g=111.0,
             open_avg_weight_pre_g=110.2, open_biomass_kg=11_100.0,
             close_count=99_990.0, close_avg_weight_g=111.0, close_biomass_kg=11_098.9,
             feed_kg_week=529.0)
    wb = openpyxl.Workbook()
    excel_io.write_weekly_report(wb, [], [], [s])
    excel_io.write_monthly_report(wb, [], [], [s], report_start=w)
    want = math.log(111.0 / 110.2) / 1.0 * 100.0
    wr = next(x for x in _table(wb["WeeklyReport"], lambda r: r[0] == "Scenario")
              if x["Batch"] == "B37")
    mr = next(x for x in _table(wb["MonthlyReport"], lambda r: r[0] == "Scenario")
              if x["Batch"] == "B37" and x["Month"] == "2024-12")
    assert abs(wr["SGR (%/day)"] - want) < 1e-3
    assert abs(mr["SGR (%/day)"] - want) < 1e-3
