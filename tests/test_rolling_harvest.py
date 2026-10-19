"""RollingHarvestPlan (report only, 2026-10-15): the operator's paste-ready
weekly harvest layout. Each check below has a negative control proving it can
fail."""
from __future__ import annotations

import copy
import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from forecast.events import Harvest
from forecast.excel_io import whole_harvest_counts, write_rolling_harvest_plan
from forecast.rolling_harvest import (BANDS_LB, band_lbs, build_rolling_rows,
                                      harvest_inputs_from_events, pct_lbs_in_range,
                                      rows_as_table, shown_band, source_tank_text)

ROOT = Path(__file__).resolve().parents[1]
MON43 = date(2026, 10, 19)          # Monday of 2026-W43


def _h(d, tank, batch, count, kg_per_fish=5.0, y=0.81):
    return {"date": d, "tank": tank, "batch": batch, "count": count,
            "hog_kg": count * kg_per_fish, "hog_yield": y}


def _two_tank_week():
    # 50,003 fish over 5 days -> 10,001 x3, 10,000 x2. Tank 61 holds 30,000,
    # so Wednesday finishes 61 (9,998) and starts 67 (3): a mixed day.
    return [_h(MON43, 61, "B44", 30000, 5.0), _h(MON43, 67, "B45", 20003, 6.0)]


def _transfers():
    return [
        {"week": "2026-W41", "batch": "B44", "from_tank": 41, "to_tank": 61, "count": 23100},
        {"week": "2026-W41", "batch": "B44", "from_tank": 33, "to_tank": 61, "count": 6900},
        {"week": "2026-W41", "batch": "B45", "from_tank": 61, "to_tank": 67, "count": 500},
        {"week": "2026-W41", "batch": "B45", "from_tank": 67, "to_tank": 67, "count": 999},
    ]


# ---- checkers (each has a negative control) --------------------------------

def _week_rows(rows):
    """The week blocks without the year summary (its row + blank line)."""
    out, skip = [], False
    for r in rows:
        if r["kind"] == "year":
            skip = True
            continue
        if skip and r["kind"] == "blank":
            skip = False
            continue
        out.append(r)
    return out


def block_structure_errors(rows):
    rows = _week_rows(rows)
    errs = []
    i = 0
    while i < len(rows):
        wk = rows[i]["week_label"]
        blk = [r for r in rows if r["week_label"] == wk]
        kinds = [r["kind"] for r in blk]
        n_day = kinds.count("day")
        if kinds != ["day"] * n_day + ["total", "blank"] or not 1 <= n_day <= 5:
            errs.append((wk, kinds))
        i += len(blk)
    return errs


def total_errors(rows, week_totals):
    errs = []
    for wk, want in week_totals.items():
        days = [r for r in rows if r["week_label"] == wk and r["kind"] == "day"]
        tot = [r for r in rows if r["week_label"] == wk and r["kind"] == "total"]
        got = sum(r["count"] for r in days)
        if got != want or not tot or tot[0]["count"] != want:
            errs.append((wk, got, want))
    return errs


def evenness_errors(rows):
    rows = _week_rows(rows)
    errs = []
    for wk in {r["week_label"] for r in rows}:
        c = [r["count"] for r in rows if r["week_label"] == wk and r["kind"] == "day"]
        if max(c) - min(c) > 1 or c != sorted(c, reverse=True):
            errs.append((wk, c))
    return errs


# ---- tests ------------------------------------------------------------------

def test_block_structure_and_negative_control():
    rows = build_rolling_rows(_two_tank_week(), report_start=date(2026, 10, 12))
    assert [r["kind"] for r in rows] == ["day"] * 5 + ["total", "blank", "year",
                                                       "blank"]
    assert block_structure_errors(rows) == []
    bad = copy.deepcopy(rows)
    del bad[5]                                   # drop the Total row
    assert block_structure_errors(bad)


def test_days_sum_to_week_total_and_negative_control():
    rows = build_rolling_rows(_two_tank_week())
    assert [r["count"] for r in rows if r["kind"] == "day"] == [10001] * 3 + [10000] * 2
    assert total_errors(rows, {"2026-W43": 50003}) == []
    assert evenness_errors(rows) == []
    bad = copy.deepcopy(rows)
    bad[1]["count"] += 1                         # alter one day's count
    assert total_errors(bad, {"2026-W43": 50003})
    assert evenness_errors(bad)


def test_kg_follows_each_tank_and_mixed_day_lists_both():
    rows = build_rolling_rows(_two_tank_week())
    wed = rows[2]
    assert wed["tank"] == "61,67" and wed["batch"] == "B44,B45"
    assert wed["kg_hog"] == pytest.approx(9998 * 5.0 + 3 * 6.0)
    assert rows[0]["tank"] == 61 and rows[0]["batch"] == "B44"
    assert rows[3]["tank"] == 67
    tot = rows[5]
    assert tot["kg_hog"] == pytest.approx(30000 * 5.0 + 20003 * 6.0)
    assert tot["tank"] == "61, 67" and tot["batch"] == "B44, B45"
    assert tot["date"] == "43 Total"
    assert rows[6]["kind"] == "blank" and rows[6]["count"] == ""   # blank line
    # negative control: a single-tank week has no mixed day
    one = build_rolling_rows([_h(MON43, 61, "B44", 50000)])
    assert all("," not in str(r["tank"]) for r in one if r["kind"] == "day")


def test_first_week_clipped_to_report_start():
    rows = build_rolling_rows(_two_tank_week(), report_start=date(2026, 10, 21))
    days = [r for r in rows if r["kind"] == "day"]
    assert [r["date"] for r in days] == [date(2026, 10, 21), date(2026, 10, 22),
                                         date(2026, 10, 23)]
    assert sum(r["count"] for r in days) == 50003
    assert block_structure_errors(rows) == []


def test_source_tank_percent_format_and_negative_control():
    evs = _two_tank_week()
    # 6N source (61->67) and retention (67->67) are excluded.
    assert source_tank_text(evs, evs, _transfers()) == "41(77%)/33(23%)"
    assert source_tank_text(evs[1:], evs, _transfers()) == ""
    single = [t for t in _transfers() if t["from_tank"] == 41]
    assert source_tank_text(evs, evs, single) == "41"
    # negative control: a transfer outside the fill window (before the tank's
    # previous harvest) does not count.
    prior = [_h(date(2026, 10, 5), 61, "B44", 100)]
    assert source_tank_text(evs[:1], evs + prior, _transfers()) == "41(77%)/33(23%)"
    late_prior = [_h(date(2026, 10, 12), 61, "B44", 100)]
    assert source_tank_text(evs[:1], evs + late_prior, _transfers()) == ""
    rows = build_rolling_rows(evs, _transfers())
    assert rows[0]["source"] == "41(77%)/33(23%)"
    assert all(r["source"] == "" for r in rows[1:])


def test_harvest_id_sequence_from_knob():
    hs = _two_tank_week() + [_h(date(2026, 10, 26), 63, "B46", 40000)]
    rows = build_rolling_rows(hs, first_id=117)
    ids = [r["harvest_id"] for r in rows if r["kind"] == "day"]
    assert ids[:5] == ["B44T61-117", "B44T61-118", "B44,B45T61,67-119",
                       "B45T67-120", "B45T67-121"]
    assert ids[5] == "B46T63-122"
    # negative control: default starts at 1
    assert build_rolling_rows(hs)[0]["harvest_id"] == "B44T61-1"


def test_weeks_knob_limits_blocks():
    hs = [_h(date.fromisocalendar(2026, w, 1), 61, "B44", 1000) for w in range(42, 54)]
    hs += [_h(date.fromisocalendar(2027, 1, 1), 61, "B44", 1000)]
    rs = date(2026, 10, 12)
    wk = lambda rows: sorted({r["week_label"] for r in _week_rows(rows)})
    assert len(wk(build_rolling_rows(hs, report_start=rs, weeks=12))) == 12
    assert len(wk(build_rolling_rows(hs, report_start=rs, weeks=3))) == 3
    assert len(wk(build_rolling_rows(hs, report_start=rs, weeks=0))) == 13
    assert wk(build_rolling_rows(hs, report_start=rs, weeks=3))[-1] == "2026-W44"


def test_event_inputs_tie_to_whole_harvest_counts():
    evs = [Harvest("B44", MON43, 61, 26592.5, 5000.0),
           Harvest("B45", MON43, 67, 26592.5, 5200.0),
           Harvest("B45", MON43, 69, 1815.0, 5200.0)]
    shown = whole_harvest_counts(evs)
    ins = harvest_inputs_from_events(evs, 0.81, {})
    assert [i["count"] for i in ins] == [shown[id(e)] for e in evs]
    rows = build_rolling_rows(ins)
    assert sum(r["count"] for r in rows if r["kind"] == "day") == 55000
    # per-tank whole fish are respected: tank 61's fish all come out
    taken = sum(r["count"] for r in rows if r["kind"] == "day" and r["tank"] == 61)
    assert taken <= shown[id(evs[0])]
    assert total_errors(rows, {"2026-W43": 55000}) == []
    assert total_errors(rows, {"2026-W43": 55001})          # negative control


def test_writer_layout_formulas_and_yield():
    evs = [Harvest("B44", MON43, 61, 30000, 6000.0),
           Harvest("B45", MON43, 67, 20003, 6100.0)]
    wb = openpyxl.Workbook()
    write_rolling_harvest_plan(wb, evs, 0.81, {"2026-W43": 0.8}, first_id=5,
                               report_start=date(2026, 10, 12))
    ws = wb["RollingHarvestPlan"]
    hdr = [ws.cell(row=4, column=c).value for c in range(1, 16)]
    assert hdr[1:] == ["Source Tank", "Year ", "Harvest ID", "Week", "Date",
                       "Tank ", "Batch", "Count", "Weight (kg HOG)",
                       "Average weight (kg HOG)", "Live weight (kg)",
                       "Origin FT Weight (kg)", "Weight (lb HOG)",
                       "Average weight (lb HOG)"]
    assert ws["I3"].value == "Kg" and ws["N3"].value == "Lb"
    assert ws["D5"].value == "B44T61-5" and ws["E5"].value == 43
    assert ws["K5"].value == "=J5/I5" and ws["L5"].value == "=K5/0.8"
    assert ws["N5"].value == "=J5*2.2046" and ws["O5"].value == "=K5*2.2046"
    assert ws["I10"].value == "=SUM(I5:I9)" and ws["F10"].value == '=E10&" Total"'
    # the line after a week's Total is blank, all the way across A..W
    assert all(ws.cell(row=11, column=c).value is None for c in range(1, 24))
    # then the year summary (values) and another blank line
    assert ws["F12"].value == "2026 Total" and ws["I12"].value == 50003
    assert ws["D12"].value == "2026-10-19 to 2026-10-23"
    assert all(ws.cell(row=13, column=c).value is None for c in range(1, 24))
    assert ws.freeze_panes == "A5"
    assert sum(ws.cell(row=r, column=9).value for r in range(5, 10)) == 50003
    tab = rows_as_table(build_rolling_rows(harvest_inputs_from_events(evs, 0.81, {})))
    assert all(len(r) == 23 and r[0] == "" and r[12] == "" for r in tab)
    # P..W: the sales weight ranges, header and values
    assert [ws.cell(row=4, column=c).value for c in range(16, 24)] == [
        "0 - 2lbs HOG", "2 - 4lbs HOG", "4 - 5.8lbs HOG", "5.8 - 8lbs HOG",
        "8 - 10lbs HOG", "10 - 12lbs HOG", "12-14lbs HOG", "14-16 lbs HOG"]
    p10 = [ws.cell(row=10, column=c).value for c in range(16, 24)]
    n10 = sum(v for v in p10 if v)
    lb = sum(30000 * 6.0 * 0.8 * 2.2046 if t == 61 else 0 for t in (61,)) + \
        20003 * 6.1 * 0.8 * 2.2046
    # the ranges stop at 16 lb, as on the sales sheet: these ~10.7 lb fish put
    # ~1.5% of their weight above it, so the ranges hold just under the total
    assert 0.95 * lb < n10 <= lb
    assert p10[0] is None                  # 0-2 lb: nothing, shown blank


# ---- the sales weight ranges (operator 2026-10-19) -------------------------

def test_ranges_reproduce_the_sales_sheet_row_and_negative_control():
    # The operator's sales sheet, B43T65-978: 10,915 fish, 78,649 lb HOG ->
    # 2-4: 29, 4-5.8: 7,565, 5.8-8: 46,638, 8-10: 21,204, 10-12: 2,986,
    # 12-14: 216 (0-2 and 14-16 under 15 lb, blank).
    lb = 78649.0
    got = band_lbs(lb, lb / 10915)
    want = [0, 29, 7565, 46638, 21204, 2986, 216, 11]
    assert all(abs(g - w) <= 1.0 for g, w in zip(got, want)), [round(g) for g in got]
    assert [shown_band(round(g)) for g in got][0] == ""
    assert shown_band(round(got[7])) == "" and shown_band(round(got[1])) == 29
    # negative control: weighting by FISH (no sigma^2 shift) is not the sheet.
    import math
    from statistics import NormalDist
    s2 = math.log(1 + 0.1805 ** 2)
    d = NormalDist(math.log(lb / 10915) - 0.5 * s2, math.sqrt(s2))
    by_fish = lb * (d.cdf(math.log(10)) - d.cdf(math.log(8)))
    assert abs(by_fish - 21204) > 1000


def test_threshold_is_more_than_15_lb():
    assert shown_band(15.0) == "" and shown_band(15.01) == 15.01
    assert shown_band(0.0) == "" and shown_band("") == ""
    assert pct_lbs_in_range(0.0, 2, 4) == 0.0
    assert pct_lbs_in_range(7.0, 4, 4) == 0.0           # empty range


def test_day_ranges_follow_each_tank_and_sum_to_the_total():
    rows = build_rolling_rows(_two_tank_week())
    days = [r for r in rows if r["kind"] == "day"]
    tot = next(r for r in rows if r["kind"] == "total")
    for bi in range(len(BANDS_LB)):
        assert tot["bands"][bi] == pytest.approx(sum(r["bands"][bi] for r in days))
    for r in days:              # the ranges stop at 16 lb: never more than the total
        assert sum(r["bands"]) <= r["lb_hog"]
    for r in days[:2]:          # 11 lb fish (tank 61): only a ~2% tail above 16 lb
        assert sum(r["bands"]) > 0.95 * r["lb_hog"]
    # Wednesday finishes 61 (9,998 at 5.0 kg) and starts 67 (3 at 6.0 kg):
    # each part spread at its own weight.
    wed = days[2]
    want = [a + b for a, b in zip(band_lbs(9998 * 5.0 * 2.2046, 5.0 * 2.2046),
                                  band_lbs(3 * 6.0 * 2.2046, 6.0 * 2.2046))]
    assert wed["bands"] == pytest.approx(want)
    # negative control: one blended average gives different ranges
    blend = band_lbs(wed["lb_hog"], wed["avg_lb_hog"])
    assert any(abs(a - b) > 1e-6 for a, b in zip(blend, want))


def test_year_summary_adds_each_years_days_and_negative_control():
    # 2026-W53 runs Mon Dec 28 .. Fri Jan 1: four days in 2026, one in 2027.
    hs = [_h(date(2026, 12, 21), 61, "B44", 10000, 3.0),
          _h(date(2026, 12, 28), 63, "B45", 10000, 3.2),
          _h(date(2027, 1, 4), 65, "B46", 10000, 3.4)]
    rows = build_rolling_rows(hs, report_start=date(2026, 12, 21))
    yrs = {r["year"]: r for r in rows if r["kind"] == "year"}
    assert sorted(yrs) == [2026, 2027]
    days = [r for r in rows if r["kind"] == "day"]
    for y, yr in yrs.items():
        mine = [d for d in days if d["date"].year == y]
        assert yr["count"] == sum(d["count"] for d in mine)
        assert yr["kg_hog"] == pytest.approx(sum(d["kg_hog"] for d in mine))
        assert yr["bands"] == pytest.approx(
            [sum(d["bands"][b] for d in mine) for b in range(len(BANDS_LB))])
    assert yrs[2026]["count"] == 10000 + 8000           # W53's Jan 1 is 2027's
    assert yrs[2027]["count"] == 2000 + 10000
    assert yrs[2026]["harvest_id"] == "2026-12-21 to 2026-12-31"
    # each year's row comes after the last week holding one of its days:
    # 2026 after W53 (Dec 28 .. Jan 1), 2027 after W01
    kinds = [(r["kind"], r["week_label"] or r["year"]) for r in rows
             if r["kind"] in ("total", "year")]
    assert kinds == [("total", "2026-W52"), ("total", "2026-W53"), ("year", 2026),
                     ("total", "2027-W01"), ("year", 2027)]
    # negative control: counting W53 by its ISO year would put 10,000 in 2026
    by_week_year = sum(r["count"] for r in rows if r["kind"] == "total"
                       and r["week_label"].startswith("2026"))
    assert by_week_year != yrs[2026]["count"]
    # the blocks still read day..Total, blank, with block_top on their first day
    assert block_structure_errors(rows) == []
    for r in rows:
        if r["kind"] in ("day", "total"):
            assert rows[r["block_top"]]["kind"] == "day"
            assert rows[r["block_top"]]["week_label"] == r["week_label"]


# ---- the whole pipeline on the fixture PR -----------------------------------

FIXTURE = ROOT / "tests" / "fixtures" / "pr_2026-09-30.xlsx"


@pytest.mark.skipif(not FIXTURE.exists(), reason="fixture PR missing")
def test_fixture_run_sheet_week_totals_match_harvest_plan(tmp_path, copy_config):
    from forecast.run import main
    cfg, scn = tmp_path / "config", tmp_path / "scenario"
    copy_config(ROOT / "config", cfg)          # never the operator's costs
    shutil.copytree(ROOT / "scenario", scn)
    out = tmp_path / "out.xlsx"
    main(str(FIXTURE), str(out), config_dir=str(cfg), scenario_dir=str(scn),
         calib_log_path="")
    out = next(tmp_path.glob("out.xls*"))     # suffix follows the input's macros
    wb = openpyxl.load_workbook(out)
    assert "RollingHarvestPlan" in wb.sheetnames
    # replaced by RollingHarvestPlan (operator 2026-10-19)
    assert "Daily Harvest Schedule" not in wb.sheetnames
    plan: dict = {}
    for row in wb["HarvestPlan"].iter_rows(min_row=5, values_only=True):
        if row and isinstance(row[0], str) and "-W" in row[0]:
            plan[row[0]] = plan.get(row[0], 0) + row[3]
    ws = wb["RollingHarvestPlan"]
    got: dict = {}
    for r in range(5, ws.max_row + 1):
        f = ws.cell(row=r, column=6).value
        if isinstance(f, str) and f.endswith('" Total"'):
            continue
        c, e = ws.cell(row=r, column=9).value, ws.cell(row=r, column=5).value
        if isinstance(c, (int, float)) and e is not None:
            wk = f"{date.fromisoformat(str(f)[:10]).isocalendar()[0]}-W{int(e):02d}"
            got[wk] = got.get(wk, 0) + c
    assert got, "no daily rows written"
    for wk, n in got.items():
        assert n == plan[wk], (wk, n, plan[wk])
    # negative control: the comparison catches a one-fish difference
    some = next(iter(got))
    assert got[some] + 1 != plan[some]
