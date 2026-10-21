"""Models sheet: the growth, feed-conversion and other biology models this run
used, spelled out in one place (operator 2026-10-19: "make sure the models are
clearly spelled out in the workbook").

BatchInputs already names each batch's growth and FCR model, and RunConfig
dumps the raw curves; neither says what a model IS. This sheet does: which
batches grow on which seawater curve and whether it was measured from the
production reports or assumed, how a measured one was made and how well it
forecast later reports (read from the growth history, never written), the
curves themselves at round fish sizes, the FCR tables, and the other
assumptions a reader needs to trust the numbers.

REPORT ONLY. Nothing here is read back by the planner, and every number is the
run's own (the BiologyTables, batches and control the run used), so the sheet
cannot disagree with the plan. Appended last, like BatchInputs, so no existing
sheet's index or style ids move.
"""
from __future__ import annotations

from pathlib import Path

from .batch_order import batch_sort_key

SHEET_NAME = "Models"
GROWTH_SIZES_G = [250, 500, 750, 1000, 1500, 2000, 2500, 3000, 3500, 4000, 4500, 5000]
FCR_SIZES_G = [100, 250, 500, 1000, 2000, 3000, 4000, 5000]


def _batches_by(batches, attr, key=lambda v: v):
    out: dict = {}
    for b in batches:
        out.setdefault(key(getattr(b, attr, "") or ""), []).append(b)
    return {k: sorted(v, key=lambda b: batch_sort_key(b.batch_id)) for k, v in out.items()}


def _ids(bs) -> str:
    return ", ".join(b.batch_id for b in bs)


def adjustment_text(vals, curve="curve") -> str:
    """A batch growth multiplier in words: 1.0 is 'none', not '1' -- a bare 1
    reads like '100% of the model' (operator 2026-10-19)."""
    out = []
    for v in sorted({round(float(v), 4) for v in vals}):
        if abs(v - 1.0) < 1e-9:
            out.append(f"none - the {curve} as it is")
        else:
            out.append(f"x{v:g} - {abs(1.0 - v) * 100:.0f}% "
                       f"{'slower' if v < 1 else 'faster'} than the {curve}")
    return "; ".join(out)


def _history_record(history_path, name):
    """(record, error) for a growth model name from the growth history
    (a measured GrowthFit or an ASSUMED DerivedModel); read only."""
    if not history_path:
        return None, "no growth history given"
    p = Path(history_path)
    if not p.exists():
        return None, f"growth history not found ({p.name})"
    try:
        from .growth_history import GrowthHistory
        return GrowthHistory(p).get_version(name), None
    except KeyError:
        return None, "not in the growth history"
    except Exception as e:  # noqa: BLE001 - a report must not sink a run
        return None, f"growth history unreadable ({type(e).__name__})"


def growth_model_rows(batches, tables, history_path=None) -> list[list]:
    """[Model, Type, Batches, Adjustment on the curve, How it was made,
    Accuracy check]."""
    from .biology import BASE_GROWTH_MODEL
    rows = []
    by = _batches_by(batches, "growth_model",
                     key=lambda v: v if v and v != BASE_GROWTH_MODEL else BASE_GROWTH_MODEL)
    for name in sorted(by, key=lambda n: (n == BASE_GROWTH_MODEL, n)):
        bs = by[name]
        corr = adjustment_text(getattr(b, "sgr_correction", 1.0) or 1.0 for b in bs)
        if name == BASE_GROWTH_MODEL:
            rows.append([name, "Company table (assumed)", _ids(bs), corr,
                         "biology.yaml sgr_sw_pct_day: the original seawater "
                         "growth table, not fitted to this site's reports", ""])
            continue
        rec, err = _history_record(history_path, name)
        if rec is None:
            rows.append([name, "From biology.yaml", _ids(bs), corr,
                         f"biology.yaml sgr_sw_by_model ({err})", ""])
        elif hasattr(rec, "explain"):                       # DerivedModel
            rows.append([name, "ASSUMED improvement", _ids(bs), corr,
                         f"{rec.explain} Built from '{rec.source}'.", ""])
        else:                                               # GrowthFit
            bp = rec.breakpoint_g
            how = (f"Measured from production reports: {rec.n} clean tank-months "
                   f"({rec.n_provisional} from mid-month reports), "
                   f"{rec.window_start} to {rec.as_of} ({rec.window_months}-month "
                   f"window). Growth = {rec.a:.4f} x weight^{rec.b:.3f} g/day up "
                   f"to {bp:,.0f} g ({rec.n_below} tank-months), then a flat "
                   f"{rec.cap_g_day:.1f} g/day above {bp:,.0f} g (median of "
                   f"{rec.n_above} tank-months).")
            if getattr(rec, "excluded_batches", None):
                how += (" Batches left out of the fit: "
                        f"{', '.join(rec.excluded_batches)}.")
            ho = rec.holdout or {}
            if ho.get("fit") and ho.get("base"):
                acc = (f"Fitted without the reports after {ho['from']}, then "
                       f"forecast those {ho['n']} tank-months: median error "
                       f"{ho['fit']['median_err_pct']:+.1f}% (base curve "
                       f"{ho['base']['median_err_pct']:+.1f}%).")
            else:
                acc = "No later reports to test it on yet."
            rows.append([name, "Measured", _ids(bs), corr, how, acc])
    return rows


def growth_check_rows(batches, tables, history_path=None, facility=None) -> list[list]:
    """[Model, Data, *growth_history.CHECK_COLUMNS] for every MEASURED growth
    model the run's batches use: SGR by fish size, measured vs the curve this
    run used (biology.yaml) vs the 100% curve -- on the tank-months the model
    was fitted on, and on the ones that ended after it (data it never saw).
    Read from the growth history at run time, so it updates as PRs are added."""
    from .biology import BASE_GROWTH_MODEL
    from . import growth_history as gh
    names = sorted({getattr(b, "growth_model", "") or "" for b in batches} - {"", BASE_GROWTH_MODEL})
    pad = [None] * (len(gh.CHECK_COLUMNS) - 1)
    recs = []
    for name in names:
        rec, err = _history_record(history_path, name)
        if rec is not None and not hasattr(rec, "explain"):     # a GrowthFit
            recs.append((name, rec))
    if not recs:
        return [["No measured growth model in this run", "", *pad, None]]
    if facility is None:
        return [["(not available: no facility given)", "", *pad, None]]
    try:
        ivs = gh.GrowthHistory(Path(history_path)).intervals(
            *gh.facility_tank_sets(facility))
    except Exception as e:  # noqa: BLE001 - a report must not sink a run
        return [[f"(growth history unreadable: {type(e).__name__})", "", *pad, None]]
    rows = []
    for name, rec in recs:
        curve = (getattr(tables, "sgr_sw_by_model", None) or {}).get(name)
        sizes = tables.sgr_size_g
        if curve is None:                                       # not in biology.yaml
            curve, sizes = rec.sgr_pct_day, rec.sizes_g
        own = gh.fit_intervals(rec, ivs)
        win = f"Fitted window {rec.window_start} .. {rec.as_of}"
        if len(own) != rec.n:
            win += (f" (fitted on {rec.n} tank-months; {len(own)} as the history "
                    f"reads now - a later report replaced some)")
        for label, sub in (
                (win, own),
                (f"After {rec.as_of} (not used in the fit)",
                 gh.later_intervals(rec, ivs))):
            body = gh.check_rows(sizes, curve, sub, tables.sgr_size_g,
                                 tables.sgr_sw_pct_day)
            if not body:
                rows.append([name, label, "no tank-months yet", *pad[1:], None])
            for r in body:
                rows.append([name, label, *r])
    return rows


def growth_curve_rows(batches, tables) -> tuple[list, list[list]]:
    """(header, rows): seawater SGR %/day and g/day by fish size for every
    model in use, plus the base curve for comparison."""
    from .biology import BASE_GROWTH_MODEL, _interp
    sizes = list(tables.sgr_size_g or [])
    used = sorted({getattr(b, "growth_model", "") or BASE_GROWTH_MODEL for b in batches})
    names = [n for n in used if n != BASE_GROWTH_MODEL] + [BASE_GROWTH_MODEL]
    curves = {n: (tables.sgr_sw_pct_day if n == BASE_GROWTH_MODEL
                  else (tables.sgr_sw_by_model or {}).get(n, [])) for n in names}
    header = ["Fish size (g)"]
    for n in names:
        tag = n if n in used else f"{n} (not used; for comparison)"
        header += [f"{tag} - SGR (%/day)", f"{tag} - growth (g/day)"]
    rows = []
    for w in GROWTH_SIZES_G:
        r = [w]
        for n in names:
            sgr = _interp(w, sizes, curves[n]) if curves[n] else None
            r += [sgr, (w * sgr / 100.0) if sgr is not None else None]
        rows.append(r)
    return header, rows


def fcr_rows(batches, tables) -> tuple[list, list[list]]:
    """(header, rows): each FCR model in use, its batches and value by size."""
    from .biology import _fcr_model_key, _interp
    by = _batches_by(batches, "fcr_model", key=lambda v: _fcr_model_key(v) if v else "")
    header = (["Model", "Type", "Batches"]
              + [f"FCR at {w:,} g" for w in FCR_SIZES_G])
    rows = []
    for key in sorted(k for k in by if k):
        curve = (tables.fcr_by_model or {}).get(key, [])
        labels = sorted({b.fcr_model for b in by[key]})
        rows.append([f"{key} ({', '.join(labels)})",
                     "Company table (assumed; not yet learned from reports)",
                     _ids(by[key])]
                    + [(_interp(w, tables.fcr_size_g, curve) if curve else None)
                       for w in FCR_SIZES_G])
    return header, rows


def other_rows(batches, tables, control, hog_overrides=None,
               forecast_start=None) -> list[list]:
    """[Item, Value, Where it comes from]. Weekly factors and yields are listed
    for the forecast's own weeks (from `forecast_start`) only."""
    from .time_grid import iso_week_label
    first = iso_week_label(forecast_start) if forecast_start is not None else ""
    rows = []
    fw = sorted({round(float(getattr(b, "fw_correction", 1.0) or 1.0), 3) for b in batches})
    rows.append(["Freshwater growth",
                 (f"the company FW curve, adjusted per batch; this run from "
                  f"{adjustment_text([fw[0]], 'FW curve')} to "
                  f"{adjustment_text([fw[-1]], 'FW curve')} (each batch in BatchInputs)"
                  if len(fw) > 1 else
                  f"the company FW curve, {adjustment_text(fw, 'FW curve')}")
                 if fw else "the company FW curve",
                 "biology.yaml sgr_fw_pct_day; batches.yaml fw_correction, "
                 + ("auto-calibrated so each batch reaches its planned transfer "
                    "weight (BatchInputs)" if getattr(control, "auto_calibrate_fw", False)
                    else "as entered")])
    og = getattr(tables, "og_sgr_by_week", None) or {}
    off = {k: v for k, v in og.items()
           if v is not None and abs(v - 1.0) > 1e-9 and k >= first}
    rows.append(["Seawater growth per-week factors",
                 ", ".join(f"{k} x{v:g}" for k, v in sorted(off.items())) or
                 "none in the forecast weeks (every week x1.0)",
                 "scenario/limits.yaml sgr_correction_og - multiplies every OG "
                 "batch's growth in that week"])
    rows.append(["6N purge tanks",
                 "fish grow and eat in 6N" if getattr(control, "sixn_growth", False)
                 else "no growth and no feed while in a 6N purge tank; harvested "
                      "2 weeks after moving in",
                 "control.yaml sixn_growth"])
    mw = list(zip(tables.mortality_week_from_input or [], tables.mortality_pct_weekly or []))
    rows.append(["Mortality (% per week, by weeks since egg input)",
                 "; ".join(f"from wk {w}: {m:g}%" for w, m in mw) or "none",
                 "biology.yaml mortality_pct_weekly - each value holds until the next"])
    rows.append(["Freshwater culling",
                 "; ".join(f"day {d}: {p:g}%" for d, p in (tables.culling or [])) or "none",
                 "biology.yaml culling - days since input"])
    rows.append(["Handling mortality",
                 f"{getattr(control, 'handling_mortality_pct', 0) or 0:g}% of fish moved",
                 "control.yaml handling_mortality_pct"])
    ov = len([k for k in (hog_overrides or {}) if k >= first])
    rows.append(["HOG yield (head-on gutted / live)",
                 f"{getattr(control, 'default_hog_yield', 0.81):g}"
                 + (f" (other values in {ov} week(s))" if ov else ""),
                 "control.yaml default_hog_yield; scenario/limits.yaml hog_yield rows"])
    rows.append(["Grader efficiency", f"{getattr(control, 'grade_efficiency', 1.0):g}",
                 "control.yaml grade_efficiency (1.0 = a perfect size sort)"])
    from .rolling_harvest import MODEL_CV
    rows.append(["Harvest size spread (sales weight ranges)",
                 f"lognormal, CV {MODEL_CV:g}, by biomass",
                 "RollingHarvestPlan columns P-W (the sales sheet's PCT_LBS_IN_RANGE)"])
    return rows


def write_models_sheet(wb, *, batches, tables, control, history_path=None,
                       hog_overrides=None, forecast_start=None, facility=None,
                       sheet_name: str = SHEET_NAME) -> None:
    """Append the Models sheet (replacing a stale one)."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from . import excel_format as xf
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    batches = sorted(batches, key=lambda b: batch_sort_key(b.batch_id))
    ws.append(["MODELS USED IN THIS FORECAST"])
    ws.append(["The growth, feed-conversion and other biology models this run used: "
               "what each is, which batches use it, and the numbers behind it. "
               "Per-batch detail: BatchInputs. Report only."])
    heads, tables_at = [], []

    def section(title, header, rows, fmt=None):
        ws.append([])
        ws.append([title]); heads.append(ws._current_row)
        ws.append(header); tables_at.append(ws._current_row)
        for r in rows:
            ws.append(r)
            if fmt:
                for c, f in fmt.items():
                    cell = ws.cell(ws._current_row, c)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = f(c) if callable(f) else f

    section("SEAWATER GROWTH MODELS (OG) - which curve each batch grows on",
            ["Model", "Type", "Batches",
             "Adjustment on the curve (batches.yaml sgr_correction)",
             "How it was made", "Accuracy check"],
            growth_model_rows(batches, tables, history_path))
    gh, gr = growth_curve_rows(batches, tables)
    section("SEAWATER GROWTH BY FISH SIZE - the curves above",
            gh, gr, {c: ("0.000" if c % 2 == 0 else "0.0") for c in range(2, len(gh) + 1)}
            | {1: "#,##0"})
    from .growth_history import CHECK_COLUMNS
    section("SEAWATER GROWTH CHECKED AGAINST THE DATA - SGR %/day by fish size "
            "(100% curve = biology.yaml sgr_sw_pct_day, the original table)",
            ["Model", "Data", *CHECK_COLUMNS],
            growth_check_rows(batches, tables, history_path, facility),
            {5: "0.000", 6: "0.000", 7: "0.000", 8: '0"%"', 9: '0"%"'})
    fh, fr = fcr_rows(batches, tables)
    section("FEED CONVERSION (FCR) MODELS - kg feed per kg growth, by fish size",
            fh, fr, {c: "0.00" for c in range(4, len(fh) + 1)})
    section("OTHER BIOLOGY AND HARVEST ASSUMPTIONS",
            ["Item", "Value", "Where it comes from"],
            other_rows(batches, tables, control, hog_overrides, forecast_start))

    ws["A1"].font = Font(bold=True, size=14)
    for r in heads:
        ws.cell(r, 1).font = xf.F_HDR
        ws.cell(r, 1).fill = xf.C_HDR
    for r in tables_at:
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value is not None:
                ws.cell(r, c).font = Font(bold=True)
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            if isinstance(cell.value, str) and len(cell.value) > 40:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {1: 30, 2: 28, 3: 34, 4: 16, 5: 70, 6: 50}
    for c in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 16)
    ws.sheet_properties.tabColor = xf.TAB_INPUT
