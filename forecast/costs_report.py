"""The CostsAndProfit sheet: cash-view costs and profit in the Run workbook.

CASH VIEW (operator ruling, 2026-09-10/11). Each month's SPEND is set against
that month's SALES. No cost is carried with the fish, so a month with little
harvest shows a loss while the fish in the tanks gain value. The arithmetic
is forecast.costs, the ONE cost function. This module only gathers the
drivers from the run and lays them out.

Where each driver comes from (all of it is already in the run):
  feed kg by type, by month   excel_io._feed_by_type_week, then
                              excel_io._feed_by_type_month (the calendar-day
                              month split): the two calls FeedForecastMonthly
                              builds its matrix from. So every figure TIES to
                              that sheet (before its round).
                              It covers realized OG/SW feed (STARVE eats
                              nothing), the 6N move-in add-back and the
                              projected FW/EGG feed.
  eggs stocked                BatchInput.input_count of each batch whose
                              input_date falls in [report_start, horizon
                              end], booked in the input date's month. Eggs
                              bought before the report opens are not this
                              horizon's spend.
  fixed                       fixed_monthly x days covered / days in month,
                              from report_start to the end of the last
                              week the run reports.
  revenue, HOG kg             the HarvestPlan rows (analysis.harvest_rows_
                              from_ws). Each week is booked whole in its
                              MONDAY's month (time_grid.iso_week_month_split,
                              clipped to report_start: the MonthlyReport
                              sales convention) and each month is priced by
                              analysis.revenue_for (the Analyze pricing).

NEUTRALITY. run.main drops a stale copy of this sheet right after loading
the input. It calls write_costs_sheet ONLY when config/costs.yaml exists, as
the last write before the save. The sheet is APPENDED LAST and styled on its
own (never apply_workbook_formatting). The HarvestPlan read creates no cells.
So every existing sheet's XML is byte-identical with or without costs. A bad
costs.yaml, or any other failure here, writes a NOT COMPUTED sheet that says
why and prints a NOTE: the run never fails because of costs.
"""
from __future__ import annotations

import calendar
import datetime as dt
import math
import numbers
import re
from collections import defaultdict
from pathlib import Path
from typing import Optional

from . import costs as _costs

COSTS_SHEET = "CostsAndProfit"
TITLE = "COSTS AND PROFIT (cash view)"
HEADER_ROW = 4
NOT_COMPUTED = "NOT COMPUTED"
FEED_BLOCK = "FEED BY TYPE"
KIND_MONTH = "Month"
KIND_PRE_START = "Month (before report opens)"
NONE_TEXT = "—"                 # a figure that does not exist (never a 0)
UNPRICED_TEXT = "(unpriced)"
NO_CURRENCY = "currency not set"
_DASH = " — "

# (header, field) in sheet column order; "{cur}" becomes the currency.
COLUMNS = (
    ("Period", "period"),
    ("Kind", "kind"),
    ("Feed (kg)", "feed_kg"),
    ("Feed ({cur})", "feed"),
    ("Shipping ({cur})", "shipping"),
    ("Oxygen ({cur})", "oxygen"),
    ("Chemicals ({cur})", "chemicals"),
    ("Eggs stocked", "eggs_n"),
    ("Eggs ({cur})", "eggs"),
    ("Fixed ({cur})", "fixed"),
    ("Total cost ({cur})", "total"),
    ("HOG harvested (kg)", "hog_kg"),
    ("Revenue ({cur})", "revenue"),
    ("Profit ({cur})", "profit"),
    # Cash view: the period's SPEND over the HOG kg SOLD in it - not the cost
    # of producing the fish sold (eggs and feed for fish still in the tanks
    # are in the spend, earlier spend on the fish sold is not).
    ("Spend per kg HOG sold ({cur})", "cost_per_kg_hog"),
    ("Unpriced feed (kg)", "unpriced_kg"),
)
FEED_COLUMNS = (
    ("Feed type", "feed_type"),
    ("Item", "item"),
    ("Price per kg ({cur})", "price_per_kg"),
    ("Feed (kg)", "feed_kg"),
    ("Feed ({cur})", "feed"),
)
_MONEY, _WHOLE, _PER_KG, _PRICE = "#,##0", "#,##0", "#,##0.00", "#,##0.000"
_FORMATS = {"feed_kg": _WHOLE, "feed": _MONEY, "shipping": _MONEY,
            "oxygen": _MONEY, "chemicals": _MONEY, "eggs_n": _WHOLE,
            "eggs": _MONEY, "fixed": _MONEY, "total": _MONEY,
            "hog_kg": _WHOLE, "revenue": _MONEY, "profit": _MONEY,
            "cost_per_kg_hog": _PER_KG, "unpriced_kg": _WHOLE}
_FEED_FORMATS = {"price_per_kg": _PRICE, "feed_kg": _WHOLE, "feed": _MONEY}


def _headers(cols, cur: str) -> list:
    return [h.format(cur=cur) for h, _f in cols]


# --------------------------------------------------------------------------- #
# The stale copy
# --------------------------------------------------------------------------- #
def drop_stale_costs_sheet(wb) -> bool:
    """Delete a CostsAndProfit sheet the input workbook carried in (an
    earlier planned workbook re-used as input). True when one was dropped.
    run.main calls it right after loading, before any writer or the
    formatting pass sees the workbook, so a costs-free run never carries an
    old sheet forward and every other sheet is untouched."""
    if COSTS_SHEET in wb.sheetnames:
        del wb[COSTS_SHEET]
        return True
    return False


# --------------------------------------------------------------------------- #
# Drivers
# --------------------------------------------------------------------------- #
def _day(value, what: str = "date") -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raise ValueError(f"{what} {value!r} is not a date")


def _month_of(d: dt.date) -> str:
    return f"{d.year}-{d.month:02d}"


def _monday(label) -> dt.date:
    s = str(label)
    try:
        return dt.date.fromisocalendar(int(s[:4]), int(s[6:8]), 1)
    except (ValueError, TypeError):
        raise ValueError(f"HarvestPlan week {label!r} is not an ISO week "
                         f"label 'YYYY-Www'") from None


class _Values:
    """A worksheet's values, read WITHOUT creating cells.

    openpyxl's iter_rows on a writable sheet creates every empty cell it
    passes, and an empty row then gains a <row> element in the saved XML.
    This writer reads HarvestPlan just before the save, when every existing
    sheet must stay byte-identical, so it reads the cells that exist and
    nothing else. The rows it yields are the ones iter_rows(values_only=True)
    would give."""

    def __init__(self, ws):
        self._ws = ws

    def iter_rows(self, values_only: bool = True):
        if not values_only:
            raise ValueError("_Values yields values only")
        cells = self._ws._cells
        if not cells:
            return
        n_rows, n_cols = self._ws.max_row, self._ws.max_column
        for r in range(1, n_rows + 1):
            yield tuple(cells[(r, c)].value if (r, c) in cells else None
                        for c in range(1, n_cols + 1))


def _harvest_rows(wb) -> list:
    """HarvestPlan's rows, read without creating a cell. A read-only
    workbook's sheet creates none, so it is read as it is. A writable sheet
    is read through _Values, and one without the `_cells` it needs (an
    openpyxl that renamed it) RAISES, so the sheet goes NOT COMPUTED: falling
    back to plain iter_rows would add cells to HarvestPlan and change its
    saved XML without a word."""
    from .analysis import harvest_rows_from_ws
    if "HarvestPlan" not in wb.sheetnames:
        raise ValueError("the workbook has no HarvestPlan sheet, so revenue "
                         "and HOG cannot be read")
    ws = wb["HarvestPlan"]
    if getattr(wb, "read_only", False):
        return harvest_rows_from_ws(ws)
    if not isinstance(getattr(ws, "_cells", None), dict):
        raise RuntimeError("HarvestPlan cannot be read without adding cells "
                           "to it: this openpyxl's Worksheet has no _cells "
                           "map, and reading it the plain way would change "
                           "the saved sheet")
    return harvest_rows_from_ws(_Values(ws))


def feed_by_type_month(batch_locations, states_by_batch, tables, batch_by_id,
                       sixn_move_in_feed, report_start) -> tuple[dict, list]:
    """({"YYYY-MM": {feed type: kg}}, [every week start the feed covers]).

    FeedForecastMonthly's matrix, unrounded: the same two calls it makes
    (excel_io._feed_by_type_week, then excel_io._feed_by_type_month, the one
    calendar-day month roll-up, with report_start), so each figure rounds to
    that sheet's cell. A week wholly before report_start is left unclipped,
    as it is there."""
    from .excel_io import _feed_by_type_month, _feed_by_type_week
    ftw, wk_start = _feed_by_type_week(
        batch_locations, states_by_batch, tables, batch_by_id,
        sixn_move_in_feed=sixn_move_in_feed)
    ftm, months = _feed_by_type_month(ftw, wk_start, report_start)
    out: dict = {m: {} for m in sorted(months)}
    for (name, mo), kg in ftm.items():
        out[mo][name] = kg
    return out, [_day(w, "week start") for w in wk_start.values()]


def eggs_by_month(batches, start: dt.date, end: dt.date) -> dict:
    """{"YYYY-MM": eggs stocked} over batches whose input_date is in
    [start, end], in the input date's month. An input_count that is not a
    finite count >= 0 raises, with the message ideal_engine.eggs_by_year
    gives."""
    out: dict = defaultdict(float)
    for b in batches:
        d = _day(b.input_date, f"batch {b.batch_id} input_date")
        if not start <= d <= end:
            continue
        n = b.input_count
        if (isinstance(n, bool) or not isinstance(n, numbers.Real)
                or not math.isfinite(n) or n < 0):
            raise ValueError(f"batch {b.batch_id}: input_count {n!r} is not "
                             f"a number of eggs")
        out[_month_of(d)] += float(n)
    return dict(out)


def harvest_by_month(rows, report_start: dt.date) -> dict:
    """{"YYYY-MM": [harvest rows]}: each week whole in its Monday's month,
    clipped to report_start (the MonthlyReport sales convention)."""
    from .time_grid import iso_week_month_split
    out: dict = {}
    for r in rows:
        (ym,) = iso_week_month_split(_monday(r["week"]),
                                     clip_start=report_start)
        out.setdefault(f"{ym[0]}-{ym[1]:02d}", []).append(r)
    return out


_span_months = _costs.span_months


def month_days(months, start: dt.date, end: dt.date) -> dict:
    """{"YYYY-MM": (days of [start, end] in the month, days in the month)}:
    forecast.costs.month_days, the one day split (the Ideal year's fixed
    cost uses it too, so the two surfaces charge the same days alike)."""
    return _costs.month_days(months, start, end)


def cost_drivers(*, batch_locations, states_by_batch, tables, batch_by_id,
                 batches, sixn_move_in_feed, report_start,
                 harvest_rows) -> dict:
    """Every quantity the sheet prices, gathered from the run (no prices).

    The horizon runs from report_start to the end (Monday + 6 days) of the
    last week any feed row or harvest row names. `pre_start` holds the
    months that end before the report opens: only a week wholly before
    report_start, which calendar_day_month_split leaves unclipped, can put
    feed there."""
    start = _day(report_start, "report_start")
    feed, week_starts = feed_by_type_month(
        batch_locations, states_by_batch, tables, batch_by_id,
        sixn_move_in_feed, start)
    harvest = harvest_by_month(harvest_rows, start)
    ends = week_starts + [_monday(r["week"]) for r in harvest_rows]
    end = (max(ends) + dt.timedelta(days=6)) if ends else start - \
        dt.timedelta(days=1)
    eggs = eggs_by_month(batches, start, end)
    span = _span_months(start, end) if end >= start else []
    months = sorted(set(feed) | set(harvest) | set(eggs) | set(span))
    pre = {m for m in months
           if dt.date(int(m[:4]), int(m[5:7]),
                      calendar.monthrange(int(m[:4]), int(m[5:7]))[1])
           < start}
    return dict(start=start, end=end, months=months, feed=feed,
                harvest=harvest, eggs=eggs,
                month_days=month_days(months, start, end), pre_start=pre)


# --------------------------------------------------------------------------- #
# The rows (pure)
# --------------------------------------------------------------------------- #
def _economics(config_dir) -> tuple[Optional[dict], str]:
    """(economics or None, the sheet's revenue sentence). A missing or broken
    economics.yaml leaves revenue unpriced and says so; costs still price."""
    from .analysis import load_economics
    try:
        econ = load_economics(config_dir)
    except Exception as exc:                                  # noqa: BLE001
        return None, (f"Revenue: NOT PRICED{_DASH}economics.yaml: "
                      f"{type(exc).__name__}: {exc}.")
    if econ is None:
        return None, (f"Revenue: NOT PRICED{_DASH}economics.yaml is missing "
                      f"or has no price bands.")
    return econ, "Revenue: economics.yaml price bands (the Analyze pricing)."


def _cell_value(v):
    return NONE_TEXT if v is None else v


def build_costs_rows(drivers: dict, costs: dict, economics: Optional[dict], *,
                     feed_types, sha: str, revenue_note: str,
                     economics_sha: Optional[str] = None
                     ) -> tuple[list, dict]:
    """(rows to append, layout) for the sheet. Pure: the P&L is
    costs.monthly_pl, each month's revenue analysis.revenue_for over that
    month's harvest rows. `feed_types` is the model feed-type order (size
    order); a type the feed names that is not in it is listed after.

    Row 2 (the note) names the pricing and the first 8 hex of costs.yaml's
    signature (`sha`) and, when given, economics.yaml's (`economics_sha`),
    so a reader can tell when either changed since the run. It says so
    when feed of a type with no price was left out of the cost: the cost
    is then understated and the profit overstated."""
    from .analysis import revenue_for
    harvest = drivers["harvest"]
    hog = {m: sum(r["hog_kg"] for r in rows) for m, rows in harvest.items()}
    rev, unpriced_rev = None, 0.0
    if economics is not None:
        rev = {}
        for m, rows in harvest.items():
            got = revenue_for(rows, economics)
            rev[m] = got["total"]
            unpriced_rev += got["unpriced_kg"]
    pl = _costs.monthly_pl(drivers["feed"], drivers["eggs"], rev, hog,
                           drivers["month_days"], costs)
    cur = economics["currency"] if economics is not None else NO_CURRENCY
    note = (f"Each month's spend against that month's sales{_DASH}a month "
            f"with little harvest shows a loss. {revenue_note} Costs: "
            f"config/{_costs.COSTS_FILE} (sha {sha}).")
    if economics_sha is not None:
        note += f" Price bands: config/economics.yaml (sha {economics_sha})."
    note += f" Currency: {cur}."
    if unpriced_rev > 0.5:
        note += (f" {unpriced_rev:,.0f} kg of harvest "
                 f"({economics.get('basis', 'hog')} basis) falls outside "
                 f"every price band and is unpriced.")
    unpriced_feed = pl["total"]["unpriced_kg"]
    if unpriced_feed > 0:
        note += (f" {unpriced_feed:,.0f} kg of feed ("
                 f"{', '.join(pl['total']['unpriced_types'])}) has no price "
                 f"in config/{_costs.COSTS_FILE} and is left out of the feed "
                 f"cost: total cost is understated and profit overstated.")
    rows: list = [[TITLE], [note], [], _headers(COLUMNS, cur)]
    for r in pl["months"]:
        kind = KIND_PRE_START if r["period"] in drivers["pre_start"] \
            else KIND_MONTH
        rows.append([_cell_value(dict(r, kind=kind)[f]) for _h, f in COLUMNS])
    for r in pl["years"] + [pl["total"]]:
        rows.append([_cell_value(r[f]) for _h, f in COLUMNS])
    layout = {"body_last": len(rows)}

    kg: dict = defaultdict(float)
    for m in sorted(drivers["feed"]):
        for name, v in drivers["feed"][m].items():
            kg[name] += v
    names = list(dict.fromkeys(feed_types))
    names += sorted(n for n in kg if n not in set(names))
    prices = costs["feed_prices"]
    rows += [[], [FEED_BLOCK], _headers(FEED_COLUMNS, cur)]
    layout["feed_title"] = len(rows) - 1
    layout["feed_header"] = len(rows)
    tot_kg = tot_feed = 0.0
    for name in names:
        p, v = prices.get(name), kg.get(name, 0.0)
        tot_kg += v
        if p is None:
            rows.append([name, None, UNPRICED_TEXT, v, NONE_TEXT])
        else:
            tot_feed += v * p["price_per_kg"]
            rows.append([name, p["item"] or None, p["price_per_kg"], v,
                         v * p["price_per_kg"]])
    rows.append(["Total", None, None, tot_kg, tot_feed])
    layout["feed_last"] = len(rows)
    return rows, dict(layout, pl=pl, currency=cur, note=note)


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #
def _style(ws, layout: dict) -> None:
    """Style THIS sheet only, with the workbook's own table grammar."""
    from openpyxl.utils import get_column_letter
    from .excel_format import (A_HDR, B_HDR, C_HDR, F_HDR, TAB_REPORT,
                               _format_table)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 28
    _format_table(ws, HEADER_ROW)
    # One filter over the P&L body only; the feed block below the blank row
    # is a second table.
    ws.auto_filter.ref = (f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}"
                          f"{layout['body_last']}")

    def _fmt(r, cols, formats):
        for c, (_h, f) in enumerate(cols, start=1):
            cell = ws.cell(r, c)
            fmt = formats.get(f)
            if fmt and isinstance(cell.value, (int, float)) \
                    and not isinstance(cell.value, bool):
                cell.number_format = fmt

    for r in range(HEADER_ROW + 1, layout["body_last"] + 1):
        _fmt(r, COLUMNS, _FORMATS)
    ws.cell(layout["feed_title"], 1).font = F_HDR
    for c in range(1, len(FEED_COLUMNS) + 1):
        cell = ws.cell(layout["feed_header"], c)
        cell.fill, cell.font, cell.alignment, cell.border = (C_HDR, F_HDR,
                                                             A_HDR, B_HDR)
    for r in range(layout["feed_header"] + 1, layout["feed_last"] + 1):
        _fmt(r, FEED_COLUMNS, _FEED_FORMATS)
    ws.sheet_properties.tabColor = TAB_REPORT


def _write_not_computed(wb, reason: str) -> dict:
    from .excel_format import C_TITLE, F_BAD, F_TITLE, TAB_REPORT
    drop_stale_costs_sheet(wb)
    ws = wb.create_sheet(COSTS_SHEET)
    ws.append([TITLE])
    ws.append([f"{NOT_COMPUTED}{_DASH}{reason}"])
    ws["A1"].fill, ws["A1"].font = C_TITLE, F_TITLE
    ws["A2"].font = F_BAD
    ws.sheet_properties.tabColor = TAB_REPORT
    print(f"  NOTE: {COSTS_SHEET} NOT COMPUTED - {reason}")
    return {"status": "not_computed", "error": reason}


def _costs_error(exc: Exception) -> str:
    msg = str(exc)
    lead = f"{_costs.COSTS_FILE}: "
    return msg[len(lead):] if msg.startswith(lead) else msg


def write_costs_sheet(wb, *, config_dir, batch_locations, states_by_batch,
                      tables, batch_by_id, batches, sixn_move_in_feed,
                      report_start) -> dict:
    """Append the CostsAndProfit sheet LAST (dropping any stale copy first).

    Does nothing when config_dir has no costs.yaml ({"status": "absent"}):
    run.main only calls it when the file exists, and this keeps the rule
    true for any other caller. A costs.yaml that is empty or invalid, or any
    failure while pricing, writes a NOT COMPUTED sheet naming the reason and
    prints a NOTE; it never raises. Every other sheet is left exactly as it
    was: HarvestPlan is read without creating a cell and only this sheet is
    styled. -> {"status": "ok" | "not_computed" | "absent", ...}"""
    drop_stale_costs_sheet(wb)
    if not (Path(config_dir) / _costs.COSTS_FILE).is_file():
        return {"status": "absent"}
    where = f"config/{_costs.COSTS_FILE}"
    try:
        costs = _costs.load_costs(config_dir)
    except (ValueError, OSError) as exc:
        return _write_not_computed(wb, f"{where}: {_costs_error(exc)}")
    if costs is None:
        return _write_not_computed(
            wb, f"{where}: the file is empty{_DASH}costs not set")
    try:
        sha = _costs.costs_sig(config_dir)[:8]
        econ_sha = _costs.economics_sig(config_dir)[:8]
        economics, revenue_note = _economics(config_dir)
        drivers = cost_drivers(
            batch_locations=batch_locations, states_by_batch=states_by_batch,
            tables=tables, batch_by_id=batch_by_id, batches=batches,
            sixn_move_in_feed=sixn_move_in_feed, report_start=report_start,
            harvest_rows=_harvest_rows(wb))
        ftypes = ([n for _s, n in sorted(tables.feed_types,
                                         key=lambda x: x[0])]
                  if tables is not None and getattr(tables, "feed_types",
                                                    None) else [])
        rows, layout = build_costs_rows(drivers, costs, economics,
                                        feed_types=ftypes, sha=sha,
                                        revenue_note=revenue_note,
                                        economics_sha=econ_sha)
    except Exception as exc:                                  # noqa: BLE001
        return _write_not_computed(wb, f"{type(exc).__name__}: {exc}")
    ws = wb.create_sheet(COSTS_SHEET)
    for row in rows:
        ws.append(row)
    _style(ws, layout)
    print(f"  {COSTS_SHEET}: {len(drivers['months'])} month(s) priced from "
          f"{where} (sha {sha})")
    return {"status": "ok", "sha": sha, "months": list(drivers["months"]),
            "body_last": layout["body_last"],
            "feed_header": layout["feed_header"],
            "feed_last": layout["feed_last"]}


# --------------------------------------------------------------------------- #
# Reading it back (the app's Run page)
# --------------------------------------------------------------------------- #
def _at(row, i):
    return row[i] if row is not None and i < len(row) else None


def _back(v):
    return None if v == NONE_TEXT else v


def _blank(row) -> bool:
    return row is None or all(v is None or v == "" for v in row)


def read_costs_sheet(ws) -> dict:
    """The sheet as data:
      {"status": "ok", "sha", "economics_sha", "currency", "note", "months",
       "years", "total", "feed_by_type", "feed_total"}
                                                 row dicts keyed by the
                                                 COLUMNS / FEED_COLUMNS
                                                 fields, "—" read as None
      {"status": "not_computed", "error", "note"}
      {"status": "error", "error"}              not a sheet this module wrote
    `sha` is the first 8 hex of costs.costs_sig at run time and
    `economics_sha` of costs.economics_sig ("none" when economics.yaml was
    missing; None when the note does not carry it), so a caller can tell
    when costs.yaml or the price bands have changed since the run."""
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    if not rows or _at(rows[0], 0) != TITLE:
        return {"status": "error",
                "error": f"{COSTS_SHEET}: row 1 is not {TITLE!r}"}
    note = _at(rows[1], 0) if len(rows) > 1 else None
    note = "" if note is None else str(note)
    if note.startswith(NOT_COMPUTED):
        return {"status": "not_computed", "note": note,
                "error": note[len(NOT_COMPUTED):].strip().lstrip("—").strip()}
    hdr = rows[HEADER_ROW - 1] if len(rows) >= HEADER_ROW else ()
    m = re.fullmatch(r"Feed \((.+)\)", str(_at(hdr, 3) or ""))
    cur = m.group(1) if m else None
    if cur is None or [_at(hdr, i) for i in range(len(COLUMNS))] \
            != _headers(COLUMNS, cur):
        return {"status": "error",
                "error": f"{COSTS_SHEET}: row {HEADER_ROW} is not the "
                         f"expected header"}
    sha = re.search(r"config/costs\.yaml \(sha ([0-9a-f]{8})\)", note)
    esha = re.search(r"config/economics\.yaml \(sha ([0-9a-f]{8}|none)\)",
                     note)
    out = {"status": "ok", "sha": sha.group(1) if sha else None,
           "economics_sha": esha.group(1) if esha else None,
           "currency": cur, "note": note, "months": [], "years": [],
           "total": None, "feed_by_type": [], "feed_total": None}
    i = HEADER_ROW
    while i < len(rows) and not _blank(rows[i]):
        d = {f: _back(_at(rows[i], c)) for c, (_h, f) in enumerate(COLUMNS)}
        kind = str(d["kind"] or "")
        if kind.startswith(KIND_MONTH):
            out["months"].append(d)
        elif kind == "Year":
            out["years"].append(d)
        elif kind == "Total":
            out["total"] = d
        i += 1
    while i < len(rows) and _at(rows[i], 0) != FEED_BLOCK:
        i += 1
    i += 2                                   # the block title, its header
    while i < len(rows) and not _blank(rows[i]):
        d = {f: _back(_at(rows[i], c))
             for c, (_h, f) in enumerate(FEED_COLUMNS)}
        if d["item"] is None:
            d["item"] = ""
        if d["feed_type"] == "Total":
            out["feed_total"] = d
        else:
            out["feed_by_type"].append(d)
        i += 1
    return out
