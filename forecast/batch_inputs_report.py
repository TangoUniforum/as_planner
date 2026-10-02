"""BatchInputs sheet: what the run was TOLD about each batch, next to what it DID.

Operator request, 2026-10-02: "add an output tab to the workbook that includes
the inputs from the config files, like the input dates for each batch, the
transfer dates, the counts input and transferred to OG, and the models used per
batch." RunConfig already carries the same facts, but as a raw YAML dump that
cannot be read per batch, and the Batch Plan is the journey, not the inputs.

REPORT ONLY. Nothing here is read back by the planner. Every input column comes
from the scenario as loaded (a copy taken BEFORE FW auto-calibration rewrote
`fw_correction` in place); every "run" column comes from the run's own objects
-- the TranOG events the TransferPlan prints and the calibration records the
ValidationLog narrates -- so the sheet reconciles with those sheets by
construction rather than by a second calculation.

Written AFTER the presentation pass and after CostsAndProfit, appended last, so
no existing sheet's index or style ids move (same rule as the costs sheet).
"""
from __future__ import annotations

from datetime import date, datetime

from .batch_order import batch_sort_key
from .time_grid import iso_week_label

SHEET_NAME = "BatchInputs"

# One row per batch. Header text drives excel_format's number formats, so the
# units live in the header (e.g. "(g)" -> whole grams, "count" -> #,##0).
COLUMNS = [
    ("Batch", "Batch id."),
    ("Source", "Where the run got this batch: in OG at the PR close, in "
               "freshwater at the PR close, split across both, scheduled "
               "(not yet in the PR), or pre-start (transferred before the "
               "PR and no longer on site)."),
    ("Input date", "batches.yaml input_date (egg input)."),
    ("Input count (eggs)", "batches.yaml input_count."),
    ("Tran SF date", "batches.yaml tran_sf_date (start-feed transfer)."),
    ("Tran OG date (input)", "batches.yaml tran_og_date - the planned "
                             "freshwater -> OG transfer date."),
    ("TranOG count (input, fish)", "batches.yaml tran_og_count - the fish the "
                                   "scenario expects to transfer to OG."),
    ("TranOG avg wt (input, g)", "batches.yaml tran_og_avg_wt_g - the target "
                                 "pre-cull weight at transfer."),
    ("TranOG CV (%)", "batches.yaml tran_og_cv - size spread at transfer."),
    ("FCR model", "batches.yaml fcr_model, and in brackets the biology.yaml "
                  "fcr_by_model curve it selects."),
    ("FW correction (input)", "batches.yaml fw_correction - multiplier on the "
                              "freshwater SGR curve."),
    ("FW correction (used)", "The fw_correction the run actually used. Differs "
                             "from the input when auto_calibrate_fw rewrote it "
                             "to land the target weight at transfer."),
    ("FW calibration", "What auto-calibration did (rewritten / clamped / did "
                       "not converge / off / not applicable)."),
    ("SGR correction (SW)", "batches.yaml sgr_correction - multiplier on the "
                            "seawater SGR curve."),
    ("OG entry week (run)", "ISO week of the batch's freshwater -> OG transfer "
                            "in this plan (TransferPlan TranOG rows)."),
    ("OG entry date (run)", "Date of that transfer."),
    ("Count into OG (run, fish)", "Fish the plan transfers into OG - the sum of "
                                  "this batch's TranOG rows on the TransferPlan."),
    ("Avg wt at OG entry (run, g)", "Count-weighted average weight of those "
                                    "TranOG rows."),
    ("Entry tanks (run)", "OG tanks the batch enters."),
    ("Manual fw_to_og", "Set when the transfer was scripted in the manual "
                        "override window (window week and ISO week)."),
    ("Notes", "batches.yaml notes."),
]


def _d(v):
    """Plain date for a date/datetime/None (no time-of-day in this project)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return v


def _source(bid, b, *, pr_og, pr_fw, forecast_start) -> str:
    """Judged on the ProductionReport itself, NOT on the post-window state: a
    window that harvests a batch out or scripts its fw_to_og must not change
    where the batch was when the PR closed."""
    tog = _d(b.tran_og_date)
    fs = _d(forecast_start)
    if bid in pr_og and bid in pr_fw:
        return "In-flight: split (SW part in OG + FW part in freshwater)"
    if bid in pr_og:
        return "In-flight: in OG at PR close"
    if bid in pr_fw:
        if tog is not None and fs is not None and tog < fs:
            return ("In-flight: in freshwater at PR close "
                    "(transfer date already passed)")
        return "In-flight: in freshwater at PR close"
    if tog is not None and fs is not None and tog < fs:
        return "Pre-start: transferred before the PR, not on site"
    return "Scheduled: batches.yaml (not yet in the PR)"


def _fcr_label(fcr_model, tables) -> str:
    from .biology import _fcr_model_key
    if not fcr_model:
        return ""
    key = _fcr_model_key(fcr_model)
    known = key in (getattr(tables, "fcr_by_model", None) or {})
    return f"{fcr_model} ({key}{'' if known else ' - NOT in biology.yaml'})"


def _calib_text(rec, control, bid, in_og, manual) -> str:
    if not getattr(control, "auto_calibrate_fw", False):
        return "off (auto_calibrate_fw = false)"
    if rec is None:
        if in_og:
            return "not applicable (already in OG)"
        if manual:
            return "not applied (transfer scripted by manual fw_to_og)"
        return "not applied"
    if not rec.get("converged", True):
        return "did not converge - kept the input"
    if rec.get("clamped"):
        return (f"clamped to {rec.get('applied')} (needed {rec.get('solved')}; "
                f"bounds {rec.get('clamp_lo')}-{rec.get('clamp_hi')})")
    if (rec.get("configured") is not None and rec.get("applied") is not None
            and abs(rec["configured"] - rec["applied"]) < 1e-9):
        return "unchanged (input already on target)"
    tgt = rec.get("target_wt_g")
    return (f"rewritten to land {tgt:,.0f} g at transfer" if tgt
            else "rewritten")


def batch_input_rows(*, batches_input, batch_by_id_used, tranog_events,
                     control, tables, pr_og=(), pr_fw=(), forecast_start=None,
                     calib_records=(), manual_fw_weeks=None) -> list[list]:
    """The table body, one list per batch in batch-number order.

    `batches_input` is the scenario as loaded (before auto-calibration);
    `batch_by_id_used` the same batches as the run left them. Pure: no
    workbook access, so a test can check the arithmetic directly.
    """
    by_batch_tranog: dict = {}
    for ev in tranog_events or []:
        by_batch_tranog.setdefault(ev.batch_id, []).append(ev)
    calib = {}
    for r in calib_records or []:
        calib[r.get("batch")] = r        # last rewrite wins (it is the one used)
    manual_fw_weeks = manual_fw_weeks or {}

    rows = []
    for b in sorted(batches_input, key=lambda x: batch_sort_key(x.batch_id)):
        bid = b.batch_id
        used = batch_by_id_used.get(bid, b)
        evs = sorted(by_batch_tranog.get(bid, []),
                     key=lambda e: (_d(e.event_date), e.batch_id))
        dests = [d for e in evs for d in e.destinations]
        n_in = sum(float(d.count) for d in dests)
        wt_in = (sum(float(d.count) * float(d.avg_wt_g) for d in dests) / n_in
                 if n_in > 0 else None)
        weeks = []
        for e in evs:
            wk = iso_week_label(e.event_date)
            if wk not in weeks:
                weeks.append(wk)
        tanks = []
        for d in dests:
            if d.tank_id not in tanks:
                tanks.append(d.tank_id)
        in_og = bid in pr_og and bid not in pr_fw
        if evs:
            entry_week = ", ".join(weeks)
            entry_date = _d(evs[0].event_date)
        elif in_og:
            entry_week, entry_date = "in OG before the forecast", None
        else:
            entry_week, entry_date = "", None
        mw = manual_fw_weeks.get(bid)
        manual = ""
        if mw is not None:
            manual = (f"Yes - window week {mw[0]} ({mw[1]})" if mw[1]
                      else f"Yes - window week {mw[0]}")
        rows.append([
            bid,
            _source(bid, b, pr_og=pr_og, pr_fw=pr_fw,
                    forecast_start=forecast_start),
            _d(b.input_date),
            b.input_count,
            _d(b.tran_sf_date),
            _d(b.tran_og_date),
            b.tran_og_count,
            b.tran_og_avg_wt_g,
            b.tran_og_cv,
            _fcr_label(b.fcr_model, tables),
            b.fw_correction,
            used.fw_correction,
            _calib_text(calib.get(bid), control, bid, in_og, mw is not None),
            b.sgr_correction,
            entry_week,
            entry_date,
            (round(n_in) if evs else None),
            (round(wt_in, 1) if wt_in is not None else None),
            ", ".join(str(t) for t in tanks),
            manual,
            b.notes or "",
        ])
    return rows


def _settings_rows(control, pr_closing, forecast_start, horizon_weeks,
                   window_weeks) -> list[list]:
    def g(name):
        return getattr(control, name, None)
    rows = [
        ["PR closing date", _d(pr_closing), "ProductionReport closing date"],
        ["Forecast start", _d(forecast_start), "PR closing + 1 day"],
        ["Horizon (weeks)", horizon_weeks, "control.yaml horizon_weeks"],
    ]
    if window_weeks:
        rows.append(["Manual override window (weeks)", window_weeks,
                     "last scripted week in scenario/manual_events"])
    rows += [
        ["scenario_name", g("scenario_name"), "control.yaml"],
        ["placement_method", g("placement_method"), "control.yaml"],
        ["tran_og_default_tanks", g("tran_og_default_tanks"),
         "control.yaml - OG1/2 tanks a transfer lands in at minimum"],
        ["handling_mortality_pct", g("handling_mortality_pct"),
         "control.yaml - mortality applied on handling"],
        ["default_hog_yield", g("default_hog_yield"),
         "control.yaml - head-on gutted yield"],
        ["grade_efficiency", g("grade_efficiency"),
         "control.yaml - 1.0 = perfect grader"],
        ["auto_calibrate_fw", g("auto_calibrate_fw"),
         "control.yaml - rewrite fw_correction to land tran_og_avg_wt_g"],
        ["auto_calibrate_fw_min", g("auto_calibrate_fw_min"), "control.yaml"],
        ["auto_calibrate_fw_max", g("auto_calibrate_fw_max"), "control.yaml"],
    ]
    return rows


def _model_rows(tables) -> list[list]:
    rows = []
    sz = list(getattr(tables, "sgr_size_g", []) or [])
    fw = [v for v in (getattr(tables, "sgr_fw_pct_day", []) or []) if v is not None]
    sw = [v for v in (getattr(tables, "sgr_sw_pct_day", []) or []) if v is not None]
    if sz:
        rows.append(["SGR curves (size -> %/day)",
                     f"{len(sz)} size points, {min(sz):g}-{max(sz):g} g",
                     f"biology.yaml sgr_fw_pct_day ({len(fw)} values), "
                     f"sgr_sw_pct_day ({len(sw)} values)"])
    fcr = getattr(tables, "fcr_by_model", None) or {}
    if fcr:
        rows.append(["FCR models available", ", ".join(sorted(fcr)),
                     "biology.yaml fcr_by_model - each batch's fcr_model "
                     "selects one (column FCR model)"])
    mw = list(getattr(tables, "mortality_week_from_input", []) or [])
    if mw:
        rows.append(["Mortality table",
                     f"{len(mw)} rows, weeks {min(mw)}-{max(mw)} from input",
                     "biology.yaml mortality_pct_weekly - shared by all batches"])
    cull = list(getattr(tables, "culling", []) or [])
    if cull:
        rows.append(["FW culling schedule",
                     "; ".join(f"day {d}: {p:g}%" for d, p in cull),
                     "biology.yaml culling - days since input"])
    ft = list(getattr(tables, "feed_types", []) or [])
    if ft:
        rows.append(["Feed types", f"{len(ft)} size brackets",
                     "biology.yaml feed_types"])
    og = getattr(tables, "og_sgr_by_week", None) or {}
    off = {k: v for k, v in og.items() if v is not None and abs(v - 1.0) > 1e-9}
    if off:
        wks = sorted(off)
        rows.append(["OG SGR per-week factors",
                     f"{len(off)} week(s) != 1.0, {wks[0]}-{wks[-1]}, "
                     f"range {min(off.values()):g}-{max(off.values()):g}",
                     "scenario/limits.yaml sgr_correction_og"])
    else:
        rows.append(["OG SGR per-week factors", "none (all weeks 1.0)",
                     "scenario/limits.yaml sgr_correction_og"])
    return rows


def write_batch_inputs_sheet(wb, *, batches_input, batch_by_id_used,
                             tranog_events, control, tables, pr_closing=None,
                             forecast_start=None, horizon_weeks=None,
                             window_weeks=0, pr_og=(), pr_fw=(),
                             calib_records=(),
                             manual_fw_weeks=None,
                             sheet_name: str = SHEET_NAME) -> int:
    """Append the BatchInputs sheet (replacing a stale one). Returns rows written."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    rows = batch_input_rows(
        batches_input=batches_input, batch_by_id_used=batch_by_id_used,
        tranog_events=tranog_events, control=control, tables=tables,
        pr_og=pr_og, pr_fw=pr_fw, forecast_start=forecast_start,
        calib_records=calib_records, manual_fw_weeks=manual_fw_weeks)

    ws.append(["BATCH INPUTS - what the scenario told the run about each "
               "batch, next to what the run did with it"])
    ws.append(["Input columns are scenario/batches.yaml as loaded; '(used)' "
               "and '(run)' columns are this plan (TransferPlan TranOG rows, "
               "ValidationLog FW calibration). Column meanings, run settings "
               "and the biology models in force are listed below the table."])
    ws.append([])
    # Row numbers from the append cursor, not max_row: an appended BLANK row
    # creates no cells, so max_row would point the header one row too high.
    ws.append([c for c, _m in COLUMNS])
    header_row = ws._current_row
    for r in rows:
        ws.append(r)
    last_batch_row = ws._current_row

    ws.append([])
    sec_rows = []
    ws.append(["RUN SETTINGS THAT SHAPE BATCH INPUTS"]); sec_rows.append(ws._current_row)
    for r in _settings_rows(control, pr_closing, forecast_start,
                            horizon_weeks, window_weeks):
        ws.append(r)
    ws.append([])
    ws.append(["BIOLOGY MODELS IN FORCE (shared by all batches; per-batch "
               "choices are the FCR model and the two corrections)"])
    sec_rows.append(ws._current_row)
    for r in _model_rows(tables):
        ws.append(r)
    ws.append([])
    ws.append(["COLUMN MEANINGS"]); sec_rows.append(ws._current_row)
    for name, meaning in COLUMNS:
        ws.append([name, meaning])

    _style(ws, header_row, last_batch_row, sec_rows)
    return len(rows)


def _style(ws, header_row, last_batch_row, sec_rows) -> None:
    """Same look as the other table sheets (excel_format), then keep the
    filter on the batch table only and mark the sections below it."""
    try:
        from openpyxl.styles import Alignment
        from . import excel_format as xf
        xf._format_table(ws, header_row)
        ws.auto_filter.ref = (f"A{header_row}:"
                              f"{ws.cell(header_row, len(COLUMNS)).column_letter}"
                              f"{last_batch_row}")
        ws.sheet_properties.tabColor = xf.TAB_INPUT
        for r in sec_rows:
            ws.cell(r, 1).font = xf.F_HDR
            ws.cell(r, 1).fill = xf.C_HDR
        widths = {"A": 30, "B": 44, "C": 12, "J": 26, "M": 34, "O": 18,
                  "S": 14, "T": 26, "U": 30}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        # The corrections are multipliers: 4 places, not a raw float tail.
        corr_cols = [i for i, (h, _m) in enumerate(COLUMNS, start=1)
                     if "correction" in h.lower()]
        for r in range(header_row + 1, last_batch_row + 1):
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
            for c in corr_cols:
                if isinstance(ws.cell(r, c).value, (int, float)):
                    ws.cell(r, c).number_format = "0.0000"
        # Below the table the B column holds settings values; keep the
        # meanings readable rather than clipped.
        for r in range(last_batch_row + 2, ws.max_row + 1):
            for c in (2, 3):
                ws.cell(r, c).alignment = Alignment(wrap_text=False)
    except Exception:  # noqa: BLE001 - presentation only, never fail a run
        pass
