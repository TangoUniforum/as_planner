"""Growth history: the seawater growth the Production Reports actually show, kept
as every PR is added, and the ACTUAL growth model fitted from it.

Operator plan 2026-10-18 ("as we add new PRs it builds on the history to improve
the actual growth model"). Three pieces, all data, none of them read by the
engine -- a fitted curve only reaches a plan when the operator copies a frozen
version into biology.yaml sgr_sw_by_model and a batch picks it (growth_model):

1. SNAPSHOTS. One per Production Report added: closing date, kind and the OG
   tank rows (batch, tank, count, mean weight). Kind is FULL when the workbook
   has the accumulated (Acc) tab -- the full monthly report -- else MID-MONTH
   (the PR sheet only).
2. THE CHAIN. One snapshot per closing date, chosen by a fixed rule, and only
   one source per month: a full report supersedes every mid-month PR of its
   month; in a month with no full report yet, only the latest mid-month PR is
   used and its interval is PROVISIONAL. Consecutive chain snapshots give the
   CLEAN TANK INTERVALS: the same batch in the same tank at both ends, no fish
   added and under 5% lost, not a 6N tank. Anything else -- a harvest, grade
   or move -- changes the mean weight without growth and is left out.
3. THE FIT (the hindcast winner of 2026-10-17, "model B"): grams per day by
   size, AGR = a x w^b below a breakpoint (2 kg) and a flat cap above it,
   fitted on the intervals that end in the rolling window (24 months = one full
   life cycle), each weighted by its days (a 4-day mid-month interval counts
   4/30 of a month). It is converted to an SGR table on the biology size rows.
   Each fit is saved as a FROZEN version: adding PRs later never changes a saved
   version, it makes the next one.
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import math
import os
import re
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Optional

FULL, MID = "full", "mid-month"
STORE_FORMAT = 1
# A tank interval losing more than this share of its fish was harvested,
# graded or moved, so its mean-weight change is not growth.
CLEAN_LOSS = 0.05
SIXN_SYSTEM = "OG6N"
DEFAULT_WINDOW_MONTHS = 24
DEFAULT_BREAKPOINT_G = 2000.0
DEFAULT_SW_MIN_G = 150.0          # biology rows below this keep the base value
MIN_INTERVALS = 20                # per side of the breakpoint
HOLDOUT_MONTHS = 3


# --------------------------------------------------------------------------- #
# Snapshots
# --------------------------------------------------------------------------- #
@dataclass
class Snapshot:
    closing: str                  # ISO date
    kind: str                     # FULL or MID
    source: str                   # file name as added
    sha1: str
    added: str                    # ISO timestamp, the order of adding
    og: list = field(default_factory=list)   # [batch, tank, count, avg_wt_g]
    modified: str = ""            # the file's last-modified time (ISO)

    @property
    def closing_date(self) -> date:
        return date.fromisoformat(self.closing)

    @property
    def prelim(self) -> bool:
        return "prelim" in self.source.lower()


def _sha1(path: Path) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_snapshot(path, *, added: Optional[str] = None) -> Snapshot:
    """One Production Report -> Snapshot. Raises ValueError (naming the file)
    when the workbook has no readable ProductionReport sheet."""
    import openpyxl
    from .accumulated import find_acc_sheet
    from .production_report import read_production_report
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            try:
                closing, og, _fw = read_production_report(wb)
            except Exception as e:  # noqa: BLE001 - reported by name below
                raise ValueError(f"{path.name}: no readable ProductionReport "
                                 f"({type(e).__name__}: {str(e)[:120]})") from e
            kind = FULL if find_acc_sheet(wb) is not None else MID
    finally:
        wb.close()
    if closing is None:
        raise ValueError(f"{path.name}: the ProductionReport has no closing date")
    closing = closing.date() if isinstance(closing, datetime) else closing
    return Snapshot(
        closing=closing.isoformat(), kind=kind, source=path.name,
        sha1=_sha1(path), added=added or datetime.now().isoformat(timespec="seconds"),
        modified=datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
        og=[[r.batch_id, int(r.tank_id), float(r.closing_count),
             float(r.closing_avg_wt_g)] for r in og])


# A report with fewer tank rows than this share of the fullest one for the
# same closing date is an extract (a few batches), never the whole report.
EXTRACT_SHARE = 0.8


def _rank(s: Snapshot, most_rows: int) -> tuple:
    """Higher wins for one closing date: full > mid-month, final > prelim, a
    whole report > a few-batch extract, then the most recently modified file
    (V3 over V2), then the later add."""
    return (s.kind == FULL, not s.prelim, len(s.og) >= EXTRACT_SHARE * most_rows,
            s.modified, s.added)


def choose_chain(snapshots) -> tuple[list, list]:
    """(chain, superseded): one snapshot per closing date, one source per month.

    Superseded entries are (snapshot, reason) so the operator can see why a
    report was not used."""
    by_date: dict[str, list] = {}
    for s in snapshots:
        by_date.setdefault(s.closing, []).append(s)
    best: dict[str, Snapshot] = {}
    superseded = []
    for closing, group in by_date.items():
        most = max(len(s.og) for s in group)
        ranked = sorted(group, key=lambda s: _rank(s, most), reverse=True)
        best[closing] = ranked[0]
        for s in ranked[1:]:
            superseded.append((s, f"same closing date as {ranked[0].source}"))
    full_months = {s.closing[:7] for s in best.values() if s.kind == FULL}
    latest_mid: dict[str, Snapshot] = {}
    chain = []
    for s in sorted(best.values(), key=lambda x: x.closing):
        m = s.closing[:7]
        if s.kind == MID and m in full_months:
            superseded.append((s, "the full report for its month is in"))
            continue
        if s.kind == MID:
            prev = latest_mid.get(m)
            if prev is not None:
                superseded.append((prev, f"a later mid-month PR ({s.source}) "
                                         f"is in for {m}"))
                chain.remove(prev)
            latest_mid[m] = s
        chain.append(s)
    return chain, superseded


# --------------------------------------------------------------------------- #
# Clean tank intervals
# --------------------------------------------------------------------------- #
@dataclass
class Interval:
    batch: str
    tank: int
    d0: str
    d1: str
    days: int
    c0: float
    c1: float
    w0: float
    w1: float
    provisional: bool             # an end is a mid-month PR

    @property
    def wmid(self) -> float:
        return math.sqrt(self.w0 * self.w1)

    @property
    def agr(self) -> float:
        return (self.w1 - self.w0) / self.days

    @property
    def sgr(self) -> float:
        return 100.0 * math.log(self.w1 / self.w0) / self.days

    @property
    def weight(self) -> float:
        """Days covered, capped at one month: a 4-day interval counts 4/30."""
        return min(self.days, 30) / 30.0


def clean_intervals(chain, sixn_tanks, og_tanks=None) -> list:
    """Clean tank intervals between consecutive chain snapshots.

    `sixn_tanks`: tank ids of the 6N purge/harvest system (excluded).
    `og_tanks`: if given, only these tank ids count (unknown ids are skipped).
    """
    sixn = set(sixn_tanks)
    out = []
    for s0, s1 in zip(chain, chain[1:]):
        days = (s1.closing_date - s0.closing_date).days
        if days <= 0:
            continue
        a = {(b, t): (c, w) for b, t, c, w in s0.og}
        b_ = {(b, t): (c, w) for b, t, c, w in s1.og}
        for key in sorted(set(a) & set(b_)):
            batch, tank = key
            if tank in sixn or (og_tanks is not None and tank not in og_tanks):
                continue
            (c0, w0), (c1, w1) = a[key], b_[key]
            if w0 <= 0 or w1 <= 0 or c0 <= 0:
                continue
            if c1 > c0 * 1.001 or c1 < c0 * (1.0 - CLEAN_LOSS):
                continue
            out.append(Interval(batch, tank, s0.closing, s1.closing, days,
                                c0, c1, w0, w1,
                                provisional=(s0.kind == MID or s1.kind == MID)))
    return out


def facility_tank_sets(facility) -> tuple[set, set]:
    """(sixn_tank_ids, og_tank_ids) from a FacilityConfig."""
    sixn = {t.tank_id for t in facility.tanks if t.system_id == SIXN_SYSTEM}
    og = {t.tank_id for t in facility.tanks if t.type == "OG"}
    return sixn, og


# --------------------------------------------------------------------------- #
# The fit
# --------------------------------------------------------------------------- #
def _months_before(d: date, months: int) -> date:
    y, m = divmod(d.year * 12 + (d.month - 1) - months, 12)
    m += 1
    last = [31, 29 if (y % 4 == 0 and (y % 100 or y % 400 == 0)) else 28,
            31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(y, m, min(d.day, last))


def _wmedian(vals, wts) -> float:
    pairs = sorted(zip(vals, wts))
    half = sum(w for _, w in pairs) / 2.0
    acc = 0.0
    for v, w in pairs:
        acc += w
        if acc >= half:
            return v
    return pairs[-1][0]


@dataclass
class GrowthFit:
    name: str
    as_of: str
    window_months: int
    window_start: str
    breakpoint_g: float
    a: float
    b: float
    cap_g_day: float
    n: int
    n_provisional: int
    n_below: int
    n_above: int
    sizes_g: list
    sgr_pct_day: list             # one per biology size row
    holdout: Optional[dict] = None
    sources: list = field(default_factory=list)
    created: str = ""
    # Batches the operator left out of the fit (e.g. a disease episode). Their
    # tank-months are dropped before the window, the fit and the holdout.
    excluded_batches: list = field(default_factory=list)

    def agr_g_day(self, w: float) -> float:
        return min(self.a * w ** self.b, self.cap_g_day)


def _fit_params(ivs, breakpoint_g):
    lo = [i for i in ivs if i.wmid < breakpoint_g and i.agr > 0]
    hi = [i for i in ivs if i.wmid >= breakpoint_g]
    if len(lo) < MIN_INTERVALS or len(hi) < MIN_INTERVALS:
        raise ValueError(
            f"not enough clean tank intervals to fit: {len(lo)} below "
            f"{breakpoint_g:,.0f} g and {len(hi)} above (need {MIN_INTERVALS} each)")
    xs = [math.log(i.wmid) for i in lo]
    ys = [math.log(i.agr) for i in lo]
    ws = [i.weight for i in lo]
    sw = sum(ws)
    mx = sum(w * x for w, x in zip(ws, xs)) / sw
    my = sum(w * y for w, y in zip(ws, ys)) / sw
    b = (sum(w * (x - mx) * (y - my) for w, x, y in zip(ws, xs, ys))
         / sum(w * (x - mx) ** 2 for w, x in zip(ws, xs)))
    a = math.exp(my - b * mx)
    cap = _wmedian([i.agr for i in hi], [i.weight for i in hi])
    return a, b, cap, len(lo), len(hi)


def _agr_to_sgr(agr: float, w: float) -> float:
    return 100.0 * math.log(1.0 + agr / w)


def _predict(w0: float, days: int, sgr_fn) -> float:
    w = w0
    for _ in range(days):
        w *= 1.0 + sgr_fn(w) / 100.0
    return w


def _holdout(ivs, as_of: date, breakpoint_g, sizes, base_sw) -> Optional[dict]:
    """Fit without the last HOLDOUT_MONTHS of intervals, predict them, and score
    the end weight -- the same for the base curve. None if too little data."""
    from .biology import _interp
    cut = _months_before(as_of, HOLDOUT_MONTHS).isoformat()
    train = [i for i in ivs if i.d1 <= cut]
    test = [i for i in ivs if i.d1 > cut]
    if not test:
        return None
    try:
        a, b, cap, _, _ = _fit_params(train, breakpoint_g)
    except ValueError:
        return None
    fit_fn = lambda w: _agr_to_sgr(min(a * w ** b, cap), w)       # noqa: E731
    base_fn = lambda w: _interp(w, sizes, base_sw)                 # noqa: E731
    out = {"n": len(test), "from": cut}
    for label, fn in (("fit", fit_fn), ("base", base_fn)):
        errs = [100.0 * (_predict(i.w0, i.days, fn) - i.w1) / i.w1 for i in test]
        errs.sort()
        absd = sorted(abs(e) for e in errs)
        out[label] = {"median_err_pct": round(errs[len(errs) // 2], 2),
                      "median_abs_err_pct": round(absd[len(absd) // 2], 2)}
    return out


def fit_actual(intervals, *, sizes, base_sw, as_of, name=None,
               window_months=DEFAULT_WINDOW_MONTHS,
               breakpoint_g=DEFAULT_BREAKPOINT_G,
               sw_min_g=DEFAULT_SW_MIN_G, sources=(),
               exclude_batches=()) -> GrowthFit:
    """Fit the actual growth model on the intervals ending inside the window.

    `sizes`/`base_sw`: biology.yaml sgr_size_g / sgr_sw_pct_day. Rows below
    `sw_min_g` keep the base value (no seawater fish are that small).
    `exclude_batches`: batch ids whose tank-months are left out."""
    as_of = as_of if isinstance(as_of, date) else date.fromisoformat(str(as_of))
    start = _months_before(as_of, window_months)
    out = {str(b).strip() for b in exclude_batches if str(b).strip()}
    ivs = [i for i in intervals
           if start.isoformat() < i.d1 <= as_of.isoformat() and i.batch not in out]
    a, b, cap, n_lo, n_hi = _fit_params(ivs, breakpoint_g)
    table = []
    for w, base in zip(sizes, base_sw):
        if w < sw_min_g:
            table.append(base)
        else:
            table.append(round(_agr_to_sgr(min(a * w ** b, cap), w), 4))
    return GrowthFit(
        name=name or f"Actual {as_of.isoformat()}", as_of=as_of.isoformat(),
        window_months=window_months, window_start=start.isoformat(),
        breakpoint_g=float(breakpoint_g), a=a, b=b, cap_g_day=cap,
        n=len(ivs), n_provisional=sum(i.provisional for i in ivs),
        n_below=n_lo, n_above=n_hi, sizes_g=list(sizes), sgr_pct_day=table,
        holdout=_holdout(ivs, as_of, breakpoint_g, list(sizes), list(base_sw)),
        sources=list(sources),
        created=datetime.now().isoformat(timespec="seconds"),
        excluded_batches=sorted(out, key=lambda b: (int(re.sub(r"\D", "", b) or 0), b)))


# --------------------------------------------------------------------------- #
# Check against the data (operator 2026-10-20: "is there a table that shows the
# result from the historical relative to the 100% model ... as % SGR/day")
# --------------------------------------------------------------------------- #
CHECK_BANDS_G = (0, 500, 1000, 1500, 2000, 2500, 3000, 3500, 4000)
CHECK_COLUMNS = ["Fish size", "Tank-months", "Measured SGR (%/day)",
                 "Model SGR (%/day)", "100% curve SGR (%/day)",
                 "Model as % of 100%", "Measured as % of 100%"]


def _excluded(fit) -> set:
    return set(getattr(fit, "excluded_batches", None) or [])


def fit_intervals(fit, intervals) -> list:
    """The tank-months a saved fit was made from: its window, without the
    batches it left out. The same rule as fit_actual."""
    out = _excluded(fit)
    return [i for i in intervals
            if fit.window_start < i.d1 <= fit.as_of and i.batch not in out]


def later_intervals(fit, intervals) -> list:
    """Tank-months that ended AFTER the fit's as-of date: data the frozen fit
    has never seen. This set grows as PRs are added, so the check of a saved
    version keeps testing it. Batches it left out stay out."""
    out = _excluded(fit)
    return [i for i in intervals if i.d1 > fit.as_of and i.batch not in out]


def _sgr_over(w0: float, days: int, sgr_fn) -> float:
    """The SGR (%/day) a curve gives over `days` from `w0`: daily steps, the
    same way the engine grows a tank."""
    return 100.0 * math.log(_predict(w0, days, sgr_fn) / w0) / days


def check_rows(model_sizes, model_sgr, intervals, sizes, base_sw,
               bands=CHECK_BANDS_G) -> list[list]:
    """One row per fish-size band (by each tank-month's mid weight), in
    CHECK_COLUMNS order: tank-months; the measured SGR; the SGR the model and
    the 100% curve (biology.yaml sgr_sw_pct_day) give over the same tank-months
    from the same start weight; the model and the measurement as % of the 100%
    curve. Medians weighted by days, as in the fit. Bands with no tank-months
    are left out. Report only."""
    from .biology import _interp
    model_fn = lambda w: _interp(w, list(model_sizes), list(model_sgr))  # noqa: E731
    base_fn = lambda w: _interp(w, list(sizes), list(base_sw))           # noqa: E731
    edges = list(bands) + [math.inf]
    rows = []
    for lo, hi in zip(edges, edges[1:]):
        b = [i for i in intervals
             if lo <= i.wmid < hi and i.days > 0 and i.w0 > 0 and i.w1 > 0]
        if not b:
            continue
        wts = [i.weight for i in b]
        meas = _wmedian([100.0 * math.log(i.w1 / i.w0) / i.days for i in b], wts)
        mod = _wmedian([_sgr_over(i.w0, i.days, model_fn) for i in b], wts)
        base = _wmedian([_sgr_over(i.w0, i.days, base_fn) for i in b], wts)
        label = f"{lo:,}-{hi:,} g" if hi != math.inf else f"{lo:,}+ g"
        pct = (lambda v: round(100.0 * v / base)) if base > 0 else (lambda v: None)
        rows.append([label, len(b), round(meas, 3), round(mod, 3), round(base, 3),
                     pct(mod), pct(meas)])
    return rows


# --------------------------------------------------------------------------- #
# The store
# --------------------------------------------------------------------------- #
class GrowthHistory:
    """The snapshots and the frozen fits, in one JSON file.

    There is no default path on purpose: every caller names the file, so a
    test or a scratch run can never write the operator's history by accident.
    """

    def __init__(self, path):
        self.path = Path(path)
        self.snapshots: list[Snapshot] = []
        self.fits: list[GrowthFit] = []
        self.derived: list[DerivedModel] = []
        if self.path.exists():
            d = json.loads(self.path.read_text(encoding="utf-8"))
            if d.get("format") != STORE_FORMAT:
                raise ValueError(f"{self.path}: unknown growth history format "
                                 f"{d.get('format')!r}")
            self.snapshots = [Snapshot(**s) for s in d.get("snapshots", [])]
            self.fits = [GrowthFit(**f) for f in d.get("fits", [])]
            self.derived = [DerivedModel(**m) for m in d.get("derived", [])]

    def add(self, snap: Snapshot) -> str:
        """'duplicate' (same file content already in), else 'added'."""
        if any(s.sha1 == snap.sha1 for s in self.snapshots):
            return "duplicate"
        self.snapshots.append(snap)
        return "added"

    def chain(self):
        return choose_chain(self.snapshots)

    def intervals(self, sixn_tanks, og_tanks=None):
        return clean_intervals(self.chain()[0], sixn_tanks, og_tanks)

    def add_fit(self, fit: GrowthFit) -> None:
        """Frozen: a name is never overwritten with different values."""
        if any(m.name == fit.name for m in self.derived):
            raise ValueError(f"{fit.name!r} is already an ASSUMED version - "
                             f"fit under a new name")
        for f in self.fits:
            if f.name == fit.name:
                if f.sgr_pct_day == fit.sgr_pct_day:
                    return
                raise ValueError(
                    f"growth model version {fit.name!r} already exists with "
                    f"different values - versions are frozen; fit under a new name")
        self.fits.append(fit)

    def get_fit(self, name: str) -> GrowthFit:
        for f in self.fits:
            if f.name == name:
                return f
        raise KeyError(f"no growth model version {name!r} "
                       f"(have: {[f.name for f in self.fits]})")

    def add_derived(self, m: "DerivedModel") -> None:
        """Frozen like the fits; its source must be a saved measured version."""
        self.get_fit(m.source)
        if any(f.name == m.name for f in self.fits):
            raise ValueError(f"{m.name!r} is already a measured version")
        for o in self.derived:
            if o.name == m.name:
                if o.sgr_pct_day == m.sgr_pct_day:
                    return
                raise ValueError(
                    f"ASSUMED version {m.name!r} already exists with different "
                    f"values - versions are frozen; save under a new name")
        self.derived.append(m)

    def versions(self) -> list:
        """Every saved version, measured first, then ASSUMED."""
        return list(self.fits) + list(self.derived)

    def get_version(self, name: str):
        for v in self.versions():
            if v.name == name:
                return v
        raise KeyError(f"no growth model version {name!r}")

    def save(self) -> None:
        d = {"format": STORE_FORMAT,
             "snapshots": [asdict(s) for s in self.snapshots],
             "fits": [asdict(f) for f in self.fits]}
        if self.derived:
            d["derived"] = [asdict(m) for m in self.derived]
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=self.path.name + ".tmp-",
                                   dir=str(self.path.parent))
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(d, fh, indent=1)
        os.replace(tmp, self.path)


def biology_entry(fit, sizes) -> tuple[str, list]:
    """(name, values) for biology.yaml sgr_sw_by_model, from a measured or an
    ASSUMED version. Refuses one made on other size rows -- its values would
    land on the wrong sizes."""
    if [float(x) for x in fit.sizes_g] != [float(x) for x in sizes]:
        raise ValueError(f"{fit.name!r} was fitted on different biology size rows")
    return fit.name, list(fit.sgr_pct_day)


# --------------------------------------------------------------------------- #
# Expected improvement (step 4) -- ASSUMED, never measured
# --------------------------------------------------------------------------- #
# The history measures how fish HAVE grown. Planning better growth for fish
# still to come is an assumption the operator makes; these helpers turn one
# into a separate frozen version whose name says ASSUMED, built from a measured
# version. Both work in grams per day (what was measured), size row by size
# row; rows below DEFAULT_SW_MIN_G keep the source value. The engine holds one
# curve per batch for its whole life, so each version is made for the fish
# GROWING IN a stated month: pick it for the batches growing then.
ASSUMED = "ASSUMED"


@dataclass
class DerivedModel:
    name: str
    source: str                   # the measured version it is built from
    method: str                   # "pct_per_year" | "blend"
    params: dict
    growing_in: str               # the date the improvement is reckoned to
    factor: float                 # pct: the AGR multiplier; blend: share of the way
    explain: str                  # the plain sentence shown everywhere
    sizes_g: list
    sgr_pct_day: list
    created: str = ""
    assumed: bool = True


def _sgr_to_agr(sgr: float, w: float) -> float:
    return w * (math.exp(sgr / 100.0) - 1.0)


def _as_date(d) -> date:
    return d if isinstance(d, date) else date.fromisoformat(str(d)[:10])


def improve_pct_per_year(fit: GrowthFit, pct: float, growing_in, *, name=None,
                         sw_min_g=DEFAULT_SW_MIN_G) -> DerivedModel:
    """Grams per day x (1 + pct/100) ^ years from the version's as-of date to
    `growing_in`. A date before the as-of date gives factor 1 (no improvement
    backwards)."""
    if not -50.0 < float(pct) < 50.0:
        raise ValueError(f"{pct}% per year is outside -50..50")
    g = _as_date(growing_in)
    years = max(0.0, (g - _as_date(fit.as_of)).days / 365.25)
    k = (1.0 + float(pct) / 100.0) ** years
    table = [s if (w < sw_min_g or s is None)
             else round(_agr_to_sgr(_sgr_to_agr(s, w) * k, w), 4)
             for w, s in zip(fit.sizes_g, fit.sgr_pct_day)]
    explain = (f"{ASSUMED}: '{fit.name}' (measured to {fit.as_of}) with growth "
               f"{float(pct):+g}% per year, {years:.2f} years to {g.isoformat()} "
               f"-> grams per day x {k:.3f} at every seawater size")
    return DerivedModel(
        name=name or f"{fit.name} {float(pct):+g}%/yr for {g.strftime('%Y-%m')} ({ASSUMED})",
        source=fit.name, method="pct_per_year", params={"pct_per_year": float(pct)},
        growing_in=g.isoformat(), factor=round(k, 6), explain=explain,
        sizes_g=list(fit.sizes_g), sgr_pct_day=table,
        created=datetime.now().isoformat(timespec="seconds"))


def blend_to_target(fit: GrowthFit, target_name: str, target_sgr, target_date,
                    growing_in, *, name=None,
                    sw_min_g=DEFAULT_SW_MIN_G) -> DerivedModel:
    """Grams per day moved from the measured version toward a target curve (on
    the same size rows), in a straight line: none of the way at the version's
    as-of date, all of it at `target_date`, held there after."""
    t0, t1, g = _as_date(fit.as_of), _as_date(target_date), _as_date(growing_in)
    if t1 <= t0:
        raise ValueError(f"the target date {t1} must be after {fit.name!r}'s "
                         f"as-of date {t0}")
    if len(target_sgr) != len(fit.sizes_g):
        raise ValueError(f"target {target_name!r} has {len(target_sgr)} values, "
                         f"{fit.name!r} has {len(fit.sizes_g)} size rows")
    f = min(1.0, max(0.0, (g - t0).days / (t1 - t0).days))
    table = []
    for w, s, t in zip(fit.sizes_g, fit.sgr_pct_day, target_sgr):
        if w < sw_min_g or s is None or t is None:
            table.append(s)
            continue
        a0, a1 = _sgr_to_agr(s, w), _sgr_to_agr(t, w)
        table.append(round(_agr_to_sgr(a0 + f * (a1 - a0), w), 4))
    explain = (f"{ASSUMED}: '{fit.name}' (measured to {t0}) moving in a straight "
               f"line to '{target_name}' by {t1}; fish growing in {g.isoformat()} "
               f"are {f:.0%} of the way there")
    return DerivedModel(
        name=name or (f"{fit.name} -> {target_name} {f:.0%} for "
                      f"{g.strftime('%Y-%m')} ({ASSUMED})"),
        source=fit.name, method="blend",
        params={"target": target_name, "target_date": t1.isoformat()},
        growing_in=g.isoformat(), factor=round(f, 6), explain=explain,
        sizes_g=list(fit.sizes_g), sgr_pct_day=table,
        created=datetime.now().isoformat(timespec="seconds"))


# --------------------------------------------------------------------------- #
# Per-batch factor from its own history (step 4)
# --------------------------------------------------------------------------- #
# How a batch grew in its clean tank-months against what the measured version
# predicts for the same start weight and days: 1.10 = it put on 10% more
# weight than the fleet curve. Pulled toward 1.0 by BATCH_PRIOR_MONTHS: a
# batch with that many tank-months of evidence counts half, so one good month
# in one tank does not set a batch's whole future.
BATCH_PRIOR_MONTHS = 12.0


@dataclass
class BatchFactor:
    batch: str
    tank_months: float            # summed interval weights (a 30-day interval = 1)
    n: int
    n_provisional: int
    measured: float               # weighted median, not pulled toward 1.0
    suggested: float              # pulled toward 1.0
    first: str
    last: str


def batch_factors(intervals, fit: GrowthFit, *, prior_months=BATCH_PRIOR_MONTHS,
                  sw_min_g=DEFAULT_SW_MIN_G) -> list:
    """BatchFactor per batch with clean seawater tank-months in `fit`'s window,
    in batch order. Batches with none are not listed (fleet only = 1.0)."""
    from .biology import _interp
    sgr = lambda w: _interp(w, fit.sizes_g, fit.sgr_pct_day)        # noqa: E731
    by: dict = {}
    for i in intervals:
        if not (fit.window_start < i.d1 <= fit.as_of) or i.w0 < sw_min_g:
            continue
        pred = _predict(i.w0, i.days, sgr)
        if pred <= i.w0:
            continue
        by.setdefault(i.batch, []).append(
            (math.log(i.w1 / i.w0) / math.log(pred / i.w0), i))
    out = []
    for b, rows in by.items():
        wts = [i.weight for _, i in rows]
        m = _wmedian([r for r, _ in rows], wts)
        tm = sum(wts)
        out.append(BatchFactor(
            batch=b, tank_months=round(tm, 2), n=len(rows),
            n_provisional=sum(i.provisional for _, i in rows),
            measured=round(m, 3),
            suggested=round(1.0 + (m - 1.0) * tm / (tm + prior_months), 3),
            first=min(i.d0 for _, i in rows), last=max(i.d1 for _, i in rows)))
    key = lambda f: (int(re.sub(r"\D", "", f.batch) or 0), f.batch)  # noqa: E731
    return sorted(out, key=key)


_PR_NAME = re.compile(r"production report|(^|[ _])pr\.xlsx$", re.IGNORECASE)


def is_report_name(name: str) -> bool:
    """A file name that looks like a Production Report or a mid-month PR."""
    low = name.lower()
    return (low.endswith((".xlsx", ".xlsm")) and not low.startswith("~$")
            and "template" not in low and bool(_PR_NAME.search(name)))


# --------------------------------------------------------------------------- #
# App / run helpers (step 3)
# --------------------------------------------------------------------------- #
STORE_NAME = "growth_history.json"


def live_store_path() -> Path:
    """The operator's growth history, beside fw_calibration_history.jsonl.

    AS_GROWTH_HISTORY overrides it -- only so an app test can point the app at
    a throwaway file instead of the operator's."""
    env = os.environ.get("AS_GROWTH_HISTORY")
    if env:
        return Path(env)
    return Path(__file__).resolve().parent.parent / STORE_NAME


def record_pr(pr_path, store_path) -> str:
    """Add one Production Report to the history at `store_path` and save.

    Returns what happened, in words, for the run log. Hashes first, so a PR
    that is already in (the usual case: the same PR run again) is not re-read.
    """
    pr_path = Path(pr_path)
    store = GrowthHistory(store_path)
    sha = _sha1(pr_path)
    if any(s.sha1 == sha for s in store.snapshots):
        return f"already in ({pr_path.name})"
    snap = read_snapshot(pr_path)
    store.add(snap)
    store.save()
    return f"added {snap.closing} {snap.kind} ({pr_path.name})"


def describe_batch_growth(batch, tables) -> str:
    """The plain sentence for one batch: which curve, times which factor."""
    from .biology import BASE_GROWTH_MODEL
    name = getattr(batch, "growth_model", "") or BASE_GROWTH_MODEL
    known = (name == BASE_GROWTH_MODEL
             or name in (getattr(tables, "sgr_sw_by_model", None) or {}))
    corr = float(getattr(batch, "sgr_correction", 1.0) or 1.0)
    s = f"{name}" + ("" if known else " (NOT in biology.yaml - the run will stop)")
    if abs(corr - 1.0) > 1e-9:
        s += f" x {corr:g} (sgr_correction)"
    return s


def set_batch_growth(rows, first, last, *, model=None, sw_adjustment=None,
                     fw_adjustment=None):
    """Set growth for every batch from `first` to `last` (batch number order,
    both ends included) in batches.yaml-shaped rows: the seawater growth model
    (BASE_GROWTH_MODEL = blank), sgr_correction (SW) and fw_correction (FW).
    None leaves that field as it is. Returns (new rows, batch ids changed); the
    input rows are not modified. One click for what was a dropdown per batch
    (operator 2026-10-20: in-flight batches on the measured model, future ones
    on the base curve with an offset)."""
    from .batch_order import batch_sort_key
    from .biology import BASE_GROWTH_MODEL
    lo, hi = batch_sort_key(first), batch_sort_key(last)
    if lo > hi:
        raise ValueError(f"'{first}' comes after '{last}' - pick the first batch "
                         f"first")
    for what, v in (("SW adjustment", sw_adjustment),
                    ("FW adjustment", fw_adjustment)):
        if v is not None and not float(v) > 0:
            raise ValueError(f"{what} must be above 0 (1 = none), not {v}")
    out, changed = [], []
    for r in rows:
        r = dict(r)
        bid = r.get("batch_id")
        if bid not in (None, "") and lo <= batch_sort_key(bid) <= hi:
            if model is not None:
                r["growth_model"] = "" if model == BASE_GROWTH_MODEL else model
            if sw_adjustment is not None:
                r["sgr_correction"] = float(sw_adjustment)
            if fw_adjustment is not None:
                r["fw_correction"] = float(fw_adjustment)
            changed.append(bid)
        out.append(r)
    return out, changed


def add_growth_model_to_biology(config_dir, name: str, values) -> str:
    """Put a frozen version into biology.yaml sgr_sw_by_model (biology.yaml
    only). Plans do not change until a batch picks it. Returns 'added' or
    'already there'; refuses a name that holds different values."""
    from .config_io import dump_biology, load_biology_tables
    tables = load_biology_tables(config_dir)
    if len(values) != len(tables.sgr_size_g):
        raise ValueError(f"{name!r} has {len(values)} values but biology.yaml has "
                         f"{len(tables.sgr_size_g)} size rows")
    have = (tables.sgr_sw_by_model or {}).get(name)
    if have is not None:
        if list(have) == list(values):
            return "already there"
        raise ValueError(f"biology.yaml already has a different growth model named "
                         f"{name!r}")
    tables.sgr_sw_by_model = dict(tables.sgr_sw_by_model or {})
    tables.sgr_sw_by_model[name] = list(values)
    dump_biology(config_dir, tables)
    return "added"


def _monthly_harvest(out_path) -> dict:
    """{YYYY-MM: [fish, HOG kg]} from a run's HarvestReport sheet."""
    import openpyxl
    wb = openpyxl.load_workbook(out_path, read_only=True, data_only=True)
    m: dict = {}
    try:
        for r in wb["HarvestReport"].iter_rows(min_row=5, values_only=True):
            if not (r and isinstance(r[0], int) and hasattr(r[1], "strftime")):
                continue
            e = m.setdefault(r[1].strftime("%Y-%m"), [0.0, 0.0])
            e[0] += float(r[6] or 0)
            e[1] += float(r[8] or 0)
    finally:
        wb.close()
    return m


def compare_plan(pr_path, config_dir, scenario_dir, name: str, values, work_dir,
                 batches=None) -> list:
    """Before/after: the scenario as it is, and the same with `batches` (all
    when None) on growth model `name` -- both run on COPIES, recording nothing.
    Returns rows (month, fish before, HOG t before, fish after, HOG t after)."""
    import shutil
    import yaml
    from .run import main as run_main
    work_dir = Path(work_dir)
    rows_by = {}
    for arm in ("before", "after"):
        d = work_dir / arm
        if d.exists():
            shutil.rmtree(d)
        shutil.copytree(config_dir, d / "config",
                        ignore=shutil.ignore_patterns("costs.yaml*"))
        shutil.copytree(scenario_dir, d / "scenario")
        if arm == "after":
            bp = d / "config" / "biology.yaml"
            bd = yaml.safe_load(bp.read_text(encoding="utf-8")) or {}
            bd.setdefault("sgr_sw_by_model", {})[name] = list(values)
            bp.write_text(yaml.safe_dump(bd, sort_keys=False), encoding="utf-8")
            sp = d / "scenario" / "batches.yaml"
            sd = yaml.safe_load(sp.read_text(encoding="utf-8")) or {}
            for b in sd.get("batches", []):
                if batches is None or b.get("batch_id") in batches:
                    b["growth_model"] = name
            sp.write_text(yaml.safe_dump(sd, sort_keys=False), encoding="utf-8")
        out = d / ("out" + Path(pr_path).suffix)
        inp = d / ("pr" + Path(pr_path).suffix)
        shutil.copy(pr_path, inp)
        with contextlib.redirect_stdout(io.StringIO()):
            rc = run_main(str(inp), str(out), config_dir=str(d / "config"),
                          scenario_dir=str(d / "scenario"), calib_log_path="",
                          growth_history_path="")
        if rc != 0:
            raise RuntimeError(f"the {arm} run failed (rc={rc})")
        if not out.exists():
            out = next(iter(d.glob("out.*")))
        rows_by[arm] = _monthly_harvest(out)
    months = sorted(set(rows_by["before"]) | set(rows_by["after"]))
    rows = []
    for m in months:
        b = rows_by["before"].get(m, [0.0, 0.0])
        a = rows_by["after"].get(m, [0.0, 0.0])
        rows.append((m, round(b[0]), round(b[1] / 1000, 1), round(a[0]), round(a[1] / 1000, 1)))
    return rows


def scan_reports(store: GrowthHistory, dirs, since: str = "2024-11-01") -> dict:
    """Add every report under `dirs` closing on/after `since` (files already in
    are recognised by content hash and not re-read). Saves the store.
    Returns {'added': [...], 'already': n, 'before_since': n, 'unreadable': [...]}"""
    known = {s.sha1 for s in store.snapshots}
    out = {"added": [], "already": 0, "before_since": 0, "unreadable": []}
    for d in dirs:
        for p in sorted(Path(d).rglob("*")):
            if not (p.is_file() and is_report_name(p.name)):
                continue
            try:
                if _sha1(p) in known:
                    out["already"] += 1
                    continue
                s = read_snapshot(p)
            except (ValueError, OSError) as e:
                out["unreadable"].append(str(e))
                continue
            if s.closing < since:
                out["before_since"] += 1
                continue
            store.add(s)
            known.add(s.sha1)
            out["added"].append(f"{s.closing} {s.kind} {s.source}")
    store.save()
    return out
