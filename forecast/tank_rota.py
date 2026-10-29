"""Tank rota (fixed standard schedule): every batch handled the same way, in steps you choose.

You give the steps, e.g. entry 4 tanks from week 0, grow 6 tanks from week 21, +1 tank at week 32. The tool takes one
batch's life from a finished run (growth, mortality, 6N staging by week of life - the same harvest), and:

1. quick_check  - a tank count in a second: can these steps fit the facility at all? (a necessary condition only)
2. design_batch - one batch's tank plan, all rules hard (CP-SAT): fish per tank, the step moves, the 6N drain
3. assign_tanks - physical tanks for a batch every `cadence` weeks: entry by rule, grow by CP-SAT within the system
                  biomass and feed caps (x buffer, as the engine and the checker)
4. evaluate     - 3 cycles of steady state laid out week by week and judged by least_moves.check (the same
                  independent checker as the least-moves layout), plus a companion workbook

Steady state only: getting from today's plan into the rota is not worked out here. Changes nothing else.
Tiers follow least_moves: entry = OG1N/1S/2N/2S, grow = every other OG tank in the endpoint (OG6N is never in it).
"""
from __future__ import annotations

import collections
import math
import statistics

from forecast import least_moves as lm

UNIT = 100.0                  # fish per solver unit at arrival (mortality rides in a per-week factor)


# ------------------------------------------------------------------------------------------------ 0. inputs
def parse_steps(text: str) -> list[dict]:
    """'0:4, 21:6, 32:+1' -> [{'week': 0, 'tanks': 4}, {'week': 21, 'tanks': 6}, {'week': 32, 'add': 1}].
    Step 1 = entry (week 0), step 2 = the whole batch moves to grow-out, later steps add grow tanks."""
    steps = []
    for part in [p.strip() for p in text.replace(';', ',').split(',') if p.strip()]:
        w, k = [v.strip() for v in part.split(':')]
        steps.append({'week': int(w), 'add': int(k[1:])} if k.startswith('+') else {'week': int(w), 'tanks': int(k)})
    return steps


def validate_steps(steps: list[dict], n_entry: int) -> list[str]:
    errs = []
    if len(steps) < 2:
        errs.append('give at least two steps: entry (week 0) and the move to grow-out')
        return errs
    if steps[0]['week'] != 0 or 'tanks' not in steps[0]:
        errs.append('step 1 is the entry: week 0 and a number of tanks (e.g. 0:4)')
    if 'tanks' not in steps[1]:
        errs.append('step 2 is the move to grow-out: a number of tanks (e.g. 21:6)')
    if any('tanks' in s for s in steps[2:]):
        errs.append('steps after the move to grow-out add tanks: write them as +n (e.g. 32:+1)')
    if any(b['week'] <= a['week'] for a, b in zip(steps, steps[1:])):
        errs.append('step weeks must increase')
    if 'tanks' in steps[0] and not 1 <= steps[0]['tanks'] <= n_entry:
        errs.append(f'entry tanks must be 1..{n_entry}')
    if any(s.get('tanks', 1) < 1 or s.get('add', 1) < 1 for s in steps):
        errs.append('every step needs at least one tank')
    return errs


def cadence(D: dict) -> int:
    """Weeks between arrivals in the plan (median gap)."""
    W = D['weeks']; first = []
    for b, rec in D['batch'].items():
        i = next((i for i, w in enumerate(W) if rec[w]['arr'] > 0), None)
        if i is not None:
            first.append(i)
    gaps = [b - a for a, b in zip(sorted(first), sorted(first)[1:]) if b > a]
    return int(statistics.median(gaps)) if gaps else 7


def complete_batches(D: dict) -> list[dict]:
    """Batches that arrive and leave inside the plan, in arrival order: batch, arrival week, weeks of life."""
    W = D['weeks']; out = []
    for b, rec in D['batch'].items():
        i0 = next((i for i, w in enumerate(W) if rec[w]['arr'] > 0), None)
        if i0 is None:
            continue
        end = next((i for i in range(i0, len(W)) if rec[W[i]]['n'] <= 0.5 and rec[W[i]]['arr'] <= 0), None)
        if end is not None:
            out.append(dict(batch=b, arrival=W[i0], weeks=end - i0 + 1, end_index=end, fish=rec[W[i0]]['arr']))
    return sorted(out, key=lambda d: (d['arrival'], d['batch']))


def batch_profile(D: dict, batch: str | None = None, fish: float | None = None) -> dict:
    """One batch's life by week of life (n, wt, sfr, 6N staging), from arrival to the last fish staged. Default batch:
    the latest arrival whose whole life is inside the re-planned weeks. `fish` scales counts and staging (weights,
    feed rate and harvest timing stay)."""
    W = D['weeks']
    def life(b):
        rec = D['batch'][b]
        i0 = next((i for i, w in enumerate(W) if rec[w]['arr'] > 0), None)
        if i0 is None:
            return None
        ages = []
        for w in W[i0:]:
            ages.append(dict(rec[w], week=w))
            if rec[w]['n'] <= 0.5 and rec[w]['arr'] <= 0:
                return ages
        return None                                        # still alive at the end of the horizon
    if batch is None:
        full = complete_batches(D)
        if not full:
            raise ValueError('no batch lives its whole life inside the plan: run the forecast further out')
        # the latest one that is done a cadence before the plan ends (the last ones carry end-of-horizon tails)
        early = [f for f in full if f['end_index'] <= len(W) - 1 - cadence(D)]
        batch = (early or full)[-1]['batch']
    ages = life(batch)
    if ages is None:
        raise ValueError(f'{batch} does not arrive and leave inside the plan')
    k = (fish / ages[0]['arr']) if fish else 1.0
    out = [dict(n=a['n'] * k, wt=a['wt'], sfr=a['sfr'], stage=a['stage'] * k, arr=a['arr'] * k, week=a['week'])
           for a in ages]
    cap = D['dens_cap'].get(ages[-1]['week']) or min(t[3] for t in D['tanks'])
    return dict(batch=batch, ages=out, fish=out[0]['arr'], cap=float(cap), cadence=cadence(D))


def _facility(D: dict) -> dict:
    SYS = {t[0]: t[1] for t in D['tanks']}; VOL = {t[0]: t[2] for t in D['tanks']}
    entry = sorted(t for t in SYS if SYS[t] in lm.OG12); grow = sorted(t for t in SYS if t not in entry)
    return dict(SYS=SYS, VOL=VOL, entry=entry, grow=grow, vol=min(VOL.values()))


def _units(prof: dict) -> dict:
    """Per week of life: units alive after staging (Ua), fish per unit (cs*UNIT), staged units (su) - as least_moves."""
    ages = prof['ages']; A = len(ages)
    U0 = round(ages[0]['arr'] / UNIT); c = ages[0]['arr'] / (U0 * UNIT); U = U0
    Ua, cs, su = [], [], []
    for a, r in enumerate(ages):
        s_ = min(U, round(r['stage'] / (c * UNIT))) if (a > 0 and r['stage'] > 0) else 0
        U -= s_
        if a == A - 1 or (U > 0 and r['n'] < 0.5):
            s_ += U; U = 0
        c = (r['n'] / (U * UNIT)) if U > 0 else c
        Ua.append(U); cs.append(c); su.append(s_)
    return dict(A=A, Ua=Ua, cs=cs, su=su)


def _tanks_by_age(steps: list[dict], A: int) -> list[int]:
    """Grow tanks the steps provide at each week of life (0 while in entry)."""
    out = []
    for a in range(A):
        g = 0
        if a >= steps[1]['week']:
            g = steps[1]['tanks'] + sum(s['add'] for s in steps[2:] if s['week'] <= a)
        out.append(g)
    return out


# ------------------------------------------------------------------------------------------------ 1. quick check
def quick_check(D: dict, prof: dict, steps: list[dict]) -> dict:
    """Necessary conditions only, in a second: density at every step, entry room, and a lower bound on grow tanks in
    the busiest week (each batch needs at least ceil(biomass / tank capacity) tanks while it holds fish)."""
    F = _facility(D); ages = prof['ages']; A = len(ages); cad = prof['cadence']
    cap_kg = prof['cap'] * F['vol']; E = steps[0]['tanks']; X = steps[1]['week']
    lines = []; ok = True
    bio = [a['n'] * a['wt'] for a in ages]
    first_h = next((a for a in range(A) if ages[a]['stage'] > 0), A)
    if X >= first_h:
        return dict(ok=False, lines=[f'the move to grow-out (week {X}) must come before 6N staging starts '
                                     f'(week {first_h})'], peak=None)
    # entry
    groups = len(F['entry']) // E
    need_e = bio[X - 1] / (E * cap_kg)
    if need_e > 1:
        ok = False
        lines.append(f'Entry: {E} tanks hold {E * cap_kg / 1e3:,.0f} t, the batch is {bio[X - 1] / 1e3:,.0f} t at '
                     f'week {X - 1} - too dense. Fewer weeks in entry or more entry tanks.')
    if X > groups * cad:
        ok = False
        lines.append(f'Entry: {len(F["entry"])} entry tanks / {E} per batch = {groups} batches at a time, so a batch '
                     f'can stay {groups * cad} weeks; the steps keep it {X}.')
    if ok:
        lines.append(f'Entry fits: {E} tanks, weeks 0-{X - 1}, fullest {100 * need_e:.0f}% of the density cap.')
    # grow density per step
    g = _tanks_by_age(steps, A)
    for a in range(X, first_h):
        if bio[a] > g[a] * cap_kg:
            ok = False
            lines.append(f'Grow: at week {a} the batch is {bio[a] / 1e3:,.0f} t, {g[a]} tanks hold '
                         f'{g[a] * cap_kg / 1e3:,.0f} t - add a tank earlier.')
            break
    # busiest-week lower bound
    lb = [0] * A
    for a in range(X, A):
        if ages[a]['n'] > 0.5:
            lb[a] = min(g[a], max(1, math.ceil(bio[a] / cap_kg))) if a >= first_h else g[a]
    phase = [sum(lb[a] for a in range(c, A, cad) if a >= X) for c in range(cad)]
    peak = max(phase); have = len(F['grow'])
    if peak > have:
        ok = False
        lines.append(f'Grow: in the busiest week the batches need at least {peak} grow tanks, the facility has '
                     f'{have}. It cannot fit.')
    else:
        lines.append(f'Grow count passes: the busiest week needs at least {peak} of {have} grow tanks.')
    ev = E + steps[1]['tanks'] - 1
    if ev > D['max_moves']:
        ok = False
        lines.append(f'The move to grow-out alone is {ev} move events, over the {D["max_moves"]} per week budget.')
    return dict(ok=ok, lines=lines, peak=peak, have=have)


# ------------------------------------------------------------------------------------------------ 2. one batch
def design_batch(D: dict, prof: dict, steps: list[dict], secs: float = 60, workers: int = 8) -> dict:
    """One batch's tank plan with every rule hard. Fewest grow tank-weeks while harvesting, then fewest 6N sources."""
    from ortools.sat.python import cp_model
    F = _facility(D); u = _units(prof); A = u['A']; Ua, cs, su = u['Ua'], u['cs'], u['su']
    ages = prof['ages']; cap_kg = prof['cap'] * F['vol']
    X = steps[1]['week']; G0 = steps[1]['tanks']
    start = [X] * G0
    for s in steps[2:]:
        start += [s['week']] * s['add']
    R = len(start); adds = sorted({s['week'] for s in steps[2:]})
    maxu = [int(cap_kg / (cs[a] * UNIT * ages[a]['wt'])) if ages[a]['wt'] > 0 else 10 ** 6 for a in range(A)]
    MM = [math.ceil(D['min_move'] / (cs[a] * UNIT)) for a in range(A)]
    MT = [math.ceil(D['min_tank'] / (cs[a] * UNIT)) for a in range(A)]
    HAR = [a for a in range(A) if su[a] > 0]
    U0 = max(Ua) + max(su)
    m = cp_model.CpModel()
    n = {}; occ = {}; s = {}; ys = {}; out = {}; yo = {}; recv = {}
    for r in range(R):
        for a in range(start[r], A):
            n[(r, a)] = m.NewIntVar(0, U0, ''); occ[(r, a)] = m.NewBoolVar('')
            m.Add(n[(r, a)] <= maxu[a] * occ[(r, a)])
            m.Add(n[(r, a)] >= MT[a] * occ[(r, a)])                       # a tank in use holds >= 7k
            if a > start[r]:
                m.AddImplication(occ[(r, a)], occ[(r, a - 1)])            # never refilled once empty
            if a in HAR:
                s[(r, a)] = m.NewIntVar(0, su[a], ''); ys[(r, a)] = m.NewBoolVar('')
                m.Add(s[(r, a)] <= su[a] * ys[(r, a)])
                m.Add(s[(r, a)] >= min(MM[a], su[a])).OnlyEnforceIf([ys[(r, a)], occ[(r, a)]])
        if start[r] > X:
            recv[r] = m.NewIntVar(0, U0, '')
            m.Add(recv[r] >= MM[start[r]])                                  # received piece >= 7k
    for w in adds:
        for r in range(R):
            if start[r] < w:
                out[(r, w)] = m.NewIntVar(0, U0, ''); yo[(r, w)] = m.NewBoolVar('')
                m.Add(out[(r, w)] <= U0 * yo[(r, w)]); m.Add(out[(r, w)] >= MM[w] * yo[(r, w)])
                m.Add(n[(r, w)] >= MM[w]).OnlyEnforceIf(yo[(r, w)])         # a sender keeps >= 7k (partial move)
        m.Add(sum(out[(r, w)] for r in range(R) if (r, w) in out) == sum(recv[r] for r in range(R) if start[r] == w))
    for a in range(X, A):
        m.Add(sum(n[(r, a)] for r in range(R) if (r, a) in n) == Ua[a])
        for r in range(R):
            if (r, a) not in n:
                continue
            st_ = s.get((r, a), 0)
            if a == start[r]:
                m.Add(occ[(r, a)] == 1)                                     # every tank of a step is used
                if r in recv:
                    m.Add(n[(r, a)] == recv[r] - st_)
            else:
                m.Add(n[(r, a)] == n[(r, a - 1)] - out.get((r, a), 0) - st_)
    for a in HAR:
        if a < X:
            raise ValueError(f'6N staging starts at week {a}, before the move to grow-out (week {X})')
        m.Add(sum(s[(r, a)] for r in range(R) if (r, a) in s) == su[a])
    TW = sum(occ[(r, a)] for (r, a) in occ if a >= (HAR[0] if HAR else A))
    m.Minimize(100 * TW + sum(ys.values()) + sum(yo.values()))
    sv = cp_model.CpSolver(); sv.parameters.max_time_in_seconds = secs; sv.parameters.num_workers = workers
    stt = sv.Solve(m)
    res = dict(status=sv.StatusName(stt), ok=stt in (cp_model.OPTIMAL, cp_model.FEASIBLE))
    if not res['ok']:
        return res
    V = sv.Value
    roles = []
    E = steps[0]['tanks']
    for e in range(E):
        roles.append(dict(name=f'E{e + 1}', tier='E', start=0,
                          fish={a: Ua[a] * cs[a] * UNIT / E for a in range(X)}, inn={}, out={}, st={}))
    for r in range(R):
        f = {a: V(n[(r, a)]) * cs[a] * UNIT for a in range(start[r], A) if V(n[(r, a)]) > 0}
        st_ = {a: V(s[(r, a)]) * cs[a] * UNIT for a in HAR if (r, a) in s and V(s[(r, a)])}
        inn = {start[r]: f.get(start[r], 0) + st_.get(start[r], 0)}
        o = {w: V(out[(r, w)]) * cs[w] * UNIT for w in adds if (r, w) in out and V(out[(r, w)])}
        roles.append(dict(name=f'G{r + 1}', tier='G', start=start[r], fish=f, inn=inn, out=o, st=st_))
    res.update(roles=roles, units=u, X=X, A=A)
    return res


# ------------------------------------------------------------------------------------------------ 3. tanks
def assign_tanks(D: dict, prof: dict, design: dict, secs: float = 300, workers: int = 8) -> dict:
    """Entry by rule (batch i takes entry group i mod groups, a group = one tank per entry system in turn); grow roles
    on physical tanks for a cycle of `slots` batches, with one batch per tank and system biomass/feed <= cap x buffer
    every week. Tries a cycle of ceil(life/cadence) batches, then twice that."""
    from ortools.sat.python import cp_model
    F = _facility(D); cad = prof['cadence']; ages = prof['ages']; A = design['A']
    E = sum(1 for r in design['roles'] if r['tier'] == 'E')
    groups = len(F['entry']) // E
    by_pos = sorted(F['entry'], key=lambda t: (sorted(x for x in F['entry'] if F['SYS'][x] == F['SYS'][t]).index(t),
                                                lm.OG12.index(F['SYS'][t])))
    egroup = [by_pos[g * E:(g + 1) * E] for g in range(groups)]
    grow_roles = [i for i, r in enumerate(design['roles']) if r['tier'] == 'G']
    load = {(i, a): (f * ages[a]['wt'], f * ages[a]['wt'] * ages[a]['sfr'])
            for i in grow_roles for a, f in design['roles'][i]['fish'].items()}
    BUF = 1 + D['buffer']
    base = math.ceil(A / cad)
    for L in (base, 2 * base):
        P = L * cad
        m = cp_model.CpModel(); z = {}
        for j in range(L):
            for i in grow_roles:
                for t in F['grow']:
                    z[(j, i, t)] = m.NewBoolVar('')
                m.AddExactlyOne(z[(j, i, t)] for t in F['grow'])
        at = collections.defaultdict(list)
        for j in range(L):
            for i in grow_roles:
                for a in design['roles'][i]['fish']:
                    at[(cad * j + a) % P].append((j, i, a))
        for c in range(P):
            for t in F['grow']:
                v = [z[(j, i, t)] for j, i, a in at[c]]
                if len(v) > 1:
                    m.AddAtMostOne(v)
            for sname in sorted({F['SYS'][t] for t in F['grow']}):
                bc, fc = D['sys_cap'].get(sname, [0, 0])
                ts = [t for t in F['grow'] if F['SYS'][t] == sname]
                if bc:
                    m.Add(sum(int(load[(i, a)][0]) * z[(j, i, t)] for j, i, a in at[c] for t in ts) <= int(bc * BUF))
                if fc:
                    m.Add(sum(int(load[(i, a)][1]) * z[(j, i, t)] for j, i, a in at[c] for t in ts) <= int(fc * BUF))
        sv = cp_model.CpSolver(); sv.parameters.max_time_in_seconds = secs; sv.parameters.num_workers = workers
        stt = sv.Solve(m)
        if stt in (cp_model.OPTIMAL, cp_model.FEASIBLE):
            grow = {(j, i): next(t for t in F['grow'] if sv.Value(z[(j, i, t)])) for j in range(L) for i in grow_roles}
            return dict(ok=True, status=sv.StatusName(stt), slots=L, egroup=egroup, grow=grow)
        last = sv.StatusName(stt)
    return dict(ok=False, status=last,
                proven=(last == 'INFEASIBLE'), slots=L)


def tank_of(design: dict, tanks: dict, i: int, role: int) -> int:
    r = design['roles'][role]
    if r['tier'] == 'E':
        return tanks['egroup'][i % len(tanks['egroup'])][int(r['name'][1:]) - 1]
    return tanks['grow'][(i % tanks['slots'], role)]


# ------------------------------------------------------------------------------------------------ 4. judge
def evaluate(D: dict, prof: dict, design: dict, tanks: dict, cycles: int = 3) -> dict:
    """Lay `cycles` cycles of steady state out week by week (with the batches already in the tanks at week 0) and judge
    every cycle after the first with least_moves.check - the same independent checker as the least-moves layout."""
    F = _facility(D); cad = prof['cadence']; ages = prof['ages']; A = design['A']; X = design['X']
    P = tanks['slots'] * cad
    weeks = [f'R{c + 1:04d}' for c in range(P * cycles)]
    plan = []; recs = collections.defaultdict(dict); harvested = {}
    ids = range(-(A // cad) - 1, (P * cycles) // cad + 1)
    for c, w in enumerate(weeks):
        rows = []; hv = 0.0
        for i in ids:
            b = f'R{i:+04d}'; a = c - cad * i
            rec = dict(n=0.0, kg=0.0, wt=0.0, arr=0.0, stage=0.0, sfr=0.0)
            if 0 <= a < A:
                rec.update(wt=ages[a]['wt'], sfr=ages[a]['sfr'])
                for ri, r in enumerate(design['roles']):
                    t = tank_of(design, tanks, i, ri); f = r['fish'].get(a, 0.0)
                    inn = outv = stg = arr = 0.0
                    if r['tier'] == 'E':
                        arr = f if a == 0 else 0.0
                        outv = r['fish'].get(a - 1, 0.0) if a == X else 0.0
                    else:
                        inn = r['inn'].get(a, 0.0); outv = r['out'].get(a, 0.0); stg = r['st'].get(a, 0.0)
                    if f > 0.5 or inn or outv or stg or arr:
                        rows.append(dict(b=b, t=t, n=f, inn=inn, out=outv, st=stg, arr=arr,
                                         dens=f * ages[a]['wt'] / F['VOL'][t]))
                    rec['n'] += f; hv += stg; rec['stage'] += stg
                rec['kg'] = rec['n'] * rec['wt']
            recs[b][w] = rec
        harvested[w] = hv
        plan.append(dict(week=w, moved=sum(r['inn'] for r in rows), rows=rows))
    E2 = dict(source='tank rota', weeks=weeks, tanks=D['tanks'], sys_cap=D['sys_cap'], buffer=D['buffer'],
              min_tank=D['min_tank'], min_move=D['min_move'], move_rules=True, max_moves=D['max_moves'],
              dens_cap={w: prof['cap'] for w in weeks}, batch=dict(recs), harvested=harvested)
    steady = [p for p in plan if p['week'] >= weeks[P]]
    chk = lm.check(E2, steady)
    hv = sum(harvested[p['week']] for p in steady)
    tpf = sum(r['inn'] for p in steady for r in p['rows']) / hv if hv else 0.0
    sysbio = collections.defaultdict(float)
    for p in steady:
        for r in p['rows']:
            sysbio[(p['week'], F['SYS'][r['t']])] += r['dens'] * F['VOL'][r['t']]
    bc = D['sys_cap'].get('OG3N', [0, 0])[0]
    return dict(check=chk, tpf=tpf, max_moves=max(chk['moves'].values()) if chk['moves'] else 0,
                busiest_t=max(sysbio.values()) / 1e3 if sysbio else 0.0, sys_cap_t=bc / 1e3,
                needs_buffer=bool(bc) and max(sysbio.values()) > bc, cycle=steady[:P], D=E2)


# ------------------------------------------------------------------------------------------------ 5. all of it
def run(D: dict, steps: list[dict], batch: str | None = None, fish: float | None = None, secs_batch: float = 60,
        secs_tanks: float = 300, workers: int = 8, log=print) -> dict:
    F = _facility(D)
    errs = validate_steps(steps, len(F['entry']))
    if errs:
        return dict(stage='steps', ok=False, lines=errs)
    prof = batch_profile(D, batch, fish)
    log(f"batch {prof['batch']}: {prof['fish']:,.0f} fish, {len(prof['ages'])} weeks of life, one every "
        f"{prof['cadence']} weeks, density cap {prof['cap']:g} kg/m3")
    q = quick_check(D, prof, steps)
    for ln in q['lines']:
        log('  ' + ln)
    out = dict(profile=prof, quick=q, steps=steps)
    if not q['ok']:
        return dict(out, stage='quick', ok=False)
    log('designing one batch ...')
    des = design_batch(D, prof, steps, secs_batch, workers)
    out['design'] = des
    if not des['ok']:
        return dict(out, stage='design', ok=False,
                    lines=[f"No single-batch plan meets every rule ({des['status'].lower()}) - the drain or a step "
                           f"is too tight; try a tank more or a step a week earlier."])
    log('placing on tanks ...')
    tk = assign_tanks(D, prof, des, secs_tanks, workers)
    out['tanks'] = tk
    if not tk['ok']:
        why = ('proven: no placement meets the system caps' if tk.get('proven')
               else f"none found in the time given ({tk['status'].lower()})")
        return dict(out, stage='tanks', ok=False, lines=[f'Tanks: {why}.'])
    log('checking 3 cycles ...')
    ev = evaluate(D, prof, des, tk)
    out['eval'] = ev
    return dict(out, stage='done', ok=ev['check']['ok'])


def write_workbook(res: dict, out_path: str, source: str = '') -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    prof, des, ev, steps = res['profile'], res['design'], res['eval'], res['steps']
    ages = prof['ages']; A = des['A']; cad = prof['cadence']; B_ = Font(bold=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Summary'
    stxt = f"entry {steps[0]['tanks']} -> grow {steps[1]['tanks']} at week {steps[1]['week']}" + ''.join(
        f", +{s['add']} at week {s['week']}" for s in steps[2:])
    rows = [['TANK ROTA (standard schedule) - steady state'],
            [f"One {prof['fish']:,.0f}-fish batch every {cad} weeks, every batch handled the same way: {stxt}."],
            [f"Growth, mortality and 6N staging by week of life: batch {prof['batch']} of {source or 'the source plan'}."],
            [],
            ['Result (independent rule check, cycles 2-3 of steady state)'],
            ['Transfers per harvested fish', round(ev['tpf'], 2)],
            ['Move events per week (max)', ev['max_moves']],
            [f"Busiest system (t; cap {ev['sys_cap_t']:.0f} x buffer)", round(ev['busiest_t'])],
            ['Rule breaks', ', '.join(f'{k} {v}' for k, v in ev['check']['issues'].items()) or 'none'],
            [],
            ['Limits'],
            ['Steady state only: getting from today\'s plan into this rota is not worked out here.'],
            ['Tank handovers: an emptied tank can take the next batch in the same week (entry tanks always do).']]
    if ev['needs_buffer']:
        rows.append([f"Uses the system tolerance: a system goes over {ev['sys_cap_t']:.0f} t (within the buffer)."])
    for r in rows:
        ws.append(r)
    for i in (1, 5, 11):
        ws.cell(i, 1).font = B_
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 100
    wr = wb.create_sheet('Batch recipe')
    names = [r['name'] for r in des['roles']]
    wr.append(['Week of life', 'Avg wt (kg)', 'Fish'] + names + ['Action'])
    for a in range(A):
        act = []
        if a == 0: act.append(f"arrive into {steps[0]['tanks']} entry tanks")
        if a == des['X']: act.append(f"move to {steps[1]['tanks']} grow tanks")
        for r in des['roles']:
            if r['tier'] == 'G' and r['start'] == a and a != des['X']:
                act.append(f"fill {r['name']} from " + ', '.join(f"{o['name']} {o['out'][a]:,.0f}"
                                                                for o in des['roles'] if a in o['out'] and o['out'][a]))
        st_ = [f"{r['name']} {r['st'][a]:,.0f}" for r in des['roles'] if r['st'].get(a)]
        if st_: act.append('6N ' + ', '.join(st_))
        wr.append([a, round(ages[a]['wt'], 2), round(sum(r['fish'].get(a, 0) for r in des['roles']))]
                  + [round(r['fish'].get(a, 0)) or None for r in des['roles']] + ['; '.join(act)])
    for c in wr[1]: c.font = B_
    wr.column_dimensions[get_column_letter(len(names) + 4)].width = 80
    wc = wb.create_sheet('Tank calendar')
    cyc = ev['cycle']; F = _facility(ev['D'])
    wc.append(['Tank', 'System'] + [f'wk {k + 1}' for k in range(len(cyc))])
    lab = {(r['t'], k): (r['b'], r['dens']) for k, p in enumerate(cyc) for r in p['rows'] if r['n'] > 0.5}
    bn = sorted({v[0] for v in lab.values()})
    pal = ['FFD966', '9DC3E6', 'A9D18E', 'F4B183', 'C9C9C9', 'D5A6BD', '8FAADC', 'FFE699', 'B4C7E7', 'C5E0B4']
    for t in sorted(F['SYS']):
        wc.append([t, F['SYS'][t]] + [(f'{lab[(t, k)][0]} {lab[(t, k)][1]:.0f}' if (t, k) in lab else '')
                                      for k in range(len(cyc))])
        for k in range(len(cyc)):
            if (t, k) in lab:
                wc.cell(wc.max_row, k + 3).fill = PatternFill('solid', fgColor=pal[bn.index(lab[(t, k)][0]) % len(pal)])
    wc.freeze_panes = 'C2'
    for k in range(len(cyc)): wc.column_dimensions[get_column_letter(k + 3)].width = 9
    wc.append([]); wc.append(['Cell = batch (by place in the rota) and density (kg/m3) at the end of the week. '
                              'One cycle; it repeats.'])
    wt_ = wb.create_sheet('Transfers')
    wt_.append(['Cycle week', 'Batch', 'From_Tank', 'To_Tank', 'Count (fish)'])
    ix = {p['week']: k + 1 for k, p in enumerate(cyc)}
    for g in lm.pair_moves(cyc):
        wt_.append([ix[g['week']], g['batch'], g['frm'], g['to'], round(g['fish'])])
    for c in wt_[1]: c.font = B_
    wb.save(out_path)
