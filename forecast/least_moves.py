"""Least-moves layout ("Global v2"): the same plan, re-laid-out in the tanks with as few fish moves as possible.

WHAT IT DOES
------------
A normal run decides WHAT happens to every batch: how many fish it has each week, how heavy they are, when they
arrive from freshwater and when they go into 6N for harvest. It decides WHERE (which tank) week by week, reacting
to each week's density, and that reshuffling is where most transfers come from.

This step keeps the WHAT exactly (the "endpoint": fish per batch per week outside 6N, arrivals, 6N staging -> the
same harvest, HOG and biomass by construction) and re-plans only the WHERE, looking ahead several weeks at once, so
fish are put where they can stay. It never changes the run's workbook; it writes a companion workbook.

    1. read_endpoint()   the endpoint + the engine's own layout, from a finished run's workbook
    2. solve()           CP-SAT in rolling windows (default: look 10 weeks ahead, keep the first 5, move on)
    3. check()           an independent rule check of any layout (used on both the engine's and the new one)
    4. pair_moves()      turns each week's tank in/out amounts into from -> to transfers
    5. write_workbook()  Summary / Transfers / Layout / Batches sheets

RULES (HARD = never broken; SOFT = heavily penalised and reported if ever used)
    HARD: fish conservation per batch and tank; one batch per tank; TranOG arrivals only into entry tanks (OG1/2);
          no move back into OG1/2; no split inside OG1/2 at >= 1 kg; 6N staging only from grow-out tanks.
    SOFT: tank density cap (that week's og_density row or facility.yaml); system biomass and feed x (1 + buffer);
          the weekly move budget; min fish per occupied tank; 7k move-size floors (partial moves, received pieces,
          6N staging pieces). The check() report counts every one of these on the finished layout.
Objective: fish moved between OG tanks, then move events (LAMBDA fish-equivalents each), after the penalties.
"""
from __future__ import annotations

import collections
import json
import math
import time
from pathlib import Path

import yaml

OG12 = ('OG1N', 'OG1S', 'OG2N', 'OG2S')
SIXN = {61, 63, 65, 67, 69, 71}
UNIT = 100.0                  # the solver counts fish in units of 100 (mortality rides in a per-week factor)


# ------------------------------------------------------------------------------------------------ 1. endpoint
def _sheet_rows(wb, sheet, key='Week'):
    hdr = None
    for r in wb[sheet].iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == key:
                hdr = list(r)
            continue
        if r and (str(r[0]).startswith('20') or (key != 'Week' and len(r) > 1 and str(r[1]).startswith('20'))):
            yield dict(zip(hdr, r))


def default_start_week(wb, weeks: list[str]) -> str:
    """First week after the manual starting events (those weeks are the operator's, not the planner's): the run's
    RunConfig row 'Manual override window (weeks)'."""
    n = 0
    for r in wb['RunConfig'].iter_rows(values_only=True):
        if r and r[0] == 'Manual override window (weeks)':
            n = int(r[1] or 0)
            break
    return weeks[min(n, len(weeks) - 1)]


def read_endpoint(workbook: str, config_dir: str, scenario_dir: str, start_week: str | None = None) -> dict:
    """The endpoint the new layout must reproduce, plus the engine's layout (for the warm start and the comparison).
    config_dir / scenario_dir: the config/ and scenario/ folders the run used (as forecast.run.main takes them)."""
    import openpyxl
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    loc = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0, 0.0]))   # week -> batch -> [n, kg]
    layout = collections.defaultdict(dict)                                                # week -> tank -> (b, n)
    weeks = set()
    for d in _sheet_rows(wb, 'BatchLocations'):
        t = int(d['Tank']); w = d['Week']; weeks.add(w)
        if t in SIXN:
            continue
        loc[w][d['Batch']][0] += float(d['Count (fish)'] or 0)
        loc[w][d['Batch']][1] += float(d['Biomass (kg)'] or 0)
        layout[w][t] = (d['Batch'], float(d['Count (fish)'] or 0), float(d['Density (kg/m3)'] or 0))
    weeks = sorted(weeks)
    arr = collections.defaultdict(float); arr_wt = {}
    stage = collections.defaultdict(float); moved = collections.defaultdict(float)
    eng_flow = collections.defaultdict(lambda: [0.0, 0.0, 0.0])        # (week, batch, tank) -> [in, out, staged]
    for d in _sheet_rows(wb, 'TransferPlan'):
        w, b = d['Week'], d['Batch']; c = float(d['Count (fish)'] or 0)
        if d['Type'] == 'TranOG':
            arr[(w, b)] += c; arr_wt[(w, b)] = float(d['Avg_Weight (kg)'] or 0)
            continue
        try:
            s, t = int(d['From_Tank']), int(d['To_Tank'])
        except (TypeError, ValueError):
            continue
        if s == t or s in SIXN:
            continue
        if t in SIXN:
            stage[(w, b)] += c; eng_flow[(w, b, s)][2] += c
        else:
            moved[(w, b)] += c; eng_flow[(w, b, s)][1] += c; eng_flow[(w, b, t)][0] += c
    harvested = collections.defaultdict(float)
    for d in _sheet_rows(wb, 'HarvestPlan'):
        harvested[d['Week']] += float(d['Count (fish)'] or 0)
    sfr = {}
    for d in _sheet_rows(wb, 'WeeklyReport', key='Scenario'):
        if d['Batch'] not in ('TOTAL', None) and d.get('SFR (%/day)') is not None:
            sfr[(d['Week'], d['Batch'])] = float(d['SFR (%/day)']) / 100.0
    lim = yaml.safe_load(open(Path(scenario_dir) / 'limits.yaml', encoding='utf-8')) or {}
    dens_row = {r['week']: float(r['value']) for r in lim.get('facility', []) or [] if r['metric'] == 'og_density_kg_m3'}
    fac = yaml.safe_load(open(Path(config_dir) / 'facility.yaml', encoding='utf-8'))['tanks']
    tanks = sorted((t['tank_id'], t['system_id'], t['volume_m3'], t['max_density_kg_m3']) for t in fac
                   if t['system_id'].startswith('OG') and t['system_id'] != 'OG6N')
    sysd = lim.get('system_defaults', {}) or {}
    cap = {s: (float(v.get('biomass', 0) or 0), float(v.get('feed_per_day', 0) or 0)) for s, v in sysd.items()}
    ctl = yaml.safe_load(open(Path(config_dir) / 'control.yaml', encoding='utf-8'))
    start = start_week or default_start_week(wb, weeks)
    W = [w for w in weeks if w >= start]
    prev = weeks[weeks.index(W[0]) - 1]
    batches = sorted({b for w in W for b in loc[w]} | {b for (w, b) in arr if w in W})
    D = {'source': str(workbook), 'weeks': W, 'start_from': prev, 'tanks': tanks, 'sys_cap': cap,
         'buffer': float(ctl.get('global_buffer_pct', 0.1) or 0.1),
         'min_tank': float(ctl.get('min_tank_control', 7000) or 0),
         'min_move': float(ctl.get('min_transfer_count', 7000) or 0),
         'move_rules': bool(ctl.get('move_size_rules', False)),
         'max_moves': int(ctl.get('max_transfers_per_week', 15) or 0),
         'dens_cap': {w: dens_row.get(w) for w in W},
         'start_layout': {str(t): [v[0], v[1]] for t, v in layout[prev].items()},
         'batch': {}, 'harvested': {w: harvested[w] for w in W},
         'engine_moved': {w: sum(v for (ww, b), v in moved.items() if ww == w) for w in W},
         'engine_layout': {w: {str(t): [v[0], v[1]] for t, v in layout[w].items()} for w in W}}
    for b in batches:
        D['batch'][b] = {w: {'n': loc[w][b][0], 'kg': loc[w][b][1],
                             'wt': (loc[w][b][1] / loc[w][b][0]) if loc[w][b][0] > 0 else arr_wt.get((w, b), 0),
                             'arr': arr[(w, b)], 'stage': stage[(w, b)], 'sfr': sfr.get((w, b), 0.0)}
                         for w in W}
    vol = {t[0]: t[2] for t in tanks}
    D['engine_plan'] = [{'week': w, 'rows': [          # the engine's own tank densities (its tanks are graded)
        {'b': b, 't': t, 'n': n, 'inn': eng_flow[(w, b, t)][0], 'out': eng_flow[(w, b, t)][1],
         'st': eng_flow[(w, b, t)][2], 'dens': dn}
        for t, (b, n, dn) in layout[w].items()] + [
        {'b': b, 't': t, 'n': 0.0, 'inn': f[0], 'out': f[1], 'st': f[2], 'dens': 0.0}
        for (ww, b, t), f in eng_flow.items() if ww == w and layout[w].get(t, (None,))[0] != b and t in vol]}
        for w in W]
    return D


def open_wt(B: dict, b: str, w: str, weeks: list) -> float:
    """A batch's weight when the week opens (= the end of the week before), which is when moves happen; the
    week's own weight when there is no week before or the batch was not yet in OG."""
    i = weeks.index(w)
    if i > 0 and B[b][weeks[i - 1]]['n'] > 0:
        return B[b][weeks[i - 1]]['wt']
    return B[b][w]['wt']


def sys_cap(D: dict, week: str, system: str) -> tuple:
    """(biomass kg, feed kg/day) caps of a system in a week: the week's own caps when the endpoint carries them
    (built inside a run, from limits.yaml rows), else the system defaults."""
    v = (D.get('sys_cap_w') or {}).get(week, {}).get(system)
    return tuple(v) if v else tuple(D['sys_cap'].get(system, (0, 0)))


# ------------------------------------------------------------------------------------------------ 2. solve
def solve(D: dict, horizon: int = 10, commit: int = 5, secs: float = 360, lam: float = 2000.0,
          workers: int = 8, retries: int = 3, warm_start: bool = True, log=print) -> dict:
    """Rolling-window CP-SAT. Each window looks `horizon` weeks ahead, keeps the first `commit` weeks, and the next
    window starts from where those end. A window that finds no answer or uses hard-rule slack is retried with
    double the time (up to `retries` times)."""
    from ortools.sat.python import cp_model
    W = D['weeks']; TK = D['tanks']; T = [t[0] for t in TK]
    SYS = {t[0]: t[1] for t in TK}; VOL = {t[0]: t[2] for t in TK}; CAP0 = {t[0]: t[3] for t in TK}
    ENTRY = [t for t in T if SYS[t] in OG12]; GROW = [t for t in T if SYS[t] not in OG12]
    BUF = 1 + D['buffer']; MAXMV = D['max_moves']; B = D['batch']
    _wi = {w: i for i, w in enumerate(W)}

    def open_wt(b, w):                                   # weight when the week opens (moves happen then)
        i = _wi[w]
        return B[b][W[i - 1]]['wt'] if i > 0 and B[b][W[i - 1]]['n'] > 0 else B[b][w]['wt']

    # per-batch unit bookkeeping: units change only by arrivals and 6N staging; cs = fish per unit / UNIT
    units = {}; cs = {}; stage_u = {}; arr_u = {}
    start = {int(t): v for t, v in D['start_layout'].items()}
    start_n = collections.defaultdict(float)
    for t, (b, n) in start.items():
        start_n[b] += n
    for b, rec in B.items():
        U = None; c = None
        for i, w in enumerate(W):
            r = rec[w]
            if U is None:
                if start_n.get(b, 0) > 0 and i == 0:
                    U = round(start_n[b] / UNIT); c = start_n[b] / (U * UNIT)
                elif r['arr'] > 0:
                    U = 0; c = 1.0
                else:
                    units[(b, w)] = 0; cs[(b, w)] = 1.0; stage_u[(b, w)] = 0; arr_u[(b, w)] = 0
                    continue
            su = min(U, round(r['stage'] / (c * UNIT))) if r['stage'] > 0 else 0
            au = round(r['arr'] / UNIT) if r['arr'] > 0 else 0
            U = U - su + au
            if U > 0 and r['n'] < 0.5:      # the batch has left OG this week: rounding leftovers go to 6N with it
                su += U; U = 0
            c = (r['n'] / (U * UNIT)) if U > 0 else c
            units[(b, w)] = U; cs[(b, w)] = c; stage_u[(b, w)] = su; arr_u[(b, w)] = au
    state = {t: None for t in T}
    for b in {v[0] for v in start.values()}:
        ts = sorted(t for t, v in start.items() if v[0] == b)
        want = round(start_n[b] / UNIT); got = 0
        for t in ts:
            u = round(start[t][1] / UNIT); state[t] = (b, u); got += u
        if ts and got != want:
            t0 = max(ts, key=lambda t: state[t][1]); state[t0] = (b, state[t0][1] + want - got)

    def bkg(b, w):                      # kg per unit at the end of week w, x10
        return int(round(cs[(b, w)] * UNIT * B[b][w]['wt'] * 10))

    def cap_kg(t, w):                   # tank capacity kg, x10
        return int(round((D['dens_cap'].get(w) or CAP0[t]) * VOL[t] * 10))

    plan = []; report = collections.Counter(); t0 = time.time(); k = 0
    while k < len(W):
        win = W[k:k + horizon]
        alive = sorted({b for w in win for b in B if units.get((b, w), 0) > 0 or arr_u.get((b, w), 0) > 0}
                       | {v[0] for v in state.values() if v})
        m = cp_model.CpModel()
        n = {}; x = {}; inn = {}; out = {}; st = {}; av = {}; yi = {}; yo = {}; ys = {}
        pen = []
        for b in alive:
            Umax = max([units.get((b, w), 0) for w in win] + [sum(v[1] for v in state.values() if v and v[0] == b)] + [1])
            for w in win:
                for t in T:
                    key = (b, t, w)
                    n[key] = m.NewIntVar(0, Umax, ''); x[key] = m.NewBoolVar('')
                    inn[key] = m.NewIntVar(0, Umax, ''); out[key] = m.NewIntVar(0, Umax, '')
                    yi[key] = m.NewBoolVar(''); yo[key] = m.NewBoolVar('')
                    m.Add(inn[key] <= Umax * yi[key]); m.Add(out[key] <= Umax * yo[key]); m.Add(n[key] <= Umax * x[key])
                    if t in GROW:
                        st[key] = m.NewIntVar(0, Umax, ''); ys[key] = m.NewBoolVar('')
                        m.Add(st[key] <= Umax * ys[key])
                    if t in ENTRY and arr_u.get((b, w), 0) > 0:
                        av[key] = m.NewIntVar(0, arr_u[(b, w)], '')
        for i, w in enumerate(win):
            for b in alive:
                for t in T:
                    key = (b, t, w)
                    if i == 0:
                        p = state.get(t); prev_n = p[1] if (p and p[0] == b) else 0
                    else:
                        prev_n = n[(b, t, win[i - 1])]
                    m.Add(n[key] == prev_n + inn[key] - out[key] - st.get(key, 0) + av.get(key, 0))
                    m.Add(out[key] + st.get(key, 0) <= prev_n)
                m.Add(sum(n[(b, t, w)] for t in T) == units.get((b, w), 0))
                m.Add(sum(inn[(b, t, w)] for t in T) == sum(out[(b, t, w)] for t in T))
                m.Add(sum(st.get((b, t, w), 0) for t in GROW) == stage_u.get((b, w), 0))
                if arr_u.get((b, w), 0) > 0:
                    m.Add(sum(av[(b, t, w)] for t in ENTRY) == arr_u[(b, w)])
                # the rule applies to the weight when the fish move; with entry tanks fixed to the source plan its own
                # (per-tank, graded) entry moves stand, so only "never back into OG1/2" is kept
                if open_wt(b, w) >= 1.0 and not D.get('fix_entry'):
                    for t in ENTRY:
                        m.Add(inn[(b, t, w)] == 0)                               # no split inside OG1/2 at >= 1 kg
                else:
                    m.Add(sum(inn[(b, t, w)] for t in ENTRY) <= sum(out[(b, t, w)] for t in ENTRY))   # never back
                bk = bkg(b, w) if units.get((b, w), 0) > 0 else 0
                minu = (math.ceil(D['min_tank'] / (cs[(b, w)] * UNIT))
                        if units.get((b, w), 0) * cs[(b, w)] * UNIT >= D['min_tank'] else 0)
                for t in T:
                    key = (b, t, w)
                    if bk:
                        s_ = m.NewIntVar(0, cap_kg(t, w), ''); pen.append((s_, 10 ** 4, 'density'))
                        m.Add(n[key] * bk <= cap_kg(t, w) * x[key] + s_)
                    if minu:
                        s2 = m.NewIntVar(0, minu, ''); pen.append((s2, 10 ** 6, 'min_tank'))
                        m.Add(n[key] + s2 >= minu * x[key])
                    if D['move_rules'] and D['min_move'] > 0 and units.get((b, w), 0) > 0:
                        mmu = math.ceil(D['min_move'] / (cs[(b, w)] * UNIT))
                        s3 = m.NewIntVar(0, mmu, ''); pen.append((s3, 10 ** 6, 'move_size'))
                        m.Add(out[key] + s3 >= mmu).OnlyEnforceIf([yo[key], x[key]])       # partial move
                        s3b = m.NewIntVar(0, mmu, ''); pen.append((s3b, 10 ** 6, 'move_size'))
                        m.Add(inn[key] + s3b >= mmu).OnlyEnforceIf(yi[key])               # received piece
                        if key in st and stage_u.get((b, w), 0) > 0:                       # partial 6N staging
                            smu = min(mmu, stage_u[(b, w)])
                            s3c = m.NewIntVar(0, smu, ''); pen.append((s3c, 10 ** 6, 'move_size'))
                            m.Add(st[key] + s3c >= smu).OnlyEnforceIf([ys[key], x[key]])
            for t in T:
                m.Add(sum(x[(b, t, w)] for b in alive) <= 1)                          # one batch per tank
            for sname in sorted({SYS[t] for t in T}):
                bc, fc = sys_cap(D, w, sname)
                ts = [t for t in T if SYS[t] == sname]
                if bc:
                    s4 = m.NewIntVar(0, int(bc * 10), ''); pen.append((s4, 10 ** 4, 'sys_bio'))
                    m.Add(sum(n[(b, t, w)] * bkg(b, w) for b in alive for t in ts if units.get((b, w), 0) > 0)
                          <= int(bc * BUF * 10) + s4)
                if fc:
                    s5 = m.NewIntVar(0, int(fc * 1000), ''); pen.append((s5, 10 ** 2, 'sys_feed'))
                    m.Add(sum(n[(b, t, w)] * int(round(bkg(b, w) * B[b][w]['sfr'] * 100)) for b in alive for t in ts
                              if units.get((b, w), 0) > 0) <= int(fc * BUF * 1000) + s5)
            if MAXMV:                    # per moving batch: senders + receivers - 1 moves; each 6N source is a move
                zs = []
                for b in alive:
                    z = m.NewBoolVar('')
                    for t in T:
                        m.AddImplication(yo[(b, t, w)], z)
                    m.Add(z <= sum(yo[(b, t, w)] for t in T)); zs.append(z)
                s6 = m.NewIntVar(0, 200, ''); pen.append((s6, 10 ** 9, 'moves'))
                m.Add(sum(yo[(b, t, w)] + yi[(b, t, w)] for b in alive for t in T) - sum(zs)
                      + sum(ys[(b, t, w)] for b in alive for t in GROW) <= MAXMV + s6)
        def targets(lay, b, w):                          # a layout's fish for batch b, in units summing to the batch's
            want = units.get((b, w), 0); f = cs[(b, w)] * UNIT
            ts = [(int(t), v[1] / f) for t, v in lay.items() if v[0] == b]
            fl = {t: int(math.floor(q)) for t, q in ts}
            for t, q in sorted(ts, key=lambda v: v[1] - math.floor(v[1]), reverse=True)[:max(0, want - sum(fl.values()))]:
                fl[t] += 1
            return fl
        def near(v, u):                                  # same occupancy; fish within one unit (the unit rounding)
            if u <= 0:
                m.Add(v == 0)
            else:
                m.Add(v >= max(1, u - 1)); m.Add(v <= u + 1)
        if D.get('fix_entry'):                           # entry tanks as the source plan, every week
            for w in win:
                for b in alive:
                    tg = targets(D['engine_layout'].get(w, {}), b, w)
                    for t in ENTRY:
                        near(n[(b, t, w)], tg.get(t, 0))
        if D.get('end_layout') and W[-1] in win:         # pinned last week: exactly the given tanks
            wl = W[-1]
            for b in alive:
                tg = targets(D['end_layout'], b, wl)
                for t in T:
                    near(n[(b, t, wl)], tg.get(t, 0))
        if warm_start and D.get('engine_layout'):       # start from the engine's own layout for these weeks
            for w in win:
                lay = D['engine_layout'].get(w, {})
                for b in alive:
                    for t in T:
                        v = lay.get(str(t)); on = bool(v and v[0] == b)
                        m.AddHint(x[(b, t, w)], on)
                        m.AddHint(n[(b, t, w)], int(round(v[1] / (cs[(b, w)] * UNIT))) if on and cs[(b, w)] > 0 else 0)
        fish_moved = sum(inn[(b, t, w)] * int(round(cs[(b, w)] * UNIT)) for b in alive for t in T for w in win)
        events = sum(yi[(b, t, w)] + yo[(b, t, w)] for b in alive for t in T for w in win)
        m.Minimize(sum(wg * v for v, wg, _ in pen) + 1000 * fish_moved + int(lam * 1000) * events)
        s_now = secs
        for attempt in range(retries + 1):
            sv = cp_model.CpSolver(); sv.parameters.num_workers = workers; sv.parameters.max_time_in_seconds = s_now
            stt = sv.Solve(m)
            if stt not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
                if attempt == retries:
                    break
                log(f'   no answer yet ({sv.StatusName(stt)}) - retrying with {s_now * 2:.0f}s')
                s_now *= 2
                continue
            slack = collections.Counter()
            for v, _, kind in pen:
                slack[kind] += sv.Value(v)
            hard_used = any(slack[h] > 0 for h in ('density', 'sys_bio', 'moves', 'min_tank', 'move_size'))
            if not (hard_used or slack['sys_feed'] > 5000) or attempt == retries:
                break
            log(f'   rule slack {dict(slack)} - retrying with {s_now * 2:.0f}s from this answer')
            m.ClearHints()
            for key in n:
                m.AddHint(n[key], sv.Value(n[key])); m.AddHint(x[key], sv.Value(x[key]))
            s_now *= 2
        if stt not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
            raise RuntimeError(f'least-moves: window starting {win[0]} found no layout ({sv.StatusName(stt)}); '
                               f'try more seconds per window')
        keep = win[:commit] if k + horizon < len(W) else win
        for w in keep:
            rows = []
            for b in alive:
                for t in T:
                    key = (b, t, w); V = sv.Value
                    nv = V(n[key]); f = cs[(b, w)] * UNIT
                    # staged pieces add up to the plan's own 6N staging (the unit factor is the one after the week)
                    fs = B[b][w]['stage'] / stage_u[(b, w)] if stage_u.get((b, w), 0) and B[b][w]['stage'] > 0 else f
                    if nv or V(inn[key]) or V(out[key]) or (key in st and V(st[key])) or (key in av and V(av[key])):
                        rows.append(dict(b=b, t=t, u=nv, iu=V(inn[key]), ou=V(out[key]),
                                         su=V(st[key]) if key in st else 0, au=V(av[key]) if key in av else 0,
                                         n=nv * f, inn=V(inn[key]) * f, out=V(out[key]) * f,
                                         st=V(st[key]) * fs if key in st else 0, arr=V(av[key]) * f if key in av else 0,
                                         dens=nv * bkg(b, w) / 10 / VOL[t] if units.get((b, w), 0) else 0))
            plan.append(dict(week=w, moved=sum(r['inn'] for r in rows), rows=rows))
        lastw = keep[-1]
        state = {t: None for t in T}
        for b in alive:
            for t in T:
                v = sv.Value(n[(b, t, lastw)])
                if v:
                    state[t] = (b, v)
        pv = collections.Counter()
        for v, _, kind in pen:
            pv[kind] += sv.Value(v)
        report.update(pv)
        log(f'weeks {win[0]}..{keep[-1]} done ({len(alive)} batches, {sv.StatusName(stt).lower()}, '
            f'{sum(p["moved"] for p in plan[-len(keep):]) / 1e3:,.0f}k fish moved) [{time.time() - t0:.0f}s]')
        k += len(keep)
    return dict(plan=plan, slack=dict(report), seconds=time.time() - t0,
                settings=dict(horizon=horizon, commit=commit, secs=secs, lam=lam, workers=workers, retries=retries))


# ------------------------------------------------------------------------------------------------ 3. check
def check(D: dict, plan: list[dict]) -> dict:
    """Independent rule check of a layout, week by week, from the committed rows only (not the solver's own
    bookkeeping). Works on the engine's layout too (D['engine_plan']), so both are judged by the same rules."""
    SYS = {t[0]: t[1] for t in D['tanks']}; VOL = {t[0]: t[2] for t in D['tanks']}; CAP0 = {t[0]: t[3] for t in D['tanks']}
    BUF = 1 + D['buffer']; mm = D['min_move']
    issues = collections.Counter(); worst = collections.defaultdict(float); moves = {}
    for p in plan:
        w = p['week']; rows = p['rows']
        occ = collections.defaultdict(set); sb = collections.defaultdict(float); sf = collections.defaultdict(float)
        btot = collections.defaultdict(float); by_b = collections.defaultdict(lambda: [0, 0]); stg = 0
        for r in rows:
            t = int(r['t']); b = r['b']; rec = D['batch'].get(b, {}).get(w)
            if rec is None:
                continue
            if r['n'] > 0.5:
                occ[t].add(b); kg = r['dens'] * VOL[t]; btot[b] += r['n']
                sb[SYS[t]] += kg; sf[SYS[t]] += kg * rec['sfr']
                cap = D['dens_cap'].get(w) or CAP0[t]
                if r['dens'] > cap + 0.05:
                    issues['density over cap (tank-weeks)'] += 1; worst['density_over_kg_m3'] = max(worst['density_over_kg_m3'], r['dens'] - cap)
                if r['n'] < D['min_tank'] - 1 and rec['n'] >= D['min_tank']:
                    issues['tank under min fish'] += 1
            if SYS[t] in OG12 and r['inn'] > 0.5 and open_wt(D['batch'], b, w, D['weeks']) >= 1.0:
                issues['entry split at >= 1 kg'] += 1
            if SYS[t] in OG12 and r['st'] > 0.5:
                issues['6N staging from an entry tank'] += 1
            if D['move_rules']:
                if r['out'] > 0.5 and r['n'] > 0.5 and r['out'] < mm - 1:
                    issues['partial move under min'] += 1
                if r['st'] > 0.5 and r['n'] > 0.5 and r['st'] < min(mm, rec['stage']) - 1:
                    issues['partial 6N staging under min'] += 1
                if r['inn'] > 0.5 and r['inn'] < mm - 1:
                    issues['received piece under min'] += 1
            if r['out'] > 0.5: by_b[b][0] += 1
            if r['inn'] > 0.5: by_b[b][1] += 1
            if r['st'] > 0.5: stg += 1
        issues['two batches in a tank'] += sum(1 for bs in occ.values() if len(bs) > 1)
        for s, kg in sb.items():
            bc, fc = sys_cap(D, w, s)
            if bc and kg > bc * BUF + 1:
                issues['system biomass over cap x buffer'] += 1; worst['sys_bio_over_t'] = max(worst['sys_bio_over_t'], (kg - bc * BUF) / 1e3)
            if fc and sf[s] > fc * BUF + 1:
                issues['system feed over cap x buffer'] += 1; worst['sys_feed_over_kg_d'] = max(worst['sys_feed_over_kg_d'], sf[s] - fc * BUF)
        for b, rec in D['batch'].items():
            if abs(btot.get(b, 0) - rec[w]['n']) > max(150, 0.002 * rec[w]['n']):
                issues['batch count differs from the plan'] += 1
        ev = stg + sum(max(0, s + rr - 1) for s, rr in by_b.values())
        moves[w] = ev
        if ev > D['max_moves']:
            issues['weeks over the move budget'] += 1
    return dict(issues={k: v for k, v in issues.items() if v}, worst=dict(worst), moves=moves,
                ok=not any(issues.values()))


# ------------------------------------------------------------------------------------------------ 4. transfers
def pair_moves(plan: list[dict]) -> list[dict]:
    """Each week, per batch: senders' out amounts are matched to receivers' in amounts, biggest first, into
    from -> to transfers (the fewest legs that move those amounts). 6N staging rows are listed as To_Tank '6N'."""
    legs = []
    for p in plan:
        by_b = collections.defaultdict(lambda: ([], []))
        for r in p['rows']:
            if r['out'] > 0.5: by_b[r['b']][0].append([r['t'], r['out']])
            if r['inn'] > 0.5: by_b[r['b']][1].append([r['t'], r['inn']])
            if r['st'] > 0.5:
                legs.append(dict(week=p['week'], batch=r['b'], frm=r['t'], to='6N', fish=r['st']))
            if r.get('arr', 0) > 0.5:
                legs.append(dict(week=p['week'], batch=r['b'], frm='FW', to=r['t'], fish=r['arr']))
        for b, (snd, rcv) in sorted(by_b.items()):
            snd.sort(key=lambda v: -v[1]); rcv.sort(key=lambda v: -v[1])
            i = j = 0
            while i < len(snd) and j < len(rcv):
                q = min(snd[i][1], rcv[j][1])
                legs.append(dict(week=p['week'], batch=b, frm=snd[i][0], to=rcv[j][0], fish=q))
                snd[i][1] -= q; rcv[j][1] -= q
                if snd[i][1] < 0.5: i += 1
                if rcv[j][1] < 0.5: j += 1
    return legs


# ------------------------------------------------------------------------------------------------ 5. workbook
def transfers_per_fish(D: dict, plan: list[dict]) -> dict:
    """Fish moved between OG tanks / fish harvested, overall and per year (the operator's yardstick)."""
    mv = collections.Counter(); hv = collections.Counter()
    for p in plan:
        y = p['week'][:4]
        mv[y] += sum(r['inn'] for r in p['rows']); hv[y] += D['harvested'][p['week']]
    out = {y: (mv[y] / hv[y] if hv[y] else 0.0) for y in sorted(hv)}
    out['all'] = sum(mv.values()) / sum(hv.values()) if sum(hv.values()) else 0.0
    return out


def write_workbook(D: dict, R: dict, out_path: str) -> None:
    import openpyxl
    from openpyxl.styles import Font
    eng_plan = D['engine_plan']; new_plan = R['plan']
    weeks = {p['week'] for p in new_plan}
    eng_plan = [p for p in eng_plan if p['week'] in weeks]
    ce, cn = check(D, eng_plan), check(D, new_plan)
    te, tn = transfers_per_fish(D, eng_plan), transfers_per_fish(D, new_plan)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Summary'; bold = Font(bold=True)
    ws.append(['LEAST-MOVES LAYOUT (Global v2)']); ws['A1'].font = Font(bold=True, size=14)
    ws.append(['Same plan as the source workbook (fish per batch per week, arrivals, 6N staging, harvest, HOG), '
               'laid out in the tanks with as few moves as possible. The source workbook is unchanged.'])
    ws.append(['Source workbook', D['source']])
    ws.append(['Weeks re-planned', f"{new_plan[0]['week']} .. {new_plan[-1]['week']} ({len(new_plan)} weeks; "
                                   f"weeks before {D['weeks'][0]} are the manual starting events and stay as they are)"])
    s = R['settings']
    ws.append(['Solver settings', f"look ahead {s['horizon']} wk, keep {s['commit']} wk per step, {s['secs']:.0f} s per "
                                  f"step, {s['workers']} threads; run time {R['seconds'] / 60:.0f} min"])
    ws.append([])
    ws.append(['Transfers per harvested fish', 'Engine plan', 'Least-moves', 'Change']); ws.cell(ws.max_row, 1).font = bold
    for y in te:
        ws.append([('All weeks' if y == 'all' else y), round(te[y], 2), round(tn[y], 2), round(tn[y] - te[y], 2)])
    me, mn = list(ce['moves'].values()), list(cn['moves'].values())
    ws.append(['Move events per week (mean / max)', f'{sum(me) / len(me):.1f} / {max(me)}', f'{sum(mn) / len(mn):.1f} / {max(mn)}'])
    ws.append([])
    ws.append(['Rule check (same checker on both)', 'Engine plan', 'Least-moves']); ws.cell(ws.max_row, 1).font = bold
    for k in sorted(set(ce['issues']) | set(cn['issues']) | {'density over cap (tank-weeks)', 'weeks over the move budget'}):
        ws.append([k, ce['issues'].get(k, 0), cn['issues'].get(k, 0)])
    for k in sorted(set(ce['worst']) | set(cn['worst'])):
        ws.append([f'worst {k}', round(ce['worst'].get(k, 0), 1), round(cn['worst'].get(k, 0), 1)])
    ws.append([])
    ws.append(['How to read']); ws.cell(ws.max_row, 1).font = bold
    for line in ('Transfers: every move this layout needs, week by week (FW = TranOG arrival placement, 6N = staging for '
                 'harvest; the 6N tank itself is as in the source plan).',
                 'Layout: which batch is in which tank at the end of each week, with fish and density.',
                 'Batches: transfers per harvested fish for each batch, engine plan vs this layout.',
                 'Harvest, HOG, biomass and feed are the source workbook\'s - this layout changes none of them.'):
        ws.append([line])
    ws.column_dimensions['A'].width = 44
    for c in 'BCD':
        ws.column_dimensions[c].width = 16
    wt = wb.create_sheet('Transfers'); wt.append(['Week', 'Batch', 'From_Tank', 'To_Tank', 'Count (fish)'])
    for g in pair_moves(new_plan):
        wt.append([g['week'], g['batch'], g['frm'], g['to'], round(g['fish'])])
    wl = wb.create_sheet('Layout'); wl.append(['Week', 'Tank', 'System', 'Batch', 'Count (fish)', 'AvgWt (kg)', 'Density (kg/m3)'])
    SYS = {t[0]: t[1] for t in D['tanks']}
    for p in new_plan:
        for r in sorted(p['rows'], key=lambda r: r['t']):
            if r['n'] > 0.5:
                wl.append([p['week'], r['t'], SYS[r['t']], r['b'], round(r['n']), round(D['batch'][r['b']][p['week']]['wt'], 2),
                           round(r['dens'], 1)])
    wbat = wb.create_sheet('Batches')
    wbat.append(['Batch', 'Fish harvested (re-planned weeks)', 'Transfers per fish - engine', 'Transfers per fish - least-moves'])
    mv_e = collections.Counter(); mv_n = collections.Counter(); st = collections.Counter()
    for p in eng_plan:
        for r in p['rows']:
            mv_e[r['b']] += r['inn']
    for p in new_plan:
        for r in p['rows']:
            mv_n[r['b']] += r['inn']; st[r['b']] += r['st']
    for b in sorted(D['batch'], key=lambda b: int(b[1:]) if b[1:].isdigit() else 0):
        if st[b] > 0:
            wbat.append([b, round(st[b]), round(mv_e[b] / st[b], 2), round(mv_n[b] / st[b], 2)])
    wb.save(out_path)


def run(workbook: str, config_dir: str, scenario_dir: str, out_path: str | None = None, start_week: str | None = None,
        log=print, **solve_kw) -> dict:
    """read -> solve -> check -> write. Returns the summary numbers. The companion workbook is written next to the
    source as '<source name>_least_moves.xlsx' unless out_path is given."""
    D = read_endpoint(workbook, config_dir, scenario_dir, start_week)
    log(f"least-moves: {len(D['weeks'])} weeks {D['weeks'][0]}..{D['weeks'][-1]}, {len(D['batch'])} batches, "
        f"{len(D['tanks'])} tanks")
    R = solve(D, log=log, **solve_kw)
    out_path = out_path or str(Path(workbook).with_name(Path(workbook).stem + '_least_moves.xlsx'))
    write_workbook(D, R, out_path)
    weeks = {p['week'] for p in R['plan']}
    eng = [p for p in D['engine_plan'] if p['week'] in weeks]
    res = dict(out=out_path, tpf=transfers_per_fish(D, R['plan']), engine_tpf=transfers_per_fish(D, eng),
               check=check(D, R['plan'])['issues'], engine_check=check(D, eng)['issues'])
    log(f"transfers per harvested fish: engine {res['engine_tpf']['all']:.2f} -> least-moves {res['tpf']['all']:.2f}; "
        f"rule check {res['check'] or 'clean'}; saved {out_path}")
    return res
