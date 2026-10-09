"""The readable `RunConfig` tab: the run's config + scenario as formatted tables.

Operator, 2026-10-02: "make the RunConfig output more visually pleasing". The
sheet used to be ~1,800 lines of raw YAML in column A. The YAML still travels,
unchanged, on the `RunConfig YAML` tab (forecast.config_snapshot) -- that is
the record import reads. This tab is for people: run identity, then one
section per file, each rendered as a Setting | Value | Meaning list or a
proper table, with the big tables outlined so they collapse.

PRESENTATION ONLY, with three load-bearing rules:
  * A1 keeps the KIND_SNAPSHOT_MARK so run_config_kind still says "snapshot".
  * No cell ever holds a `# ===== key =====` block marker, so the old-layout
    fallback reader can never parse this tab as YAML.
  * Control settings sit as key in A / value in B, the dialect
    analysis.convergence_review already reads, and per-week limit METRICS are
    column headers (never a key cell followed by a number), so a per-week
    `min_harvest_per_week` row cannot shadow the control value there.

Meanings are taken from what the project already documents -- the
docs/USER_GUIDE.md settings tables, first sentence -- never invented here. A
key the guide does not document gets a blank meaning.
"""
from __future__ import annotations

import re
from pathlib import Path

import yaml

# Section/table look borrowed from excel_format so the tab matches the others.
_MAX_COLS = 12
_COLLAPSE_AT = 15        # tables longer than this start collapsed

_MEANINGS_CACHE: dict | None = None


def _meanings() -> dict:
    """{setting: first sentence of its USER_GUIDE description}. Empty when the
    guide is not alongside the code (the view then shows no meanings)."""
    global _MEANINGS_CACHE
    if _MEANINGS_CACHE is not None:
        return _MEANINGS_CACHE
    out: dict = {}
    try:
        guide = Path(__file__).resolve().parent.parent / "docs" / "USER_GUIDE.md"
        for m in re.finditer(r"^\| `([a-z_0-9]+)` \| (.+?) \|",
                             guide.read_text(encoding="utf-8"), re.M):
            key, text = m.group(1), m.group(2)
            if key in out:
                continue                      # first (settings-table) wins
            out[key] = _first_sentence(text)
    except Exception:  # noqa: BLE001 - meanings are a courtesy
        out = {}
    _MEANINGS_CACHE = out
    return out


def _first_sentence(md: str, limit: int = 160) -> str:
    t = re.sub(r"\*\*|`|\*", "", md).strip()
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    m = re.search(r"(.+?[.;:])(\s|$)", t)
    t = m.group(1) if m else t
    t = t.rstrip(".;:")
    return t if len(t) <= limit else t[: limit - 1].rstrip() + "…"


# --------------------------------------------------------------------------- #
# Value rendering
# --------------------------------------------------------------------------- #
def _scalar(v):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float, str)):
        return v
    return _flow(v)


def _flow(v) -> str:
    """Compact one-line text for a nested value."""
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_flow(x)}" for k, x in v.items())
    if isinstance(v, list):
        if v and all(isinstance(x, dict) and "tank" in x for x in v):
            parts = []
            for d in v:
                s = str(d.get("tank"))
                if d.get("count") is not None:
                    s += f":{d['count']:,.0f}" if isinstance(d["count"], (int, float)) else f":{d['count']}"
                if d.get("size_class"):
                    s += f" ({d['size_class']})"
                parts.append(s)
            return ", ".join(parts)
        return ", ".join(_flow(x) for x in v)
    if isinstance(v, float):
        return f"{v:,.4g}" if abs(v) < 1000 else f"{v:,.0f}"
    if v is None:
        return "—"
    return str(v)


def _number_format(v) -> str | None:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    if isinstance(v, int) or float(v).is_integer():
        return "#,##0" if abs(v) >= 1000 else "0"
    return "#,##0.00" if abs(v) >= 1000 else "General"


# --------------------------------------------------------------------------- #
# The writer
# --------------------------------------------------------------------------- #
class _View:
    def __init__(self, ws):
        from . import excel_format as xf
        self.ws, self.xf = ws, xf
        self.row = 0
        self.sections: list[int] = []
        self.headers: list[int] = []
        self.notes: list[int] = []

    def put(self, values, *, note=False):
        self.row += 1
        for c, v in enumerate(values, start=1):
            if v is None:
                continue
            cell = self.ws.cell(self.row, c, v)
            fmt = "yyyy-mm-dd" if hasattr(v, "year") else _number_format(v)
            if fmt:
                cell.number_format = fmt
        if note:
            self.notes.append(self.row)
        return self.row

    def blank(self):
        self.row += 1

    def section(self, title, note=None):
        self.blank()
        self.sections.append(self.put([title]))
        if note:
            self.put([note], note=True)

    def table(self, header, rows, collapse=None, label=None):
        rows = list(rows)
        if collapse is None:
            collapse = len(rows) > _COLLAPSE_AT
        if label is not None:
            self.put([f"{label} — {len(rows)} rows, collapsed: click + in the "
                      f"left margin to expand" if collapse and rows else label])
        self.headers.append(self.put(header))
        first = self.row + 1
        for r in rows:
            self.put([_scalar(v) for v in r])
        last = self.row
        if last >= first and collapse:
            self.ws.row_dimensions.group(first, last, outline_level=1,
                                         hidden=True)

    def settings(self, items, meanings):
        self.headers.append(self.put(["Setting", "Value", "Meaning (USER_GUIDE)"]))
        for k, v in items:
            self.put([k, _scalar(v), meanings.get(k, "")])


def write_readable_view(ws, dirs, *, files, run_info=None, pr_closing=None,
                        yaml_sheet="RunConfig YAML", title="",
                        excluded_line="") -> None:
    v = _View(ws)
    meanings = _meanings()
    v.put([title])
    v.put([f"The run's config and scenario as tables. The exact YAML (what "
           f"Configure > 'Template & import' > 'Import config' restores) is "
           f"on the '{yaml_sheet}' tab. "
           f"Rows marked + are collapsed: click + in the left margin to "
           f"expand."], note=True)
    v.put([excluded_line], note=True)
    link_row = v.put([f"→ {yaml_sheet} (re-importable YAML)",
                      None, "→ BatchInputs (per-batch inputs)"])
    ws.cell(link_row, 1).hyperlink = f"#'{yaml_sheet}'!A1"
    ws.cell(link_row, 3).hyperlink = "#'BatchInputs'!A1"

    if run_info:
        v.section("RUN")
        v.headers.append(v.put(["Item", "Value", "Meaning"]))
        for label, value, meaning in run_info:
            v.put([label, _scalar(value) if not hasattr(value, "year") else value,
                   meaning])

    for name, kind in files:
        base = dirs.get(kind)
        if base is None:
            continue
        p = Path(base) / name
        if not p.exists():
            continue
        try:
            doc = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
        except Exception as exc:  # noqa: BLE001
            v.section(f"{kind}/{name}",
                      f"Could not be parsed for display ({type(exc).__name__}); "
                      f"see the '{yaml_sheet}' tab.")
            continue
        render = _RENDERERS.get(name, _render_generic)
        render(v, f"{kind}/{name}", doc, meanings)

    if pr_closing is not None and dirs.get("scenario"):
        _render_manual_events(v, dirs["scenario"], pr_closing)

    _style(v)


# --------------------------------------------------------------------------- #
# Per-file renderers
# --------------------------------------------------------------------------- #
def _render_control(v, key, doc, meanings):
    v.section(f"CONTROL — {key}",
              "Run settings: caps, horizon and planner knobs. forecast_start "
              "is a seed; the run derives it from the ProductionReport.")
    v.settings(list(doc.items()), meanings)


def _render_biology(v, key, doc, meanings):
    v.section(f"BIOLOGY — {key}", "Growth, feed-conversion, mortality, feed "
              "and culling curves shared by every batch.")
    used = set()

    def cols(size_key, series):
        n = len(doc.get(size_key) or [])
        return [[doc[size_key][i]] + [((doc.get(k) or [None] * n) + [None] * n)[i]
                                      for k, _h in series] for i in range(n)]

    if "sgr_size_g" in doc:
        series = [(k, h) for k, h in (("sgr_fw_pct_day", "SGR FW (%/day)"),
                                      ("sgr_sw_pct_day", "SGR SW (%/day)"))
                  if k in doc]
        v.table(["Size (g)"] + [h for _k, h in series], cols("sgr_size_g", series),
                label="SGR curves (by fish size)")
        used |= {"sgr_size_g"} | {k for k, _h in series}
    if "fcr_size_g" in doc:
        models = doc.get("fcr_by_model") or {}
        n = len(doc["fcr_size_g"])
        rows = [[doc["fcr_size_g"][i]] + [(list(models[m]) + [None] * n)[i]
                                          for m in models] for i in range(n)]
        v.table(["Size (g)"] + [f"FCR {m}" for m in models], rows,
                label="FCR curves (by fish size, one column per model)")
        used |= {"fcr_size_g", "fcr_by_model"}
    if "mortality_week_from_input" in doc:
        series = [("mortality_pct_weekly", "Mortality (% per week)")]
        v.table(["Week from input", "Mortality (% per week)"],
                cols("mortality_week_from_input", series),
                label="Mortality (by week since input)")
        used |= {"mortality_week_from_input", "mortality_pct_weekly"}
    if "feed_types" in doc:
        v.table(["Up to size (g)", "Feed"], [list(r) for r in doc["feed_types"]],
                label="Feed types (largest fish size each feed is used up to)")
        used.add("feed_types")
    if "culling" in doc:
        v.table(["Days since input", "Cull (%)"], [list(r) for r in doc["culling"]],
                label="Freshwater culling schedule")
        used.add("culling")
    rest = {k: x for k, x in doc.items() if k not in used}
    if rest:
        _render_items(v, rest, meanings)


def _render_facility(v, key, doc, meanings):
    tanks = doc.get("tanks") or []
    v.section(f"FACILITY — {key}", f"{len(tanks)} tanks: volume, density cap "
              f"and feed cap per tank.")
    _render_items(v, doc, meanings)


def _render_batches(v, key, doc, meanings):
    bs = doc.get("batches") or []
    v.section(f"BATCHES — {key}", f"{len(bs)} batches. Their inputs, side by "
              f"side with what this run did with them, are on the BatchInputs "
              f"tab (one row per batch); the exact YAML is on the YAML tab.")
    rest = {k: x for k, x in doc.items() if k != "batches"}
    if rest:
        _render_items(v, rest, meanings)


def _render_limits(v, key, doc, meanings):
    v.section(f"LIMITS — {key}", "System capacities (standing defaults, then "
              "one-off weekly exceptions) and per-week facility rows.")
    sd = doc.get("system_defaults") or {}
    if sd:
        metrics = []
        for d in sd.values():
            for m in (d or {}):
                if m != "modes" and m not in metrics:
                    metrics.append(m)
        rows = []
        for sysid, d in sd.items():
            d = d or {}
            rows.append([sysid] + [d.get(m) for m in metrics]
                        + [_flow(d["modes"]) if d.get("modes") else ""])
        v.table(["System"] + metrics + ["By mode"], rows, collapse=False,
                label="System defaults (standing capacity per system)")
    for part, label, index in (("system", "System exceptions (one row per week, "
                                          "system, metric)", ("week", "system")),
                               ("facility", "Facility rows per week (absent = "
                                            "the control default)", ("week",))):
        rows = doc.get(part)
        if rows is None:
            continue
        if not rows:
            v.put([label])
            v.put(["(none)"])
            continue
        _pivot(v, rows, index, label)
    rest = {k: x for k, x in doc.items()
            if k not in ("system_defaults", "system", "facility")}
    if rest:
        _render_items(v, rest, meanings)


def _pivot(v, rows, index, label=None):
    """week/metric/value rows -> one row per index, one column per metric.
    The metric is a HEADER, never a key cell beside a number (see module
    docstring)."""
    if not all(isinstance(r, dict) and {"metric", "value"} <= set(r) for r in rows):
        _table_of_dicts(v, rows, label)
        return
    metrics, grid = [], {}
    for r in rows:
        m = str(r["metric"])
        if m not in metrics:
            metrics.append(m)
        k = tuple(str(r.get(i, "")) for i in index)
        grid.setdefault(k, {})[m] = r["value"]
    body = [list(k) + [grid[k].get(m) for m in metrics]
            for k in sorted(grid)]
    v.table([i.capitalize() for i in index] + metrics, body, label=label)


def _render_generic(v, key, doc, meanings):
    v.section(key)
    if isinstance(doc, dict):
        _render_items(v, doc, meanings)
    else:
        v.put([_flow(doc)])


def _render_items(v, doc, meanings):
    """Scalars as a settings list; each list/dict as its own table."""
    scalars = [(k, x) for k, x in doc.items()
               if not isinstance(x, (list, dict))]
    if scalars:
        v.settings(scalars, meanings)
    for k, x in doc.items():
        if isinstance(x, list):
            if not x:
                v.put([k])
                v.put(["(none)"])
            elif all(isinstance(r, dict) for r in x):
                _table_of_dicts(v, x, k)
            elif all(isinstance(r, list) for r in x):
                width = max(len(r) for r in x)
                v.table([f"{k} [{i + 1}]" for i in range(width)], x, label=k)
            else:
                v.put([k])
                v.put(["values", _flow(x)])
        elif isinstance(x, dict):
            v.put([k])
            if not x:
                v.put(["(none)"])
            elif all(isinstance(d, dict) for d in x.values()):
                cols = []
                for d in x.values():
                    for c in d:
                        if c not in cols:
                            cols.append(c)
                v.table([k] + cols, [[n] + [d.get(c) for c in cols]
                                     for n, d in x.items()], label=None)
            else:
                v.settings(list(x.items()), meanings)


def _table_of_dicts(v, rows, label=None):
    cols = []
    for r in rows:
        for c in r:
            if c not in cols:
                cols.append(c)
    v.table(cols, [[r.get(c) for c in cols] for r in rows], label=label)


def _render_manual_events(v, scenario_dir, pr_closing):
    try:
        from .manual_events import events_file_for
        p = events_file_for(scenario_dir, pr_closing)
    except Exception:  # noqa: BLE001
        return
    v.section(f"MANUAL EVENTS — scenario/manual_events/{p.name}",
              "The operator-scripted window for this PR. Also on the YAML "
              "tab: an import restores it (the replaced file is kept as "
              "*.yaml.before-import).")
    if not p.exists():
        v.put(["(no events file for this PR)"])
        return
    try:
        evs = (yaml.safe_load(p.read_text(encoding="utf-8")) or {}).get("events") or []
    except Exception as exc:  # noqa: BLE001
        v.put([f"(could not be parsed: {type(exc).__name__})"])
        return
    if not evs:
        v.put(["(no events)"])
        return
    cols = ["week", "type", "batch", "from_tank", "destinations", "count",
            "mode", "notes"]
    extra = [c for e in evs for c in e if c not in cols]
    cols += [c for i, c in enumerate(extra) if c not in extra[:i]]
    rows = sorted(evs, key=lambda e: (int(e.get("week") or 0),))
    v.table(["Week", "Type", "Batch", "From tank", "Destinations (tank:count)",
             "Count", "Mode", "Notes"] + cols[8:],
            [[e.get(c) for c in cols] for e in rows], collapse=False)


_RENDERERS = {
    "control.yaml": _render_control,
    "biology.yaml": _render_biology,
    "facility.yaml": _render_facility,
    "batches.yaml": _render_batches,
    "limits.yaml": _render_limits,
}


# --------------------------------------------------------------------------- #
# Look
# --------------------------------------------------------------------------- #
def _style(v) -> None:
    from openpyxl.styles import Alignment
    xf, ws = v.xf, v.ws
    for c in range(1, _MAX_COLS + 1):
        ws.cell(1, c).fill = xf.C_TITLE
    ws.cell(1, 1).font = xf.F_TITLE
    for r in v.notes:
        ws.cell(r, 1).font = xf.F_NOTE
    for r in v.sections:
        for c in range(1, _MAX_COLS + 1):
            ws.cell(r, c).fill = xf.C_HDR
        ws.cell(r, 1).font = xf.F_HDR
    for r in v.headers:
        for c in range(1, _MAX_COLS + 1):
            cell = ws.cell(r, c)
            if cell.value is None:
                continue
            cell.font = xf.F_HDR
            cell.border = xf.B_HDR
            cell.alignment = xf.A_HDR
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = "A5"
    widths = {"A": 34, "B": 22, "C": 60}
    for i in range(4, _MAX_COLS + 1):
        widths[chr(64 + i)] = 14
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for r in range(5, v.row + 1):
        cell = ws.cell(r, 3)
        if isinstance(cell.value, str) and len(cell.value) > 60:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        b = ws.cell(r, 2)
        if b.value is not None and r not in v.headers:
            b.alignment = Alignment(horizontal="left")
    ws.sheet_view.zoomScale = 100
