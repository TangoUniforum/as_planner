"""Judge a knob (or a method) across the WHOLE PR corpus, not on one PR.

WHY THIS EXISTS. On 2026-09-19 density_relief_pct 0.85 looked like the best
lever of the session on a single PR -- fewer moves, lower harvest CV, and it
fixed a processing-cap breach. Across the 17 corpus PRs that both arms could
plan it was a wash or worse: HOG mean -0.41% (better on 6 of 17), weeks below
the harvest floor worse on 6 and better on 2, and it made one PR INFEASIBLE
that plans fine today. The planner is mode-discontinuous -- tiny knob nudges
swing months by 100+ t -- so a single-PR win is not evidence, it is the
discontinuity.

Runs are seconds, so the corpus costs minutes. This turns "remember to check
the corpus" into one command:

    python -m tools.knob_corpus --set density_relief_pct=0.85
    python -m tools.knob_corpus --method controller-lns --identical

--identical is the DELETION check: it asks whether the variant produces a
BYTE-IDENTICAL plan on every PR, which is what has to be true before a knob or
a method can be removed as dead.

Nothing here touches production files: every run is an isolated temp copy of
config + scenario via forecast.methods.run_method, and costs.yaml is never
copied.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import statistics
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

# Sheets that define the PLAN. A knob that leaves these identical on every PR
# changed nothing an operator acts on.
PLAN_SHEETS = ("BatchLocations", "TransferPlan", "HarvestPlan")


def parse_set(pairs) -> dict:
    """knob=value pairs to a dict, read as YAML scalars so 3200 stays an int,
    0.85 a float, false a bool and an empty value None."""
    out = {}
    for p in pairs or ():
        if "=" not in p:
            raise ValueError("--set wants knob=value, got %r" % (p,))
        k, v = p.split("=", 1)
        out[k.strip()] = yaml.safe_load(v) if v.strip() != "" else None
    return out


def write_config(src_config: Path, dst: Path, knobs: dict) -> Path:
    """A temp config with `knobs` applied to control.yaml. costs.yaml is never
    copied -- no corpus comparison needs it, and it holds real cost data."""
    shutil.rmtree(dst, ignore_errors=True)
    shutil.copytree(src_config, dst, ignore=shutil.ignore_patterns("costs.yaml*"))
    p = dst / "control.yaml"
    doc = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    doc.update(knobs)
    p.write_text(yaml.safe_dump(doc, sort_keys=True), encoding="utf-8")
    return dst


def plan_hash(workbook) -> str:
    """The plan's identity: every row of the sheets an operator acts on."""
    wb = openpyxl.load_workbook(str(workbook), data_only=True, read_only=True)
    try:
        h = hashlib.sha256()
        for name in PLAN_SHEETS:
            if name not in wb.sheetnames:
                continue
            h.update(name.encode())
            for row in wb[name].iter_rows():
                h.update(repr([c.value for c in row]).encode())
        return h.hexdigest()[:16]
    finally:
        wb.close()


def _table(wb, sheet, key="Week"):
    rows = [[c.value for c in r] for r in wb[sheet].iter_rows()]
    for i, r in enumerate(rows):
        if r and str(r[0]).strip() == key:
            hdr = [str(v).strip() if v else "" for v in r]
            cols = {n.split(" (")[0]: j for j, n in enumerate(hdr) if n}
            return cols, rows[i + 1:]
    return {}, []


def score(workbook, floor: float, cap: float) -> dict:
    """The few numbers a knob is judged on, from the workbook it wrote."""
    wb = openpyxl.load_workbook(str(workbook), data_only=True, read_only=True)
    try:
        cols, body = _table(wb, "YearlySummary", "Year")
        hog = 0.0
        if cols:
            for r in body:
                if str(r[0]).strip().isdigit():
                    hog += float(r[cols["Harvest_HOG"]] or 0)
        cols, body = _table(wb, "HarvestPlan")
        per: dict = {}
        for r in body:
            w = r[cols["Week"]] if cols else None
            if w and str(w).startswith("20"):
                per[str(w)] = per.get(str(w), 0.0) + float(r[cols["Count"]] or 0)
        weekly = list(per.values())
        cols, body = _table(wb, "TransferPlan")
        moves = 0
        for r in body:
            if not cols or not r[cols["Week"]]:
                continue
            if str(r[cols["Week"]]).startswith("20") and \
                    str(r[cols["Type"]]).strip() == "Transfer":
                moves += 1
        mean = statistics.mean(weekly) if weekly else 0.0
        return {
            "hog_t": hog,
            "moves": moves,
            "harvest_cv": (statistics.pstdev(weekly) / mean) if mean else 0.0,
            "weeks_below_floor": sum(1 for x in weekly if x < floor),
            "weeks_over_cap": sum(1 for x in weekly if x > cap),
        }
    finally:
        wb.close()


def summarise(rows: list, identical_mode: bool) -> dict:
    """rows: per-PR dicts carrying base/variant scores and plan hashes."""
    ok = [r for r in rows if r.get("base") and r.get("variant")]
    out = {
        "n": len(ok),
        "failed": [r["pr"] for r in rows
                   if r.get("base") and not r.get("variant")],
        "identical": sum(1 for r in ok if r["base_hash"] == r["variant_hash"]),
    }
    if identical_mode or not ok:
        return out
    d_hog = [100.0 * (r["variant"]["hog_t"] - r["base"]["hog_t"]) / r["base"]["hog_t"]
             for r in ok if r["base"]["hog_t"]]
    d_mv = [r["variant"]["moves"] - r["base"]["moves"] for r in ok]
    out.update({
        "hog_mean_pct": statistics.mean(d_hog) if d_hog else 0.0,
        "hog_better": sum(1 for v in d_hog if v > 0),
        "moves_mean": statistics.mean(d_mv) if d_mv else 0.0,
        "moves_fewer": sum(1 for v in d_mv if v < 0),
        "floor_worse": sum(1 for r in ok if r["variant"]["weeks_below_floor"]
                           > r["base"]["weeks_below_floor"]),
        "floor_better": sum(1 for r in ok if r["variant"]["weeks_below_floor"]
                            < r["base"]["weeks_below_floor"]),
        "cap_worse": sum(1 for r in ok if r["variant"]["weeks_over_cap"]
                         > r["base"]["weeks_over_cap"]),
    })
    return out


def verdict(summary: dict, identical_mode: bool) -> str:
    """One line an operator can act on."""
    n = summary["n"]
    if summary["failed"]:
        return ("REJECT -- the variant made %d PR(s) INFEASIBLE that plan fine "
                "today: %s" % (len(summary["failed"]),
                               ", ".join(summary["failed"])))
    if not n:
        return "NO EVIDENCE -- no PR ran on both arms."
    if identical_mode:
        if summary["identical"] == n:
            return "DEAD -- identical plan on every PR; safe to delete."
        return ("NOT DEAD -- the plan differs on %d of %d PRs."
                % (n - summary["identical"], n))
    if summary["identical"] == n:
        return "INERT -- identical plan on every PR; it changes nothing."
    parts = ["HOG %+.2f%% mean (better on %d/%d)"
             % (summary["hog_mean_pct"], summary["hog_better"], n),
             "moves %+.1f (fewer on %d/%d)"
             % (summary["moves_mean"], summary["moves_fewer"], n)]
    if summary["floor_worse"] or summary["cap_worse"]:
        parts.append("floor worse on %d, cap breaches worse on %d"
                     % (summary["floor_worse"], summary["cap_worse"]))
    return " | ".join(parts)


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--set", action="append", metavar="KNOB=VALUE",
                    help="control.yaml knob to change (repeatable)")
    ap.add_argument("--method", help="a methods.py key to run as the variant")
    ap.add_argument("--corpus", default=str(ROOT / "pr_corpus"))
    ap.add_argument("--config-dir", default=str(ROOT / "config"))
    ap.add_argument("--scenario-dir", default=str(ROOT / "scenario"))
    ap.add_argument("--base-method", default="controller")
    ap.add_argument("--identical", action="store_true",
                    help="deletion check: is the plan byte-identical on EVERY PR?")
    ap.add_argument("--limit", type=int, default=0, help="first N PRs only")
    a = ap.parse_args()
    if not a.set and not a.method:
        ap.error("give --set KNOB=VALUE and/or --method")

    from forecast import methods as M
    knobs = parse_set(a.set)
    cfg = Path(a.config_dir)
    ctl = yaml.safe_load((cfg / "control.yaml").read_text(encoding="utf-8")) or {}
    floor = float(ctl.get("min_harvest_per_week") or 0)
    cap = float(ctl.get("max_harvest_per_week") or 10 ** 9)
    prs = sorted(Path(a.corpus).glob("*.xls*"))
    if a.limit:
        prs = prs[:a.limit]
    head = "%d PRs | baseline %s" % (len(prs), a.base_method)
    if knobs:
        head += " | knobs %s" % (knobs,)
    if a.method:
        head += " | method %s" % (a.method,)
    print(head, flush=True)

    work = Path(tempfile.mkdtemp(prefix="knobcorpus_"))
    try:
        base_cfg = write_config(cfg, work / "cfg_base", {})
        var_cfg = write_config(cfg, work / "cfg_var", knobs)
        rows = []
        for pr in prs:
            row = {"pr": pr.stem}
            arms = (("base", base_cfg, a.base_method),
                    ("variant", var_cfg, a.method or a.base_method))
            for tag, cdir, meth in arms:
                out = work / ("%s_%s.xlsm" % (pr.stem, tag))
                # The two failures are DIFFERENT and must never be conflated:
                # the engine refusing to plan a PR is evidence about the knob,
                # while this tool failing to read the workbook is a bug in this
                # tool. Reporting the second as the first is how a scoring typo
                # became "NO EVIDENCE -- no PR ran on both arms" on 2026-09-21,
                # when in fact every run had succeeded.
                try:
                    M.run_method(M.REGISTRY[meth], str(pr), str(out),
                                 str(cdir), a.scenario_dir, quiet=True)
                except Exception as exc:                      # noqa: BLE001
                    row[tag] = None
                    row[tag + "_err"] = ("%s: %s" % (type(exc).__name__, exc))[:90]
                    continue
                try:
                    row[tag] = score(out, floor, cap)
                    row[tag + "_hash"] = plan_hash(out)
                except Exception as exc:                      # noqa: BLE001
                    raise RuntimeError(
                        "the run succeeded but this tool could not score %s (%s): "
                        "%s: %s" % (pr.stem, tag, type(exc).__name__, exc)) from exc
            rows.append(row)
            b, v = row.get("base"), row.get("variant")
            if b and v:
                same = row["base_hash"] == row["variant_hash"]
                # format(), not %-formatting: "%+,.0f" is not valid there and
                # this branch only runs when a PR DIFFERS, so the crash hid the
                # one result that mattered (2026-09-21).
                extra = "" if same else (
                    "  HOG %s t  moves %+d"
                    % (format(v["hog_t"] - b["hog_t"], "+,.0f"),
                       v["moves"] - b["moves"]))
                print("  %-12s %s%s" % (row["pr"], "same" if same else "DIFFERS",
                                        extra), flush=True)
            else:
                which = "both arms" if not b else "the variant"
                print("  %-12s FAILED (%s)" % (row["pr"], which), flush=True)
        s = summarise(rows, a.identical)
        print("\nran on both arms: %d of %d | identical plan: %d"
              % (s["n"], len(prs), s["identical"]))
        print("VERDICT: %s" % verdict(s, a.identical))
        with open(ROOT / "knob_corpus_last.json", "w", encoding="utf-8") as fh:
            json.dump(rows, fh, indent=1)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
